# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР4. В текст отчёта не входит."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT.parent / "ЛР3" / "ЛР3_Отчет_Буренков_СТЗ.docx"
OUT = ROOT / "ЛР4_Отчет_Буренков_СТЗ.docx"
RES = ROOT / "results"
BLACK = RGBColor(0, 0, 0)

HEADINGS = {
    "ВВЕДЕНИЕ": "PV",
    "1 Объект и исходное изображение": "P1",
    "2 Границы методом Кэнни": "P2",
    "3 Контуры детали": "P3",
    "4 Ключевые точки": "P4",
    "5 Цвет и текстура": "P5",
    "6 Заливка связной области": "P6",
    "7 Сравнение и признаки для разметки": "P7",
    "ЗАКЛЮЧЕНИЕ": "PZ",
}


def rfonts(run, name):
    run.font.name = name
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def style_run(run, size=14, bold=False, font="Times New Roman"):
    rfonts(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def set_spacing(paragraph, before=0, after=0, line=360):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line / 240.0


def set_indent(paragraph, first=None, left=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))
    if left is not None:
        ind.set(qn("w:left"), str(left))


def add_text(doc, text, *, first=709, before=0, after=0, line=360, center=False, bold=False, size=14):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before, after, line)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if not center and first:
        set_indent(paragraph, first=first)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold)
    return paragraph


def add_heading(doc, text, *, center=False, break_before=False):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=12, line=360)
    paragraph.paragraph_format.keep_with_next = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph, first=709)
    if break_before:
        run_br = paragraph.add_run()
        style_run(run_br, bold=True)
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_br._r.append(br)
    run = paragraph.add_run(text)
    style_run(run, bold=True)
    return paragraph


def add_toc(doc, title, token):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, line=360)
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    ppr.append(tabs)
    run = paragraph.add_run(title)
    style_run(run)
    paragraph.add_run().add_tab()
    num = paragraph.add_run(token)
    style_run(num)


def add_code(doc, line):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_indent(paragraph, first=0, left=420)
    run = paragraph.add_run(line if line else " ")
    style_run(run, size=10, font="Courier New")
    return paragraph


def set_cell(cell, text, *, bold=False, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_spacing(paragraph, before=2, after=2, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    style_run(run, size=12, bold=bold)
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    if header:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        tcpr.append(shd)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)


def add_table(doc, caption, headers, rows, widths):
    inset = max(0, (9355 - sum(widths)) // 2)
    add_text(doc, caption, first=0, before=6, after=6, line=240)
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = True
    set_indent(cap, first=0, left=inset)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, header=True)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], value)


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=6, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    add_text(doc, caption, first=0, before=6, after=10, line=360, center=True)


def replace_title(doc):
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 4",
        "Предобработка, фильтрация, преобразование изображений и видеокадров на производстве":
            "Обнаружение и выделение признаков на изображениях и видеокадрах для решения производственных задач",
    }
    found = set()
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
            found.add(node.text)
    if found != set(replacements.values()):
        raise SystemExit(f"Титул заменён не полностью: {found}")


def fill_pages(path):
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(path))
    doc.Repaginate()
    pages = {}
    for para in doc.Paragraphs:
        text = para.Range.Text.replace("\r", "").replace("\x07", "").strip()
        if text in HEADINGS:
            pages[text] = int(para.Range.Information(3))
    doc.Close(False)
    word.Quit()
    return pages


def write_page_numbers(path, pages):
    doc = Document(str(path))
    tokens = {token: str(pages[title]) for title, token in HEADINGS.items()}
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in tokens:
            node.text = tokens[node.text]
    doc.save(str(path))


def build():
    doc = Document(str(SRC))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replace_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ", center=True)
    for title, token in HEADINGS.items():
        add_toc(doc, title, token)

    add_heading(doc, "ВВЕДЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "Цель работы — научиться выделять признаки объекта на кадре: границы, контур, ключевые точки, цвет и текстуру. Отдельным этапом реализована заливка связной области. Она отделяет пиксели детали от фона до разметки и обучения сети.",
    )
    add_text(
        doc,
        "Общая часть одинакова для всех вариантов: детектор Кэнни, характеристики контура, точки ORB на исходном и повёрнутом кадре, цветовая гистограмма и HOG. Вариант 1 добавляет заливку с самостоятельной проверкой нижнего смежного пикселя.",
    )

    add_heading(doc, "1 Объект и исходное изображение", break_before=True)
    add_text(
        doc,
        "Вариант 1. Производственный объект — стальной глухарь: крепёжная деталь с шестигранной головкой и резьбой по дереву. На головке читаются грани, на стержне — витки и остриё. Фон кадра чёрный, поэтому силуэт отделяется от поля по яркости.",
    )
    add_text(
        doc,
        "Исходный кадр имеет размер 952×492 пикселя, три канала, тип uint8. Для расчётов длинная сторона приведена к 960 пикселям с сохранением пропорций, рабочий кадр — 960×496. Масштаб 1,008, форма детали не растянута. Параметры объектива неизвестны, дисторсия не исправлялась. Кадр показан на рисунке 1.",
    )
    add_picture(doc, RES / "original.jpg", "Рисунок 1 — Исходный кадр глухаря", 15.2)
    add_text(
        doc,
        "Перевод в оттенки серого для границ и точек выполнен по чувствительности глаза: Y = 0,2126·R + 0,7152·G + 0,0722·B. Дальше пороги и детекторы работают с этим полутоновым кадром.",
    )

    add_heading(doc, "2 Границы методом Кэнни", break_before=True)
    add_text(
        doc,
        "Детектор Кэнни сглаживает кадр, считает градиент яркости по двум осям, подавляет немаксимумы и оставляет пиксель границы, если градиент проходит между нижним и верхним порогом. Сравнены две пары порогов (рисунок 2, таблица 1).",
    )
    add_picture(doc, RES / "canny.png", "Рисунок 2 — Границы при порогах 40/110 и 90/180", 16.2)
    add_table(
        doc,
        "Таблица 1 — Число пикселей границы",
        ["Нижний порог", "Верхний порог", "Пикселей границы", "Что остаётся на кадре"],
        [
            ["40", "110", "26265", "Силуэт, витки и рябь на стержне"],
            ["90", "180", "13280", "Кромка головки и витки резьбы"],
        ],
        [1700, 1900, 2300, 3455],
    )
    add_text(
        doc,
        "Пара 90 и 180 уменьшает число пикселей границы почти вдвое: с 26265 до 13280. На стержне пропадает мелкая рябь покрытия, контур головки и периодические кромки витков сохраняются. Для описания формы детали взята эта пара. Пара 40 и 110 полезна, если нужно смотреть текстуру поверхности, а не силуэт.",
        before=8,
    )

    add_heading(doc, "3 Контуры детали", break_before=True)
    add_text(
        doc,
        "Полутоновый кадр бинаризован порогом 60: металл светлее фона. Внешний контур найден по этой маске, контуры площадью меньше 300 пикселей отброшены. На кадре остался один контур — сам глухарь (рисунок 3, таблица 2). Зелёная линия идёт по пикселям силуэта, красный прямоугольник построен по этой линии и целиком её вмещает.",
    )
    add_picture(doc, RES / "contours.png", "Рисунок 3 — Контур глухаря и ограничивающий прямоугольник", 15.2)
    add_table(
        doc,
        "Таблица 2 — Характеристики контура",
        ["Площадь, пикс.", "Периметр, пикс.", "x", "y", "Ширина", "Высота", "Выпуклость"],
        [["143714,5", "2442,0", "33", "58", "914", "306", "0,729"]],
        [1600, 1600, 900, 900, 1300, 1300, 1755],
    )
    add_text(
        doc,
        "Площадь 143714,5 пикселя посчитана по зелёному контуру, а не по площади прямоугольника 914×306. Выпуклость 0,729 — отношение площади контура к площади выпуклой оболочки. Впадины резьбы в оболочку не входят, поэтому отношение меньше единицы. Это признак развитой кромки: у гладкого стержня без витков выпуклость была бы ближе к единице.",
        before=8,
    )

    add_heading(doc, "4 Ключевые точки", break_before=True)
    add_text(
        doc,
        "Ключевая точка должна повторяться при сдвиге, повороте и смене освещения. На кадре это углы шестигранника и изломы витков. Детектор ORB ограничен 120 точками. Чтобы поворот на 25° не обрезал остриё и головку, кадр вписан в чёрный квадрат со стороной 1089 пикселей. Точки на исходном и повёрнутом кадре показаны на рисунке 4.",
    )
    add_picture(doc, RES / "orb.png", "Рисунок 4 — Точки ORB до поворота и после поворота на 25°", 16.2)
    add_text(
        doc,
        "На обоих кадрах найдено по 120 точек. Сопоставление дескрипторов по расстоянию Хэмминга дало 73 пары. Для каждой пары исходная точка повёрнута на те же 25° известным преобразованием и сравнена с найденной точкой. У 71 пары расхождение меньше 12 пикселей. Точки сидят на витках и на гранях головки. Гладкий участок стержня почти пуст: там нет устойчивого угла.",
    )
    add_text(
        doc,
        "Для сравнения на том же полутоновом кадре посчитан детектор Харриса. Окно 2×2, апертура Собеля 3, коэффициент k = 0,04. Ответ R = det(M) − k·(trace(M))². Оставлены локальные максимумы в окне 7×7, у которых R выше 0,04 от максимума по кадру. Таких углов 97 (рисунок 5). Они тоже собираются на гребнях резьбы и на кромке головки. Единичные отклики на чёрном поле — шум порога, в признак детали они не входят.",
    )
    add_picture(doc, RES / "harris.png", "Рисунок 5 — Углы, найденные детектором Харриса", 15.2)

    add_heading(doc, "5 Цвет и текстура", break_before=True)
    add_text(
        doc,
        "Цвет посчитан по пикселям внутри контура, чёрное поле в гистограмму не входит (рисунок 6). Средние значения каналов: B = 143,7, G = 147,1, R = 154,1. Металл чуть теплее нейтрального серого. Основная масса пикселей лежит около уровня 190. Узкий пик у нуля дают тёмные впадины под головкой, они остаются внутри внешнего контура.",
    )
    add_picture(doc, RES / "hist_rgb.png", "Рисунок 6 — Гистограмма каналов B, G и R внутри контура", 14.5)
    add_text(
        doc,
        "HOG построен по ограничивающему прямоугольнику, приведённому к 256×256. Ячейка 16×16, 9 корзин направления на полукруг. Штрих на рисунке 7 показывает направление сильнейшей корзины. Вдоль головки штрихи повторяют её контур, на резьбе они идут поперёк витков. На чёрных клетках поля направление неустойчиво: градиента там нет. Для разметки берётся рисунок штрихов на металле.",
    )
    add_picture(doc, RES / "hog.png", "Рисунок 7 — Область детали и её HOG-представление", 15.5)

    add_heading(doc, "6 Заливка связной области", break_before=True)
    add_text(
        doc,
        "Заливка собирает связную область из пикселей, близких по цвету. Начальная точка выбрана внутри детали, на головке, с координатами (98; 73). Карта данных отмечает уже просмотренные пиксели. В очередь попадает сосед, если он ещё не отмечен и по каждому каналу отличается от текущего пикселя не больше чем на 48 уровней. Допуск подобран так, чтобы область выросла по металлу и не вышла на чёрный фон.",
    )
    add_text(
        doc,
        "По варианту 1 отдельно записана проверка нижнего соседа. Проверки верхнего, левого и правого соседа устроены так же: другая координата и та же граница кадра. Основной цикл снимает пиксель с головы очереди и ставит в хвост принятых соседей. Цикл кончается на пустой очереди.",
    )
    for line in [
        "def check_bottom_neighbor(field, img, row, col, n_rows, lo, up):",
        "    if row + 1 >= n_rows:",
        "        return None",
        "    if field[row + 1, col] == 1:",
        "        return None",
        "    if close_color(img[row, col], img[row + 1, col], lo, up):",
        "        field[row + 1, col] = 1",
        "        return (row + 1, col)",
        "    return None",
        "",
        "queue = deque([(seed_row, seed_col)])",
        "while queue:",
        "    row, col = queue.popleft()",
        "    for nxt in (",
        "        check_bottom_neighbor(field, img, row, col, n_rows, lo, up),",
        "        check_top_neighbor(field, img, row, col, lo, up),",
        "        check_left_neighbor(field, img, row, col, lo, up),",
        "        check_right_neighbor(field, img, row, col, n_cols, lo, up)):",
        "        if nxt is not None:",
        "            queue.append(nxt)",
    ]:
        add_code(doc, line)
    add_text(
        doc,
        "Та же точка и тот же допуск ±48 переданы во встроенную функцию floodFill. Своя заливка и встроенная покрыли одинаковые 143754 пикселя, доля совпадения равна 1. Результат своей заливки показан на рисунке 8: слева начальная точка, справа залитая область.",
        before=8,
    )
    add_picture(doc, RES / "flood.png", "Рисунок 8 — Начальная точка заливки и залитая область", 16.2)

    add_heading(doc, "7 Сравнение и признаки для разметки", break_before=True)
    add_text(
        doc,
        "Площадь контура 143714,5 пикселя, заливка покрыла 143754 пикселя. Пересечение с силуэтом контура даёт меру IoU 0,993. Разница в несколько десятков пикселей возникает потому, что площадь контура считается по ломаной, а заливка считает клетки маски. Форма совпадает: головка, стержень, витки и остриё есть и на зелёном контуре, и на жёлтой маске.",
    )
    add_text(
        doc,
        "Для этого кадра форму детали точнее задаёт контур. Он сразу даёт периметр, прямоугольник и выпуклость, а заливка при допуске 48 повторяет тот же силуэт: IoU равен 0,993. Пороги Кэнни 90 и 180 оставляют для признака «кромка витка». Набор для разметки в следующей работе собран в таблице 3.",
    )
    add_table(
        doc,
        "Таблица 3 — Признаки глухаря для последующей разметки",
        ["Признак", "Значение на кадре", "Зачем он в разметке"],
        [
            ["Площадь контура", "143714,5 пикс.", "Отличить деталь от мелкого мусора"],
            ["Периметр", "2442,0 пикс.", "Оценить развитость резьбы"],
            ["Прямоугольник", "914×306, левый верх (33; 58)", "Положение детали на кадре"],
            ["Выпуклость", "0,729", "Впадины резьбы против гладкого стержня"],
            ["ORB при повороте 25°", "71 из 73 пар ближе 12 пикс.", "Узнать тот же крепёж на соседнем кадре"],
            ["HOG области", "штрихи вдоль головки и поперёк витков", "Текстура кромки при разметке"],
            ["Средний цвет B, G, R", "143,7; 147,1; 154,1", "Отделить металл от чёрного фона"],
        ],
        [2800, 3300, 3255],
    )

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "На кадре стального глухаря выделены признаки для дальнейшей разметки. Детектор Кэнни с порогами 90 и 180 оставляет кромку головки и витки: 13280 пикселей границы против 26265 при порогах 40 и 110. Внешний контур при пороге яркости 60 имеет площадь 143714,5, периметр 2442,0 и выпуклость 0,729. ORB даёт 120 точек, после поворота на 25° совпали 71 из 73 пар. Харрис на том же кадре отмечает 97 углов на резьбе и головке. Заливка с допуском 48 и проверкой нижнего соседа покрыла 143754 пикселя и полностью совпала со встроенной функцией. Маска заливки и контур описывают одну и ту же деталь, IoU равен 0,993.",
    )

    doc.save(str(OUT))
    pages = fill_pages(OUT)
    missing = [title for title in HEADINGS if title not in pages]
    if missing:
        raise SystemExit(f"Нет страниц у разделов: {missing}")
    write_page_numbers(OUT, pages)
    print(pages)
    print(OUT)


if __name__ == "__main__":
    build()
