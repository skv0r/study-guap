# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР3. В текст отчёта не входит."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT.parent / "ЛР2" / "ЛР2_Отчет_Буренков_СТЗ.docx"
OUT = ROOT / "ЛР3_Отчет_Буренков_СТЗ.docx"
RES = ROOT / "results"
BLACK = RGBColor(0, 0, 0)


def rfonts(run, name):
    run.font.name = name
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def style_run(run, size=14, bold=False, font="Times New Roman"):
    rfonts(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def set_spacing(paragraph, before=0, after=0, line=360):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line / 240.0


def set_indent(paragraph, first=None, left=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))
    if left is not None:
        ind.set(qn("w:left"), str(left))


def add_text(doc, text, *, first=709, before=0, after=0, line=360, center=False, bold=False, size=14):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before, after, line)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if not center and first:
        set_indent(paragraph, first=first)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold)
    return paragraph


def add_heading(doc, text, *, center=False, break_before=False):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=12, line=360)
    paragraph.paragraph_format.keep_with_next = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph, first=709)
    if break_before:
        run_br = paragraph.add_run()
        style_run(run_br, bold=True)
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_br._r.append(br)
    run = paragraph.add_run(text)
    style_run(run, bold=True)
    return paragraph


def add_toc(doc, title, token):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, line=360)
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    ppr.append(tabs)
    run = paragraph.add_run(title)
    style_run(run)
    paragraph.add_run().add_tab()
    num = paragraph.add_run(token)
    style_run(num)


def set_cell(cell, text, *, bold=False, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_spacing(paragraph, before=2, after=2, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    style_run(run, size=12, bold=bold)
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    if header:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        tcpr.append(shd)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)


def add_table(doc, caption, headers, rows, widths):
    inset = max(0, (9355 - sum(widths)) // 2)
    add_text(doc, caption, first=0, before=6, after=6, line=240)
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = True
    set_indent(cap, first=0, left=inset)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, header=True)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], value)


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=6, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    add_text(doc, caption, first=0, before=6, after=10, line=360, center=True)


def replace_title(doc):
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3",
        "Подключение и настройка видеокамеры. Захват изображений и видеокадров на производстве":
            "Предобработка, фильтрация, преобразование изображений и видеокадров на производстве",
    }
    found = set()
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
            found.add(node.text)
    if found != set(replacements.values()):
        raise SystemExit(f"Титул заменён не полностью: {found}")


def build():
    doc = Document(str(SRC))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replace_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ", center=True)
    toc = [
        ("ВВЕДЕНИЕ", "PV"),
        ("1 Объект и исходные данные", "P1"),
        ("2 Параметры кадра и гистограмма", "P2"),
        ("3 Цветовые пространства", "P3"),
        ("4 Коррекция яркости и контраста", "P4"),
        ("5 Сравнение фильтров", "P5"),
        ("6 Приведение к размеру 640×640", "P6"),
        ("7 Итоговый конвейер", "P7"),
        ("ЗАКЛЮЧЕНИЕ", "PZ"),
    ]
    for title, token in toc:
        add_toc(doc, title, token)

    add_heading(doc, "ВВЕДЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "Цель работы — освоить предварительную обработку кадра перед обнаружением объекта: оценить яркость, сравнить цветовые представления, коррекцию контраста и сглаживающие фильтры, затем привести кадры к одному размеру без искажения формы.",
    )
    add_text(
        doc,
        "Производственный объект тот же, что во второй лабораторной работе: оператор в наушниках. Съёмка для этой работы взята с веб-камеры на рабочем месте. Для следующих работ полезен контур головы и наушников на фоне помещения.",
    )

    add_heading(doc, "1 Объект и исходные данные", break_before=True)
    add_text(
        doc,
        "Исходная последовательность снята с веб-камеры. В ролике 597 кадров, частота 25 кадр/с, длительность 23,9 с. Разрешение кадра 640×480 пикселей, три цветовых канала, тип данных uint8. В поле зрения оператор в наушниках в профиль, за ним шкаф, столешница и плита. Стена и потолок светлые, наушники и волосы заметно темнее фона.",
    )
    add_text(
        doc,
        "Из ролика с шагом 11 кадров сохранены 55 изображений. Номер в имени файла совпадает с номером кадра в ролике, поэтому положение кадра во времени восстанавливается. Параметры калибровки объектива неизвестны, исправление дисторсии не выполнялось.",
    )

    add_heading(doc, "2 Параметры кадра и гистограмма", break_before=True)
    add_text(
        doc,
        "Дальнейшие сравнения сделаны по кадру № 297, середина ролика (рисунок 1). Ширина 640 пикселей, высота 480 пикселей, каналов 3, тип uint8. Частота потока 25 кадр/с, поэтому на один кадр приходится около 40 мс, если обработку нужно успевать без пропуска.",
    )
    add_picture(doc, RES / "original.jpg", "Рисунок 1 — Исходный кадр № 297", 11)
    add_picture(doc, RES / "hist_original.png", "Рисунок 2 — Гистограмма яркости кадра № 297", 14)
    add_text(
        doc,
        "Гистограмма яркости приведена на рисунке 2. Средняя яркость 154,18 уровня, среднеквадратическое отклонение 60,80. Пятый перцентиль равен 45, девяносто пятый — 243. Масса пикселей смещена к светлым уровням: это стена и потолок, у 245 есть узкий пик. Тёмный хвост гистограммы дают наушники и волосы. Кадр в целом светлый. Мешает перепад между ярким фоном и тёмной головой, а не общая недодержка.",
    )

    add_heading(doc, "3 Цветовые пространства", break_before=True)
    add_text(
        doc,
        "Кадр переведён в оттенки серого и в HSV (рисунок 3). На сером изображении и на канале яркости V читаются контур головы, дужка наушников, шкаф и плита. Канал тона шумный и устойчивой маски лица не даёт. Насыщенность выше на полках и на лице, чем на светлой стене.",
    )
    add_picture(doc, RES / "color_spaces.png", "Рисунок 3 — Оттенки серого и каналы HSV", 15.5)
    add_text(
        doc,
        "Главный признак оператора на этом кадре — контур по яркости, а не тон. Цвет при этом оставлен: плита, дерево шкафа и чёрные наушники в кадре различаются. В набор для следующих работ кадр сохраняется с тремя каналами. Коррекция контраста выполняется по яркости, чтобы не перекрашивать сцену.",
    )

    add_heading(doc, "4 Коррекция яркости и контраста", break_before=True)
    add_text(
        doc,
        "Сравнены три способа (рисунок 4, таблица 1). Линейное преобразование взято с коэффициентом контраста 1,15 и сдвигом яркости 10. Гамма-коррекция выполнена при γ = 0,85. CLAHE применён к каналу L пространства Lab с ограничением контраста 2,0 и сеткой 8×8, после чего кадр возвращён в цвет.",
    )
    add_picture(doc, RES / "correction.png", "Рисунок 4 — Исходный кадр и три способа коррекции", 15.5)
    add_table(
        doc,
        "Таблица 1 — Яркость кадра после коррекции",
        ["Способ", "Среднее", "СКО", "95-й перцентиль"],
        [
            ["Исходный кадр", "154,18", "60,80", "243"],
            ["Линейная, α = 1,15, β = 10", "182,11", "63,37", "255"],
            ["Гамма, γ = 0,85", "163,82", "57,17", "244"],
            ["CLAHE, 2,0, сетка 8×8", "154,58", "57,19", "238"],
        ],
        [3600, 1800, 1800, 2155],
    )
    add_text(
        doc,
        "Линейная коррекция осветляет весь кадр одинаково и обрезает светлую стену: девяносто пятый перцентиль доходит до 255. Гамма 0,85 немного поднимает тени, разброс яркости при этом падает. CLAHE оставляет среднее почти исходным и усиливает локальный контраст волос и кромки наушников, не заливая стену белым. Для конвейера выбран CLAHE.",
        before=8,
    )

    add_heading(doc, "5 Сравнение фильтров", break_before=True)
    add_text(
        doc,
        "На исходном кадре сравнены гауссов фильтр с ядром 5×5, медианный фильтр с ядром 5×5 и билатеральный фильтр с диаметром 9 и σ = 75 (рисунок 5). Зернистости и отдельных выбитых точек на кадре нет, поэтому три результата почти не отличаются от исходного. Включать сглаживание в конвейер из-за такого кадра нет причины: границы наушников и так чёткие.",
    )
    add_picture(doc, RES / "filters_clean.png", "Рисунок 5 — Фильтры на исходном кадре", 15.5)
    add_text(
        doc,
        "Чтобы увидеть разницу методов, на копию кадра добавлен импульсный шум: около 1,5 % пикселей заменены на чёрные и белые точки (рисунок 6). Эта копия в итоговый набор не входит. Гауссов фильтр размазывает точки и смягчает контур. Билатеральный фильтр часть точек оставляет. Медианный фильтр 5×5 убирает точки и сохраняет границу наушников.",
    )
    add_picture(doc, RES / "filters_noise.png", "Рисунок 6 — Фильтры на копии с импульсным шумом", 15.5)
    add_text(
        doc,
        "Время измерено на исходном кадре 640×480, по 30 повторам каждого вызова. В таблицу 2 записано среднее.",
    )
    add_table(
        doc,
        "Таблица 2 — Среднее время фильтрации",
        ["Метод", "Параметры", "Среднее время, мс"],
        [
            ["Гауссов фильтр", "ядро 5×5", "0,090"],
            ["Медианный фильтр", "ядро 5×5", "1,108"],
            ["Билатеральный фильтр", "d = 9, σ = 75", "4,447"],
        ],
        [3400, 3200, 2755],
    )
    add_text(
        doc,
        "При 25 кадр/с бюджет кадра около 40 мс. Все три фильтра в него укладываются, билатеральный самый медленный. На реальных кадрах импульсного шума нет, поэтому в конвейер фильтр не входит. Если такой шум появится, выбирается медианный фильтр 5×5: он убирает точки и не размывает контур так, как гауссов.",
        before=8,
    )

    add_heading(doc, "6 Приведение к размеру 640×640", break_before=True)
    add_text(
        doc,
        "Кадр приводится к квадрату 640×640 без растяжения (рисунок 7). Ширина уже равна 640, поэтому масштаб равен 1, высота остаётся 480 пикселей. Свободные полосы сверху и снизу по 80 пикселей заполняются значением 114. Лицо и наушники не вытягиваются по вертикали.",
    )
    add_picture(doc, RES / "prepared_640.jpg", "Рисунок 7 — Кадр № 297 после CLAHE и дополнения до 640×640", 11)

    add_heading(doc, "7 Итоговый конвейер", break_before=True)
    add_text(
        doc,
        "В конвейер вошли только те шаги, которые меняют пригодность кадра (таблица 3). Сглаживание исключено. Последовательность: загрузка кадра, CLAHE по каналу яркости, дополнение до 640×640. Цвет сохраняется. Среднее время этой цепочки на кадре 640×480 — 1,695 мс при 30 повторах.",
    )
    add_table(
        doc,
        "Таблица 3 — Параметры итогового конвейера",
        ["Этап", "Параметры", "Назначение"],
        [
            ["Исходный кадр", "640×480, 3 канала, uint8, 25 кадр/с", "Кадр с веб-камеры"],
            ["Коррекция", "CLAHE по L, ограничение 2,0, сетка 8×8", "Контраст головы на светлой стене"],
            ["Фильтрация", "не применяется", "На кадре нет шума, который она убирает"],
            ["Размер", "масштаб 1, поля по 80 пикселей, значение 114", "Квадрат 640×640 без растяжения"],
        ],
        [2200, 4000, 3155],
    )
    add_text(
        doc,
        "Теми же параметрами обработаны все 55 кадров. На рисунке 8 сверху исходные кадры № 0, № 297 и № 594, снизу — результат. Положение головы меняется вдоль ролика, пропорции сохраняются, серые поля одинаковые.",
        before=8,
    )
    add_picture(doc, RES / "pairs.png", "Рисунок 8 — Три кадра до обработки и после неё", 15.2)

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "С веб-камеры снят ролик оператора в наушниках длительностью 23,9 с. Контур головы читается по яркости, цвет кадра при этом сохранён. Контраст выравнивается через CLAHE по яркости, с ограничением 2,0 и сеткой 8×8. Линейная коррекция обрезает светлую стену до уровня 255. Гауссов, медианный и билатеральный фильтры на исходном кадре почти ничего не меняют и в конвейер не входят. Кадр дополняется до 640×640 со значением 114, форма объекта не растягивается. Среднее время цепочки 1,695 мс, при 25 кадр/с запас на обнаружение остаётся. Эти 55 кадров и те же параметры используются в следующих работах.",
    )

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
