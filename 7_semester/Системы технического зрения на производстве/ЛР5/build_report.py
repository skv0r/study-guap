# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР5. В текст отчёта не входит."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SRC = ROOT.parent / "ЛР4" / "ЛР4_Отчет_Буренков_СТЗ.docx"
OUT = ROOT / "ЛР5_Отчет_Буренков_СТЗ.docx"
RES = ROOT / "results"
BLACK = RGBColor(0, 0, 0)

HEADINGS = {
    "ВВЕДЕНИЕ": "PV",
    "1 Вариант, объект и классы": "P1",
    "2 Набор данных": "P2",
    "3 Размеченные кадры": "P3",
    "4 Выборки": "P4",
    "5 Правила разметки": "P5",
    "6 Проверка малых рамок": "P6",
    "7 Анализ и версия набора": "P7",
    "ЗАКЛЮЧЕНИЕ": "PZ",
}


def rfonts(run, name):
    run.font.name = name
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def style_run(run, size=14, bold=False, font="Times New Roman"):
    rfonts(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def set_spacing(paragraph, before=0, after=0, line=360):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line / 240.0


def set_indent(paragraph, first=None, left=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))
    if left is not None:
        ind.set(qn("w:left"), str(left))


def add_text(doc, text, *, first=709, before=0, after=0, line=360, center=False, bold=False, size=14):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before, after, line)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    if not center and first:
        set_indent(paragraph, first=first)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold)
    return paragraph


def add_heading(doc, text, *, center=False, break_before=False):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=12, line=360)
    paragraph.paragraph_format.keep_with_next = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph, first=709)
    if break_before:
        run_br = paragraph.add_run()
        style_run(run_br, bold=True)
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_br._r.append(br)
    run = paragraph.add_run(text)
    style_run(run, bold=True)
    return paragraph


def add_toc(doc, title, token):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, line=360)
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    ppr.append(tabs)
    run = paragraph.add_run(title)
    style_run(run)
    paragraph.add_run().add_tab()
    num = paragraph.add_run(token)
    style_run(num)


def add_code(doc, line):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_indent(paragraph, first=0, left=420)
    run = paragraph.add_run(line if line else " ")
    style_run(run, size=10, font="Courier New")
    return paragraph


def set_cell(cell, text, *, bold=False, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_spacing(paragraph, before=2, after=2, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    style_run(run, size=12, bold=bold)
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    if header:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        tcpr.append(shd)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)


def add_table(doc, caption, headers, rows, widths):
    inset = max(0, (9355 - sum(widths)) // 2)
    add_text(doc, caption, first=0, before=6, after=6, line=240)
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = True
    set_indent(cap, first=0, left=inset)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, header=True)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], value)


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=6, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    add_text(doc, caption, first=0, before=6, after=10, line=360, center=True)


def replace_title(doc):
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 4": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 5",
        "Обнаружение и выделение признаков на изображениях и видеокадрах для решения производственных задач":
            "Разметка изображений и видеокадров производственных объектов",
    }
    found = set()
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
            found.add(node.text)
    if found != set(replacements.values()):
        raise SystemExit(f"Титул заменён не полностью: {found}")


def fill_pages(path):
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(path))
    doc.Repaginate()
    pages = {}
    for para in doc.Paragraphs:
        text = para.Range.Text.replace("\r", "").replace("\x07", "").strip()
        if text in HEADINGS:
            pages[text] = int(para.Range.Information(3))
    doc.Close(False)
    word.Quit()
    return pages


def write_page_numbers(path, pages):
    doc = Document(str(path))
    tokens = {token: str(pages[title]) for title, token in HEADINGS.items()}
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in tokens:
            node.text = tokens[node.text]
    doc.save(str(path))


def build():
    doc = Document(str(SRC))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replace_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ", center=True)
    for title, token in HEADINGS.items():
        add_toc(doc, title, token)

    add_heading(doc, "ВВЕДЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "Цель работы — собрать размеченный набор кадров производственных деталей, записать рамки в формате YOLO, разбить набор на обучающую, валидационную и тестовую выборки и проверить качество разметки программой. Проверяется слишком малая рамка: площадь w·h меньше заданной доли кадра.",
    )
    add_text(
        doc,
        "Общая часть одинакова для всех вариантов: не меньше сорока кадров, у каждого класса не меньше десяти экземпляров, правила разметки, нормированные рамки и разбиение 70/20/10. Вариант 1 добавляет восстановление рамок на кадре и проверку малых рамок.",
    )

    add_heading(doc, "1 Вариант, объект и классы", break_before=True)
    add_text(
        doc,
        "Вариант 1. Производственный объект — детали крепежа и зубчатая звёздочка, те классы, которые дальше размечают на конвейере. Классов три. gear — зубчатая звёздочка: зубцы по окружности и отверстия. washer — плоская шайба с отверстием. nut — шестигранная гайка. Имена классов записаны латиницей, как в строке аннотации.",
    )
    add_text(
        doc,
        "Индивидуальное задание — алгоритм 1, проверка (а). Порог площади MIN_AREA = 0,001, то есть одна тысячная площади кадра. Рамка меньше порога чаще всего появляется от случайного щелчка и в обучение не годится. Остальные проверки — выход за кадр, пустой файл и дубликат — для этого варианта не обязательны.",
    )

    add_heading(doc, "2 Набор данных", break_before=True)
    add_text(
        doc,
        "Источник — фотографии трёх деталей: звёздочки, шайбы и гайки. Из них собрано 50 кадров размером 640×480 пикселей. Кадры разложены в пять серий по десять: video01–video05. Серия держит один характер фона, кадры внутри серии отличаются положением деталей. Имя файла начинается с имени серии, поэтому серию можно отделить и при другом разбиении.",
    )
    add_text(
        doc,
        "Условия от кадра к кадру меняются. Фон — поле с градиентом и лёгким шумом, цвет и яркость разные. Деталь поворачивается на произвольный угол и ставится в новом масштабе. На восьми кадрах звёздочка перекрывает часть гайки: video01_0001, video01_0009, video01_0010, video02_0007, video03_0005, video04_0003, video05_0001 и video05_0009. Пустых кадров нет.",
    )

    add_heading(doc, "3 Размеченные кадры", break_before=True)
    add_text(
        doc,
        "Тип разметки — обнаружение объектов. В проекте три класса: gear, washer, nut. Строка аннотации хранит номер класса и рамку в долях кадра: xc, yc, w, h. Центр и размеры нормированы на ширину и высоту, поэтому одна и та же запись годится для кадра другого размера. На рисунке 1 программа восстановила рамки из файлов аннотаций и подписала класс. Слева кадр с перекрытием, в центре звёздочка и шайба, справа одна шайба.",
    )
    add_picture(
        doc,
        RES / "marked.jpg",
        "Рисунок 1 — Три кадра с рамками, восстановленными из разметки",
        16.2,
    )

    add_heading(doc, "4 Выборки", break_before=True)
    add_text(
        doc,
        "Набор перемешан и разрезан в отношении 70/20/10. Из 50 кадров получилось 35 обучающих, 10 валидационных и 5 тестовых. В каждой выборке есть все три класса. Число кадров и рамок собрано в таблице 1. Всего рамок 111: gear — 37, washer — 39, nut — 35. У каждого класса больше десяти экземпляров.",
    )
    add_table(
        doc,
        "Таблица 1 — Кадры и рамки по выборкам",
        ["Выборка", "Кадров", "gear", "washer", "nut"],
        [
            ["train", "35", "26", "27", "25"],
            ["valid", "10", "9", "8", "5"],
            ["test", "5", "2", "4", "5"],
            ["Всего", "50", "37", "39", "35"],
        ],
        [1871, 1871, 1871, 1871, 1871],
    )
    add_text(
        doc,
        "Описание экспорта YOLOv8 лежит в файле data.yaml. В нём указаны каталоги train, valid и test и те же три имени классов: 0 — gear, 1 — washer, 2 — nut.",
        before=8,
    )

    add_heading(doc, "5 Правила разметки", break_before=True)
    add_text(
        doc,
        "Правило 1. Размечается каждый объект трёх классов. На кадре video04_0002 шайба одна, и в файле одна строка. Кадров без объекта в наборе нет, пустой файл был бы пропуском.",
    )
    add_text(
        doc,
        "Правило 2. Рамка покрывает деталь целиком. Для перекрытой гайки в рамку входит и та часть, которую закрыла звёздочка, а не только видимый край.",
    )
    add_text(
        doc,
        "Правило 3. Рамка тесная: её стороны проходят по крайним пикселям детали, лишнее поле фона в рамку не берётся. На рисунке 1 это видно по шайбе справа и по звёздочке в центре.",
    )
    add_text(
        doc,
        "Правило 4. Перекрытый объект размечается так, как если бы он был виден полностью. На кадре video01_0001 звёздочка закрывает гайку. У гайки своя рамка на весь шестигранник, у звёздочки своя. Если оставить гайке только видимый кусок, сеть выучит обрезанную деталь.",
    )
    add_text(
        doc,
        "Правило 5. Имена классов записаны латиницей и совпадают со словарём в data.yaml. В строке стоит номер класса, на кадре программа подписывает имя.",
    )
    add_text(
        doc,
        "Перевод рамки из пикселей в строку YOLO: xc = (x1 + x2) / (2W), yc = (y1 + y2) / (2H), w = (x2 − x1) / W, h = (y2 − y1) / H. Обратный перевод стоит в листинге визуализации: доля умножается на ширину или высоту кадра, затем восстанавливаются левый верхний и правый нижний углы.",
    )

    add_heading(doc, "6 Проверка малых рамок", break_before=True)
    add_text(
        doc,
        "Сначала строка разбирается: в ней должно быть пять чисел, класс — из диапазона имён, все доли лежат от 0 до 1. Затем рамка рисуется на кадре. Отдельно считается площадь w·h. Если она меньше 0,001, файл попадает в список замечаний.",
    )
    for line in [
        "MIN_AREA = 0.001",
        "",
        "def check_label_line(line, n_classes):",
        "    parts = line.split()",
        "    if len(parts) != 5:",
        '        return "ошибка: должно быть 5 чисел"',
        "    cls = int(parts[0])",
        "    vals = [float(value) for value in parts[1:]]",
        "    if not 0 <= cls < n_classes:",
        '        return "ошибка: неизвестный класс"',
        "    if any(value < 0 or value > 1 for value in vals):",
        '        return "ошибка: значение вне [0, 1]"',
        '    return "ok"',
        "",
        "def draw_boxes(image, lines, names):",
        "    height, width = image.shape[:2]",
        "    for line in lines:",
        '        if check_label_line(line, len(names)) != "ok":',
        "            continue",
        "        cls, xc, yc, bw, bh = line.split()",
        "        xc, bw = float(xc) * width, float(bw) * width",
        "        yc, bh = float(yc) * height, float(bh) * height",
        "        x1, y1 = int(xc - bw / 2), int(yc - bh / 2)",
        "        x2, y2 = int(xc + bw / 2), int(yc + bh / 2)",
        "        cv2.rectangle(image, (x1, y1), (x2, y2), (40, 170, 40), 2)",
        "        cv2.putText(image, names[int(cls)], (x1, max(16, y1 - 6)),",
        "                    cv2.FONT_HERSHEY_SIMPLEX, 0.55, (40, 170, 40), 2)",
        "",
        "def audit_small(label_dir):",
        "    problems = []",
        '    for path in sorted(label_dir.glob("*.txt")):',
        "        rows = [line.split() for line in path.read_text().splitlines()",
        "                if line.strip()]",
        "        for row in rows:",
        "            bw, bh = float(row[3]), float(row[4])",
        "            if bw * bh < MIN_AREA:",
        '                problems.append((path.name, "слишком малая рамка", bw * bh))',
        "    return problems",
    ]:
        add_code(doc, line)
    add_text(
        doc,
        "Чтобы увидеть, что порог срабатывает, в копию файла video03_0001.txt добавлена строка класса nut: 2 0.500000 0.500000 0.012000 0.012000. Стороны рамки по 0,012 доли кадра, площадь 0,000144. Это меньше 0,001. На кадре 640×480 рамка занимает примерно 8 пикселей по ширине и 6 по высоте и сидит в центре (рисунок 2). Программа вернула одно замечание. Вывод записан в таблице 2.",
        before=8,
    )
    add_picture(
        doc,
        RES / "small_box.jpg",
        "Рисунок 2 — Кадр video03_0001 с добавленной мелкой рамкой в центре",
        14.0,
    )
    add_table(
        doc,
        "Таблица 2 — Вывод проверки малых рамок",
        ["Файл", "Замечание", "Площадь w·h"],
        [
            ["video03_0001.txt (копия)", "слишком малая рамка", "0,000144"],
            ["рабочий каталог", "замечаний нет", "—"],
        ],
        [3600, 3200, 2555],
    )

    add_heading(doc, "7 Анализ и версия набора", break_before=True)
    add_text(
        doc,
        "Замечание относилось только к копии. В рабочем каталоге лишней строки не было. После проверки копия не используется. Повторный просмотр рабочих файлов замечаний не дал: формат строк принят, рамок с площадью меньше 0,001 нет.",
    )
    add_text(
        doc,
        "На восьми кадрах с перекрытием рамка гайки оставлена полной. Это не ошибка проверки (а): площадь такой рамки намного больше порога, а правило 4 как раз требует размечать скрытую часть. Тесность рамок на рисунке 1 совпадает с крайними пикселями деталей.",
    )
    add_text(
        doc,
        "Классы в выборках не выпали. В тесте меньше всего звёздочек — две рамки, в валидации меньше всего гаек — пять. Для оценки после обучения этого достаточно, чтобы каждый класс хотя бы раз встретился вне обучающей выборки. Обучающая выборка несёт 26, 27 и 25 рамок, перекоса в одну деталь нет.",
    )
    add_text(
        doc,
        "Окончательная версия для лабораторной работы №6 — этот каталог после удаления мелкой рамки. В нём 50 кадров, разбиение 35 / 10 / 5 и файл data.yaml с классами gear, washer, nut. Каталоги выборок: train, valid, test. Именно эта версия идёт в следующую работу.",
    )

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "Собран набор из 50 кадров 640×480 по фотографиям звёздочки, шайбы и гайки. Классы gear, washer и nut содержат 37, 39 и 35 рамок. Кадры лежат в пяти сериях, условия меняются по углу, масштабу, яркости и фону. На восьми кадрах звёздочка перекрывает гайку, и рамка гайки записана целиком. Набор разрезан 35 / 10 / 5, описание экспорта — data.yaml. Проверка малых рамок на копии файла video03_0001 нашла рамку площадью 0,000144. В рабочем каталоге таких рамок нет. Эта версия набора подготовлена для лабораторной работы №6.",
    )

    doc.save(str(OUT))
    pages = fill_pages(OUT)
    missing = [title for title in HEADINGS if title not in pages]
    if missing:
        raise SystemExit(f"Нет страниц у разделов: {missing}")
    write_page_numbers(OUT, pages)
    print(pages)
    print(OUT)


if __name__ == "__main__":
    build()
