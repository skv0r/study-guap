#!/usr/bin/env python3
"""Отчёт ЛР1 по СТЗ на производстве: бланк ГУАП, ГОСТ 7.32 / ГОСТ 2.105."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
BLANK = BASE / "guap_blanks" / "lab.docx"
OUT = BASE / "ЛР1_Отчет_Буренков_СТЗ.docx"
PDF = BASE / "ЛР1_Отчет_Буренков_СТЗ.pdf"
FIG = BASE / "figures" / "scheme.png"

TOC_DEF = [
    ("ВВЕДЕНИЕ", "intro", 0),
    ("1 Постановка задачи", "s1", 0),
    ("2 Паспорт подсистем СТЗ", "s2", 0),
    ("3 Классификационный профиль системы", "s3", 0),
    ("4 Анализ типов данных", "s4", 0),
    ("5 Расчёт параметров информационного процесса", "s5", 0),
    ("6 Выбор интерфейса и проверка временного цикла", "s6", 0),
    ("7 Структурная схема системы", "s7", 0),
    ("ЗАКЛЮЧЕНИЕ", "conc", 0),
    ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "src", 0),
]

SEARCH = {
    "intro": "ВВЕДЕНИЕ",
    "s1": "1 Постановка задачи",
    "s2": "2 Паспорт подсистем СТЗ",
    "s3": "3 Классификационный профиль системы",
    "s4": "4 Анализ типов данных",
    "s5": "5 Расчёт параметров информационного процесса",
    "s6": "6 Выбор интерфейса и проверка временного цикла",
    "s7": "7 Структурная схема системы",
    "conc": "ЗАКЛЮЧЕНИЕ",
    "src": "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
}


def set_run_font(run, size=14, bold=False, name="Times New Roman", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(
    p,
    *,
    first_indent=True,
    align="justify",
    space_after=0,
    space_before=0,
    left_indent=0,
    line_spacing=1.5,
):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]


def add_body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p)
    set_run_font(p.add_run(text))
    return p


def _no_hyphens(p):
    pPr = p._p.get_or_add_pPr()
    el = OxmlElement("w:suppressAutoHyphens")
    el.set(qn("w:val"), "true")
    pPr.append(el)


def add_section_h(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_before=0, space_after=12)
    p.paragraph_format.keep_with_next = True
    _no_hyphens(p)
    set_run_font(p.add_run(text), bold=True)


def add_struct(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_after=18)
    p.paragraph_format.keep_with_next = True
    _no_hyphens(p)
    set_run_font(p.add_run(text), bold=True)


def add_toc_line(doc, title, page, *, indent=0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=0.75 * indent)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    p.add_run("\t")
    set_run_font(p.add_run(str(page)))


def add_caption_table(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=10, space_after=4)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def add_figure(doc, path, caption, w=16.0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=10)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(w))
    c = doc.add_paragraph()
    format_paragraph(c, first_indent=False, align="center", space_before=6, space_after=10)
    set_run_font(c.add_run(caption))


def _set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_shading(cell, fill="D9D9D9"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, headers, rows, widths_cm=None, size=10):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    usable = Cm(16.5)
    if widths_cm:
        total = sum(widths_cm)
        widths = [int(usable * (w / total)) for w in widths_cm]
    else:
        widths = [int(usable / len(headers))] * len(headers)
    grid = t._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        if i < len(widths):
            child.set("w", str(widths[i]))
    for row in t.rows:
        for cell in row.cells:
            _set_cell_border(cell)
            for i, w in enumerate(widths):
                if cell._tc is row.cells[i]._tc:
                    cell.width = Emu(w)
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        set_cell_shading(cell)
        p = cell.paragraphs[0]
        format_paragraph(p, first_indent=False, align="center", line_spacing=1.15)
        set_run_font(p.add_run(h), bold=True, size=size)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            align = "center" if j == 0 else "left"
            format_paragraph(p, first_indent=False, align=align, line_spacing=1.15)
            set_run_font(p.add_run(v), size=size)
    return t


def add_page_field(paragraph):
    run = paragraph.add_run()
    for kind, val in [("begin", None), ("instr", " PAGE "), ("end", None)]:
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    set_run_font(run)


def put_cell(cell, text, size=14):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if not p.runs:
        run = p.add_run(text)
    else:
        run = p.runs[0]
        run.text = text
    set_run_font(run, size=size)


def fill_title(doc: Document):
    for p in doc.paragraphs:
        if p.text.startswith("КАФЕДРА"):
            if len(p.runs) >= 2 and p.runs[1].text.strip("_") == "":
                p.runs[1].text = "42"
            else:
                p.clear()
                r = p.add_run("КАФЕДРА № 42")
                set_run_font(r, size=14)
        elif "20__" in p.text:
            for r in p.runs:
                if "20__" in r.text:
                    r.text = r.text.replace("20__", "2026")

    t_prep, t_work, t_stud = doc.tables
    put_cell(t_prep.rows[0].cells[0], "преподаватель")
    put_cell(t_prep.rows[0].cells[4], "М.Д. Яушкина")

    put_cell(t_work.rows[0].cells[0], "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1")
    put_cell(
        t_work.rows[1].cells[0],
        "Первичное обследование состава системы технического зрения на производстве. "
        "Рассмотрение особенностей типов данных",
    )
    put_cell(t_work.rows[2].cells[0], "по курсу: Системы технического зрения на производстве")

    put_cell(t_stud.rows[0].cells[1], "4321")
    put_cell(t_stud.rows[0].cells[5], "Г.В. Буренков")


def setup_body_section(doc: Document):
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    title_sec, body_sec = doc.sections[0], doc.sections[1]

    for sec in (title_sec, body_sec):
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)

    title_sec.footer.is_linked_to_previous = False
    for p in title_sec.footer.paragraphs:
        p.clear()

    body_sec.footer.is_linked_to_previous = False
    body_sec.header.is_linked_to_previous = False

    sect_pr = body_sec._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "2")
    sect_pr.append(pg)

    fp = body_sec.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def make_scheme():
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    FIG.parent.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(10.2, 6.4), dpi=160)
    ax.set_xlim(0, 12.3)
    ax.set_ylim(0, 7.4)
    ax.axis("off")

    items = [
        (0.4, 5.5, "Объект\nсталь, DataMatrix 12×12\nFOV 60×60 мм"),
        (3.2, 5.5, "Освещение\nтёмное поле LED VIS\nполяризаторы"),
        (6.0, 5.5, "Оптика\nобычный объектив\nбез телецентричности"),
        (8.8, 5.5, "Камера CMOS\n1600×1200 Mono12\n3,84 МБ/кадр"),
        (0.4, 2.2, "GigE Vision\n4,80 МБ/с\nLкаб = 3 м"),
        (3.2, 2.2, "Вычислитель\nсмарт-камера, DPM\nкод / No-Read"),
        (6.0, 2.2, "ПЛК / MES\nстрока кода\ntIO = 2 мс"),
        (8.8, 2.2, "Исполнитель\nNo-Read: сброс\ntисп ≈ 50 мс"),
    ]
    for x, y, text in items:
        ax.add_patch(
            FancyBboxPatch(
                (x, y),
                2.55,
                1.7,
                boxstyle="round,pad=0.03,rounding_size=0.08",
                facecolor="#f2f2f2",
                edgecolor="black",
                linewidth=1.0,
            )
        )
        ax.text(x + 1.275, y + 0.85, text, ha="center", va="center", fontsize=8.5)
    for x in (2.95, 5.75, 8.55):
        ax.annotate("", xy=(x + 0.25, 6.35), xytext=(x, 6.35), arrowprops=dict(arrowstyle="-|>", color="black", lw=1.1))
    for x in (2.95, 5.75, 8.55):
        ax.annotate("", xy=(x + 0.25, 3.05), xytext=(x, 3.05), arrowprops=dict(arrowstyle="-|>", color="black", lw=1.1))
    ax.plot([11.35, 11.7, 11.7, 1.675], [6.35, 6.35, 4.35, 4.35], color="black", lw=1.1)
    ax.annotate("", xy=(1.675, 3.9), xytext=(1.675, 4.35), arrowprops=dict(arrowstyle="-|>", color="black", lw=1.1))
    ax.text(6.0, 0.55, "Обратная связь: триггер такта (фиксация детали) → экспозиция", ha="center", fontsize=9)
    fig.tight_layout(pad=0.3)
    fig.savefig(FIG, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_body(doc: Document, toc_pages: dict):
    add_struct(doc, "СОДЕРЖАНИЕ", new_page=False)
    for title, key, ind in TOC_DEF:
        add_toc_line(doc, title, toc_pages.get(key, "…"), indent=ind)

    add_struct(doc, "ВВЕДЕНИЕ")
    add_body(
        doc,
        "Целью работы является освоение первичного обследования системы технического "
        "зрения (СТЗ) на производстве: описание состава подсистем, типов данных и "
        "расчёт базовых информационных характеристик процесса [1].",
    )
    add_body(
        doc,
        "Работа выполнена для варианта 2. Объект контроля — стальная деталь; "
        "контролируемый параметр — чтение лазерного DataMatrix 12 × 12. Расчёты "
        "выполнены по правилам методических указаний [1]. Численный подбор "
        "фокусного расстояния и глубины резкости не требуется и не выполнялся.",
    )

    add_section_h(doc, "1 Постановка задачи")
    add_body(
        doc,
        "Изделие — стальная деталь с лазерной маркировкой DataMatrix ECC200 размера "
        "12 × 12 модулей. Контролируется читаемость кода и совпадение считанного "
        "идентификатора с ожидаемым. Задача относится к функции идентификации "
        "(прослеживаемость). Система ставится на межоперационном контроле после "
        "лазерной маркировки либо перед сборкой. Пропуск непрочитанного или неверно "
        "считанного кода на этом этапе разрушает прослеживаемость партии и может "
        "привести к установке не той детали в узел.",
    )
    add_body(
        doc,
        "Количественный критерий брака. Идентификация считается неуспешной, если "
        "декодер не восстанавливает DataMatrix 12 × 12 с модулем 0,3 мм (габарит "
        "кода 3,6 × 3,6 мм) с валидным ECC либо считанная строка не совпадает с "
        "эталоном партии. Допустимое число No-Read и Misread на годной детали — 0. "
        "Особенности сцены: зеркальные блики стали и низкий контраст лазерной "
        "маркировки.",
    )
    add_body(
        doc,
        "Исходные данные варианта: поле зрения FOV = 60 × 60 мм; минимальный "
        "признак dmin = 0,3 мм (модуль кода); такт Tтакт = 0,8 с; объект неподвижен "
        "(v = 0); длина линии Lкаб = 3 м. Режим расчёта частоты съёмки — А "
        "(дискретная съёмка по такту). Проверка смаза не выполняется: деталь "
        "зафиксирована [1].",
    )

    add_section_h(doc, "2 Паспорт подсистем СТЗ")
    add_body(
        doc,
        "Требования к шести подсистемам сведены в таблицу 1. Разрешение камеры, "
        "частота съёмки и интерфейс указаны после расчёта раздела 5.",
    )
    add_caption_table(doc, "Таблица 1 — Паспорт подсистем СТЗ")
    add_table(
        doc,
        ["№", "Подсистема", "Назначение в данной системе", "Ключевой параметр", "Принятое значение / тип", "Обоснование выбора"],
        [
            [
                "1",
                "Освещение",
                "Контраст модуля DataMatrix на зеркальной стали без засветки бликами",
                "геометрия, спектр, режим",
                "тёмное поле (кольцевая низкоугловая LED VIS); кросс-поляризация; непрерывный режим",
                "Лазерная риска рассеивает свет, полированная поверхность отражает его мимо объектива. Объект неподвижен — стробоскоп не требуется",
            ],
            [
                "2",
                "Оптика",
                "Построение изображения зоны маркировки на матрице в пределах заданного FOV",
                "FOV, тип объектива, телецентричность",
                "FOV 60 × 60 мм; обычный объектив; телецентричность не требуется",
                "Задача — чтение кода, а не прецизионное измерение. f, WD и ГРИП без паспорта камеры не рассчитываются [1]",
            ],
            [
                "3",
                "Фотоприемник",
                "Цифровой кадр сцены в такте 0,8 с с различимым низкоконтрастным кодом",
                "тип матрицы, затвор; разрешение и FPS — после этапа 5",
                "CMOS, VIS, Mono12; бегущий затвор допустим; 1600 × 1200; 1,25 кадр/с",
                "Деталь неподвижна, смаз отсутствует. 12 бит — для слабоконтрастной маркировки и бликов [1]",
            ],
            [
                "4",
                "Передача данных",
                "Доставка кадра (или результата декодирования) к вычислителю / ПЛК",
                "интерфейс, длина линии",
                "GigE Vision; Lкаб = 3 м",
                "Поток 4,80 МБ/с с запасом > 30 %. USB3 Vision допустим по длине (3 м < 5 м), но запас по кабелю в цехе мал",
            ],
            [
                "5",
                "Вычислитель",
                "Локализация кода, DPM-декодирование, сравнение с эталоном, сигнал код/No-Read",
                "архитектура, ускорение",
                "смарт-камера; классическое CV (декодер DataMatrix); GPU не требуется",
                "Одна специализированная задача идентификации. tобр до 30 мс укладывается в такт 800 мс [1]",
            ],
            [
                "6",
                "Исполнительное устройство",
                "Передача идентификатора в MES; сброс детали при No-Read/Misread",
                "тип, шина, задержка",
                "Ethernet к ПЛК/MES; пневмоотбраковщик при отказе чтения; tисп ≈ 50 мс",
                "Выход — строка кода, а не только Годен/Брак. Задержка механики в TСТЗ не входит [1]",
            ],
        ],
        widths_cm=[0.9, 2.2, 3.4, 2.6, 3.2, 4.2],
        size=9,
    )

    add_section_h(doc, "3 Классификационный профиль системы")
    add_body(
        doc,
        "Класс проектируемой СТЗ по признакам методических указаний [1] приведён в таблице 2.",
    )
    add_caption_table(doc, "Таблица 2 — Класс системы")
    add_table(
        doc,
        ["Признак", "Возможные категории", "Выбор", "Обоснование"],
        [
            ["Решаемая задача", "контроль / измерение / идентификация / позиционирование", "идентификация", "Нужна строка DataMatrix и признак успешного чтения, а не размер"],
            ["Размерность данных", "1D / 2D / 2.5D / 3D", "2D", "DataMatrix — двумерный код; признак проявляется как контраст модулей в плоскости"],
            ["Спектральный диапазон", "UV / VIS / NIR-SWIR / Thermal / X-Ray", "VIS", "Лазерная маркировка различима в видимом свете при тёмном поле"],
            ["Степень интеграции", "оптический датчик / смарт-камера / PC-based / встраиваемая / распределенная", "смарт-камера", "Одна задача DPM-чтения; декодер размещается на борту камеры"],
            ["Число камер и схема обзора", "однокамерная / многокамерная / стерео / круговая / сканирующая", "однокамерная", "FOV 60 × 60 мм покрывает зону маркировки одним кадром"],
            ["Подвижность", "стационарная / роботизированная / мобильная", "стационарная", "Камера над приспособлением; деталь неподвижна в такте"],
            ["Способ освещения", "внешнее / активное контролируемое / структурированное / эмиссионное", "активное контролируемое", "Нужны тёмное поле и поляризаторы; цеховой свет даёт блики"],
            ["Характер обработки", "пороговая / классическая CV / ML / DL", "классическая CV", "Штатный декодер DataMatrix / DPM; нейросеть не требуется"],
            ["Степень участия человека", "автоматическая / автоматизированная / поддержка решений / регистрация", "автоматическая", "Чтение и отбраковка No-Read в такте без оператора"],
            ["Условия эксплуатации", "лабораторные / цеховые / агрессивные / взрывоопасные / экстремальные", "цеховые", "Механообработка/сборка; корпус камеры IP67"],
        ],
        widths_cm=[3.2, 4.6, 2.4, 6.3],
        size=9,
    )

    add_section_h(doc, "4 Анализ типов данных")
    add_body(
        doc,
        "Минимально достаточная размерность — 2D. DataMatrix по определению является "
        "двумерным кодом: одномерный профиль вдоль строки не восстанавливает сетку "
        "12 × 12 модулей произвольной ориентации. Переход к 2.5D или 3D даёт карту "
        "глубины лазерной риски, но декодирование выполняется по контрасту яркости; "
        "объём данных и стоимость растут без изменения критерия успешного чтения.",
    )
    add_body(
        doc,
        "Спектральный диапазон — видимый. Признак не требует UV, тепловизора или "
        "рентгена: контраст модуля формируется геометрией тёмного поля. Формат — "
        "Mono12 (unpacked, B = 2 байт/пикс): сортировка по цвету не задана, а "
        "маркировка слабоконтрастна, на стали остаются блики. По [1] 8 бит "
        "недостаточно, если нужно устойчиво различать перепад яркости менее 0,4 % "
        "динамического диапазона. RGB8 увеличил бы поток в 1,5 раза относительно "
        "Mono12 без новой информации о коде.",
    )
    add_body(
        doc,
        "Объёмы кадра для одной и той же сцены 1600 × 1200 пикс. приведены в таблице 3. "
        "Основное сжатие данных происходит не на передаче, а на этапе экстракции: "
        "кадр 3,84 МБ сводится к строке идентификатора и флагу успешного чтения. "
        "В смарт-камере полный кадр, как правило, не покидает устройство. "
        "Выбранное 2D Mono12 даёт 3,84 МБ/кадр: цвет увеличивает объём без информации "
        "для декодера, облако точек дороже примерно в шесть раз.",
    )
    add_caption_table(doc, "Таблица 3 — Сравнение типов данных для одной сцены")
    add_table(
        doc,
        ["Представление", "Формат", "Байт/точку", "Объём кадра, МБ", "Что даёт", "Чего не даёт"],
        [
            ["1D профиль (строка)", "Mono8", "1", "0,00160", "Профиль яркости вдоль линии", "Сетку DataMatrix вне линии"],
            ["2D полутоновое", "Mono8", "1", "1,92", "Геометрию кода при высоком контрасте", "Запас по слабому контрасту и бликам"],
            ["2D цветное", "RGB8", "3", "5,76", "Цвет поверхности", "Не требуется заданием; поток втрое больше Mono8"],
            ["2.5D карта высот", "Coord3D C16", "2", "3,84", "Рельеф лазерной риски", "Избыточно для декодирования по яркости"],
            ["3D облако точек", "Coord3D ABC32f", "12", "23,04", "Полную геометрию детали", "Не нужно декодеру; объём ×12 к Mono8"],
        ],
        widths_cm=[3.0, 2.4, 1.8, 2.2, 3.6, 3.5],
        size=9,
    )

    add_section_h(doc, "5 Расчёт параметров информационного процесса")
    add_body(
        doc,
        "Вариант относится к режиму А. Для чтения 2D-кода принято n = 6 пикселей на "
        "модуль (допустимый диапазон 5–8) [1]. Формат Mono12, B = 2 байт/пикс. "
        "Ход расчёта совпадает с таблицей 4.",
    )
    add_body(
        doc,
        "Требуемый масштабный коэффициент kтреб = dmin / n = 0,3 / 6 = 0,05 мм/пикс. "
        "Требуемое разрешение Wтреб = 60 / 0,05 = 1200 пикс, Hтреб = 60 / 0,05 = 1200 пикс. "
        "Ряд 1280 × 1024 не проходит по высоте (1024 < 1200). По учебному ряду "
        "матричных камер принимается первое значение не ниже обоих порогов: "
        "1600 × 1200 пикс.",
    )
    add_body(
        doc,
        "Фактический масштаб kx = 60 / 1600 = 0,0375 мм/пикс, ky = 60 / 1200 = 0,05 мм/пикс, "
        "k = max(kx, ky) = 0,05 мм/пикс. На модуль 0,3 мм приходится 6 пикселей, "
        "весь код 3,6 мм занимает 72 пикселя. Объём кадра V = 1600 · 1200 · 2 / 10^6 = "
        "3,84 МБ. FPSтреб = 1 / 0,8 = 1,25 кадр/с. Поток D = 3,84 · 1,25 = 4,80 МБ/с.",
    )
    add_caption_table(doc, "Таблица 4 — Расчёт параметров информационного процесса")
    add_table(
        doc,
        ["Параметр", "Обозначение", "Формула", "Результат"],
        [
            ["Требуемый масштабный коэффициент", "kтреб", "dmin / n = 0,3 / 6", "0,05 мм/пикс"],
            ["Требуемое разрешение", "Wтреб × Hтреб", "FOVx / kтреб; FOVy / kтреб", "1200 × 1200 пикс"],
            ["Принятое разрешение камеры", "W × H", "первый стандарт из ряда А.7 [1]", "1600 × 1200 пикс"],
            ["Фактический масштаб", "kx; ky; k", "FOVx/W; FOVy/H; max(kx, ky)", "0,0375; 0,05; 0,05 мм/пикс"],
            ["Объём кадра", "V", "W · H · B / 10^6, B = 2", "3,84 МБ"],
            ["Требуемая частота съёмки", "FPSтреб", "режим А: 1 / Tтакт", "1,25 кадр/с"],
            ["Поток данных", "D", "V · FPS", "4,80 МБ/с"],
            ["Время передачи кадра", "tпер", "1000 · V / Cэфф, Cэфф = 110 МБ/с", "34,9 мс"],
            ["Предельное время экспозиции", "tэксп, max", "объект неподвижен, v = 0", "не ограничено смазом"],
            ["Принятая экспозиция", "tэксп", "по свету слабоконтрастной маркировки", "10 мс"],
            ["Доступное время для решения", "Tдоп", "режим А: Tтакт", "800 мс"],
            ["Латентность СТЗ", "TСТЗ", "tэксп + tсч + tпер + tобр + tлог + tIO", "79,9 мс"],
            ["Проверка по времени", "—", "TСТЗ ≤ Tдоп", "выполнено"],
        ],
        widths_cm=[4.4, 2.6, 5.8, 3.7],
        size=9,
    )
    add_body(
        doc,
        "Учебные составляющие латентности по [1]: tсч = 2 мс, tлог = 1 мс, tIO = 2 мс; "
        "для классического CV принята верхняя оценка tобр = 30 мс. Тогда "
        "TСТЗ = 10 + 2 + 34,9 + 30 + 1 + 2 = 79,9 мс, что меньше Tдоп = 800 мс.",
    )

    add_section_h(doc, "6 Выбор интерфейса и проверка временного цикла")
    add_body(
        doc,
        "Сравнение каналов для потока 4,80 МБ/с и длины 3 м приведено в таблице 5. "
        "Запас по пропускной способности должен быть не менее 30 % [1]. Требуемый "
        "порог с запасом: 4,80 · 1,30 = 6,24 МБ/с.",
    )
    add_caption_table(doc, "Таблица 5 — Выбор интерфейса передачи данных")
    add_table(
        doc,
        ["Интерфейс", "Эффективная скорость, МБ/с", "Макс. длина", "Запас к 4,80 МБ/с", "Вывод"],
        [
            ["USB3 Vision", "350–400", "до 5 м", "многократный", "3 м < 5 м, но запас по кабелю в цехе мал"],
            ["GigE Vision (1G)", "около 110", "до 100 м", "более 20 раз", "Принят: запас и длина"],
            ["10 GigE Vision", "около 1100", "до 100 м", "избыточен", "Не требуется"],
            ["Camera Link (Full)", "около 850", "до 10 м", "избыточен", "Нужна плата захвата"],
            ["CoaXPress CXP-12", "около 1250 на канал", "до 40 м", "избыточен", "Не требуется"],
        ],
        widths_cm=[3.4, 3.4, 2.4, 3.2, 4.1],
        size=9,
    )
    add_body(
        doc,
        "Принят GigE Vision 1G. Условие по времени выполняется: 79,9 мс < 800 мс. "
        "Полный цикл с механикой отбраковки No-Read Tполн = TСТЗ + tисп = "
        "79,9 + 50 = 129,9 мс. Деталь неподвижна в приспособлении в течение такта "
        "0,8 с, поэтому условие Lисп / v не применяется. Конвейеризация алгоритма "
        "не нужна: запас по такту составляет более чем порядок. Выделение ROI вокруг "
        "ожидаемого положения кода (габарит 3,6 мм в поле 60 мм) снизило бы объём "
        "данных, но для выполнения цикла не требуется.",
    )
    add_body(
        doc,
        "Проверка смаза не выполняется: v = 0, tэксп,max не ограничено смещением "
        "0,5 пикс. Глобальный затвор и стробоскоп по условию tэксп,max < 1 мс не "
        "требуются [1]. Бегущий затвор допустим. Импульсная подсветка может быть "
        "добавлена только для подавления внешнего цехового света, а не из-за движения.",
    )

    add_section_h(doc, "7 Структурная схема системы")
    add_body(
        doc,
        "Структура СТЗ от объекта до исполнительного устройства с указанием типа "
        "данных на переходах показана на рисунке 1. Обратная связь — сигнал такта "
        "фиксации детали на запуск экспозиции.",
    )
    add_figure(doc, FIG, "Рисунок 1 — Структурная схема СТЗ чтения DataMatrix", w=16.0)
    add_body(
        doc,
        "На входе — оптическое изображение зоны маркировки. После камеры данные "
        "имеют вид кадра Mono12 объёмом 3,84 МБ. Расчётный поток по GigE Vision — "
        "4,80 МБ/с; в смарт-камере кадр обрабатывается на борту, наружу уходит "
        "строка кода и флаг чтения. Выход вычислителя — идентификатор (или No-Read) "
        "на ПЛК/MES; при отказе чтения исполнитель сбрасывает деталь. Ключевой "
        "признак СТЗ соблюдён: результатом является управляющее воздействие в "
        "заданном такте, а не изображение для оператора [1].",
    )

    add_struct(doc, "ЗАКЛЮЧЕНИЕ")
    add_body(
        doc,
        "Выявляемый признак проявляется оптически: лазерные модули DataMatrix дают "
        "контраст при освещении тёмным полем. Рентген, ультразвук или вихретоковый "
        "контроль код не читают. Если контраст маркировки окажется недостаточным "
        "даже после смены геометрии света, следует менять технологию маркировки "
        "(глубина лазера, иглоударная маркировка), а не тип датчика.",
    )
    add_body(
        doc,
        "Для данной задачи дороже ошибка II рода в форме Misread: неверный "
        "идентификатор, принятый как годный, ломает прослеживаемость и может "
        "поставить не ту деталь в сборку. No-Read годного кода (ошибка I рода) "
        "дешевле — деталь уходит на повторное чтение или ручной сканер. Порог и "
        "проверка ECC смещаются в сторону отказа при сомнении: лучше лишний "
        "No-Read, чем ложно положительное чтение.",
    )
    add_body(
        doc,
        "Риски эксплуатации: остаточные блики стали, масляная плёнка на коде, пыль "
        "на объективе, вибрации приспособления, нестабильный внешний свет. Меры: "
        "тёмное поле и поляризаторы, IP67, укрытие зоны съёмки, крепление с "
        "виброизоляцией, контроль чистоты маркировки.",
    )
    add_body(
        doc,
        "Количественный расчёт окупаемости не выполняется: в варианте нет стоимости "
        "оборудования, брака и объёма выпуска. Для срока окупаемости нужны цена СТЗ, "
        "трудозатраты ручного сканирования, стоимость перепутывания детали и годовой "
        "выпуск. Качественно внедрение оправдано прослеживаемостью, снижением "
        "ручного считывания и исключением сборки с неверной деталью.",
    )
    add_body(
        doc,
        "Принятые допущения: n = 6; учебный ряд разрешений; tэксп = 10 мс; "
        "tсч = 2 мс, tлог = 1 мс, tIO = 2 мс, tобр = 30 мс; Cэфф GigE = 110 МБ/с; "
        "tисп = 50 мс; f, WD и ГРИП не рассчитывались. Условие TСТЗ ≤ Tдоп "
        "выполняется с запасом (79,9 мс против 800 мс).",
    )

    add_struct(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    sources = [
        "Первичное обследование состава системы технического зрения на производстве. "
        "Рассмотрение особенностей типов данных : метод. указания к лабораторной работе № 1. "
        "– Санкт-Петербург : ГУАП, 2026.",
    ]
    for i, src in enumerate(sources, 1):
        p = doc.add_paragraph()
        format_paragraph(p)
        set_run_font(p.add_run(f"{i}. {src}"))


def build(toc_pages: dict) -> Path:
    make_scheme()
    shutil.copy(BLANK, OUT)
    doc = Document(str(OUT))
    fill_title(doc)
    setup_body_section(doc)
    build_body(doc, toc_pages)
    doc.save(str(OUT))
    return OUT


def convert():
    src = str(OUT.resolve())
    dst = str(PDF.resolve())
    script = f'''
tell application "Microsoft Word"
  repeat with d in (get documents)
    close d saving no
  end repeat
  delay 0.4
  set theDoc to open file name POSIX file "{src}"
  delay 0.4
  save as theDoc file name POSIX file "{dst}" file format format PDF
  close theDoc saving no
end tell
'''
    subprocess.run(["osascript", "-e", script], check=True)


def measure() -> dict[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(str(PDF))
    pages: dict[str, int] = {}
    for i, page in enumerate(reader.pages):
        if i < 2:
            continue
        text = page.extract_text() or ""
        printed = i + 1
        for key, needle in SEARCH.items():
            if key not in pages and needle in text:
                pages[key] = printed
    return pages


def main():
    build({k: "…" for k in SEARCH})
    try:
        convert()
        pages = measure()
        print("pass1", pages)
        if pages:
            build(pages)
            convert()
            print("pass2", measure())
    except Exception as exc:
        print("PDF skipped:", exc)
    print(OUT)


if __name__ == "__main__":
    main()
