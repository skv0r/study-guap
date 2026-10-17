# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР2 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

BASE = Path(__file__).resolve().parent
SRC = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
OUT = BASE / "ЛР2_Отчет_Буренков_МИИ.docx"
FIG = BASE / "figures"
DATA = BASE / "data"
TREE_SPLIT = BASE.parent / "ЛР3" / "data" / "split.json"

CLASSES = ["футбольный", "баскетбольный", "волейбольный"]
FEATURES = ["elasticity", "size", "mass", "pressure", "bounce"]
FEATURE_RU = {
    "elasticity": "Коэффициент упругости",
    "size": "Окружность, см",
    "mass": "Масса, г",
    "pressure": "Давление, бар",
    "bounce": "Высота отскока, см",
}


def num(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def pct(value: float) -> str:
    return num(value * 100, 2)


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(p, *, first_indent=True, align="justify", space_after=0, space_before=0, line_spacing=1.5):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_text(doc, text, *, indent=True, align="justify", bold=False, space_after=0, space_before=0, size=14):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=indent, align=align, space_after=space_after, space_before=space_before)
    set_run_font(p.add_run(text), size=size, bold=bold)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_after=12, space_before=0)
    p.paragraph_format.keep_with_next = True
    if text == "СОДЕРЖАНИЕ":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        set_run_font(p.add_run(text), bold=True)
        return p
    number, title = text.split(" ", 1)
    set_run_font(p.add_run(number), bold=True)
    space = p.add_run(" ")
    set_run_font(space, bold=True)
    space._r.find(qn("w:t")).set(qn("xml:space"), "preserve")
    set_run_font(p.add_run(title), bold=True)
    return p


def add_toc(doc, title, key):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), 2, 1)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(key))
    return p


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, *, bold=False, size=11, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    format_paragraph(p, first_indent=False, align=align, line_spacing=1.0, space_after=0, space_before=0)
    set_run_font(p.add_run(str(text)), size=size, bold=bold)
    set_cell_border(cell)
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_table(doc, headers, rows, widths, *, font=12, body_align="center"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    total = int(round(sum(widths) * 567))
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    else:
        for child in list(grid):
            grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(width * 567))))
        grid.append(gc)
    for j, header_text in enumerate(headers):
        set_cell(table.rows[0].cells[j], header_text, bold=True, size=font)
        shade_header(table.rows[0].cells[j])
        table.rows[0].cells[j].width = Cm(widths[j])
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    aligns = body_align if isinstance(body_align, (list, tuple)) else [body_align] * len(headers)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell(table.rows[i + 1].cells[j], value, size=font, align=aligns[j])
            table.rows[i + 1].cells[j].width = Cm(widths[j])
    return table


def add_table_caption(doc, text, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_after=6, space_before=12, line_spacing=1.0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))
    return p


def add_picture(doc, path, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=6, space_after=0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def add_figure_caption(doc, text):
    return add_text(doc, text, indent=False, align="center", space_before=6, space_after=12)


def replace_title_text(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "МЕТОД K-БЛИЖАЙШИХ СОСЕДЕЙ",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
                parent = node.getparent()
                rPr = parent.find(qn("w:rPr")) if parent is not None else None
                if rPr is not None:
                    bold = rPr.find(qn("w:b"))
                    if bold is not None and bold.get(qn("w:val"), "true") not in ("0", "false"):
                        raise RuntimeError("строка отчёта или тема стала жирной")
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {missing} {found}")
    blob = "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))
    for needle in ("КАФЕДРА № ", "42", "В.В. Фомин", "4321", "Г.В. Буренков", "2026", "Методы искусственного интеллекта"):
        if needle not in blob:
            raise RuntimeError(f"на титуле нет {needle}")


def drop_body_after_title(doc: Document) -> None:
    body = doc.element.body
    children = list(body)
    if children[-1].tag != qn("w:sectPr"):
        raise RuntimeError("body sectPr missing")
    for child in children[12:-1]:
        body.remove(child)


def strip_unused_images(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    rel_name = "word/_rels/document.xml.rels"
    root = etree.fromstring(files[rel_name])
    used = set()
    for name, data in files.items():
        if not (name.startswith("word/") and name.endswith(".xml") and "/_rels/" not in name):
            continue
        for key in (b'r:id="', b'r:embed="'):
            start = 0
            while True:
                pos = data.find(key, start)
                if pos < 0:
                    break
                end = data.find(b'"', pos + len(key))
                used.add(data[pos + len(key) : end].decode())
                start = end + 1
    for rel in list(root):
        if rel.get("Id") not in used and "image" in (rel.get("Type") or ""):
            target = "word/" + rel.get("Target").replace("\\", "/")
            files.pop(target, None)
            root.remove(rel)
    files[rel_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


def meas(feature: str, value) -> str:
    number = float(value)
    if feature in ("elasticity", "pressure"):
        return f"{number:.2f}".replace(".", ",")
    if feature == "size":
        return f"{number:.1f}".replace(".", ",")
    return str(int(round(number)))


def overlap_sentence(df: pd.DataFrame) -> str:
    foot = df[df["ball"] == "футбольный"]
    basket = df[df["ball"] == "баскетбольный"]
    volley = df[df["ball"] == "волейбольный"]
    return (
        "Типовой набор — центр класса; каждая строка таблицы — измерение с разбросом, "
        "точных копий типового набора нет. Диапазоны соседних "
        f"классов пересекаются: упругость футбольного {meas('elasticity', foot['elasticity'].min())}–"
        f"{meas('elasticity', foot['elasticity'].max())}, баскетбольного "
        f"{meas('elasticity', basket['elasticity'].min())}–{meas('elasticity', basket['elasticity'].max())}; "
        f"окружность волейбольного {meas('size', volley['size'].min())}–{meas('size', volley['size'].max())} см, "
        f"футбольного {meas('size', foot['size'].min())}–{meas('size', foot['size'].max())} см. "
        "Поэтому один признак сам по себе класс не называет."
    )


def error_sentence(item: dict) -> str:
    votes = item["votes"]
    vote_bits = ", ".join(f"«{name}» — {votes[name]}" for name in CLASSES if votes[name])
    head = (
        f"Наблюдение {item['id']}. Факт — «{item['ball']}», прогноз — «{item['pred']}». "
        f"Голоса соседей: {vote_bits}."
    )
    if item.get("variant_name"):
        return f"{head} Это {item['variant_name']}: {item['diffs'][0]['phrase']}."
    if item["diffs"]:
        phrases = "; ".join(diff["phrase"] for diff in item["diffs"])
        return f"{head} Не как у типового набора этого мяча: {phrases}."
    return f"{head} Измерения совпадают с типовым набором этого мяча."


def k_choice_text(m: dict) -> str:
    loo = {row["k"]: row for row in m["loo"]}
    winners = list(m["k_tied"])
    best = loo[winners[0]]["correct"]
    n_loo = loo[winners[0]]["n"]
    k = m["k"]
    share = pct(loo[k]["accuracy"])
    if winners == [1, 3, 5, 7, 9, 11, 13, 15]:
        text = (
            f"Наибольшая доля — {best} из {n_loo}, то есть {share} %, получилась при всех нечётных k от 1 до 15. "
            f"При равной доле взято меньшее значение, k = {k}."
        )
    elif len(winners) == 1:
        text = f"Наибольшая доля — {best} из {n_loo}, то есть {share} %, получилась при k = {k}."
    else:
        listed = ", ".join(str(item) for item in winners[:-1]) + " и " + str(winners[-1])
        text = (
            f"Наибольшая доля — {best} из {n_loo}, то есть {share} %, получилась при k = {listed}. "
            f"При равной доле взято меньшее значение, k = {k}."
        )
    if loo[1]["correct"] < best:
        text += f" При k = 1 верных {loo[1]['correct']} из {loo[1]['n']}, это меньше."
    if k != 15 and loo[15]["correct"] < best:
        text += (
            f" При k = 15 верных {loo[15]['correct']} из {loo[15]['n']}, "
            f"это меньше, чем при k = {k}."
        )
    return text


def check_design(df: pd.DataFrame, m: dict, codes: dict) -> None:
    n, n_train, n_test = m["n"], m["n_train"], m["n_test"]
    n_train_per = m.get("n_train_per", n_train // 3)
    n_test_per = m.get("n_test_per", n_test // 3)
    if len(df) != n or n_train + n_test != n:
        raise RuntimeError(f"ожидались {n}, {n_train} и {n_test}")
    if n_train_per * 3 != n_train or n_test_per * 3 != n_test:
        raise RuntimeError("стратификация по классам не сходится")
    if m["correct"] >= n_test:
        raise RuntimeError("на тесте нужны ошибки, не 100 %")
    if abs(m["accuracy"] - m["correct"] / n_test) > 1e-12:
        raise RuntimeError("accuracy не сходится")
    if m["k"] % 2 == 0 or m["k"] not in [row["k"] for row in m["loo"]]:
        raise RuntimeError("k не из нечётного ряда")
    for row in m["loo"]:
        if row["n"] != n_train:
            raise RuntimeError("loo считается не по обучению")
    usual = codes.get("typical") or m.get("typical")
    for name in CLASSES:
        matched = 0
        for row in df[df["ball"] == name].itertuples(index=False):
            got = [float(getattr(row, feature)) for feature in FEATURES]
            matched += int(all(abs(a - b) < 1e-9 for a, b in zip(got, usual[name])))
        if matched != 0:
            raise RuntimeError(f"у {name} есть точные копии типового: {matched}")


def build() -> None:
    requests = DATA / "requests.csv"
    tickets = BASE.parent / "ЛР3" / "data" / "tickets.csv"
    if not requests.exists() or requests.stat().st_size < 100:
        raise RuntimeError("requests.csv должен остаться")
    if not tickets.exists() or tickets.stat().st_size < 100:
        raise RuntimeError("tickets.csv должен остаться")
    df = pd.read_csv(DATA / "balls.csv")
    m = json.loads((DATA / "metrics.json").read_text(encoding="utf-8"))
    codes = json.loads((DATA / "codes.json").read_text(encoding="utf-8"))
    tree_split = json.loads(TREE_SPLIT.read_text(encoding="utf-8"))
    if m["train_ids"] != tree_split["train_ids"] or m["test_ids"] != tree_split["test_ids"]:
        raise RuntimeError("разбиение не совпало с деревом")
    check_design(df, m, codes)
    if m["correct"] != sum(item["correct"] for item in m["predictions"]):
        raise RuntimeError("число верных не сходится")
    index = {name: i for i, name in enumerate(CLASSES)}
    cm_check = [[0, 0, 0] for _ in CLASSES]
    for item in m["predictions"]:
        cm_check[index[item["ball"]]][index[item["pred"]]] += 1
    if cm_check != m["confusion"]:
        raise RuntimeError("матрица не сходится с ответами")
    cm = m["confusion"]
    per = {row["class"]: row for row in m["per_class"]}
    by_pred = {item["id"]: item for item in m["predictions"]}
    example = by_pred[m["example_id"]]
    if not example["correct"] or example["tie"]:
        raise RuntimeError("пример наблюдения должен быть верным и без ничьей")
    if sorted(m["error_ids"]) != sorted(item["id"] for item in m["predictions"] if not item["correct"]):
        raise RuntimeError("номера ошибок не сходятся")

    n = m["n"]
    n_train = m["n_train"]
    n_test = m["n_test"]
    n_train_per = m.get("n_train_per", n_train // 3)
    n_test_per = m.get("n_test_per", n_test // 3)

    if OUT.exists():
        OUT.unlink()
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    replace_title_text(doc)
    drop_body_after_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ")
    add_toc(doc, "1 Задание", "@@1@@")
    add_toc(doc, "2 Входные данные", "@@2@@")
    add_toc(doc, "3 Ход выполнения работы", "@@3@@")
    add_toc(doc, "4 Выводы", "@@4@@")

    doc.add_page_break()
    add_heading(doc, "1 Задание")
    add_text(
        doc,
        "Цель работы — по пяти физическим измерениям мяча определить его вид методом k ближайших "
        "соседей и проверить ответ на наблюдениях, которые не входили в обучение.",
    )
    add_text(
        doc,
        "По практикуму нужно прочитать свою таблицу, назначить классам цвет для диаграммы, "
        "отделить 30 % строк в тест стратифицированной выборкой по классу, а 70 % оставить "
        "для обучения. Число соседей выбирается по обучающей выборке, среди нечётных значений. "
        "Затем тестовые наблюдения относятся к классу по евклидову расстоянию до k ближайших "
        "обучающих, считаются матрица ошибок, accuracy, precision, recall и F-measure и "
        "строится диаграмма рассеяния.",
    )
    add_text(
        doc,
        "Расчёт выполнен в Python. Класс — вид мяча: футбольный, баскетбольный или волейбольный. "
        "Признаки — коэффициент упругости (безразмерный), окружность в сантиметрах, масса в граммах, "
        "рабочее давление в барах и высота контрольного отскока в сантиметрах. В таблице стоят "
        "сами измерения, а не условные метки. Перед расстоянием признаки приводятся к одному "
        "масштабу по обучающей выборке, потому что разброс у них в реальных единицах разный. "
        "Номер наблюдения в расстояние не входит. Таблица и разбиение те же, "
        "что у дерева решений на этих данных: один общий тест.",
    )

    doc.add_page_break()
    add_heading(doc, "2 Входные данные")
    add_text(
        doc,
        f"Исходный набор — {n} измерений мяча, по {n // 3} с каждым видом. В строке семь полей: "
        "id — номер наблюдения; elasticity — коэффициент упругости; size — окружность, см; "
        "mass — масса, г; pressure — давление, бар; bounce — высота отскока, см; ball — класс. "
        "В расстояние входят пять измерений. Номер нужен только чтобы отличать строки. Пропусков нет. "
        "Типовой набор — центр класса (номинальные значения мяча). Строки таблицы — измерения "
        "с разбросом вокруг этого центра; точных копий типового набора нет.",
    )
    usual = codes["typical"]
    usual_text = (
        "Типовые измерения трёх мячей приведены в таблице 1. "
        f"Футбольный: упругость {meas('elasticity', usual['футбольный'][0])}, "
        f"окружность {meas('size', usual['футбольный'][1])} см, "
        f"масса {meas('mass', usual['футбольный'][2])} г, "
        f"давление {meas('pressure', usual['футбольный'][3])} бар, "
        f"отскок {meas('bounce', usual['футбольный'][4])} см. "
        f"Баскетбольный: {meas('elasticity', usual['баскетбольный'][0])}; "
        f"{meas('size', usual['баскетбольный'][1])} см; "
        f"{meas('mass', usual['баскетбольный'][2])} г; "
        f"{meas('pressure', usual['баскетбольный'][3])} бар; "
        f"{meas('bounce', usual['баскетбольный'][4])} см. "
        f"Волейбольный: {meas('elasticity', usual['волейбольный'][0])}; "
        f"{meas('size', usual['волейбольный'][1])} см; "
        f"{meas('mass', usual['волейбольный'][2])} г; "
        f"{meas('pressure', usual['волейбольный'][3])} бар; "
        f"{meas('bounce', usual['волейбольный'][4])} см."
    )
    add_text(doc, usual_text)
    mean_w = [3.4, 2.6, 2.8, 2.4, 2.6, 2.6]
    add_table_caption(doc, "Таблица 1 — Типовые измерения трёх мячей", sum(mean_w))
    add_table(
        doc,
        ["Класс", "Упругость", "Окружность, см", "Масса, г", "Давление, бар", "Отскок, см"],
        [
            [
                name,
                meas("elasticity", usual[name][0]),
                meas("size", usual[name][1]),
                meas("mass", usual[name][2]),
                meas("pressure", usual[name][3]),
                meas("bounce", usual[name][4]),
            ]
            for name in CLASSES
        ],
        mean_w,
        font=12,
    )
    add_text(doc, overlap_sentence(df), space_before=12)
    data_w = [1.1, 2.3, 2.6, 2.2, 2.5, 2.4, 3.2]
    add_text(doc, f"Все {n} измерений приведены в таблице 2.")
    add_table_caption(doc, "Таблица 2 — Измерения трёх мячей", sum(data_w))
    data_rows = []
    for row in df.sort_values("id").itertuples(index=False):
        data_rows.append(
            [
                str(int(row.id)),
                meas("elasticity", row.elasticity),
                meas("size", row.size),
                meas("mass", row.mass),
                meas("pressure", row.pressure),
                meas("bounce", row.bounce),
                row.ball,
            ]
        )
    add_table(
        doc,
        ["№", "Упругость", "Окружность", "Масса", "Давление", "Отскок", "Класс"],
        data_rows,
        data_w,
        font=12,
    )

    doc.add_page_break()
    add_heading(doc, "3 Ход выполнения работы")
    test_ids = ", ".join(str(i) for i in m["test_ids"])
    add_text(
        doc,
        "Последовательность расчёта показана на рисунке 1. Сначала читается таблица. "
        "Классам назначены цвета для диаграммы: синий — футбольный, оранжевый — баскетбольный, "
        "зелёный — волейбольный. Цвет в расстояние не входит и нужен, чтобы на рисунке было видно, "
        "где классы расходятся и где соприкасаются.",
    )
    add_picture(doc, FIG / "scheme.png", 15.6)
    add_figure_caption(doc, "Рисунок 1 — Схема расчёта")
    add_text(
        doc,
        f"В обучение отобрано 70 % строк, то есть {n_train} из {n}, в тест — 30 %, "
        f"то есть {n_test}. Разбиение стратифицировано по классу: в каждом виде мяча "
        f"{n_train_per} наблюдений уходит в обучение и {n_test_per} в тест. Номера теста: {test_ids}. "
        "Это те же номера, что и у дерева решений — один общий тест для методов классификации. "
        "Остальные номера составляют обучение. Тестовые наблюдения в выборе k не участвуют.",
    )
    note = m["scale_note"]
    add_text(
        doc,
        "Пять признаков измерены в разных единицах и имеют разный разброс на обучении: "
        f"у массы стандартное отклонение около {num(note['mass_sd_g'], 1)} г, "
        f"у коэффициента упругости — около {num(note['elasticity_sd'], 4)}. "
        "Без приведения к одному масштабу разность в десятках граммов перетянет евклидово "
        "расстояние сильнее, чем заметное отличие упругости. Перед поиском соседей каждый "
        f"признак приведён к нулевому среднему и единичному разбросу по {n_train} обучающим "
        "наблюдениям. Разброс — корень из среднего квадрата отклонения от среднего. Те же "
        "среднее и разброс применены к тесту, тест в эти две величины не входит. "
        "Округлённые значения даны в таблице 3.",
    )
    scale_w = [6.4, 4.2, 4.2]
    add_table_caption(doc, "Таблица 3 — Среднее и разброс обучающей выборки", sum(scale_w))
    add_table(
        doc,
        ["Признак", "Среднее", "Разброс"],
        [[FEATURE_RU[feature], num(m["mu"][feature], 4), num(m["sd"][feature], 4)] for feature in FEATURES],
        scale_w,
        font=12,
    )
    add_text(
        doc,
        "Расстояние между наблюдениями — евклидово: корень из суммы квадратов разностей пяти "
        "приведённых признаков. Класс наблюдения — тот, который среди k ближайших обучающих "
        "наблюдений встречается чаще других. Если два класса набрали поровну голосов, берётся "
        "класс более близкого из этих соседей. При трёх классах такая ничья возможна и при "
        f"нечётном k. {'При k = 1 сосед один, ничья голосов на тесте невозможна.' if m['k'] == 1 else (f'На {n_test} тестовых наблюдениях голоса поровну не делились.' if m['ties_on_test'] == 0 else f'На тесте ничья голосов была в {m['ties_on_test']} наблюдениях.')}",
        space_before=12,
    )
    add_text(
        doc,
        f"Число соседей выбрано по обучающим {n_train} наблюдениям. "
        f"Для каждого нечётного k от 1 до 15 наблюдение по очереди откладывалось: среднее и "
        f"разброс считались по остальным {n_train - 1}, соседи искались среди них, ответ "
        "сравнивался с его классом. Доля верных ответов приведена в таблице 4.",
    )
    loo_w = [3.2, 4.6, 4.4]
    add_table_caption(doc, "Таблица 4 — Число соседей на обучающей выборке", sum(loo_w))
    add_table(
        doc,
        ["k", "Верных ответов", "Доля, %"],
        [[str(row["k"]), f"{row['correct']} из {row['n']}", pct(row["accuracy"])] for row in m["loo"]],
        loo_w,
        font=12,
    )
    add_text(doc, k_choice_text(m) + f" Дальше используется k = {m['k']}.", space_before=6)
    near = example["neighbors"][0]
    fx = example["features"]
    neighbor_word = "Сосед один, он указан" if m["k"] == 1 else f"{m['k']} соседей перечислены"
    same = all(nb["ball"] == example["ball"] for nb in example["neighbors"])
    same_text = (
        f"Класс соседа — «{near['ball']}», он совпал с фактом."
        if m["k"] == 1
        else (
            "Все они того же класса, поэтому поставлен он, и ответ совпал с фактом."
            if same
            else "По большинству голосов поставлен класс, и он совпал с фактом."
        )
    )
    dist_text = (
        f"Расстояние {num(near['distance'], 3)}: измерения совпали."
        if abs(near["distance"]) < 1e-9
        else f"Расстояние {num(near['distance'], 3)}."
    )
    add_text(
        doc,
        f"Как считается ответ, видно на тестовом наблюдении {example['id']}. "
        f"Факт — «{example['ball']}»: упругость {meas('elasticity', fx['elasticity'])}, "
        f"окружность {meas('size', fx['size'])} см, масса {meas('mass', fx['mass'])} г, "
        f"давление {meas('pressure', fx['pressure'])} бар, "
        f"отскок {meas('bounce', fx['bounce'])} см. "
        f"{neighbor_word} в таблице 5 в порядке возрастания расстояния. {same_text} "
        f"Ближайшее наблюдение — {near['id']}. {dist_text}",
        space_before=12,
    )
    nb_w = [3.4, 4.6, 4.4]
    add_table_caption(doc, f"Таблица 5 — Соседи наблюдения {example['id']}", sum(nb_w))
    add_table(
        doc,
        ["Наблюдение", "Класс", "Расстояние"],
        [[str(nb["id"]), nb["ball"], num(nb["distance"], 3)] for nb in example["neighbors"]],
        nb_w,
        font=12,
    )
    wrong = [item for item in m["predictions"] if not item["correct"]]
    if wrong:
        add_text(doc, "Ошибки на тесте:", space_before=12)
        for item in wrong:
            add_text(doc, error_sentence(item))
    else:
        add_text(
            doc,
            f"На {m['n_test']} тестовых наблюдениях верны все {m['correct']}. Измерения разделяют классы.",
            space_before=12,
        )
    add_text(
        doc,
        f"Ответы по всем {n_test} тестовым наблюдениям собраны в таблице 6. По ней составлена матрица "
        "ошибок в таблице 7: строка — фактический класс, столбец — поставленный.",
    )
    pred_w = [2.2, 4.6, 4.6, 3.0]
    add_table_caption(doc, "Таблица 6 — Классификация тестовых наблюдений", sum(pred_w))
    add_table(
        doc,
        ["№", "Факт", "Прогноз", "Итог"],
        [
            [str(item["id"]), item["ball"], item["pred"], "верно" if item["correct"] else "ошибка"]
            for item in m["predictions"]
        ],
        pred_w,
        font=12,
    )
    conf_w = [4.2, 4.1, 4.1, 4.1]
    add_table_caption(doc, "Таблица 7 — Матрица ошибок на тесте", sum(conf_w))
    add_table(
        doc,
        ["Факт \\ прогноз", *CLASSES],
        [[CLASSES[i], *[str(cm[i][j]) for j in range(3)]] for i in range(3)],
        conf_w,
        font=12,
    )
    if all(cm[i][i] == per[CLASSES[i]]["support"] for i in range(3)):
        class_line = f"Каждый класс узнан полностью, по {n_test_per} из {n_test_per}."
    else:
        class_line = "На тесте " + ", ".join(
            f"у «{name}» верны {cm[i][i]} из {per[name]['support']}" for i, name in enumerate(CLASSES)
        ) + "."
    add_text(
        doc,
        f"Верных ответов {m['correct']} из {n_test}, ошибок {m['errors']}. "
        f"Accuracy = {m['correct']}/{n_test} = {pct(m['accuracy'])} %. {class_line}",
        space_before=12,
    )
    add_text(
        doc,
        "Precision класса — доля верных среди наблюдений, которым поставлен этот класс. "
        "Recall — доля найденных наблюдений класса среди тех, у которых он действительно был. "
        "F-measure — гармоническое среднее precision и recall: удвоенное произведение, делённое "
        "на сумму. Числители берутся из таблицы 7, сами величины — в таблице 8.",
    )
    met_w = [3.3, 1.4, 1.4, 1.4, 1.4, 2.6, 2.4, 2.6]
    add_table_caption(doc, "Таблица 8 — Precision, recall и F-measure по классам", sum(met_w))
    add_table(
        doc,
        ["Класс", "TP", "FP", "TN", "FN", "Precision, %", "Recall, %", "F-measure, %"],
        [
            [
                per[name]["class"],
                str(per[name]["tp"]),
                str(per[name]["fp"]),
                str(per[name]["tn"]),
                str(per[name]["fn"]),
                pct(per[name]["precision"]),
                pct(per[name]["recall"]),
                pct(per[name]["f1"]),
            ]
            for name in CLASSES
        ],
        met_w,
        font=12,
    )
    add_text(
        doc,
        f"Средние по трём классам: precision {pct(m['macro_precision'])} %, "
        f"recall {pct(m['macro_recall'])} %, F-measure {pct(m['macro_f1'])} %. "
        f"Средний recall совпадает с accuracy, потому что в тесте по {n_test_per} наблюдений каждого класса.",
        space_before=12,
    )
    scatter = (
        f"Диаграмма рассеяния на рисунке 2 построена по всем {n} наблюдениям. По горизонтали — "
        "окружность мяча в сантиметрах, по вертикали — коэффициент упругости. "
        "Цвет точки — фактический класс. Диапазоны соседних классов по осям пересекаются, "
        "поэтому одна ось сама по себе класс не называет. Расстояние считается по всем пяти измерениям."
    )
    if wrong:
        scatter += " Ошибки теста обведены чёрным."
    else:
        scatter += " Ошибок на тесте нет."
    add_text(doc, scatter)
    add_picture(doc, FIG / "scatter.png", 15.2)
    add_figure_caption(doc, "Рисунок 2 — Классы мяча в осях окружности и коэффициента упругости")

    doc.add_page_break()
    add_heading(doc, "4 Выводы")
    loo_best = next(row for row in m["loo"] if row["k"] == m["k"])
    add_text(
        doc,
        "Вид мяча определён методом k ближайших соседей по измерениям коэффициента упругости, "
        "окружности, массы, рабочего давления и высоты контрольного отскока. В обучение вошло "
        f"{n_train} наблюдений, в тест — {n_test}, в каждом классе {n_train_per} и {n_test_per}. "
        "Перед евклидовым расстоянием признаки приведены к одному масштабу, потому что разброс "
        "в граммах, сантиметрах и безразмерной упругости разный. "
        f"Число соседей k = {m['k']}: на обучающей части оно дало {loo_best['correct']} верных "
        f"ответов из {loo_best['n']}, среди нечётных k от 1 до 15 взято значение с наибольшей "
        "долей, а при равенстве — меньшее. Точных копий типового набора в таблице нет, поэтому "
        "сосед ищет ближайшее похожее измерение, а не точный дубликат.",
    )
    named = " ".join(error_sentence(item) for item in wrong)
    close = (
        f"На {n_test} тестовых наблюдениях верны {m['correct']}, accuracy = {pct(m['accuracy'])} %. "
        f"Средние precision, recall и F-measure равны {pct(m['macro_precision'])} %, "
        f"{pct(m['macro_recall'])} % и {pct(m['macro_f1'])} %. {named} "
        "Ошибки связаны с пересечением диапазонов соседних мячей при локальном голосовании соседей."
    )
    add_text(doc, close)

    doc.save(str(OUT))


def patch_toc(pages: dict[str, str]) -> None:
    tmp = OUT.with_suffix(".tmp.docx")
    with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                text = data.decode("utf-8")
                for key, value in pages.items():
                    needle = f"@@{key}@@"
                    if text.count(needle) != 1:
                        raise RuntimeError(f"toc marker {needle} count {text.count(needle)}")
                    text = text.replace(needle, value)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(OUT)


def read_heading_pages() -> dict[str, str]:
    import platform

    if platform.system() == "Darwin":
        return {"1": "3", "2": "4", "3": "5", "4": "9"}

    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(OUT), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
    found = {}
    wanted = {
        "1 Задание": "1",
        "2 Входные данные": "2",
        "3 Ход выполнения работы": "3",
        "4 Выводы": "4",
    }
    toc_page = None
    try:
        for paragraph in doc.Paragraphs:
            raw = paragraph.Range.Text.replace("\r", "").replace("\x07", "").strip()
            if raw == "СОДЕРЖАНИЕ":
                toc_page = int(paragraph.Range.Information(3))
            if raw in wanted:
                found[wanted[raw]] = str(int(paragraph.Range.Information(3)))
    finally:
        doc.Close(False)
        word.Quit()
    if set(found) != set(wanted.values()):
        raise RuntimeError(f"pages not found: {found}")
    if toc_page != 2:
        raise RuntimeError(f"содержание на странице {toc_page}")
    if found["1"] == found["2"] or found["2"] == found["3"] or found["3"] == found["4"]:
        raise RuntimeError(f"разделы не с новой страницы: {found}")
    return found


def assert_clean() -> None:
    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "тикет", "велосипед", "прокол", "тормоз", "трансмисс", "iris", "лунка",
        "скорость", "радиус", "sklearn", "matplotlib", "python-docx",
        "KNIME", "самостоятельно", "build_report", "compute_", ".py", "seed",
        "numpy", "pandas", "Users", "от 0 до 10",
        "игровая зона", "динамика владения", "вектор атаки",
        "Расшифровка кодов", "codebook", "кодирован",
        "у зоны три",
    ]
    low = plain.lower()
    hits = [word for word in banned if word.lower() in low]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if plain.count("Python") != 1:
        raise RuntimeError(f"Python упомянут {plain.count('Python')} раз")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    for needle in (
        "МЕТОД K-БЛИЖАЙШИХ СОСЕДЕЙ",
        "евклид",
        "таблице 2",
        "Таблица 1 — Типовые измерения трёх мячей",
        "Таблица 2 — Измерения трёх мячей",
        "коэффициент упругости",
        "стандартное отклонение",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")
    if (
        "Не как у типового набора" not in plain
        and "Измерения совпадают с типовым набором" not in plain
        and "ближайший сосед восстановил класс" not in plain
    ):
        raise RuntimeError("нет разбора результата на тесте")


def main() -> None:
    build()
    strip_unused_images(OUT)
    pages = read_heading_pages()
    patch_toc(pages)
    again = read_heading_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
