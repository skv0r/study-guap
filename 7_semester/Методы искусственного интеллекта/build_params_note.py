# -*- coding: utf-8 -*-
"""Отдельная справка: начальные параметры классификаторов (ЛР2–ЛР5)."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
OUT = BASE / "Начальные_параметры_классификаторов_Буренков.docx"


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(p, *, first_indent=True, align="justify", space_after=0, space_before=0, line_spacing=1.5):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_text(doc, text, *, indent=True, align="justify", bold=False, space_after=0, space_before=0, size=14):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=indent, align=align, space_after=space_after, space_before=space_before)
    set_run_font(p.add_run(text), size=size, bold=bold)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_after=12, space_before=12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    return p


def set_cell_border(cell):
    tc = cell._tc
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def set_cell(cell, text, *, bold=False, size=12, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    format_paragraph(p, first_indent=False, align=align, line_spacing=1.15, space_after=0, space_before=0)
    set_run_font(p.add_run(str(text)), size=size, bold=bold)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, w in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = Cm(w)
    for j, h in enumerate(headers):
        set_cell(table.rows[0].cells[j], h, bold=True, size=12)
        tcPr = table.rows[0].cells[j]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F2F2F2")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            align = "left" if j in (0, 2) else "center"
            set_cell(table.rows[i].cells[j], value, size=12, align=align)
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, space_after=6)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=6, space_after=6)
    set_run_font(p.add_run(text), size=12)
    return p


def build() -> None:
    knn = json.loads((BASE / "ЛР2" / "data" / "metrics.json").read_text(encoding="utf-8"))
    tree = json.loads((BASE / "ЛР3" / "data" / "metrics.json").read_text(encoding="utf-8"))
    bayes = json.loads((BASE / "ЛР4" / "data" / "metrics.json").read_text(encoding="utf-8"))
    svm = json.loads((BASE / "ЛР5" / "data" / "metrics.json").read_text(encoding="utf-8"))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    lines = [
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ АВТОНОМНОЕ ОБРАЗОВАТЕЛЬНОЕ",
        "УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ",
        "«САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ УНИВЕРСИТЕТ",
        "АЭРОКОСМИЧЕСКОГО ПРИБОРОСТРОЕНИЯ»",
        "",
        "КАФЕДРА № 42",
        "",
        "",
        "СПРАВКА",
        "Начальные параметры методов классификации",
        "(лабораторные работы № 2–5)",
        "",
        "",
        "по дисциплине: Методы искусственного интеллекта",
        "",
        "Выполнил: студент группы 4321",
        "Г.В. Буренков",
        "",
        "Преподаватель: профессор, д-р техн. наук,",
        "профессор В.В. Фомин",
        "",
        "",
        "Санкт-Петербург",
        "2026",
    ]
    for line in lines:
        add_text(doc, line, indent=False, align="center", bold=("СПРАВКА" in line or line.startswith("Начальные")))

    doc.add_page_break()

    add_heading(doc, "1 Общие параметры данных")
    add_text(
        doc,
        "Для методов лабораторных работ № 2–5 используется один и тот же набор измерений мяча "
        "и одно и то же разбиение на обучение и тест. Это обеспечивает сопоставимость результатов.",
    )
    add_caption(doc, "Таблица 1 — Общие параметры выборки")
    add_table(
        doc,
        ["Параметр", "Значение", "Обоснование"],
        [
            ["Объём выборки", "150 наблюдений", "По 50 строк каждого класса"],
            ["Классы", "футбольный, баскетбольный, волейбольный", "Задача классификации вида мяча"],
            ["Признаки", "упругость, окружность, масса, давление, отскок", "Физические измерения без кодирования"],
            ["Доля обучения / теста", "70 % / 30 %", "По лабораторному практикуму"],
            ["Объём обучения / теста", f"{knn['n_train']} / {knn['n_test']}", "По 35 и 15 наблюдений в классе"],
            ["Стратификация", "по классу", "В тесте представлены все три вида мяча"],
            ["Генератор разбиения", f"random_state = {knn.get('seed', 300)}", "Фиксация одного общего теста"],
        ],
        [5.0, 4.5, 7.0],
    )

    add_heading(doc, "2 Метод k ближайших соседей (лабораторная работа № 2)")
    add_text(
        doc,
        f"Число соседей выбрано только по обучающей выборке среди нечётных значений от 1 до 15. "
        f"Наибольшая доля верных ответов при проверке на обучении получена при k = {knn['k']} "
        f"({max(row['correct'] for row in knn['loo'])} из {knn['n_train']}).",
    )
    add_caption(doc, "Таблица 2 — Начальные параметры метода k ближайших соседей")
    add_table(
        doc,
        ["Параметр", "Значение", "Обоснование"],
        [
            ["Мера близости", "евклидово расстояние", "Стандартная мера для непрерывных признаков"],
            [
                "Масштабирование",
                "стандартизация по обучению",
                "Признаки имеют разные единицы (г, см, бар); без масштаба масса доминирует",
            ],
            ["Кандидаты k", "1, 3, 5, …, 15", "Нечётное k уменьшает вероятность ничьей голосов"],
            [
                "Выбранное k",
                str(knn["k"]),
                "Максимум верных ответов на обучении; при равенстве берётся меньшее k",
            ],
            ["Номер наблюдения", "не используется", "Идентификатор не является признаком класса"],
        ],
        [5.0, 4.5, 7.0],
    )

    add_heading(doc, "3 Дерево решений (лабораторная работа № 3)")
    add_text(
        doc,
        f"Фактическая глубина обученного дерева — {tree['actual_depth']}, число листьев — "
        f"{tree['n_leaves']}.",
    )
    add_caption(doc, "Таблица 3 — Начальные параметры дерева решений")
    add_table(
        doc,
        ["Параметр", "Значение", "Обоснование"],
        [
            ["Критерий разбиения", "индекс Джини", "Задан лабораторным практикумом"],
            [
                "Максимальная глубина",
                str(tree["max_depth_limit"]),
                "Ограничение переобучения и читаемости правил",
            ],
            [
                "Минимум наблюдений в листе",
                str(tree["min_samples_leaf"]),
                "Лист не строится по одной-двум случайным точкам",
            ],
            [
                "Фактическая глубина / листья",
                f"{tree['actual_depth']} / {tree['n_leaves']}",
                "Результат обучения на данной выборке",
            ],
        ],
        [5.0, 4.5, 7.0],
    )

    add_heading(doc, "4 Наивный байесовский классификатор (лабораторная работа № 4)")
    prior = bayes["priors"]["футбольный"]
    add_caption(doc, "Таблица 4 — Начальные параметры наивного Байеса")
    add_table(
        doc,
        ["Параметр", "Значение", "Обоснование"],
        [
            [
                "Распределение признака",
                "нормальное",
                "Непрерывные измерения; среднее и σ оцениваются по классу на обучении",
            ],
            [
                "Априорная вероятность класса",
                f"{prior:.4f}".replace(".", ",") + " (1/3)",
                "Классы в обучении сбалансированы (по 35 из 105)",
            ],
            [
                "Минимальное σ в плотности",
                "0,0001",
                "По практикуму: нулевое отклонение заменяется, чтобы плотность не обнулялась",
            ],
            [
                "Допущение независимости",
                "принято",
                "Суть «наивности» метода: совместная плотность — произведение одномерных",
            ],
        ],
        [5.0, 4.5, 7.0],
    )

    add_heading(doc, "5 Метод опорных векторов (лабораторная работа № 5)")
    gamma = svm["gamma"]
    add_caption(doc, "Таблица 5 — Начальные параметры метода опорных векторов")
    add_table(
        doc,
        ["Параметр", "Значение", "Обоснование"],
        [
            ["Ядро", "RBF (радиальная базисная функция)", "Задано практикумом; нелинейная граница классов"],
            ["Параметр сигма σ", str(svm["sigma"]).replace(".", ","), "Задан лабораторным практикумом"],
            [
                "Параметр γ = 1/(2σ²)",
                str(gamma).replace(".", ","),
                "Следует из формулы ядра RBF при выбранной сигме",
            ],
            [
                "Штраф C",
                str(int(svm["penalty"])) if float(svm["penalty"]).is_integer() else str(svm["penalty"]).replace(".", ","),
                "Задан практикумом: баланс ошибок обучения и ширины разделяющей полосы",
            ],
            [
                "Масштабирование",
                "стандартизация по обучению",
                "Те же единицы признаков, что у k ближайших соседей",
            ],
            [
                "Число опорных векторов",
                str(svm["n_support_total"]),
                "Результат обучения (не задаётся вручную)",
            ],
        ],
        [5.0, 4.5, 7.0],
    )

    add_heading(doc, "6 Заключение")
    add_text(
        doc,
        "Параметры, указанные в практикуме (критерий Джини, ограничение глубины, нормальное "
        "распределение, порог 0,0001, ядро RBF, сигма 0,5, штраф 1), приняты без изменения. "
        f"Число соседей k = {knn['k']} выбрано по результатам на обучающей выборке. Разбиение "
        "данных и масштабирование признаков согласованы между лабораторными работами № 2–5.",
    )

    if OUT.exists():
        OUT.unlink()
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
