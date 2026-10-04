# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР1 МИИ. Титул берётся из проверенного файла, текст не описывает сборку."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
TITLE = Path(
    r"C:\Users\skvor\OneDrive\Desktop\GitHub\study-guap\7_semester\_титулы\ТИТУЛ_МИИ_ЛР1.docx"
)
OUT = BASE / "ЛР1_Отчет_Буренков_МИИ.docx"
FIG = BASE / "figures"


def num(value: float, digits: int = 3) -> str:
    text = f"{value:.{digits}f}".replace(".", ",")
    if text.startswith("-"):
        text = "−" + text[1:]
    return text


def set_run_font(run, size=14, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(p, *, first_indent=True, align="justify", space_after=0, space_before=0, line_spacing=1.5):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_text(doc, text, *, indent=True, align="justify", bold=False, space_after=0, space_before=0, size=14):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=indent, align=align, space_after=space_after, space_before=space_before)
    set_run_font(p.add_run(text), size=size, bold=bold)
    return p


def add_heading(doc, text, *, align="left"):
    indented = align == "left"
    p = add_text(doc, "", indent=indented, align=align, space_after=12, space_before=0)
    p.paragraph_format.keep_with_next = True
    if indented and " " in text and text[0].isdigit():
        number, title = text.split(" ", 1)
        set_run_font(p.add_run(number), bold=True)
        space = p.add_run(" ")
        set_run_font(space, bold=True)
        space._r.find(qn("w:t")).set(qn("xml:space"), "preserve")
        set_run_font(p.add_run(title), bold=True)
    else:
        set_run_font(p.add_run(text), bold=True)
    return p


def add_toc(doc, title, page):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t" + str(page)))
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def set_cell(cell, text, *, bold=False, size=12, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    format_paragraph(p, first_indent=False, align=align, line_spacing=1.0, space_after=0)
    set_run_font(p.add_run(str(text)), size=size, bold=bold)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, *, font=12):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    for j, header in enumerate(headers):
        set_cell(table.rows[0].cells[j], header, bold=True, size=font)
        shade_header(table.rows[0].cells[j])
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            align = "left" if j == 0 and not str(value).replace(",", "").replace("−", "").replace(".", "").isdigit() else "center"
            if headers[j] in ("Показатель", "Что означает"):
                align = "left" if j == 0 or headers[j] == "Что означает" else "center"
            set_cell(table.rows[i + 1].cells[j], value, size=font, align="center" if j else "left")
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    # total width
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(widths) * 567)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    return table


def add_table_caption(doc, text, width_cm):
    """Подпись вплотную к таблице и от её левого края."""
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_after=6, space_before=0, line_spacing=1.0)
    content_cm = 16.5
    p.paragraph_format.left_indent = Cm(max(0, (content_cm - width_cm) / 2))
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))
    return p


def add_picture(doc, path, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=6, space_after=0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=14)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def main():
    metrics = json.loads((BASE / "data" / "metrics.json").read_text(encoding="utf-8"))
    orders = pd.read_csv(BASE / "data" / "courier_delivery.csv")
    pred = pd.read_csv(BASE / "data" / "predictions.csv")
    test = metrics["test"]
    rounded = metrics["test_rounded"]
    b0, b1 = metrics["intercept"], metrics["slope"]
    ids = ", ".join(str(v) for v in metrics["train_order_ids"])

    shutil.copy2(TITLE, OUT)
    doc = Document(str(OUT))
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_number(section.footer.paragraphs[0])
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "2")
    section._sectPr.append(pg)

    add_heading(doc, "СОДЕРЖАНИЕ", align="center")
    toc_items = [
        ("1 Задание", "3"),
        ("2 Входные данные", "4"),
        ("3 Ход выполнения работы", "7"),
        ("4 Выводы", "11"),
    ]
    for title, page in toc_items:
        add_toc(doc, title, page)

    doc.add_page_break()
    add_heading(doc, "1 Задание")
    add_text(
        doc,
        "Цель работы — описать зависимость времени курьерской доставки от расстояния до адреса моделью линейной регрессии и проверить, насколько эта прямая годится для прогноза на заказах, которые не участвовали в обучении.",
    )
    add_text(
        doc,
        "По практикуму нужно прочитать таблицу заказов, назначить наблюдениям цвет для последующей диаграммы, отделить 20 % строк линейной выборкой и на этой части оценить коэффициенты. По остальным строкам получить прогноз времени, посчитать показатели качества и построить диаграмму рассеяния. Прогноз дополнительно округляется до целой минуты: диспетчеру удобнее целое значение, а дробный столбец из показываемой таблицы убирается.",
    )
    add_text(
        doc,
        "Расчёт выполнен в Python методом наименьших квадратов. Отклик — время доставки в минутах, единственный признак — расстояние в километрах. Номер заказа в модель не входит.",
    )

    doc.add_page_break()
    add_heading(doc, "2 Входные данные")
    add_text(
        doc,
        f"Исходный набор — {metrics['n']} заказов службы доставки. В каждой строке три поля: order_id — номер заказа, distance_km — расстояние до адреса, time_min — фактическое время от назначения заказа до вручения. Независимая переменная X — distance_km, зависимая Y — time_min. Оба признака вещественные. Номер заказа нужен только чтобы отличать строки.",
    )
    add_text(
        doc,
        f"Расстояние лежит в интервале от {num(metrics['distance_min'], 2)} до {num(metrics['distance_max'], 2)} км, время — от {num(metrics['time_min'], 1)} до {num(metrics['time_max'], 1)} мин. Строки записаны по возрастанию расстояния. Полный набор приведён в таблице 1.",
    )
    add_table_caption(doc, "Таблица 1 — Заказы курьерской доставки", 4.5 + 5.5 + 5.0)
    data_rows = [
        (str(int(r.order_id)), num(r.distance_km, 2), num(r.time_min, 1))
        for r in orders.itertuples(index=False)
    ]
    add_table(
        doc,
        ["Номер заказа", "Расстояние, км", "Время, мин"],
        data_rows,
        [4.5, 5.5, 5.0],
        font=12,
    )

    doc.add_page_break()
    add_heading(doc, "3 Ход выполнения работы")
    add_text(
        doc,
        "Последовательность расчёта показана на рисунке 1. Сначала читается таблица и проверяются типы: номер заказа — целое число, расстояние и время — вещественные. Пропусков нет. После разбиения обучающим заказам назначен синий цвет, тестовым — оранжевый. Цвет не входит в оценку коэффициентов и нужен на диаграмме, чтобы было видно, на каких точках прямая построена.",
    )
    add_picture(doc, FIG / "scheme.png", 12.5)
    add_text(doc, "Рисунок 1 — Схема расчёта", indent=False, align="center", space_before=6, space_after=12)

    add_text(
        doc,
        f"В обучение отобрано 20 % строк, то есть {metrics['n_train']} из {metrics['n']}. Способ отбора — линейная выборка: строки берутся равномерно по всей таблице, а не подряд с начала файла. Первая и последняя строки входят в обучение. Остальные {metrics['n_test']} заказов составляют тест. Номера заказов обучения: {ids}.",
    )
    add_text(
        doc,
        f"Из-за того, что файл упорядочен по расстоянию, равномерный отбор захватывает и короткие, и длинные маршруты. В обучении расстояние от {num(metrics['train_distance_min'], 2)} до {num(metrics['train_distance_max'], 2)} км, в тесте — от {num(metrics['test_distance_min'], 2)} до {num(metrics['test_distance_max'], 2)} км. Прямая проверяется на тех же масштабах расстояния, на которых она оценивалась, а не только на ближних адресах.",
    )
    add_text(
        doc,
        "Целевой столбец модели — time_min, независимый — distance_km. Свободный член сохранён. По 20 обучающим строкам получены коэффициенты таблицы 2. Уравнение имеет вид",
    )
    add_text(
        doc,
        f"time_min = {num(b0, 3)} + {num(b1, 3)} · distance_km.",
        indent=False,
        align="center",
        space_before=6,
        space_after=6,
    )
    add_text(
        doc,
        f"Свободный член {num(b0, 3)} мин — это оценка времени, которое остаётся при нулевом расстоянии: передача заказа курьеру и вручение у двери. Коэффициент при расстоянии равен {num(b1, 3)} мин/км: каждый дополнительный километр увеличивает время доставки на {num(b1, 3)} мин. В пересчёте на скорость это около {num(60 / b1, 0)} км/ч, то есть городская поездка с остановками, а не движение по свободной трассе.",
    )
    add_table_caption(doc, "Таблица 2 — Коэффициенты линейной регрессии", 6.2 + 3.2 + 6.6)
    add_table(
        doc,
        ["Коэффициент", "Оценка", "Смысл"],
        [
            ["Свободный член", num(b0, 4), "время без учёта дороги, мин"],
            ["Коэффициент при distance_km", num(b1, 4), "добавка времени на 1 км, мин"],
        ],
        [6.2, 3.2, 6.6],
        font=12,
    )

    add_text(
        doc,
        "По уравнению для каждой тестовой строки посчитан прогноз, столбец Prediction (time_min). Показатели качества в таблице 3 посчитаны по всем 80 тестовым строкам и по этому дробному прогнозу. Коэффициент детерминации — доля разброса фактического времени, которую объясняет расстояние. Средняя абсолютная ошибка и корень из среднего квадрата ошибки измеряются в минутах. Средняя разность со знаком — это среднее значение «прогноз минус факт»: отрицательное число означает, что модель в среднем занижает время.",
        space_before=12,
    )
    add_table_caption(doc, "Таблица 3 — Показатели качества на 80 тестовых заказах", 12.0 + 4.0)
    add_table(
        doc,
        ["Показатель", "Значение"],
        [
            ["R²", num(test["R2"], 3)],
            ["Скорректированный R²", num(test["adj_R2"], 3)],
            ["Средняя абсолютная ошибка, мин", num(test["MAE"], 3)],
            ["Средний квадрат ошибки, мин²", num(test["MSE"], 3)],
            ["Корень из среднего квадрата ошибки, мин", num(test["RMSE"], 3)],
            ["Средняя разность со знаком, мин", num(test["mean_pred_minus_fact"], 3)],
            ["Средняя абсолютная процентная ошибка, %", num(test["MAPE_pct"], 2)],
        ],
        [12.0, 4.0],
        font=12,
    )
    add_text(
        doc,
        f"R² = {num(test['R2'], 3)}: расстояние объясняет {num(test['R2'] * 100, 1)} % разброса времени на заказах, которых модель не видела. Скорректированный R² = {num(test['adj_R2'], 3)} почти совпадает с R², потому что признак один, а тестовых строк 80. Типичный промах — {num(test['RMSE'], 3)} мин. Средняя абсолютная процентная ошибка равна {num(test['MAPE_pct'], 2)} %. Средняя разность со знаком {num(test['mean_pred_minus_fact'], 3)} мин: сдвиг меньше одной минуты и меньше разброса ошибок, отдельного смещения прямой на тесте нет.",
        space_before=6,
    )
    add_text(
        doc,
        f"Для показа диспетчеру прогноз округлён до ближайшей целой минуты, половина минуты округляется вверх. Дробный столбец в эту таблицу не входит. Фрагмент первых десяти тестовых заказов в порядке файла приведён в таблице 4. Если сравнивать факт уже с целыми минутами, средняя абсолютная ошибка составляет {num(rounded['MAE'], 2)} мин, R² = {num(rounded['R2'], 3)}. Отличие от таблицы 3 меньше десятой минуты.",
        space_before=12,
    )
    add_table_caption(doc, "Таблица 4 — Фрагмент прогноза, округлённого до минуты", 3.2 + 4.2 + 3.8 + 4.0)
    frag = []
    for r in pred.head(10).itertuples(index=False):
        frag.append(
            [
                str(int(r.order_id)),
                num(r.distance_km, 2),
                num(r.time_min, 1),
                str(int(r.prediction_min)),
            ]
        )
    add_table(
        doc,
        ["Заказ", "Расстояние, км", "Факт, мин", "Прогноз, мин"],
        frag,
        [3.2, 4.2, 3.8, 4.0],
        font=12,
    )
    add_text(
        doc,
        "Диаграмма рассеяния построена по всем 100 заказам. По горизонтали — расстояние, по вертикали — фактическое время. Прямая проведена по коэффициентам, оцененным только на синих точках. Рисунок 2 показывает, что время растёт вместе с расстоянием и точки обоих цветов лежат вдоль одной прямой. Отдельные отклонения вверх и вниз — это заказы, где на время повлияли не только километры: пробка, ожидание у двери, поиск подъезда. Одним признаком такие отклонения не описываются.",
    )
    add_picture(doc, FIG / "scatter.png", 15.0)
    add_text(doc, "Рисунок 2 — Зависимость времени доставки от расстояния", indent=False, align="center", space_before=6)

    doc.add_page_break()
    add_heading(doc, "4 Выводы")
    add_text(
        doc,
        f"Зависимость времени курьерской доставки от расстояния линейная. По 20 заказам, отобранным линейной выборкой, получено уравнение time_min = {num(b0, 3)} + {num(b1, 3)} · distance_km. Каждый километр пути добавляет {num(b1, 3)} мин, ещё {num(b0, 3)} мин уходит на операции, не связанные с дорогой.",
    )
    add_text(
        doc,
        f"На 80 тестовых заказах R² = {num(test['R2'], 3)}, скорректированный R² = {num(test['adj_R2'], 3)}, средняя абсолютная ошибка {num(test['MAE'], 3)} мин, средний квадрат ошибки {num(test['MSE'], 3)} мин², корень из него {num(test['RMSE'], 3)} мин, средняя абсолютная процентная ошибка {num(test['MAPE_pct'], 2)} %. Эти показатели и расположение точек вдоль прямой показывают, что для данного набора линейная регрессия уместна. Прогноз с ошибкой около двух минут пригоден как базовая оценка времени доставки. Пробки и тип здания одним расстоянием не учитываются.",
    )

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
