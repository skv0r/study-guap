# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР6 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import importlib.util
import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SRC_BUILD = BASE.parent / "ЛР2" / "build_report.py"
spec = importlib.util.spec_from_file_location("lr2_docx", SRC_BUILD)
docx = importlib.util.module_from_spec(spec)
spec.loader.exec_module(docx)

OUT = BASE / "ЛР6_Отчет_Буренков_МИИ.docx"
docx.OUT = OUT
FIG = BASE / "figures"
DATA = BASE / "data"
CLASSES = ["футбольный", "баскетбольный", "волейбольный"]
FEATURES = ["elasticity", "size", "mass", "pressure", "bounce"]


def pct(value: float) -> str:
    return f"{float(value) * 100:.2f}".replace(".", ",")


def num(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def meas(feature: str, value) -> str:
    number = float(value)
    if feature in ("elasticity", "pressure"):
        return f"{number:.2f}".replace(".", ",")
    if feature == "size":
        return f"{number:.1f}".replace(".", ",")
    return str(int(round(number)))


def replace_title(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 6",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "МНОГОСЛОЙНЫЙ ПЕРЦЕПТРОН",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {found}")


def build() -> None:
    m = json.loads((DATA / "metrics.json").read_text(encoding="utf-8"))
    df = pd.read_csv(DATA / "balls.csv")
    if len(m.get("comparison", [])) != 5:
        raise RuntimeError("нужно сравнение пяти методов")
    n = m["n"]
    n_train = m["n_train"]
    n_test = m["n_test"]
    n_train_per = m["n_train_per"]
    n_test_per = m["n_test_per"]
    if len(df) != n or n_train + n_test != n:
        raise RuntimeError(f"ожидались {n}, {n_train} и {n_test}")
    if m["correct"] >= n_test:
        raise RuntimeError("нужны ошибки на тесте")
    if m["n_iter"] >= m["max_iter"]:
        raise RuntimeError("обучение не закончилось")
    usual = m["typical"]
    for name in CLASSES:
        matched = 0
        for row in df[df["ball"] == name].itertuples(index=False):
            got = [float(getattr(row, feature)) for feature in FEATURES]
            matched += int(all(abs(a - b) < 1e-9 for a, b in zip(got, usual[name])))
        if matched != 0:
            raise RuntimeError(f"у {name} есть точные копии типового")

    src = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(src, OUT)
    doc = Document(str(OUT))
    replace_title(doc)
    docx.drop_body_after_title(doc)
    docx.add_heading(doc, "СОДЕРЖАНИЕ")
    docx.add_toc(doc, "1 Задание", "@@1@@")
    docx.add_toc(doc, "2 Входные данные", "@@2@@")
    docx.add_toc(doc, "3 Ход выполнения работы", "@@3@@")
    docx.add_toc(doc, "4 Выводы", "@@4@@")

    doc.add_page_break()
    docx.add_heading(doc, "1 Задание")
    docx.add_text(
        doc,
        "Цель работы — отнести измерение мяча к одному из трёх классов многослойным перцептроном "
        "и сравнить точность с методами лабораторных 2, 3, 4 и 5 на тех же строках одного общего теста.",
    )
    docx.add_text(
        doc,
        "По практикуму читается своя таблица и то же разбиение, что в предыдущих работах. "
        "Сеть: пять входов, десять нейронов скрытого слоя с функцией ReLU, три выхода — "
        "по одному на класс. На выходе считаются вероятности классов. Обучение — метод L-BFGS. "
        "Перед обучением признаки приводятся к одному масштабу по обучающей выборке. "
        "Затем считаются матрица ошибок, accuracy, precision, recall и F-measure.",
    )
    docx.add_text(
        doc,
        "Расчёт выполнен в Python. Класс — вид мяча: футбольный, баскетбольный или волейбольный. "
        "Признаки — коэффициент упругости, окружность в сантиметрах, масса в граммах, давление "
        "в барах и высота контрольного отскока в сантиметрах. Номер наблюдения в модель не входит.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "2 Входные данные")
    foot = df[df["ball"] == "футбольный"]
    basket = df[df["ball"] == "баскетбольный"]
    volley = df[df["ball"] == "волейбольный"]
    docx.add_text(
        doc,
        f"Таблица, классы и разбиение те же, что в лабораторных 2, 3, 4 и 5. {n} измерений мяча, "
        f"по {n // 3} футбольных, баскетбольных и волейбольных. Обучение — {n_train} строк "
        f"(по {n_train_per} в классе), тест — {n_test} (по {n_test_per}). Это один общий тест. "
        "Пять признаков: коэффициент упругости, окружность в сантиметрах, масса в граммах, "
        "давление в барах и высота контрольного отскока в сантиметрах. Типовой набор — центр "
        "класса; точных копий типового набора среди 150 измерений нет. Типовые измерения собраны в таблице 1.",
    )
    mean_w = [4.0, 2.2, 2.8, 2.2, 2.6, 2.4]
    docx.add_table_caption(doc, "Таблица 1 — Типовые измерения трёх мячей", sum(mean_w))
    docx.add_table(
        doc,
        ["Класс", "Упругость", "Окружность, см", "Масса, г", "Давление, бар", "Отскок, см"],
        [
            [
                name,
                meas("elasticity", usual[name][0]),
                meas("size", usual[name][1]),
                meas("mass", usual[name][2]),
                meas("pressure", usual[name][3]),
                meas("bounce", usual[name][4]),
            ]
            for name in CLASSES
        ],
        mean_w,
        font=14,
    )
    docx.add_text(
        doc,
        "Диапазоны соседних классов пересекаются: "
        f"упругость футбольного {meas('elasticity', foot['elasticity'].min())}–"
        f"{meas('elasticity', foot['elasticity'].max())}, баскетбольного "
        f"{meas('elasticity', basket['elasticity'].min())}–{meas('elasticity', basket['elasticity'].max())}; "
        f"окружность волейбольного {meas('size', volley['size'].min())}–"
        f"{meas('size', volley['size'].max())} см, футбольного "
        f"{meas('size', foot['size'].min())}–{meas('size', foot['size'].max())} см.",
        space_before=12,
    )

    doc.add_page_break()
    docx.add_heading(doc, "3 Ход выполнения работы")
    weights = m["first_hidden_weights"]
    docx.add_text(
        doc,
        "Схема сети показана на рисунке 1. Скрытый слой — десять нейронов, функция активации ReLU: "
        "отрицательный вход обращается в ноль, положительный проходит дальше. Три выхода "
        "переводятся в вероятности классов. Признаки стандартизованы по обучению: у каждого "
        "признака нулевое среднее и единичный разброс на обучающих строках; те же среднее и "
        "разброс применены к тесту.",
    )
    docx.add_picture(doc, FIG / "scheme.png", 15.5)
    docx.add_figure_caption(doc, "Рисунок 1 — Схема сети: пять входов, десять скрытых нейронов, три класса")
    docx.add_text(
        doc,
        f"Обучение методом L-BFGS закончилось за {m['n_iter']} итерации, значение функции потерь "
        f"{num(m['loss'], 4)}. У первого скрытого нейрона веса на стандартизованные признаки "
        f"упругость, окружность, масса, давление и отскок равны {num(weights[0])}, {num(weights[1])}, "
        f"{num(weights[2])}, {num(weights[3])} и {num(weights[4])}, смещение равно {num(m['hidden_bias'][0])}.",
    )
    example = m["example"]
    probs = example["probs"]
    fx = example["features"]
    docx.add_text(
        doc,
        f"Для строки {example['id']}: упругость {meas('elasticity', fx['elasticity'])}, "
        f"окружность {meas('size', fx['size'])} см, масса {meas('mass', fx['mass'])} г, "
        f"давление {meas('pressure', fx['pressure'])} бар, отскок {meas('bounce', fx['bounce'])} см. "
        f"Факт — «{example['ball']}», прогноз — «{example['pred']}». "
        f"Вероятности: футбольный {pct(probs['футбольный'])} %, "
        f"баскетбольный {pct(probs['баскетбольный'])} %, волейбольный {pct(probs['волейбольный'])} %.",
        space_before=12,
    )
    wrong = m["errors"]
    sample = "; ".join(
        f"№ {item['id']}: факт «{item['ball']}», прогноз «{item['pred']}»" for item in wrong
    )
    docx.add_text(
        doc,
        f"Ошибки на тесте, всего {len(wrong)}: {sample}. Матрица этих ответов приведена в таблице 2.",
        space_before=12,
    )
    cm = m["confusion"]
    cm_w = [4.2, 4.1, 4.1, 4.1]
    docx.add_table_caption(doc, f"Таблица 2 — Матрица ошибок на {n_test} тестовых строках", sum(cm_w))
    docx.add_table(
        doc,
        ["Факт \\ прогноз", *CLASSES],
        [[name, *[str(cm[i][j]) for j in range(3)]] for i, name in enumerate(CLASSES)],
        cm_w,
        font=12,
    )
    per = {row["class"]: row for row in m["per_class"]}
    met_w = [3.4, 3.2, 3.2, 3.4, 3.3]
    docx.add_text(doc, "Показатели по классам приведены в таблице 3.", space_before=12)
    docx.add_table_caption(doc, "Таблица 3 — Precision, recall и F-measure по классам", sum(met_w))
    docx.add_table(
        doc,
        ["Класс", "Precision, %", "Recall, %", "F-measure, %", "Наблюдений"],
        [
            [
                name,
                pct(per[name]["precision"]),
                pct(per[name]["recall"]),
                pct(per[name]["f1"]),
                str(per[name]["support"]),
            ]
            for name in CLASSES
        ],
        met_w,
        font=14,
    )
    docx.add_text(
        doc,
        f"Верных ответов {m['correct']} из {n_test}. Accuracy равна {pct(m['accuracy'])} %, "
        f"средние precision, recall и F-measure — {pct(m['macro_precision'])} %, "
        f"{pct(m['macro_recall'])} % и {pct(m['macro_f1'])} %. "
        "Баскетбольный класс на тесте узнан полностью: recall 100,00 %. "
        "Два футбольных измерения названы баскетбольными, три волейбольных — футбольными.",
        space_before=12,
    )
    docx.add_text(
        doc,
        f"Сравнение сделано на одном и том же тесте из {n_test} строк. Лабораторная 2 — "
        f"k ближайших соседей при k = {m['knn_k']}. Лабораторная 3 — дерево решений глубиной "
        f"{m['tree_depth']}. Лабораторная 4 — наивный Байес. Лабораторная 5 — опорные векторы. "
        "Лабораторная 6 — многослойный перцептрон. Числа собраны в таблице 4.",
        space_before=12,
    )
    cmp_w = [1.7, 3.7, 2.2, 2.3, 2.2, 2.2, 2.2]
    docx.add_table_caption(doc, "Таблица 4 — Сравнение пяти методов на одном тесте", sum(cmp_w))
    docx.add_table(
        doc,
        ["Работа", "Метод", "Верных", "Accuracy, %", "Precision, %", "Recall, %", "F-measure, %"],
        [
            [
                item["lab"],
                item["method"],
                f"{item['correct']} из {item.get('n_test', n_test)}",
                pct(item["accuracy"]),
                pct(item["precision"]),
                pct(item["recall"]),
                pct(item["f1"]),
            ]
            for item in m["comparison"]
        ],
        cmp_w,
        font=12,
    )

    doc.add_page_break()
    docx.add_heading(doc, "4 Выводы")
    ranked = sorted(m["comparison"], key=lambda item: (-item["correct"], -item["f1"], item["lab"]))
    best = ranked[0]
    rest = "; ".join(
        f"{item['method']} — {item['correct']}/{n_test} ({pct(item['accuracy'])} %)"
        for item in ranked[1:]
    )
    docx.add_text(
        doc,
        f"На одном тесте из {n_test} строк лучший результат у метода «{best['method']}»: "
        f"{best['correct']}/{n_test}, accuracy {pct(best['accuracy'])} %, F-measure {pct(best['f1'])} %; "
        f"далее {rest}. Сеть дала {m['correct']} верных ответов из {n_test}, accuracy {pct(m['accuracy'])} %. "
        f"Ключевые параметры сети: пять входов, {m['hidden']} скрытых нейронов с ReLU, три выхода, "
        f"обучение L-BFGS за {m['n_iter']} итерации, стандартизация признаков по обучающей выборке. "
        f"У соседей k = {m['knn_k']}; у дерева глубина {m['tree_depth']} и критерий Джини; "
        "у Байеса нормали по классу; у опорных векторов ядро RBF с сигмой 0,5 и штрафом 1. "
        "Стопроцентной точности нет, потому что диапазоны измерений соседних мячей пересекаются.",
    )
    doc.save(str(OUT))


def assert_clean() -> None:
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "тикет", "велосипед", "iris", "sklearn", "matplotlib", "KNIME",
        "build_report", "compute_", "seed", "codebook", "кодирован",
        "игровая зона", "RProp", "Adam",
    ]
    hits = [word for word in banned if word.lower() in plain.lower()]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    if plain.count("Python") != 1:
        raise RuntimeError(f"Python упомянут {plain.count('Python')} раз")
    for needle in (
        "МНОГОСЛОЙНЫЙ ПЕРЦЕПТРОН",
        "Таблица 4 — Сравнение пяти методов",
        "L-BFGS",
        "ReLU",
        "40 из 45",
        "лучший результат",
        "Ключевые параметры",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")


def main() -> None:
    build()
    docx.strip_unused_images(OUT)
    pages = docx.read_heading_pages()
    docx.patch_toc(pages)
    again = docx.read_heading_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
