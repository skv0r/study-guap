# -*- coding: utf-8 -*-
"""Конспект к устному ответу по ЛР6. В сам конспект этот файл не попадает."""
from __future__ import annotations

import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
FIG = BASE / "conspect_figures"
OUT = BASE / "Конспект_к_защите_ЛР6.docx"


def save(fig, name: str) -> Path:
    FIG.mkdir(exist_ok=True)
    path = FIG / name
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return path


def style_axes(ax, xlabel: str, ylabel: str, title: str) -> None:
    ax.set_xlabel(xlabel, fontsize=11)
    ax.set_ylabel(ylabel, fontsize=11)
    ax.set_title(title, fontsize=12)
    ax.axhline(0, color="#888888", lw=0.7)
    ax.axvline(0, color="#888888", lw=0.7)
    ax.grid(True, lw=0.4, alpha=0.45)
    ax.tick_params(labelsize=10)


def draw_neuron() -> Path:
    fig, ax = plt.subplots(figsize=(9.2, 3.3))
    ax.set_xlim(0, 14)
    ax.set_ylim(0, 5)
    ax.axis("off")
    for name, weight, y in (("x1", "w1", 4.0), ("x2", "w2", 2.5), ("xn", "wn", 1.0)):
        ax.annotate("", xy=(4.15, 2.5), xytext=(1.7, y), arrowprops={"arrowstyle": "->", "lw": 1.2})
        ax.text(0.15, y, name, va="center", fontsize=14)
        ax.text(2.15, (y + 2.5) / 2 + 0.2, weight, fontsize=12)
    ax.add_patch(plt.Circle((4.7, 2.5), 0.55, fill=False, lw=1.5))
    ax.text(4.7, 2.5, "Σ", ha="center", va="center", fontsize=15)
    ax.annotate("", xy=(6.55, 2.5), xytext=(5.25, 2.5), arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.text(5.9, 2.9, "P", ha="center", fontsize=13)
    ax.add_patch(plt.Rectangle((6.55, 1.85), 1.8, 1.3, fill=False, lw=1.5))
    ax.text(7.45, 2.5, "f", ha="center", va="center", fontsize=16)
    ax.annotate("", xy=(10.3, 2.5), xytext=(8.35, 2.5), arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.text(10.45, 2.5, "Y", va="center", fontsize=15)
    ax.text(7.0, 0.35, "вход xi  →  вес wi  →  сумма P  →  функция f  →  выход Y", ha="center", fontsize=11)
    return save(fig, "neuron.png")


def draw_activations() -> Path:
    p = np.linspace(-4, 4, 500)
    curves = [
        (np.where(p >= 0, 1.0, -1.0), "Пороговая: Y зависит от P ступенькой", (-4, 4), (-1.4, 1.4)),
        (p, "Линейная: Y = P, зависимость прямая", (-4, 4), (-4.5, 4.5)),
        (1 / (1 + np.exp(-p)), "Сигмоида: Y плавно растёт от 0 к 1", (-4, 4), (-0.05, 1.05)),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(11.4, 3.6))
    for ax, (y, title, xlim, ylim) in zip(axes, curves):
        ax.plot(p, y, color="#111111", lw=2.2)
        style_axes(ax, "потенциал P", "выход Y", title)
        ax.set_xlim(*xlim)
        ax.set_ylim(*ylim)
    axes[0].plot([0], [1], "o", color="#111111")
    axes[0].annotate("при P ≥ 0 выход 1", xy=(0.15, 1), fontsize=8)
    axes[2].plot([0], [0.5], "o", color="#111111")
    axes[2].annotate("при P = 0 выход 0,5", xy=(0.2, 0.5), fontsize=8)
    return save(fig, "activations.png")


def draw_steepness() -> Path:
    p = np.linspace(-4, 4, 500)
    fig, ax = plt.subplots(figsize=(6.4, 3.8))
    for k, label in ((0.5, "k = 0,5 ближе к прямой"), (1, "k = 1"), (4, "k = 4 ближе к ступеньке")):
        ax.plot(p, 1 / (1 + np.exp(-k * p)), lw=2, label=label)
    style_axes(ax, "потенциал P", "выход Y", "Одна сигмоида: чем больше k, тем круче график")
    ax.set_ylim(-0.05, 1.05)
    ax.legend(frameon=False, fontsize=9)
    return save(fig, "steepness.png")


def draw_relu() -> Path:
    p = np.linspace(-3, 3, 300)
    fig, ax = plt.subplots(figsize=(5.6, 3.6))
    ax.plot(p, np.maximum(0, p), color="#111111", lw=2.2)
    style_axes(ax, "потенциал P", "выход Y", "ReLU: слева от нуля выход 0, справа Y = P")
    ax.set_xlim(-3, 3)
    ax.set_ylim(-0.2, 3.2)
    return save(fig, "relu.png")


def draw_mlp() -> Path:
    fig, ax = plt.subplots(figsize=(8.2, 3.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")
    layers = [(1.3, [4.4, 3.0, 1.6], "вход"), (4.7, [5.0, 3.6, 2.2, 0.8], "скрытый"), (8.2, [3.7, 2.1], "выход")]
    coords = []
    for x, ys, title in layers:
        pts = []
        ax.text(x, 5.65, title, ha="center", fontsize=12)
        for y in ys:
            ax.add_patch(plt.Circle((x, y), 0.28, fill=False, lw=1.3))
            pts.append((x, y))
        coords.append(pts)
    for left, right in zip(coords, coords[1:]):
        for x1, y1 in left:
            for x2, y2 in right:
                ax.annotate(
                    "",
                    xy=(x2 - 0.32, y2),
                    xytext=(x1 + 0.32, y1),
                    arrowprops={"arrowstyle": "->", "lw": 0.7, "color": "#444444"},
                )
    ax.text(5.0, 0.15, "стрелки только вперёд", ha="center", fontsize=11)
    return save(fig, "mlp.png")


def draw_xor() -> Path:
    fig, ax = plt.subplots(figsize=(5.8, 4.4))
    ax.scatter([0, 1], [0, 1], s=80, color="#1a7f37", label="класс А", zorder=3)
    ax.scatter([0, 1], [1, 0], s=80, color="#d95f02", label="класс Б", zorder=3)
    ax.plot([-0.2, 1.2], [0.5, 0.5], ls="--", color="#333333", label="одна прямая не делит")
    ax.set_xlim(-0.3, 1.3)
    ax.set_ylim(-0.3, 1.3)
    ax.set_xlabel("признак x1")
    ax.set_ylabel("признак x2")
    ax.set_title("Исключающее ИЛИ: прямой границы нет")
    ax.legend(frameon=False, fontsize=9, loc="upper center", bbox_to_anchor=(0.5, -0.18), ncol=3)
    ax.set_aspect("equal")
    return save(fig, "xor.png")


def hopfield_run() -> tuple[list[float], list[str]]:
    pattern = np.array([1, -1, 1, -1])
    weight = np.outer(pattern, pattern).astype(float)
    np.fill_diagonal(weight, 0.0)

    def energy(state: np.ndarray) -> float:
        return float(-0.5 * state @ weight @ state)

    state = np.array([1, 1, 1, -1])
    energies = [energy(state)]
    states = [" ".join(str(int(v)) for v in state)]
    for step in (1, 0, 2, 3):
        field = weight[step] @ state
        state = state.copy()
        state[step] = 1 if field >= 0 else -1
        energies.append(energy(state))
        states.append(" ".join(str(int(v)) for v in state))
    if abs(energies[-1] - energy(pattern)) > 1e-9:
        raise RuntimeError(f"Хопфилд не дошёл до образа: {states} {energies}")
    if any(right > left + 1e-9 for left, right in zip(energies, energies[1:])):
        raise RuntimeError(f"энергия выросла: {energies}")
    return energies, states


def draw_energy(energies: list[float]) -> Path:
    fig, ax = plt.subplots(figsize=(6.2, 3.6))
    xs = list(range(len(energies)))
    ax.plot(xs, energies, marker="o", color="#111111", lw=2)
    for x, y in zip(xs, energies):
        ax.annotate(f"{y:.0f}", xy=(x, y), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=10)
    ax.set_xticks(xs)
    ax.set_xlabel("шаг")
    ax.set_ylabel("энергия")
    ax.set_title("Хопфилд: энергия падает и останавливается")
    ax.grid(True, lw=0.4, alpha=0.45)
    return save(fig, "energy.png")


def draw_recurrent() -> Path:
    fig, ax = plt.subplots(figsize=(7.4, 3.2))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")
    ax.add_patch(plt.Rectangle((3.4, 1.6), 3.0, 1.8, fill=False, lw=1.4))
    ax.text(4.9, 2.5, "слой", ha="center", va="center", fontsize=13)
    ax.annotate("", xy=(3.4, 2.5), xytext=(1.1, 2.5), arrowprops={"arrowstyle": "->", "lw": 1.3})
    ax.text(0.2, 2.7, "вход", fontsize=12)
    ax.annotate("", xy=(8.7, 2.5), xytext=(6.4, 2.5), arrowprops={"arrowstyle": "->", "lw": 1.3})
    ax.text(8.85, 2.35, "выход", fontsize=12)
    ax.annotate("", xy=(4.9, 1.6), xytext=(4.9, 0.7), arrowprops={"arrowstyle": "-", "lw": 1.2})
    ax.annotate("", xy=(2.2, 2.15), xytext=(4.9, 0.7), arrowprops={"arrowstyle": "->", "lw": 1.2})
    ax.text(5.0, 0.2, "прошлый выход возвращается на вход", ha="center", fontsize=11)
    return save(fig, "recurrent.png")


def draw_hopfield() -> Path:
    fig, ax = plt.subplots(figsize=(4.6, 4.4))
    ax.set_xlim(-1.6, 1.6)
    ax.set_ylim(-1.7, 1.55)
    ax.set_aspect("equal")
    ax.axis("off")
    pts = [(0, 1), (-1, 0), (0, -1), (1, 0)]
    for i in range(4):
        for j in range(i + 1, 4):
            ax.plot([pts[i][0], pts[j][0]], [pts[i][1], pts[j][1]], color="#444444", lw=1.1)
    for (x, y), name in zip(pts, "1234"):
        ax.add_patch(plt.Circle((x, y), 0.2, facecolor="white", edgecolor="black", lw=1.3, zorder=3))
        ax.text(x, y, name, ha="center", va="center", fontsize=12, zorder=4)
    ax.set_title("все со всеми, на себя связи нет", fontsize=11)
    return save(fig, "hopfield.png")


def kohonen_distances() -> tuple[np.ndarray, int]:
    point = np.array([1.0, 0.0])
    weights = np.array([[0.0, 1.0], [0.8, 0.2], [-1.0, 0.0]])
    dist = np.linalg.norm(weights - point, axis=1)
    winner = int(np.argmin(dist))
    if winner != 1:
        raise RuntimeError(dist)
    return dist, winner


def draw_distances(dist: np.ndarray, winner: int) -> Path:
    fig, ax = plt.subplots(figsize=(5.8, 3.6))
    labels = ["нейрон 1", "нейрон 2", "нейрон 3"]
    colors = ["#bbbbbb", "#f4d35e", "#bbbbbb"]
    bars = ax.bar(labels, dist, color=colors, edgecolor="black")
    for bar, value in zip(bars, dist):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.05, f"{value:.2f}".replace(".", ","), ha="center", fontsize=11)
    ax.set_ylabel("расстояние до входа")
    ax.set_title("Кохонен: побеждает самый низкий столбик")
    ax.set_ylim(0, max(dist) + 0.5)
    ax.tick_params(labelsize=11)
    return save(fig, "distances.png")


def draw_kohonen() -> Path:
    fig, ax = plt.subplots(figsize=(7.6, 3.4))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")
    ax.text(0.2, 2.5, "вход", fontsize=12, va="center")
    for i, y in enumerate((3.6, 2.5, 1.4)):
        ax.annotate("", xy=(3.2, y), xytext=(1.5, 2.5), arrowprops={"arrowstyle": "->", "lw": 1.0})
        face = "#f4d35e" if i == 1 else "white"
        ax.add_patch(plt.Circle((3.6, y), 0.32, facecolor=face, edgecolor="black", lw=1.2))
        ax.text(4.15, y, f"нейрон {i + 1}", va="center", fontsize=11)
    for r in range(3):
        for c in range(3):
            x, y = 7.0 + c * 0.72, 1.2 + r * 0.72
            face = "#f4d35e" if (r, c) == (1, 1) else "white"
            ax.add_patch(plt.Rectangle((x, y), 0.64, 0.64, facecolor=face, edgecolor="black", lw=0.9))
    ax.text(8.05, 3.6, "похожие входы\nсадятся рядом", ha="center", fontsize=10)
    return save(fig, "kohonen.png")


def convolution() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    image = np.array([[1, 0, 2, 1], [0, 1, 1, 0], [2, 1, 0, 2], [1, 0, 1, 1]], dtype=int)
    kernel = np.array([[1, -1], [1, -1]], dtype=int)
    out = np.zeros((3, 3), dtype=int)
    for i in range(3):
        for j in range(3):
            out[i, j] = int(np.sum(image[i : i + 2, j : j + 2] * kernel))
    if int(out[0, 0]) != 0 or int(out[0, 1]) != -2:
        raise RuntimeError(out)
    return image, kernel, out


def draw_grid(ax, matrix: np.ndarray, title: str, highlight=None) -> None:
    ax.set_xlim(-0.5, matrix.shape[1] - 0.5)
    ax.set_ylim(matrix.shape[0] - 0.5, -0.5)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title(title, fontsize=12)
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            face = "#fff3bf" if highlight and highlight[0] <= i < highlight[2] and highlight[1] <= j < highlight[3] else "white"
            ax.add_patch(plt.Rectangle((j - 0.5, i - 0.5), 1, 1, facecolor=face, edgecolor="black", lw=1.1))
            ax.text(j, i, str(int(matrix[i, j])), ha="center", va="center", fontsize=14)


def draw_conv(image: np.ndarray, kernel: np.ndarray, out: np.ndarray) -> Path:
    fig, axes = plt.subplots(1, 3, figsize=(9.2, 3.3))
    draw_grid(axes[0], image, "картинка, жёлтым первое окно", (0, 0, 2, 2))
    draw_grid(axes[1], kernel, "ядро")
    draw_grid(axes[2], out, "результат, первая клетка 0")
    return save(fig, "conv.png")


def draw_pool() -> Path:
    src = np.array([[4, 1, 0, 2], [3, 5, 1, 0], [0, 2, 8, 1], [1, 0, 3, 4]])
    pooled = np.array([[5, 2], [2, 8]])
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.4))
    draw_grid(axes[0], src, "окно 2×2, максимум 5", (0, 0, 2, 2))
    draw_grid(axes[1], pooled, "после подвыборки")
    return save(fig, "pool.png")


def draw_cnn() -> Path:
    fig, ax = plt.subplots(figsize=(9.6, 2.4))
    ax.set_xlim(0, 16)
    ax.set_ylim(0, 3)
    ax.axis("off")
    blocks = [(0.2, "картинка"), (2.8, "свёртка"), (5.4, "ReLU"), (8.0, "окно\nмаксимума"), (10.6, "свёртка"), (13.2, "класс")]
    for x, label in blocks:
        ax.add_patch(plt.Rectangle((x, 0.7), 2.3, 1.5, fill=False, lw=1.3))
        ax.text(x + 1.15, 1.45, label, ha="center", va="center", fontsize=11)
    for x in (2.5, 5.1, 7.7, 10.3, 12.9):
        ax.annotate("", xy=(x + 0.3, 1.45), xytext=(x, 1.45), arrowprops={"arrowstyle": "->", "lw": 1.1})
    return save(fig, "cnn.png")


def set_run(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_p(doc: Document, text: str, *, bold=False, center=False, size=14, space_before=0, space_after=6, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(1.25) if indent and not center else Cm(0)
    set_run(p.add_run(text), size=size, bold=bold)
    return p


def add_h(doc: Document, text: str):
    return add_p(doc, text, bold=True, center=True, space_before=14, space_after=8, indent=False)


def add_say(doc: Document, text: str):
    return add_p(doc, text, space_after=4)


def add_pic(doc: Document, path: Path, caption: str, width=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.first_line_indent = Cm(0)
    set_run(cap.add_run(caption), size=14)


def main() -> None:
    energies, states = hopfield_run()
    dist, winner = kohonen_distances()
    image, kernel, conv = convolution()
    sig = 1 / (1 + math.exp(1.5))
    if abs(sig - 0.1824) > 0.001:
        raise RuntimeError(sig)

    paths = {
        "neuron": draw_neuron(),
        "acts": draw_activations(),
        "steep": draw_steepness(),
        "relu": draw_relu(),
        "mlp": draw_mlp(),
        "xor": draw_xor(),
        "energy": draw_energy(energies),
        "rec": draw_recurrent(),
        "hop": draw_hopfield(),
        "dist": draw_distances(dist, winner),
        "koh": draw_kohonen(),
        "conv": draw_conv(image, kernel, conv),
        "pool": draw_pool(),
        "cnn": draw_cnn(),
    }

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.5)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(14)

    add_h(doc, "КАК ОТВЕЧАТЬ НА ЗАЩИТЕ ЛАБОРАТОРНОЙ 6")
    add_p(
        doc,
        "На листок пишется блок «Сказать» и перерисовывается график под ним. График — это зависимость: по горизонтали причина, по вертикали результат. Смотрю на линию и говорю, что с ней происходит. Блок «Если спросят» вслух, только когда вопрос уже задан.",
        indent=False,
    )

    add_h(doc, "1. Модель искусственного нейрона")
    add_p(doc, "Сказать.", bold=True, indent=False)
    for line in (
        "Искусственный нейрон повторяет нервную клетку в числах. Вход — это дендрит, вес — сила синапса, сумма считается в теле, выход уходит по аксону.",
        "На входы приходят числа x1, x2 и так далее. У каждого входа свой вес. Потенциал — их взвешенная сумма: P = w1x1 + w2x2 + … + wnxn.",
        "Выход — одно число Y = f(P). Буква f — передаточная функция. Обучение — это подбор весов, формула нейрона при этом не меняется.",
        "В лекции три графика зависимости Y от P. Пороговый: пока P ниже нуля, выход −1, как только P перешёл через ноль, выход сразу 1. Линейный: выход растёт ровно вместе с P, Y = kP. Сигмоида: выход плавно поднимается от 0 к 1. Крутизна k делает тот же график то положе, то ближе к ступеньке.",
    ):
        add_say(doc, line)
    add_pic(doc, paths["neuron"], "Рисунок 1 — Что нарисовать первым: вход, сумма, функция, выход")
    add_p(doc, "Глядя на рисунок 1, говорю: слева сигналы, на стрелках веса, в круге сумма P, в квадрате функция, справа один выход Y.", indent=False)
    add_pic(doc, paths["acts"], "Рисунок 2 — Три графика зависимости выхода Y от потенциала P", width=16.2)
    add_p(
        doc,
        "Глядя на рисунок 2, говорю слева направо. Первая линия ломается в нуле: это порог. Вторая — прямая через начало координат: это линейная зависимость. Третья ползёт от нуля к единице и в нуле равна 0,5: это сигмоида.",
        indent=False,
    )
    add_pic(doc, paths["steep"], "Рисунок 3 — Та же сигмоида при разной крутизне k", width=13)
    add_p(
        doc,
        "Глядя на рисунок 3, говорю: формула одна, Y = 1 / (1 + e в степени −kP). Малое k даёт пологую кривую, почти прямую. Большое k прижимает кривую к ступеньке. Поэтому сигмоида стоит между линейной функцией и порогом.",
        indent=False,
    )
    add_p(
        doc,
        "Числа, если попросят посчитать. x1 = 1, x2 = 2, w1 = 0,5, w2 = −1. Тогда P = 0,5·1 + (−1)·2 = −1,5. Порог даёт −1. Прямая при k = 1 даёт −1,5. Сигмоида при k = 1 даёт примерно "
        + f"{sig:.2f}".replace(".", ",")
        + ". Потенциал один, график функции другой — выход другой.",
    )
    add_p(doc, "Если спросят.", bold=True, indent=False, space_before=8)
    for line in (
        "Почему порог плохо учится? Ниже ступеньки выход один и тот же, сеть не видит, далеко она от порога или рядом. На сигмоиде наклон есть в каждой точке, по нему вес можно сдвинуть.",
        "Что такое вес? Насколько сильно этот вход тянет сумму. Вход — сам сигнал, вес — его важность.",
        "Чем нейрон отличается от сети? Нейрон выдаёт одно число. Сеть — много нейронов, выход одного приходит на вход другого.",
    ):
        add_say(doc, line)

    add_h(doc, "2. Четыре вида сетей")
    add_p(
        doc,
        "По лекции порядок такой: многослойный перцептрон, рекуррентная сеть вместе с Хопфилдом, сеть Кохонена, свёрточная сеть. На каждый вид — пять фраз и один рисунок. Вместе это около полутора страниц устного ответа.",
        indent=False,
    )

    add_h(doc, "2.1. Многослойный перцептрон")
    add_p(doc, "Сказать.", bold=True, indent=False)
    for line in (
        "Нейроны стоят слоями. Связь есть только с соседним слоем и только вперёд. Внутри слоя стрелок нет.",
        "Один слой умеет отделить классы одной прямой. Это предел, который показали Минский и Пейперт. Если классы прямой не делятся, одного слоя мало. Несколько слоёв строят ломаную границу.",
        "На рисунке с точками это видно сразу. Зелёные точки (0, 0) и (1, 1) — один класс, оранжевые (0, 1) и (1, 0) — другой. Любая прямая берёт не ту пару. Это исключающее ИЛИ.",
        "Учат так: показали пример, ответ верный — веса не трогаем, ответ неверный — веса сдвигаем. В нескольких слоях ошибку передают от выхода назад. Это обратное распространение.",
        "Плюс: таблица чисел с известными классами, наша лабораторная именно такая. Минус: картинку из пикселей так считать тяжело, прошлого шага сеть не помнит.",
    ):
        add_say(doc, line)
    add_pic(doc, paths["mlp"], "Рисунок 4 — Сеть прямого хода")
    add_pic(doc, paths["xor"], "Рисунок 5 — График: одной прямой два класса не разделить", width=12)
    add_p(doc, "Глядя на рисунок 5, говорю: пунктир — попытка прямой границы. Снизу оказались зелёная и оранжевая точки сразу, значит прямая классы смешала.", indent=False)
    add_pic(doc, paths["relu"], "Рисунок 6 — ReLU в скрытом слое нашей сети", width=11)
    add_p(
        doc,
        "Наша сеть в отчёте: 5 входов, 10 скрытых нейронов, 3 выхода. В скрытом слое ReLU, график на рисунке 6: отрицательный потенциал даёт ноль, положительный проходит как есть. Обучение L-BFGS кончилось за 33 шага. На тесте 40 верных из 45.",
        indent=False,
    )
    add_p(doc, "Если спросят.", bold=True, indent=False, space_before=8)
    for line in (
        "Почему не все 45? Измерения соседних мячей пересекаются. Два футбольных названы баскетбольными, три волейбольных — футбольными.",
        "RProp в практикуме и L-BFGS в отчёте — два способа подобрать те же веса. Сеть в обоих случаях многослойный перцептрон.",
        "Зачем выравнивать признаки? Масса в сотнях, упругость около единицы. Без этого масса забьёт сумму P.",
    ):
        add_say(doc, line)

    add_h(doc, "2.2. Рекуррентная сеть и Хопфилд")
    add_p(doc, "Сказать.", bold=True, indent=False)
    for line in (
        "У рекуррентной сети выход возвращается на вход. Поэтому один и тот же новый сигнал может дать разный ответ: сеть ещё помнит, что выдала раньше.",
        "В лекции конкретная модель — Хопфилд, ассоциативная память. Нейроны связаны со всеми, связь симметрична: wji = wij. На себя связи нет: wii = 0.",
        "Правило Хебба: два нейрона вспыхнули вместе — связь крепче, порознь — слабее.",
        "График энергии от шага идёт вниз и встаёт. Точка остановки — запомненный образ. Испорченный пример, если он близко, достраивается до этого образа.",
        "Плюс: восстановить знакомый узор по обрывку. Минус: связей очень много, без симметрии энергия из лекции не строится, слишком много образов сеть путает. Для пяти чисел про мяч и для большой фотографии это плохой выбор.",
    ):
        add_say(doc, line)
    add_pic(doc, paths["rec"], "Рисунок 7 — Обратная связь")
    add_pic(doc, paths["hop"], "Рисунок 8 — Четыре нейрона Хопфилда", width=10)
    add_pic(doc, paths["energy"], "Рисунок 9 — График зависимости энергии от шага", width=12.5)
    add_p(
        doc,
        "Маленький пример к рисунку 9. Запомненный образ: 1  −1  1  −1. Старт с ошибкой во втором знаке: "
        + states[0]
        + ". Энергия по шагам: "
        + ", ".join(f"{value:.0f}" for value in energies)
        + ". Уже на следующем шаге состояние "
        + states[-1]
        + ", это запомненный образ, и энергия дальше стоит на месте. Глядя на график, говорю: кривая падает и становится горизонтальной, сеть села в минимум.",
        indent=False,
    )
    add_p(doc, "Если спросят.", bold=True, indent=False, space_before=8)
    for line in (
        "Почему нельзя связь нейрона с самим собой? Он начнёт держать своё состояние независимо от остальных.",
        "Чем это не перцептрон? У перцептрона сигнал прошёл и забылся. Здесь выход или соседи возвращаются.",
        "Незнакомый образ всё равно куда-то сядет. Это может быть чужой запомненный образ, нового класса сеть так не выучит.",
    ):
        add_say(doc, line)

    add_h(doc, "2.3. Сеть Кохонена")
    add_p(doc, "Сказать.", bold=True, indent=False)
    for line in (
        "Готовых ответов нет. Нейроны соревнуются. У каждого свой вектор весов той же длины, что вход.",
        "Считаю расстояние от входа до каждого вектора. Побеждает самый близкий, на графике это самый низкий столбик. Остальные в этот момент молчат. Веса победителя чуть сдвигаются к показанному входу.",
        "Если нейроны разложены на клетчатой карте, вместе с победителем чуть двигают соседей. Похожие входы садятся в соседние клетки.",
        "Плюс: разложить примеры по группам, когда подписей классов нет. Минус: число нейронов задают заранее, а группы могут получиться бессмысленными. В нашей лабораторной классы мячей уже подписаны, поэтому здесь уместнее перцептрон.",
    ):
        add_say(doc, line)
    add_pic(doc, paths["dist"], "Рисунок 10 — График: расстояние от входа до трёх нейронов", width=12)
    dist_text = ", ".join(f"{value:.2f}".replace(".", ",") for value in dist)
    add_p(
        doc,
        f"Числа рисунка 10. Вход (1, 0). Веса нейронов (0, 1), (0,8, 0,2) и (−1, 0). Расстояния {dist_text}. Ниже всех второй столбик, победитель — нейрон {winner + 1}. Глядя на график, говорю: сравниваю высоты, минимум и есть ответ.",
        indent=False,
    )
    add_pic(doc, paths["koh"], "Рисунок 11 — Победитель и карта")
    add_p(doc, "Если спросят.", bold=True, indent=False, space_before=8)
    for line in (
        "Это с учителем? Нет. Учителя с готовым классом нет.",
        "Как выбрать победителя на доске? Вычесть вход из весов, возвести разности в квадрат, сложить, взять корень. У кого меньше, тот победил.",
        "Нейронов меньше, чем групп, — две группы слипнутся. Нейронов больше — одна группа займёт несколько клеток.",
    ):
        add_say(doc, line)

    add_h(doc, "2.4. Свёрточная сеть")
    add_p(doc, "Сказать.", bold=True, indent=False)
    for line in (
        "Это сеть прямого хода для картинки. Её собрал Ян Лекун в 1988 году. Слои свёртки чередуются со слоями, где из окна берут максимум. Учат обратным распространением. Функция после свёртки любая, обычно ReLU.",
        "По картинке едет одно маленькое ядро. Кусок картинки умножают на ядро клетка к клетке, произведения складывают, сумма становится одной клеткой результата. Ядро везде одно и то же.",
        "Дальше окно 2×2 оставляет только максимум и картинка уменьшается. Потом снова свёртка и в конце обычный слой, который называет класс.",
        "Плюс: мало весов и один и тот же край ищется в любом месте кадра. Минус: нужно много картинок, точное место после максимума теряется. Пять чисел про мяч — не картинка, свёртке там не по чему ехать.",
    ):
        add_say(doc, line)
    add_pic(doc, paths["cnn"], "Рисунок 12 — Порядок блоков")
    add_pic(doc, paths["conv"], "Рисунок 13 — Свёртка на числах, жёлтым первое окно", width=16)
    add_p(
        doc,
        "Считаю вслух первое окно рисунка 13. В окне сверху 1 и 0, снизу 0 и 1. Ядро сверху 1 и −1, снизу 1 и −1. Сумма: 1·1 + 0·(−1) + 0·1 + 1·(−1) = 0. Соседнее окно сдвигается на столбец: 0·1 + 2·(−1) + 1·1 + 1·(−1) = −2. Эти два числа — левые клетки верхнего ряда результата.",
        indent=False,
    )
    add_pic(doc, paths["pool"], "Рисунок 14 — Из окна остаётся максимум", width=13)
    add_p(
        doc,
        "Глядя на рисунок 14, говорю: в жёлтом окне числа 4, 1, 3 и 5. Максимум 5 уходит в результат, остальные три числа пропадают. Карта становится меньше.",
        indent=False,
    )
    add_p(doc, "Если спросят.", bold=True, indent=False, space_before=8)
    for line in (
        "Чем свёртка не обычный нейрон? Обычный нейрон видит весь вход сразу. Свёртка видит маленькое окно и таскает одни веса по всей картинке.",
        "Зачем ReLU после свёртки? Свёртка — сумма произведений, она линейная. Несколько линейных слоёв равны одному. ReLU ставит излом: минус обнуляется.",
        "Ядро из одних единиц просто сложит яркость окна, край так не найти.",
        "Почему мячи считал перцептрон? Признаки уже готовые пять чисел, сетки пикселей нет.",
    ):
        add_say(doc, line)

    add_h(doc, "Связный ответ примерно на полторы страницы")
    add_p(
        doc,
        "Сети отличаются соединением нейронов. Многослойный перцептрон соединён только вперёд. Один слой проводит прямую границу, несколько слоёв — ломаную, поэтому берутся и классы вроде исключающего ИЛИ. Учат по ошибке. Такая сеть хороша для таблицы с известным классом: у нас 5 входов, 10 скрытых нейронов и 3 класса мяча, 40 ответов из 45 верные. Для фотографии она уже тяжёлая.",
    )
    add_p(
        doc,
        "Рекуррентная сеть возвращает выход назад и помнит прошлое. Модель Хопфилда — память: все связаны со всеми, связи симметричны, петли на себя нет. График энергии от шага падает и встаёт в запомненном образе. Хороша, чтобы достроить знакомый узор. Плоха, когда связей слишком много или образа в памяти нет.",
    )
    add_p(
        doc,
        "Кохонен работает без подписей. Побеждает нейрон с самым маленьким расстоянием до входа, это нижний столбик на графике. На карте похожие входы ложатся рядом. Хороша для групп без учителя. Плоха, когда класс уже написан и нужна матрица ошибок.",
    )
    add_p(
        doc,
        "Свёрточная сеть возит одно ядро по картинке: окно умножается на ядро, сумма пишется в клетку, потом из соседнего окна берётся максимум. Весов мало, край ищется где угодно в кадре. Хороша для снимка. Для пяти чисел про мяч её не берут.",
    )

    add_h(doc, "Если спросят одной фразой")
    for line in (
        "ИНС — набор нейронов. У каждого несколько входов и один выход, выходы одних идут на входы других.",
        "Формула нейрона: P = сумма wi·xi, Y = f(P).",
        "Три графика лекции: ступенька, прямая, сигмоида. Крутизна k двигает сигмоиду от прямой к ступеньке.",
        "Четыре вида: перцептрон, Хопфилд, Кохонен, свёрточная сеть.",
        "Учитель есть у перцептрона и у свёрточной сети. У Хопфилда в память кладут образ. У Кохонена готового класса нет.",
        "В отчёте сеть 5—10—3, ReLU, L-BFGS, 33 шага, 40 из 45.",
    ):
        add_say(doc, line)

    doc.save(str(OUT))
    print(OUT)
    print("energy", energies)
    print("states", states)
    print("dist", [round(float(v), 4) for v in dist])


if __name__ == "__main__":
    main()
