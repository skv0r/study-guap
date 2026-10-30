# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР7 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import importlib.util
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SRC_BUILD = BASE.parent / "ЛР2" / "build_report.py"
spec = importlib.util.spec_from_file_location("lr2_docx", SRC_BUILD)
docx = importlib.util.module_from_spec(spec)
spec.loader.exec_module(docx)

OUT = BASE / "ЛР7_Отчет_Буренков_МИИ.docx"
docx.OUT = OUT
FIG = BASE / "figures"
DATA = BASE / "data"


def num(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def integer(value: float) -> str:
    number = float(value)
    if abs(number - round(number)) > 1e-9:
        raise RuntimeError(f"ожидалось целое: {value}")
    return str(int(round(number)))


def replace_title(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 7",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "ТРИ АЛГОРИТМА КЛАССИФИКАЦИИ",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {found}")


def build() -> None:
    m = json.loads((DATA / "examples.json").read_text(encoding="utf-8"))
    knn = m["knn"]
    bayes = m["bayes"]
    net = m["perceptron"]
    q = knn["query"]

    src = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(src, OUT)
    doc = Document(str(OUT))
    replace_title(doc)
    docx.drop_body_after_title(doc)
    docx.add_heading(doc, "СОДЕРЖАНИЕ")
    for title, key in (
        ("1 Задание", "@@1@@"),
        ("2 Метод ближайших соседей", "@@2@@"),
        ("3 Наивный Байес", "@@3@@"),
        ("4 Перцептрон", "@@4@@"),
        ("5 Выводы", "@@5@@"),
    ):
        docx.add_toc(doc, title, key)

    doc.add_page_break()
    docx.add_heading(doc, "1 Задание")
    docx.add_text(
        doc,
        "Цель работы — разобрать три разных алгоритма классификации на коротких числовых примерах, "
        "которые можно заново посчитать на доске. Алгоритмы выбраны разные по правилу решения: "
        "голоса соседей по расстоянию, сравнение нормальных плотностей и порог одного нейрона.",
    )
    docx.add_text(
        doc,
        "Предметная область во всех трёх примерах одна: мяч относится к футбольному, баскетбольному "
        "или волейбольному. Признаки примеров — окружность в сантиметрах и высота контрольного "
        "отскока в сантиметрах. У соседей шесть размеченных мячей, у Байеса девять измерений "
        "окружности, у перцептрона четыре объекта. У каждого алгоритма ниже записаны формула, "
        "исходные числа, ход счёта и класс, который получается в конце.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "2 Метод ближайших соседей")
    docx.add_text(
        doc,
        "Есть шесть уже размеченных мячей и одно новое измерение Н: окружность "
        f"{q['size']} см, отскок {q['bounce']} см. Класс Н ищется по {knn['k']} ближайшим объектам. "
        "Расстояние евклидово:",
    )
    docx.add_text(
        doc,
        "d = √((71 − s)² + (125 − h)²),   d² = (71 − s)² + (125 − h)².",
    )
    docx.add_text(
        doc,
        "Корень — возрастающая функция, поэтому сосед ближе ровно тогда, когда меньше d². "
        "На доске достаточно квадратов. Расположение точек показано на рисунке 1, "
        "числа расстояний — в таблице 1.",
    )
    docx.add_picture(doc, FIG / "neighbors.png", 14.5)
    docx.add_figure_caption(doc, "Рисунок 1 — Шесть размеченных мячей и новое измерение Н")
    ranked = knn["ranked"]
    cap_w = [1.6, 3.4, 2.6, 2.2, 1.8, 1.8, 1.6]
    docx.add_table_caption(doc, "Таблица 1 — Расстояния от измерения Н до шести мячей", sum(cap_w))
    docx.add_table(
        doc,
        ["Объект", "Класс", "Окружность, см", "Отскок, см", "d²", "d", "Ранг"],
        [
            [
                row["name"],
                row["ball"],
                str(row["size"]),
                str(row["bounce"]),
                integer(row["d2"]),
                num(row["d"]),
                str(row["rank"]),
            ]
            for row in ranked
        ],
        cap_w,
        font=12,
    )
    top = ranked[:3]
    docx.add_text(
        doc,
        "Квадраты считаются так. Ф1: (71 − 70)² + (125 − 126)² = 1 + 1 = 2. "
        "Ф2: (71 − 72)² + (125 − 127)² = 1 + 4 = 5. "
        "В2: (71 − 69)² + (125 − 119)² = 4 + 36 = 40. "
        "В1: (71 − 68)² + (125 − 116)² = 9 + 81 = 90. "
        "Б1: (71 − 73)² + (125 − 136)² = 4 + 121 = 125. "
        "Б2: (71 − 75)² + (125 − 140)² = 16 + 225 = 241.",
        space_before=12,
    )
    names = ", ".join(f"{row['name']} (d² = {integer(row['d2'])}, {row['ball']})" for row in top)
    votes = knn["votes"]
    docx.add_text(
        doc,
        f"Три ближайших: {names}. Голоса: футбольный — {votes.get('футбольный', 0)}, "
        f"волейбольный — {votes.get('волейбольный', 0)}, баскетбольный — {votes.get('баскетбольный', 0)}. "
        f"Большинство одно: класс измерения Н — {knn['winner']}.",
    )
    docx.add_text(
        doc,
        "При k = 1 ближайший один, это Ф1, и класс тоже футбольный. При k = 5 в пятёрку входят "
        "Ф1, Ф2, В2, В1 и Б1: футбольных два и волейбольных два. Большинства нет, поэтому для "
        "этого набора взято k = 3, где ответ определяется без жеребьёвки.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "3 Наивный Байес")
    docx.add_text(
        doc,
        "Здесь один признак — окружность, и три класса. В каждом классе три измерения, "
        "априорные вероятности равны 1/3. Плотность класса принята нормальной:",
    )
    docx.add_text(
        doc,
        "f(x) = 1 / √(2πσ²) · exp(−(x − μ)² / (2σ²)).",
    )
    docx.add_text(
        doc,
        "Среднее μ — среднее арифметическое трёх измерений. Дисперсия — выборочная, с делителем "
        "n − 1. Для волейбольного: измерения 67, 68 и 69 см, μ = (67 + 68 + 69) / 3 = 68. "
        "Сумма квадратов отклонений 1² + 0² + 1² = 2, поэтому σ² = 2 / 2 = 1. "
        "У футбольного измерения 69, 70 и 71 см, μ = 70, σ² = 1. "
        "У баскетбольного 72, 73 и 74 см, μ = 73, σ² = 1.",
    )
    order = ["волейбольный", "футбольный", "баскетбольный"]
    by_ball = {item["ball"]: item for item in bayes["classes"]}
    bayes_w = [3.4, 3.4, 2.2, 2.4, 2.4, 2.6]
    docx.add_text(
        doc,
        "Параметры плотностей и вероятности для x = 71 см собраны в таблице 2.",
    )
    docx.add_table_caption(doc, "Таблица 2 — Параметры нормальных плотностей трёх классов по окружности", sum(bayes_w))
    docx.add_table(
        doc,
        ["Класс", "Измерения, см", "μ, см", "σ²", "Вес exp", "Вероятность"],
        [
            [
                name,
                ", ".join(str(value) for value in by_ball[name]["values"]),
                integer(by_ball[name]["mu"]),
                integer(by_ball[name]["var"]),
                num(by_ball[name]["weight"], 4),
                num(by_ball[name]["prob"], 4),
            ]
            for name in order
        ],
        bayes_w,
        font=12,
    )
    foot = by_ball["футбольный"]
    basket = by_ball["баскетбольный"]
    volley = by_ball["волейбольный"]
    rounded_weights = [round(item["weight"], 4) for item in (foot, basket, volley)]
    if abs(sum(rounded_weights) - 0.7529) > 1e-9:
        raise RuntimeError(rounded_weights)
    rounded_probs = [round(item["prob"], 4) for item in (foot, basket, volley)]
    if abs(sum(rounded_probs) - 1) > 1e-9:
        raise RuntimeError(rounded_probs)
    docx.add_text(
        doc,
        f"Новое измерение x = {bayes['x']} см. Множитель 1/√(2πσ²) при σ² = 1 одинаков у всех "
        "классов, априорные вероятности тоже равны. Сравнение сводится к весу "
        "exp(−(x − μ)² / 2), затем веса делятся на их сумму.",
        space_before=12,
    )
    docx.add_text(
        doc,
        f"Футбольный: (71 − 70)² / 2 = 0,5; exp(−0,5) = {num(foot['weight'], 4)}. "
        f"Баскетбольный: (71 − 73)² / 2 = 2; exp(−2) = {num(basket['weight'], 4)}. "
        f"Волейбольный: (71 − 68)² / 2 = 4,5; exp(−4,5) = {num(volley['weight'], 4)}. "
        f"На доске веса удобно писать четырьмя знаками: {num(foot['weight'], 4)}, "
        f"{num(basket['weight'], 4)} и {num(volley['weight'], 4)}. Их сумма равна 0,7529. "
        f"Точная сумма, округлённая отдельно, равна {num(bayes['total'], 4)}: расхождение в одну единицу "
        "четвёртого знака появляется из-за округления слагаемых. Вероятности в таблице 2 получены "
        "делением точных весов.",
    )
    docx.add_text(
        doc,
        f"После деления на сумму: футбольный {num(foot['prob'], 4)}, "
        f"баскетбольный {num(basket['prob'], 4)}, волейбольный {num(volley['prob'], 4)}. "
        f"Наибольшая вероятность у класса «{bayes['winner']}».",
    )
    docx.add_text(
        doc,
        "Граница футбольного и баскетбольного при равных дисперсиях и равных априорных — середина "
        f"средних, {num(bayes['boundary'], 1)} см. В этой точке оба веса равны "
        f"{num(bayes['tied_weight'], 4)}. Измерение 71 см лежит по футбольную сторону границы: "
        "до среднего 70 см остаётся 1 см, до среднего 73 см — 2 см.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "4 Перцептрон")
    docx.add_text(
        doc,
        "Перцептрон здесь — один нейрон с двумя входами и смещением. Класс 1 — футбольный, "
        "класс 0 — волейбольный. Чтобы поправка на доске оставалась целой, входы сдвинуты:",
    )
    docx.add_text(
        doc,
        "x1 = окружность − 70,   x2 = (отскок − 120) / 2.",
    )
    docx.add_text(
        doc,
        "Сумма нейрона z = w1·x1 + w2·x2 + b. Прогноз ŷ равен 1, если z ≥ 0, и равен 0 иначе. "
        "Скорость обучения равна 1. Старт: w1 = 0, w2 = 0, b = 0. После объекта с ошибкой",
    )
    docx.add_text(
        doc,
        "w := w + (y − ŷ)·x,   b := b + (y − ŷ).",
    )
    docx.add_text(
        doc,
        "Объекты этого примера названы П1, П2, П3 и П4. Ход по ним записан в таблице 3.",
    )
    step_w = [1.5, 2.0, 1.5, 1.5, 1.3, 1.6, 1.3, 1.5, 1.5, 1.5]
    docx.add_table_caption(doc, "Таблица 3 — Четыре объекта и веса после каждого шага", sum(step_w))
    docx.add_table(
        doc,
        ["Шаг", "Объект", "x1", "x2", "y", "z", "ŷ", "w1", "w2", "b"],
        [
            [
                str(index),
                step["name"],
                integer(step["x1"]),
                integer(step["x2"]),
                integer(step["y"]),
                integer(step["z"]),
                integer(step["yhat"]),
                integer(step["w1"]),
                integer(step["w2"]),
                integer(step["b"]),
            ]
            for index, step in enumerate(net["steps"], start=1)
        ],
        step_w,
        font=12,
    )
    docx.add_text(
        doc,
        "Шаг 1, объект П1: окружность 68 см, отскок 116 см, поэтому x1 = −2 и x2 = −2, факт y = 0. "
        "При нулевых весах z = 0, прогноз ŷ = 1. Ошибка y − ŷ = −1. "
        "w1 = 0 + (−1)·(−2) = 2, w2 = 0 + (−1)·(−2) = 2, b = 0 + (−1) = −1.",
        space_before=12,
    )
    docx.add_text(
        doc,
        "Шаг 2, объект П2: окружность 71 см, отскок 126 см, x1 = 1, x2 = 3, y = 1. "
        "z = 2·1 + 2·3 − 1 = 7 ≥ 0, прогноз 1 совпадает с фактом. Веса остаются 2, 2 и −1.",
    )
    docx.add_text(
        doc,
        "Шаг 3, объект П3: окружность 69 см, отскок 118 см, x1 = −1, x2 = −1, y = 0. "
        "z = 2·(−1) + 2·(−1) − 1 = −5 < 0, прогноз 0 совпадает с фактом.",
    )
    docx.add_text(
        doc,
        "Шаг 4, объект П4: окружность 72 см, отскок 128 см, x1 = 2, x2 = 4, y = 1. "
        "z = 2·2 + 2·4 − 1 = 11 ≥ 0, прогноз 1 совпадает с фактом.",
    )
    docx.add_text(
        doc,
        f"После одного прохода w1 = {integer(net['w1'])}, w2 = {integer(net['w2'])}, "
        f"b = {integer(net['b'])}. Порог z = 0 записывается как 2x1 + 2x2 − 1 = 0, "
        "то есть x1 + x2 = 0,5.",
    )
    first, second = net["checks"]
    docx.add_text(
        doc,
        f"Новое измерение: окружность {first['size']} см, отскок {first['bounce']} см. "
        f"x1 = {integer(first['x1'])}, x2 = {integer(first['x2'])}. "
        f"z = 2·{integer(first['x1'])} + 2·{integer(first['x2'])} − 1 = {integer(first['z'])}. "
        "z ≥ 0, прогноз — футбольный.",
    )
    docx.add_text(
        doc,
        f"Второе новое измерение: окружность {second['size']} см, отскок {second['bounce']} см. "
        f"x1 = {integer(second['x1'])}, x2 = {integer(second['x2'])}. "
        f"z = 2·({integer(second['x1'])}) + 2·{integer(second['x2'])} − 1 = {integer(second['z'])}. "
        "z < 0, прогноз — волейбольный.",
    )
    docx.add_text(
        doc,
        "Веса сдвинулись на первом объекте, потому что нулевой нейрон выдал футбольный класс "
        "волейбольному мячу. Дальше три объекта уже лежат по верную сторону порога, и проход "
        "по ним веса не меняет.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "5 Выводы")
    docx.add_text(
        doc,
        "Три алгоритма решают одну задачу разными величинами. Ближайшие соседи сравнивают "
        f"евклидовы расстояния и берут большинство из k = {knn['k']}: у измерения Н "
        f"(окружность {q['size']} см, отскок {q['bounce']} см) класс {knn['winner']}, "
        "два голоса против одного. Наивный Байес сравнивает нормальные плотности одного "
        f"признака. При окружности {bayes['x']} см вероятность футбольного класса "
        f"{num(foot['prob'], 4)}, и она больше вероятностей двух других классов. "
        "Перцептрон сравнивает взвешенную сумму с нулём. После обучения w1 = 2, w2 = 2, b = −1; "
        "измерение 70 см и 122 см даёт z = 1 и футбольный класс, измерение 68 см и 120 см "
        "даёт z = −5 и волейбольный класс.",
    )
    docx.add_text(
        doc,
        "На защите каждый ответ восстанавливается из своей таблицы. Для соседей достаточно "
        "шести квадратов и подсчёта голосов. Для Байеса — трёх средних, общей дисперсии 1 "
        "и трёх экспонент. Для перцептрона — одного исправления весов на первом объекте "
        "и подстановки нового измерения в z = 2x1 + 2x2 − 1.",
    )
    doc.save(str(OUT))


def assert_clean() -> None:
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "CART", "Джини", "Gini", "доход", "iris", "трансформер", "ChatGPT",
        "sklearn", "matplotlib", "KNIME", "build_report", "compute_", "seed",
        "Python", "скрипт",
    ]
    hits = [word for word in banned if word.lower() in plain.lower()]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    for needle in (
        "ТРИ АЛГОРИТМА КЛАССИФИКАЦИИ",
        "евклидово",
        "exp(−(x − μ)² / (2σ²))",
        "x1 + x2 = 0,5",
        "класс измерения Н — футбольный",
        "Наибольшая вероятность у класса «футбольный»",
        "прогноз — волейбольный",
        "Рисунок 1",
        "Таблица 3",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")


def read_pages() -> dict[str, str]:
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(OUT), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
    wanted = {
        "1 Задание": "1",
        "2 Метод ближайших соседей": "2",
        "3 Наивный Байес": "3",
        "4 Перцептрон": "4",
        "5 Выводы": "5",
    }
    found = {}
    toc_page = None
    try:
        for paragraph in doc.Paragraphs:
            raw = paragraph.Range.Text.replace("\r", "").replace("\x07", "").strip()
            if raw == "СОДЕРЖАНИЕ":
                toc_page = int(paragraph.Range.Information(3))
            if raw in wanted:
                found[wanted[raw]] = str(int(paragraph.Range.Information(3)))
    finally:
        doc.Close(False)
        word.Quit()
    if set(found) != set(wanted.values()):
        raise RuntimeError(f"pages not found: {found}")
    if toc_page != 2:
        raise RuntimeError(f"содержание на странице {toc_page}")
    pages = [int(found[key]) for key in ("1", "2", "3", "4", "5")]
    if any(left >= right for left, right in zip(pages, pages[1:])):
        raise RuntimeError(f"разделы не с новой страницы: {found}")
    return found


def main() -> None:
    build()
    docx.strip_unused_images(OUT)
    pages = read_pages()
    docx.patch_toc(pages)
    again = read_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
