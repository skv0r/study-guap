# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР5 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import importlib.util
import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SRC_BUILD = BASE.parent / "ЛР2" / "build_report.py"
spec = importlib.util.spec_from_file_location("lr2_docx", SRC_BUILD)
docx = importlib.util.module_from_spec(spec)
spec.loader.exec_module(docx)

OUT = BASE / "ЛР5_Отчет_Буренков_МИИ.docx"
docx.OUT = OUT
FIG = BASE / "figures"
DATA = BASE / "data"
CLASSES = ["футбольный", "баскетбольный", "волейбольный"]
FEATURES = ["elasticity", "size", "mass", "pressure", "bounce"]


def pct(value: float) -> str:
    return f"{float(value) * 100:.2f}".replace(".", ",")


def num(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def meas(feature: str, value) -> str:
    number = float(value)
    if feature in ("elasticity", "pressure"):
        return f"{number:.2f}".replace(".", ",")
    if feature == "size":
        return f"{number:.1f}".replace(".", ",")
    return str(int(round(number)))


def replace_title(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 5",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "МЕТОД ОПОРНЫХ ВЕКТОРОВ",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {found}")


def error_sentence(item: dict) -> str:
    head = f"Наблюдение {item['id']}. Факт — «{item['ball']}», прогноз — «{item['pred']}»."
    if item.get("diffs"):
        phrases = "; ".join(diff["phrase"] for diff in item["diffs"])
        return f"{head} Не как у типового набора этого мяча: {phrases}."
    return f"{head} Измерения совпадают с типовым набором этого мяча."


def build() -> None:
    m = json.loads((DATA / "metrics.json").read_text(encoding="utf-8"))
    df = pd.read_csv(DATA / "balls.csv")
    if len(m.get("comparison", [])) != 4:
        raise RuntimeError("нужно сравнение четырёх методов")
    methods = {item["method"] for item in m["comparison"]}
    needed = {"k ближайших соседей", "дерево решений", "наивный Байес", "опорные векторы"}
    if methods != needed:
        raise RuntimeError(f"в сравнении не все методы: {methods}")
    n = m["n"]
    n_train = m["n_train"]
    n_test = m["n_test"]
    n_train_per = m.get("n_train_per", n_train // 3)
    n_test_per = m.get("n_test_per", n_test // 3)
    if len(df) != n or n_train + n_test != n:
        raise RuntimeError(f"ожидались {n}, {n_train} и {n_test}")
    if m["correct"] >= n_test:
        raise RuntimeError("нужны ошибки на тесте")
    usual = m["typical"]
    for name in CLASSES:
        matched = 0
        for row in df[df["ball"] == name].itertuples(index=False):
            got = [float(getattr(row, feature)) for feature in FEATURES]
            matched += int(all(abs(a - b) < 1e-9 for a, b in zip(got, usual[name])))
        if matched != 0:
            raise RuntimeError(f"у {name} есть точные копии типового")

    src = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(src, OUT)
    doc = Document(str(OUT))
    replace_title(doc)
    docx.drop_body_after_title(doc)
    docx.add_heading(doc, "СОДЕРЖАНИЕ")
    docx.add_toc(doc, "1 Задание", "@@1@@")
    docx.add_toc(doc, "2 Входные данные", "@@2@@")
    docx.add_toc(doc, "3 Ход выполнения работы", "@@3@@")
    docx.add_toc(doc, "4 Выводы", "@@4@@")

    doc.add_page_break()
    docx.add_heading(doc, "1 Задание")
    docx.add_text(
        doc,
        "Цель работы — отнести измерение мяча к одному из трёх классов методом опорных векторов "
        "и сравнить точность с методами лабораторных 2, 3 и 4 на тех же строках одного общего теста.",
    )
    docx.add_text(
        doc,
        "По практикуму читается своя таблица. Модель обучается на обучающей части; для теста "
        "считаются класс и вероятности классов. Ядро — радиальная базисная функция, штраф за "
        "неверную точку равен 1, параметр сигма равен 0,5. Перед обучением признаки приводятся "
        "к одному масштабу по обучающей выборке. Затем считаются матрица ошибок, accuracy, "
        "precision, recall и F-measure. В выводе — краткое сравнение четырёх методов на одном "
        "тесте и ключевые параметры, из‑за которых отличаются показатели.",
    )
    docx.add_text(
        doc,
        "Расчёт выполнен в Python. Класс — вид мяча: футбольный, баскетбольный или волейбольный. "
        "Признаки — коэффициент упругости, окружность в сантиметрах, масса в граммах, давление "
        "в барах и высота контрольного отскока в сантиметрах. Номер наблюдения в модель не входит.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "2 Входные данные")
    foot = df[df["ball"] == "футбольный"]
    basket = df[df["ball"] == "баскетбольный"]
    volley = df[df["ball"] == "волейбольный"]
    docx.add_text(
        doc,
        f"Таблица, классы и разбиение те же, что в лабораторных 2, 3 и 4. {n} измерений мяча, "
        f"по {n // 3} футбольных, баскетбольных и волейбольных. Обучение — {n_train} строк "
        f"(по {n_train_per} в классе), тест — {n_test} (по {n_test_per}). Это один общий тест. "
        "Пять признаков: коэффициент упругости, окружность в сантиметрах, масса в граммах, "
        "давление в барах и высота контрольного отскока в сантиметрах. Типовой набор — центр "
        "класса; точных копий типового набора в таблице нет, строки — измерения с разбросом.",
    )
    mean_w = [3.4, 2.6, 2.8, 2.4, 2.6, 2.6]
    docx.add_table_caption(doc, "Таблица 1 — Типовые измерения трёх мячей", sum(mean_w))
    docx.add_table(
        doc,
        ["Класс", "Упругость", "Окружность, см", "Масса, г", "Давление, бар", "Отскок, см"],
        [
            [
                name,
                meas("elasticity", usual[name][0]),
                meas("size", usual[name][1]),
                meas("mass", usual[name][2]),
                meas("pressure", usual[name][3]),
                meas("bounce", usual[name][4]),
            ]
            for name in CLASSES
        ],
        mean_w,
        font=12,
    )
    docx.add_text(
        doc,
        "Диапазоны соседних классов пересекаются: "
        f"упругость футбольного {meas('elasticity', foot['elasticity'].min())}–"
        f"{meas('elasticity', foot['elasticity'].max())}, баскетбольного "
        f"{meas('elasticity', basket['elasticity'].min())}–{meas('elasticity', basket['elasticity'].max())}; "
        f"окружность волейбольного {meas('size', volley['size'].min())}–"
        f"{meas('size', volley['size'].max())} см, футбольного "
        f"{meas('size', foot['size'].min())}–{meas('size', foot['size'].max())} см. "
        "Поэтому стопроцентной точности на тесте ни у одного из четырёх методов нет.",
        space_before=12,
    )

    doc.add_page_break()
    docx.add_heading(doc, "3 Ход выполнения работы")
    docx.add_text(
        doc,
        "Последовательность расчёта показана на рисунке 1. Штраф равен 1, ядро — радиальная "
        "базисная функция, сигма равна 0,5. Признаки стандартизированы по обучению: у каждого "
        "признака нулевое среднее и единичный разброс на обучающих строках; те же среднее и "
        "разброс применены к тесту.",
    )
    docx.add_picture(doc, FIG / "scheme.png", 15.5)
    docx.add_figure_caption(doc, "Рисунок 1 — Схема расчёта")
    counts = m["n_support"]
    docx.add_text(
        doc,
        f"Опорных векторов {m['n_support_total']}: футбольный — {counts['футбольный']}, "
        f"баскетбольный — {counts['баскетбольный']}, волейбольный — {counts['волейбольный']}. "
        f"При {n_train} обучающих строках почти все они оказались опорными: граница классов "
        "в зоне пересечения диапазонов держится на многих пограничных измерениях, а не только "
        "на типовом центре.",
    )
    docx.add_table_caption(doc, "Таблица 2 — Число опорных векторов", 12)
    docx.add_table(
        doc,
        ["Класс", "Опорных векторов"],
        [[name, str(counts[name])] for name in CLASSES],
        [6.0, 4.5],
    )
    example = m["example"]
    probs = example["probs"]
    fx = example["features"]
    docx.add_text(
        doc,
        f"Для строки {example['id']}: упругость {meas('elasticity', fx['elasticity'])}, "
        f"окружность {meas('size', fx['size'])} см, масса {meas('mass', fx['mass'])} г, "
        f"давление {meas('pressure', fx['pressure'])} бар, отскок {meas('bounce', fx['bounce'])} см. "
        f"Факт — «{example['ball']}», прогноз — «{example['pred']}». "
        f"Вероятности: футбольный {pct(probs['футбольный'])} %, "
        f"баскетбольный {pct(probs['баскетбольный'])} %, волейбольный {pct(probs['волейбольный'])} %.",
        space_before=12,
    )
    wrong = [item for item in m["predictions"] if not item["correct"]]
    if wrong:
        sample = "; ".join(
            f"№ {item['id']}: факт «{item['ball']}», прогноз «{item['pred']}»"
            for item in wrong[:3]
        )
        more = f", всего {len(wrong)}" if len(wrong) > 3 else ""
        docx.add_text(doc, f"Ошибки на тесте{more}: {sample}.", space_before=12)
    cm = m["confusion"]
    docx.add_table_caption(doc, f"Таблица 3 — Матрица ошибок на {n_test} тестовых строках", 12)
    docx.add_table(
        doc,
        ["Факт \\ прогноз", *CLASSES],
        [[name, *[str(cm[i][j]) for j in range(3)]] for i, name in enumerate(CLASSES)],
        [4.0, 2.8, 3.4, 3.2],
        font=11,
    )
    per = {row["class"]: row for row in m["per_class"]}
    met_w = [3.6, 3.2, 2.8, 3.4, 3.5]
    docx.add_table_caption(doc, "Таблица 4 — Precision, recall и F-measure по классам", sum(met_w))
    docx.add_table(
        doc,
        ["Класс", "Precision, %", "Recall, %", "F-measure, %", "Наблюдений"],
        [
            [
                name,
                pct(per[name]["precision"]),
                pct(per[name]["recall"]),
                pct(per[name]["f1"]),
                str(per[name]["support"]),
            ]
            for name in CLASSES
        ],
        met_w,
        font=12,
    )
    docx.add_text(
        doc,
        f"Верных ответов {m['correct']} из {n_test}. Accuracy равна {pct(m['accuracy'])} %, "
        f"средние precision, recall и F-measure — {pct(m['macro_precision'])} %, "
        f"{pct(m['macro_recall'])} % и {pct(m['macro_f1'])} %.",
        space_before=12,
    )
    docx.add_picture(doc, FIG / "scatter.png", 14)
    docx.add_figure_caption(doc, "Рисунок 2 — Классы мяча в осях окружности и коэффициента упругости")

    knn_k = m.get("knn_k", "?")
    tree_depth = m.get("tree_depth", "?")
    by_method = {item["method"]: item for item in m["comparison"]}
    docx.add_text(
        doc,
        f"Сравнение сделано на одном и том же тесте из {n_test} строк. Лабораторная 2 — "
        f"k ближайших соседей при k = {knn_k}. Лабораторная 3 — дерево решений глубиной "
        f"{tree_depth}. Лабораторная 4 — наивный Байес. Лабораторная 5 — опорные векторы. "
        "Числа собраны в таблице 5: верные ответы, accuracy, precision, recall и F-measure.",
        space_before=12,
    )
    docx.add_table_caption(doc, "Таблица 5 — Сравнение четырёх методов на одном тесте", 16.8)
    docx.add_table(
        doc,
        ["Работа", "Метод", "Верных", "Accuracy, %", "Precision, %", "Recall, %", "F-measure, %"],
        [
            [
                item["lab"],
                item["method"],
                f"{item['correct']} из {item.get('n_test', n_test)}",
                pct(item["accuracy"]),
                pct(item["precision"]),
                pct(item["recall"]),
                pct(item["f1"]),
            ]
            for item in m["comparison"]
        ],
        [1.8, 3.6, 2.2, 2.4, 2.4, 2.2, 2.4],
        font=10,
    )

    doc.add_page_break()
    docx.add_heading(doc, "4 Выводы")
    ranked = sorted(m["comparison"], key=lambda item: (-item["correct"], -item["f1"], item["lab"]))
    best = ranked[0]
    knn = by_method["k ближайших соседей"]
    tree = by_method["дерево решений"]
    bayes = by_method["наивный Байес"]
    svm = by_method["опорные векторы"]
    rest = "; ".join(
        f"{item['method']} — {item['correct']}/{n_test} ({pct(item['accuracy'])} %)"
        for item in ranked[1:]
    )
    docx.add_text(
        doc,
        f"На одном тесте из {n_test} строк лучший результат у метода «{best['method']}»: "
        f"{best['correct']}/{n_test}, accuracy {pct(best['accuracy'])} %, F-measure {pct(best['f1'])} %; "
        f"далее {rest}. Ключевые параметры: у соседей k = {knn_k} и стандартизация признаков; "
        f"у дерева глубина {tree_depth}, критерий Джини; у Байеса нормали по классу и допущение "
        f"независимости признаков; у опорных векторов ядро RBF (сигма {num(m['sigma'], 1)}, "
        f"штраф {num(m['penalty'], 0)}) и та же стандартизация. Стопроцентной точности нет, "
        "потому что диапазоны измерений соседних мячей пересекаются.",
    )
    doc.save(str(OUT))


def assert_clean() -> None:
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "тикет", "велосипед", "iris", "sklearn", "matplotlib", "KNIME",
        "build_report", "compute_", "seed", "codebook", "кодирован",
        "игровая зона", "динамика владения", "Расшифровка кодов",
    ]
    hits = [word for word in banned if word.lower() in plain.lower()]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    if plain.count("Python") != 1:
        raise RuntimeError(f"Python упомянут {plain.count('Python')} раз")
    for needle in (
        "МЕТОД ОПОРНЫХ ВЕКТОРОВ",
        "Таблица 5 — Сравнение четырёх методов",
        "дерево решений",
        "k ближайших соседей",
        "наивный Байес",
        "опорные векторы",
        "коэффициент упругости",
        "лучший результат",
        "Ключевые параметры",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")


def main() -> None:
    build()
    docx.strip_unused_images(OUT)
    pages = docx.read_heading_pages()
    docx.patch_toc(pages)
    again = docx.read_heading_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
