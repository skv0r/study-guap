# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР4 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import importlib.util
import json
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
SRC_BUILD = BASE.parent / "ЛР2" / "build_report.py"
spec = importlib.util.spec_from_file_location("lr2_docx", SRC_BUILD)
docx = importlib.util.module_from_spec(spec)
spec.loader.exec_module(docx)

OUT = BASE / "ЛР4_Отчет_Буренков_МИИ.docx"
docx.OUT = OUT
FIG = BASE / "figures"
DATA = BASE / "data"
CLASSES = ["футбольный", "баскетбольный", "волейбольный"]
FEATURES = ["elasticity", "size", "mass", "pressure", "bounce"]
FEATURE_RU = {
    "elasticity": "Коэффициент упругости",
    "size": "Окружность, см",
    "mass": "Масса, г",
    "pressure": "Давление, бар",
    "bounce": "Высота отскока, см",
}


def num(value: float, digits: int = 4) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def pct(value: float) -> str:
    return f"{float(value) * 100:.2f}".replace(".", ",")


def meas(feature: str, value) -> str:
    number = float(value)
    if feature in ("elasticity", "pressure"):
        return f"{number:.2f}".replace(".", ",")
    if feature == "size":
        return f"{number:.1f}".replace(".", ",")
    return str(int(round(number)))


def replace_title(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 4",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "МЕТОД НАИВНОГО БАЙЕСА",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {found}")


def error_sentence(item: dict) -> str:
    head = f"Наблюдение {item['id']}. Факт — «{item['ball']}», прогноз — «{item['pred']}»."
    if item.get("diffs"):
        phrases = "; ".join(diff["phrase"] for diff in item["diffs"])
        return f"{head} Не как у типового набора этого мяча: {phrases}."
    return f"{head} Измерения совпадают с типовым набором этого мяча."


def build() -> None:
    m = json.loads((DATA / "metrics.json").read_text(encoding="utf-8"))
    df = pd.read_csv(DATA / "balls.csv")
    tree_split = json.loads((BASE.parent / "ЛР3" / "data" / "split.json").read_text(encoding="utf-8"))
    if m["train_ids"] != tree_split["train_ids"] or m["test_ids"] != tree_split["test_ids"]:
        raise RuntimeError("разбиение не совпало")
    n = m["n"]
    n_train = m["n_train"]
    n_test = m["n_test"]
    n_train_per = m.get("n_train_per", n_train // 3)
    n_test_per = m.get("n_test_per", n_test // 3)
    if len(df) != n or n_train + n_test != n:
        raise RuntimeError(f"ожидались {n}, {n_train} и {n_test}")
    if m["correct"] >= n_test:
        raise RuntimeError("нужны ошибки на тесте")
    usual = m["typical"]
    for name in CLASSES:
        matched = 0
        for row in df[df["ball"] == name].itertuples(index=False):
            got = [float(getattr(row, feature)) for feature in FEATURES]
            matched += int(all(abs(a - b) < 1e-9 for a, b in zip(got, usual[name])))
        if matched != 0:
            raise RuntimeError(f"у {name} есть точные копии типового")

    src = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(src, OUT)
    doc = Document(str(OUT))
    replace_title(doc)
    docx.drop_body_after_title(doc)
    docx.add_heading(doc, "СОДЕРЖАНИЕ")
    docx.add_toc(doc, "1 Задание", "@@1@@")
    docx.add_toc(doc, "2 Входные данные", "@@2@@")
    docx.add_toc(doc, "3 Ход выполнения работы", "@@3@@")
    docx.add_toc(doc, "4 Выводы", "@@4@@")

    doc.add_page_break()
    docx.add_heading(doc, "1 Задание")
    docx.add_text(
        doc,
        "Цель работы — по пяти физическим измерениям мяча определить его вид наивным байесовским "
        "методом и проверить ответ на наблюдениях, которые не входили в обучение.",
    )
    docx.add_text(
        doc,
        "По практикуму читается своя таблица, 70 % строк — обучение, 30 % — тест со "
        "стратификацией по классу. По обучению для каждого класса и признака считаются среднее "
        "и стандартное отклонение; измерение описывается нормальным распределением. Априор "
        "класса — доля обучающих строк. Для тестовой строки мера класса — произведение пяти "
        "плотностей и априора; ответ — класс с наибольшей мерой. Затем считаются матрица ошибок, "
        "accuracy, precision, recall и F-measure.",
    )
    docx.add_text(
        doc,
        "Расчёт выполнен в Python. Класс — вид мяча; признаки — упругость, окружность, масса, "
        "давление и высота отскока. Таблица и разбиение те же, что в лабораторных 2 и 3. "
        "Метод наивный: признаки внутри класса считаются независимыми.",
    )

    doc.add_page_break()
    docx.add_heading(doc, "2 Входные данные")
    docx.add_text(
        doc,
        f"Исходный набор — {n} измерений мяча, по {n // 3} с каждым видом. В строке семь полей: "
        "id — номер наблюдения; elasticity — коэффициент упругости; size — окружность, см; "
        "mass — масса, г; pressure — давление, бар; bounce — высота отскока, см; ball — класс. "
        "В модель входят пять измерений. Пропусков нет. Типовой набор — центр класса "
        "(номинальные значения мяча). Строки таблицы — измерения с разбросом вокруг этого "
        "центра; точных копий типового набора нет. Диапазоны соседних классов пересекаются, "
        "поэтому один признак класс не называет.",
    )
    docx.add_text(doc, "Типовые центры трёх мячей приведены в таблице 1.")
    mean_w = [3.4, 2.6, 2.8, 2.4, 2.6, 2.6]
    docx.add_table_caption(doc, "Таблица 1 — Типовые измерения трёх мячей", sum(mean_w))
    docx.add_table(
        doc,
        ["Класс", "Упругость", "Окружность, см", "Масса, г", "Давление, бар", "Отскок, см"],
        [
            [
                name,
                meas("elasticity", usual[name][0]),
                meas("size", usual[name][1]),
                meas("mass", usual[name][2]),
                meas("pressure", usual[name][3]),
                meas("bounce", usual[name][4]),
            ]
            for name in CLASSES
        ],
        mean_w,
        font=12,
    )
    foot = df[df["ball"] == "футбольный"]
    basket = df[df["ball"] == "баскетбольный"]
    volley = df[df["ball"] == "волейбольный"]
    docx.add_text(
        doc,
        "Диапазоны соседних классов пересекаются: "
        f"упругость футбольного {meas('elasticity', foot['elasticity'].min())}–"
        f"{meas('elasticity', foot['elasticity'].max())}, баскетбольного "
        f"{meas('elasticity', basket['elasticity'].min())}–{meas('elasticity', basket['elasticity'].max())}, "
        f"волейбольного {meas('elasticity', volley['elasticity'].min())}–"
        f"{meas('elasticity', volley['elasticity'].max())}; "
        f"окружность волейбольного {meas('size', volley['size'].min())}–"
        f"{meas('size', volley['size'].max())} см, футбольного "
        f"{meas('size', foot['size'].min())}–{meas('size', foot['size'].max())} см, "
        f"баскетбольного {meas('size', basket['size'].min())}–"
        f"{meas('size', basket['size'].max())} см; "
        f"масса футбольного {meas('mass', foot['mass'].min())}–"
        f"{meas('mass', foot['mass'].max())} г, баскетбольного "
        f"{meas('mass', basket['mass'].min())}–{meas('mass', basket['mass'].max())} г. "
        f"В обучении {n_train} строк (по {n_train_per} в классе), в тесте {n_test} "
        f"(по {n_test_per}). Это тот же тест, что в лабораторных работах 2 и 3.",
        space_before=12,
    )

    doc.add_page_break()
    docx.add_heading(doc, "3 Ход выполнения работы")
    test_ids = ", ".join(str(i) for i in m["test_ids"])
    docx.add_text(
        doc,
        "Последовательность расчёта показана на рисунке 1. Классам назначены цвета для диаграммы: "
        "синий — футбольный, оранжевый — баскетбольный, зелёный — волейбольный. Цвет в модель "
        "не входит.",
    )
    docx.add_picture(doc, FIG / "scheme.png", 15.5)
    docx.add_figure_caption(doc, "Рисунок 1 — Схема расчёта")
    prior = m["priors"][CLASSES[0]]
    docx.add_text(
        doc,
        f"Разбиение то же, что в лабораторных 2 и 3: в тесте номера {test_ids}. "
        f"Априорная вероятность каждого класса равна {n_train_per}/{n_train}, то есть "
        f"{pct(prior)} % — одной трети, потому что обучение стратифицировано. Для числового "
        "признака внутри класса берётся нормальное распределение со средним и стандартным "
        "отклонением обучающих строк этого класса. В плотность подставляется отклонение "
        "«в расчёте»: если бы выборочное отклонение равнялось нулю, подставлялось бы "
        "0,0001, иначе — само отклонение. "
        + (
            "В этой выборке отклонение каждого признака внутри класса больше нуля, подмена не понадобилась."
            if not any(stat["floored"] for stat in m["stats"])
            else "Для признаков с нулевым отклонением в плотность подставлено 0,0001."
        ),
    )
    rows = []
    for stat in m["stats"]:
        rows.append(
            [
                stat["class"],
                FEATURE_RU[stat["feature"]],
                num(stat["mean"], 4),
                num(stat["sd"], 4),
                num(stat["sd_used"], 4),
            ]
        )
    docx.add_table_caption(doc, "Таблица 2 — Среднее и стандартное отклонение по обучающим строкам", 16.5)
    docx.add_table(
        doc,
        ["Класс", "Признак", "Среднее", "Отклонение", "В расчёте"],
        rows,
        [3.4, 4.5, 2.6, 3.0, 3.0],
        font=11,
    )
    example = m["example"]
    probs = example["probs"]
    fx = example["features"]
    docx.add_text(
        doc,
        f"Пример расчёта вероятностей — тестовая строка {example['id']}. "
        f"Упругость {meas('elasticity', fx['elasticity'])}, "
        f"окружность {meas('size', fx['size'])} см, масса {meas('mass', fx['mass'])} г, "
        f"давление {meas('pressure', fx['pressure'])} бар, отскок {meas('bounce', fx['bounce'])} см. "
        f"Факт — «{example['ball']}», прогноз — «{example['pred']}». "
        "Для каждого класса перемножаются пять нормальных плотностей признаков и априор "
        f"{pct(prior)} %, затем меры нормируются на сумму по трём классам. Полученные "
        f"вероятности: футбольный {pct(probs['футбольный'])} %, "
        f"баскетбольный {pct(probs['баскетбольный'])} %, "
        f"волейбольный {pct(probs['волейбольный'])} %. Выбирается класс с наибольшей вероятностью.",
        space_before=12,
    )
    wrong = [item for item in m["predictions"] if not item["correct"]]
    if wrong:
        sample = "; ".join(
            f"№ {item['id']}: факт «{item['ball']}», прогноз «{item['pred']}»"
            for item in wrong[:3]
        )
        more = f" (всего {len(wrong)})" if len(wrong) > 3 else ""
        docx.add_text(
            doc,
            f"На тесте ошибки{more}: {sample}.",
            space_before=12,
        )
    cm = m["confusion"]
    docx.add_text(
        doc,
        f"Ответы по всем {n_test} тестовым строкам собраны в матрице ошибок таблицы 3: "
        "строка — фактический класс, столбец — прогноз.",
        space_before=12,
    )
    docx.add_table_caption(doc, f"Таблица 3 — Матрица ошибок на {n_test} тестовых строках", 12)
    docx.add_table(
        doc,
        ["Факт \\ прогноз", *CLASSES],
        [[name, *[str(cm[i][j]) for j in range(3)]] for i, name in enumerate(CLASSES)],
        [4.0, 2.8, 3.4, 3.2],
        font=11,
    )
    per = {row["class"]: row for row in m["per_class"]}
    met_w = [3.3, 1.4, 1.4, 1.4, 1.4, 2.6, 2.4, 2.6]
    docx.add_table_caption(doc, "Таблица 4 — Precision, recall и F-measure по классам", sum(met_w))
    add_rows = []
    for name in CLASSES:
        row = per[name]
        add_rows.append(
            [
                name,
                str(row.get("tp", "")),
                str(row.get("fp", "")),
                str(row.get("tn", "")),
                str(row.get("fn", "")),
                pct(row["precision"]),
                pct(row["recall"]),
                pct(row["f1"]),
            ]
        )
    docx.add_table(
        doc,
        ["Класс", "TP", "FP", "TN", "FN", "Precision, %", "Recall, %", "F-measure, %"],
        add_rows,
        met_w,
        font=11,
    )
    class_line = "На тесте " + ", ".join(
        f"у «{name}» верны {cm[i][i]} из {per[name]['support']}" for i, name in enumerate(CLASSES)
    ) + "."
    docx.add_text(
        doc,
        f"Верных ответов {m['correct']} из {n_test}, ошибок {n_test - m['correct']}. "
        f"Accuracy = {m['correct']}/{n_test} = {pct(m['accuracy'])} %. "
        f"Средние по трём классам: precision {pct(m['macro_precision'])} %, "
        f"recall {pct(m['macro_recall'])} %, F-measure {pct(m['macro_f1'])} %. "
        f"Средний recall совпадает с accuracy, потому что в тесте по {n_test_per} наблюдений "
        f"каждого класса. {class_line}",
        space_before=12,
    )
    docx.add_text(
        doc,
        f"Диаграмма рассеяния на рисунке 2 построена по всем {n} наблюдениям. По горизонтали — "
        "окружность мяча в сантиметрах, по вертикали — коэффициент упругости. Цвет точки — "
        "фактический класс. Диапазоны соседних классов по осям пересекаются, поэтому одна ось "
        "класс не называет; плотности считаются по всем пяти измерениям. Ошибки теста обведены "
        "чёрным.",
    )
    docx.add_picture(doc, FIG / "scatter.png", 14)
    docx.add_figure_caption(doc, "Рисунок 2 — Классы мяча в осях окружности и коэффициента упругости")
    doc.add_page_break()
    docx.add_heading(doc, "4 Выводы")
    docx.add_text(
        doc,
        f"На общем разбиении {n_train}/{n_test} наивный Байес верно назвал {m['correct']} "
        f"из {n_test} (accuracy {pct(m['accuracy'])} %, F-measure {pct(m['macro_f1'])} %). "
        "Признаки описаны нормальным распределением по классу при априоре одна треть; "
        "нулевое отклонение заменялось бы на 0,0001. Ошибки связаны с пересечением диапазонов "
        "и с допущением независимости признаков: у мяча упругость, давление и отскок связаны, "
        "поэтому произведение одномерных плотностей иногда ставит чужой класс выше своего.",
    )
    doc.save(str(OUT))


def assert_clean() -> None:
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "тикет", "велосипед", "iris", "sklearn", "matplotlib", "KNIME",
        "build_report", "compute_", "seed", "numpy", "codebook", "кодирован",
        "игровая зона", "динамика владения", "Расшифровка кодов",
    ]
    hits = [word for word in banned if word.lower() in plain.lower()]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    if plain.count("Python") != 1:
        raise RuntimeError(f"Python упомянут {plain.count('Python')} раз")
    for needle in (
        "МЕТОД НАИВНОГО БАЙЕСА",
        "0,0001",
        "коэффициент упругости",
        "Таблица 1 — Типовые измерения",
        "независимост",
        "тот же тест",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")


def main() -> None:
    build()
    docx.strip_unused_images(OUT)
    pages = docx.read_heading_pages()
    docx.patch_toc(pages)
    again = docx.read_heading_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
