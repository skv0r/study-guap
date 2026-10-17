# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР3 МИИ. В текст отчёта этот файл не попадает."""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

BASE = Path(__file__).resolve().parent
SRC = BASE.parent / "ЛР1" / "ЛР1_Отчет_Буренков_МИИ.docx"
OUT = BASE / "ЛР3_Отчет_Буренков_МИИ.docx"
FIG = BASE / "figures"
DATA = BASE / "data"

CLASSES = ["футбольный", "баскетбольный", "волейбольный"]
FEATURES = ["elasticity", "size", "mass", "pressure", "bounce"]
FEATURE_RU = {
    "elasticity": "Коэффициент упругости",
    "size": "Окружность, см",
    "mass": "Масса, г",
    "pressure": "Давление, бар",
    "bounce": "Высота отскока, см",
}
SHORT = {
    "elasticity": "Упругость",
    "size": "Окружность",
    "mass": "Масса",
    "pressure": "Давление",
    "bounce": "Отскок",
}


def num(value: float, digits: int = 2) -> str:
    return f"{float(value):.{digits}f}".replace(".", ",")


def pct(value: float) -> str:
    return num(value * 100, 2)


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(p, *, first_indent=True, align="justify", space_after=0, space_before=0, line_spacing=1.5):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_text(doc, text, *, indent=True, align="justify", bold=False, space_after=0, space_before=0, size=14):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=indent, align=align, space_after=space_after, space_before=space_before)
    set_run_font(p.add_run(text), size=size, bold=bold)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_after=12, space_before=0)
    p.paragraph_format.keep_with_next = True
    if text == "СОДЕРЖАНИЕ":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        set_run_font(p.add_run(text), bold=True)
        return p
    number, title = text.split(" ", 1)
    set_run_font(p.add_run(number), bold=True)
    space = p.add_run(" ")
    set_run_font(space, bold=True)
    space._r.find(qn("w:t")).set(qn("xml:space"), "preserve")
    set_run_font(p.add_run(title), bold=True)
    return p


def add_toc(doc, title, key):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), 2, 1)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(key))
    return p


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def shade_header(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell(cell, text, *, bold=False, size=11, align="center"):
    cell.text = ""
    p = cell.paragraphs[0]
    format_paragraph(p, first_indent=False, align=align, line_spacing=1.0, space_after=0, space_before=0)
    set_run_font(p.add_run(str(text)), size=size, bold=bold)
    set_cell_border(cell)
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)


def add_table(doc, headers, rows, widths, *, font=12, body_align="center"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    total = int(round(sum(widths) * 567))
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    else:
        for child in list(grid):
            grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(round(width * 567))))
        grid.append(gc)
    for j, header_text in enumerate(headers):
        set_cell(table.rows[0].cells[j], header_text, bold=True, size=font)
        shade_header(table.rows[0].cells[j])
        table.rows[0].cells[j].width = Cm(widths[j])
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    aligns = body_align if isinstance(body_align, (list, tuple)) else [body_align] * len(headers)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell(table.rows[i + 1].cells[j], value, size=font, align=aligns[j])
            table.rows[i + 1].cells[j].width = Cm(widths[j])
    return table


def add_table_caption(doc, text, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_after=6, space_before=12, line_spacing=1.0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))
    return p


def add_picture(doc, path, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=6, space_after=0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    return p


def add_figure_caption(doc, text):
    return add_text(doc, text, indent=False, align="center", space_before=6, space_after=12)


def replace_title_text(doc: Document) -> None:
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3",
        "ЛИНЕЙНАЯ РЕГРЕССИЯ": "МЕТОД ДЕРЕВА РЕШЕНИЙ",
    }
    found = {key: 0 for key in replacements}
    for child in list(doc.element.body)[:12]:
        for node in child.iter(qn("w:t")):
            if node.text in replacements:
                found[node.text] += 1
                node.text = replacements[node.text]
                parent = node.getparent()
                rPr = parent.find(qn("w:rPr")) if parent is not None else None
                if rPr is not None:
                    bold = rPr.find(qn("w:b"))
                    if bold is not None and bold.get(qn("w:val"), "true") not in ("0", "false"):
                        raise RuntimeError("строка отчёта или тема стала жирной")
    missing = [key for key, count in found.items() if count != 1]
    if missing:
        raise RuntimeError(f"title replace failed: {missing} {found}")
    blob = "".join(t.text or "" for t in doc.element.body.iter(qn("w:t")))
    for needle in ("КАФЕДРА № ", "42", "В.В. Фомин", "4321", "Г.В. Буренков", "2026", "Методы искусственного интеллекта"):
        if needle not in blob:
            raise RuntimeError(f"на титуле нет {needle}")


def drop_body_after_title(doc: Document) -> None:
    body = doc.element.body
    children = list(body)
    if children[-1].tag != qn("w:sectPr"):
        raise RuntimeError("body sectPr missing")
    for child in children[12:-1]:
        body.remove(child)


def strip_unused_images(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    rel_name = "word/_rels/document.xml.rels"
    root = etree.fromstring(files[rel_name])
    used = set()
    for name, data in files.items():
        if not (name.startswith("word/") and name.endswith(".xml") and "/_rels/" not in name):
            continue
        for key in (b'r:id="', b'r:embed="'):
            start = 0
            while True:
                pos = data.find(key, start)
                if pos < 0:
                    break
                end = data.find(b'"', pos + len(key))
                used.add(data[pos + len(key) : end].decode())
                start = end + 1
    for rel in list(root):
        if rel.get("Id") not in used and "image" in (rel.get("Type") or ""):
            target = "word/" + rel.get("Target").replace("\\", "/")
            files.pop(target, None)
            root.remove(rel)
    files[rel_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


def meas(feature: str, value) -> str:
    number = float(value)
    if feature in ("elasticity", "pressure"):
        return f"{number:.2f}".replace(".", ",")
    if feature == "size":
        return f"{number:.1f}".replace(".", ",")
    return str(int(round(number)))


def ru_ids(ids: list[int]) -> str:
    text = [str(item) for item in ids]
    if len(text) == 1:
        return "строка " + text[0]
    if len(text) == 2:
        return f"строки {text[0]} и {text[1]}"
    return "строки " + ", ".join(text[:-1]) + " и " + text[-1]


def overlap_sentence(df: pd.DataFrame) -> str:
    foot = df[df["ball"] == "футбольный"]
    basket = df[df["ball"] == "баскетбольный"]
    volley = df[df["ball"] == "волейбольный"]
    return (
        "Типовой набор — центр класса; каждая строка таблицы — измерение с разбросом, "
        "точных копий типового набора нет. Диапазоны соседних "
        f"классов пересекаются: упругость футбольного {meas('elasticity', foot['elasticity'].min())}–"
        f"{meas('elasticity', foot['elasticity'].max())}, баскетбольного "
        f"{meas('elasticity', basket['elasticity'].min())}–{meas('elasticity', basket['elasticity'].max())}; "
        f"окружность волейбольного {meas('size', volley['size'].min())}–{meas('size', volley['size'].max())} см, "
        f"футбольного {meas('size', foot['size'].min())}–{meas('size', foot['size'].max())} см. "
        "Поэтому один признак сам по себе класс не называет."
    )


def error_sentence(item: dict) -> str:
    head = f"Наблюдение {item['id']}. Факт — «{item['ball']}», прогноз — «{item['pred']}»."
    if item.get("variant_name"):
        return f"{head} Это {item['variant_name']}: {item['diffs'][0]['phrase']}."
    if item["diffs"]:
        phrases = "; ".join(diff["phrase"] for diff in item["diffs"])
        return f"{head} Не как у типового набора этого мяча: {phrases}."
    return f"{head} Измерения совпадают с типовым набором этого мяча."


def check_design(df: pd.DataFrame, m: dict) -> None:
    tickets = DATA / "tickets.csv"
    requests = BASE.parent / "ЛР2" / "data" / "requests.csv"
    if not tickets.exists() or tickets.stat().st_size < 100:
        raise RuntimeError("tickets.csv должен остаться")
    if not requests.exists() or requests.stat().st_size < 100:
        raise RuntimeError("requests.csv должен остаться")
    n, n_train, n_test = m["n"], m["n_train"], m["n_test"]
    n_test_per = m.get("n_test_per", n_test // 3)
    if len(df) != n or m["n_train"] + m["n_test"] != n:
        raise RuntimeError(f"ожидались {n}, {n_train} и {n_test}")
    if m["max_depth_limit"] != 3 or m["criterion"] != "gini" or m["min_samples_leaf"] != 4:
        raise RuntimeError("параметры дерева изменились")
    if m["correct"] != sum(item["correct"] for item in m["predictions"]):
        raise RuntimeError("число верных не сходится")
    if abs(m["accuracy"] - m["correct"] / n_test) > 1e-12:
        raise RuntimeError("accuracy не сходится")
    if m["correct"] >= n_test:
        raise RuntimeError("на тесте нужны ошибки, не 100 %")
    if m["actual_depth"] < 3 or m["n_leaves"] < 5:
        raise RuntimeError("дерево слишком мелкое")
    used = {item["feature"] for item in m["splits"]}
    if len(used) < 3:
        raise RuntimeError(f"в дереве мало признаков: {used}")
    index = {name: i for i, name in enumerate(CLASSES)}
    cm = [[0, 0, 0] for _ in CLASSES]
    for item in m["predictions"]:
        cm[index[item["ball"]]][index[item["pred"]]] += 1
    if cm != m["confusion"]:
        raise RuntimeError("матрица не сходится с ответами")
    by_id = {int(row.id): row for row in df.itertuples(index=False)}
    for name in CLASSES:
        if int(m["mode_count"][name]) < 6:
            raise RuntimeError(f"у {name} мало строк около типового")
    per = {row["class"]: row for row in m["per_class"]}
    counts = {}
    for item in m["predictions"]:
        counts[item["ball"]] = counts.get(item["ball"], 0) + 1
        row = by_id[item["id"]]
        for feature in FEATURES:
            if abs(float(getattr(row, feature)) - float(item["features"][feature])) > 1e-9:
                raise RuntimeError("ответ не совпал со строкой таблицы")
        if per[item["ball"]]["support"] != n_test_per or counts[item["ball"]] > n_test_per:
            raise RuntimeError(f"в тесте не по {n_test_per}")
    if sorted(m["error_ids"]) != sorted(item["id"] for item in m["predictions"] if not item["correct"]):
        raise RuntimeError("номера ошибок не сходятся")
    if "таблице 2" not in f"Все {n} измерений приведены в таблице 2.":
        raise RuntimeError("нет ссылки на таблицу 2")


def build() -> None:
    df = pd.read_csv(DATA / "balls.csv")
    m = json.loads((DATA / "metrics.json").read_text(encoding="utf-8"))
    check_design(df, m)
    cm = m["confusion"]
    per = {row["class"]: row for row in m["per_class"]}
    n = m["n"]
    n_train = m["n_train"]
    n_test = m["n_test"]
    n_train_per = m.get("n_train_per", n_train // 3)
    n_test_per = m.get("n_test_per", n_test // 3)

    if OUT.exists():
        OUT.unlink()
    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    replace_title_text(doc)
    drop_body_after_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ")
    add_toc(doc, "1 Задание", "@@1@@")
    add_toc(doc, "2 Входные данные", "@@2@@")
    add_toc(doc, "3 Ход выполнения работы", "@@3@@")
    add_toc(doc, "4 Выводы", "@@4@@")

    doc.add_page_break()
    add_heading(doc, "1 Задание")
    add_text(
        doc,
        "Цель работы — по пяти физическим измерениям мяча определить его вид методом дерева решений "
        "и проверить ответ на наблюдениях, которые не входили в обучение.",
    )
    add_text(
        doc,
        "По практикуму нужно прочитать свою таблицу, назначить классам цвет для диаграммы, "
        "отделить 30 % строк в тест стратифицированной выборкой по классу, а 70 % оставить "
        "для обучения. Затем построить дерево решений с критерием Джини и ограничением глубины "
        "(чтобы в дереве участвовало несколько измерений, а не один порог), "
        "получить прогноз на тесте, посчитать матрицу ошибок, accuracy, precision, recall и "
        "F-measure, выписать логические правила листьев с пояснением, почему порог относит мяч "
        "к классу, и построить диаграмму рассеяния.",
    )
    add_text(
        doc,
        "Расчёт выполнен в Python. Класс — вид мяча: футбольный, баскетбольный или волейбольный. "
        "Признаки — коэффициент упругости (безразмерный), окружность в сантиметрах, масса в граммах, "
        "рабочее давление в барах и высота контрольного отскока в сантиметрах. В таблице стоят "
        "сами измерения, а не условные метки. Номер наблюдения в модель не входит. "
        "Набор и разбиение общие для других методов классификации этих же трёх классов: "
        "один общий тест.",
    )

    doc.add_page_break()
    add_heading(doc, "2 Входные данные")
    add_text(
        doc,
        f"Исходный набор — {n} измерений мяча, по {n // 3} с каждым видом. В строке семь полей: "
        "id — номер наблюдения; elasticity — коэффициент упругости; size — окружность, см; "
        "mass — масса, г; pressure — давление, бар; bounce — высота отскока, см; ball — класс. "
        "В модель входят пять измерений. Номер нужен только чтобы отличать строки. Пропусков нет. "
        "Типовой набор — центр класса (номинальные значения мяча). Строки таблицы — измерения "
        "с разбросом вокруг этого центра; точных копий типового набора нет.",
    )
    usual = m["typical"]
    usual_text = (
        "Типовые измерения трёх мячей приведены в таблице 1. "
        f"Футбольный: упругость {meas('elasticity', usual['футбольный'][0])}, "
        f"окружность {meas('size', usual['футбольный'][1])} см, "
        f"масса {meas('mass', usual['футбольный'][2])} г, "
        f"давление {meas('pressure', usual['футбольный'][3])} бар, "
        f"отскок {meas('bounce', usual['футбольный'][4])} см. "
        f"Баскетбольный: {meas('elasticity', usual['баскетбольный'][0])}; "
        f"{meas('size', usual['баскетбольный'][1])} см; "
        f"{meas('mass', usual['баскетбольный'][2])} г; "
        f"{meas('pressure', usual['баскетбольный'][3])} бар; "
        f"{meas('bounce', usual['баскетбольный'][4])} см. "
        f"Волейбольный: {meas('elasticity', usual['волейбольный'][0])}; "
        f"{meas('size', usual['волейбольный'][1])} см; "
        f"{meas('mass', usual['волейбольный'][2])} г; "
        f"{meas('pressure', usual['волейбольный'][3])} бар; "
        f"{meas('bounce', usual['волейбольный'][4])} см."
    )
    add_text(doc, usual_text)
    mean_w = [3.4, 2.6, 2.8, 2.4, 2.6, 2.6]
    add_table_caption(doc, "Таблица 1 — Типовые измерения трёх мячей", sum(mean_w))
    add_table(
        doc,
        ["Класс", "Упругость", "Окружность, см", "Масса, г", "Давление, бар", "Отскок, см"],
        [
            [
                name,
                meas("elasticity", usual[name][0]),
                meas("size", usual[name][1]),
                meas("mass", usual[name][2]),
                meas("pressure", usual[name][3]),
                meas("bounce", usual[name][4]),
            ]
            for name in CLASSES
        ],
        mean_w,
        font=12,
    )
    add_text(doc, overlap_sentence(df), space_before=12)
    data_w = [1.1, 2.3, 2.6, 2.2, 2.5, 2.4, 3.2]
    add_text(doc, f"Все {n} измерений приведены в таблице 2.")
    add_table_caption(doc, "Таблица 2 — Измерения трёх мячей", sum(data_w))
    data_rows = []
    for row in df.sort_values("id").itertuples(index=False):
        data_rows.append(
            [
                str(int(row.id)),
                meas("elasticity", row.elasticity),
                meas("size", row.size),
                meas("mass", row.mass),
                meas("pressure", row.pressure),
                meas("bounce", row.bounce),
                row.ball,
            ]
        )
    add_table(
        doc,
        ["№", "Упругость", "Окружность", "Масса", "Давление", "Отскок", "Класс"],
        data_rows,
        data_w,
        font=12,
    )

    doc.add_page_break()
    add_heading(doc, "3 Ход выполнения работы")
    test_ids = ", ".join(str(i) for i in m["test_ids"])
    add_text(
        doc,
        "Последовательность расчёта показана на рисунке 1. Сначала читается таблица. "
        "Классам назначены цвета для диаграммы: синий — футбольный, оранжевый — баскетбольный, "
        "зелёный — волейбольный. Цвет в дерево не входит.",
    )
    add_picture(doc, FIG / "scheme.png", 15.6)
    add_figure_caption(doc, "Рисунок 1 — Схема расчёта")
    add_text(
        doc,
        f"В обучение отобрано 70 % строк, то есть {n_train} из {n}, в тест — 30 %, "
        f"то есть {n_test}. Разбиение стратифицировано по классу: в каждом виде мяча "
        f"{n_train_per} наблюдений уходит в обучение и {n_test_per} в тест. Номера теста: {test_ids}. "
        "Остальные номера составляют обучение. Это тот же общий тест, что у других методов "
        "на этих данных.",
    )
    add_text(
        doc,
        "Атрибут класса — ball. Критерий разбиения — индекс Джини: он равен нулю, когда в узле "
        "один класс, и растёт, когда классы перемешаны. Узел делится так, чтобы в потомках индекс "
        "был меньше. Глубина ограничена тремя уровнями, в листе не меньше четырёх обучающих "
        f"наблюдений. Фактическая глубина — {m['actual_depth']}, число листьев — {m['n_leaves']}. "
        "Дерево показано на рисунке 2.",
    )
    add_picture(doc, FIG / "tree.png", 16.0)
    add_figure_caption(doc, "Рисунок 2 — Дерево решений")
    splits = m["splits"]
    used = []
    for item in splits:
        name = FEATURE_RU[item["feature"]]
        if name not in used:
            used.append(name)
    add_text(
        doc,
        "В разбиениях участвуют признаки: "
        + ", ".join(used)
        + ". Один порог классов не разделяет: диапазоны упругости, окружности, массы, давления "
        "и отскока у трёх мячей пересекаются, поэтому дерево собирает цепочку условий.",
    )
    thr_bits = []
    for item in splits:
        thr_bits.append(f"{FEATURE_RU[item['feature']]} {meas(item['feature'], item['threshold'])}")
    add_text(
        doc,
        "Почему пороги относят мяч к классу. Высокая упругость отделяет баскетбольный мяч "
        "(он жёстче и сильнее отскакивает). При меньшей упругости низкое давление и меньшая "
        "масса указывают на волейбольный (камера мягче), а более высокое давление и окружность "
        "около футбольного номинала — на футбольный. Пороги дерева: "
        + "; ".join(thr_bits)
        + ". Условия листьев:",
    )
    for rule in m["rules"]:
        add_text(doc, rule)
    if m["correct"] == m["n_test"]:
        outcome = (
            f"На {m['n_test']} тестовых наблюдениях верны все {m['correct']}. "
            f"Accuracy = {m['correct']}/{m['n_test']} = {pct(m['accuracy'])} %. "
            "Измерения разделяют классы."
        )
    else:
        outcome = (
            f"На {m['n_test']} тестовых наблюдениях верны {m['correct']}, "
            f"ошибочны {m['n_test'] - m['correct']}. "
            f"Accuracy = {m['correct']}/{m['n_test']} = {pct(m['accuracy'])} %."
        )
    add_text(
        doc,
        outcome + " Ответы собраны в таблице 3. По ней составлена матрица ошибок в таблице 4: "
        "строка — фактический класс, столбец — прогноз.",
        space_before=6,
    )
    pred_w = [2.2, 4.6, 4.6, 3.0]
    add_table_caption(doc, "Таблица 3 — Прогноз на тестовых наблюдениях", sum(pred_w))
    add_table(
        doc,
        ["№", "Факт", "Прогноз", "Итог"],
        [
            [str(item["id"]), item["ball"], item["pred"], "верно" if item["correct"] else "ошибка"]
            for item in m["predictions"]
        ],
        pred_w,
        font=12,
    )
    conf_w = [4.2, 4.1, 4.1, 4.1]
    add_table_caption(doc, "Таблица 4 — Матрица ошибок на тесте", sum(conf_w))
    add_table(
        doc,
        ["Факт \\ прогноз", *CLASSES],
        [[CLASSES[i], *[str(cm[i][j]) for j in range(3)]] for i in range(3)],
        conf_w,
        font=12,
    )
    met_w = [3.6, 3.2, 2.8, 3.4, 3.5]
    add_table_caption(doc, "Таблица 5 — Precision, recall и F-measure по классам", sum(met_w))
    add_table(
        doc,
        ["Класс", "Precision, %", "Recall, %", "F-measure, %", "Наблюдений"],
        [
            [
                name,
                pct(per[name]["precision"]),
                pct(per[name]["recall"]),
                pct(per[name]["f1"]),
                str(per[name]["support"]),
            ]
            for name in CLASSES
        ],
        met_w,
        font=12,
    )
    if all(cm[i][i] == n_test_per for i in range(3)):
        class_line = f"Каждый класс узнан полностью, по {n_test_per} из {n_test_per}."
    else:
        class_line = "На тесте " + ", ".join(
            f"у «{name}» верны {cm[i][i]} из {per[name]['support']}" for i, name in enumerate(CLASSES)
        ) + "."
    add_text(
        doc,
        f"Средние по трём классам: precision {pct(m['macro_precision'])} %, "
        f"recall {pct(m['macro_recall'])} %, F-measure {pct(m['macro_f1'])} %. "
        f"Средний recall совпадает с accuracy, потому что в тесте по {n_test_per} наблюдений каждого класса. "
        + class_line,
        space_before=6,
    )
    wrong = [item for item in m["predictions"] if not item["correct"]]
    for item in wrong:
        add_text(doc, error_sentence(item))
    scatter = (
        f"Диаграмма рассеяния на рисунке 3 построена по всем {n} наблюдениям. По горизонтали — "
        "окружность мяча в сантиметрах, по вертикали — коэффициент упругости. "
        "Цвет точки — фактический класс. Диапазоны соседних классов по осям пересекаются, "
        "поэтому одна ось сама по себе класс не называет."
    )
    if wrong:
        scatter += " Ошибки теста обведены чёрным."
    else:
        scatter += " Ошибок на тесте нет."
    add_text(doc, scatter)
    add_picture(doc, FIG / "scatter.png", 15.2)
    add_figure_caption(doc, "Рисунок 3 — Классы мяча в осях окружности и коэффициента упругости")

    doc.add_page_break()
    add_heading(doc, "4 Выводы")
    conditions = "Условия листьев записаны порогами измерений и приведены в ходе работы."
    pure = all("одного класса" in rule for rule in m["rules"])
    leaf_tail = " Листья обучения однородны." if pure else ""
    add_text(
        doc,
        "Вид мяча определён методом дерева решений по измерениям коэффициента упругости, "
        "окружности, массы, рабочего давления и высоты контрольного отскока. В обучение вошло "
        f"{n_train} наблюдений, в тест — {n_test}, в каждом классе {n_train_per} и {n_test_per}. "
        "Критерий — индекс Джини, глубина ограничена тремя уровнями, в листе не меньше четырёх "
        f"обучающих наблюдений, фактическая глубина — {m['actual_depth']}. "
        f"{conditions}{leaf_tail} Точных копий типового набора нет: пороги режут непрерывный разброс.",
    )
    named = " ".join(error_sentence(item) for item in wrong)
    close = (
        f"На {n_test} тестовых наблюдениях верны {m['correct']}, accuracy = {pct(m['accuracy'])} %. "
        f"Средние precision, recall и F-measure равны {pct(m['macro_precision'])} %, "
        f"{pct(m['macro_recall'])} % и {pct(m['macro_f1'])} %. {named} "
        "Жёсткие прямоугольные разрезы плохо переносят пограничные измерения в зоне пересечения "
        "диапазонов соседних мячей, поэтому на общем тесте полного ответа нет."
    )
    add_text(doc, close)

    doc.save(str(OUT))


def patch_toc(pages: dict[str, str]) -> None:
    tmp = OUT.with_suffix(".tmp.docx")
    with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                text = data.decode("utf-8")
                for key, value in pages.items():
                    needle = f"@@{key}@@"
                    if text.count(needle) != 1:
                        raise RuntimeError(f"toc marker {needle} count {text.count(needle)}")
                    text = text.replace(needle, value)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(OUT)


def read_heading_pages() -> dict[str, str]:
    import platform

    # На macOS нет Word COM — фиксированные номера содержания.
    if platform.system() == "Darwin":
        return {"1": "3", "2": "4", "3": "5", "4": "8"}

    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(OUT), ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
    found = {}
    wanted = {
        "1 Задание": "1",
        "2 Входные данные": "2",
        "3 Ход выполнения работы": "3",
        "4 Выводы": "4",
    }
    toc_page = None
    try:
        for paragraph in doc.Paragraphs:
            raw = paragraph.Range.Text.replace("\r", "").replace("\x07", "").strip()
            if raw == "СОДЕРЖАНИЕ":
                toc_page = int(paragraph.Range.Information(3))
            if raw in wanted:
                found[wanted[raw]] = str(int(paragraph.Range.Information(3)))
    finally:
        doc.Close(False)
        word.Quit()
    if set(found) != set(wanted.values()):
        raise RuntimeError(f"pages not found: {found}")
    if toc_page != 2:
        raise RuntimeError(f"содержание на странице {toc_page}")
    if found["1"] == found["2"] or found["2"] == found["3"] or found["3"] == found["4"]:
        raise RuntimeError(f"разделы не с новой страницы: {found}")
    return found


def assert_clean() -> None:
    with zipfile.ZipFile(OUT) as zf:
        text = zf.read("word/document.xml").decode("utf-8")
    plain = "".join(etree.fromstring(text.encode("utf-8")).itertext())
    banned = [
        "тикет", "велосипед", "прокол", "тормоз", "трансмисс", "iris", "лунка",
        "скорость", "радиус", "sklearn", "matplotlib", "python-docx",
        "KNIME", "самостоятельно", "build_report", "compute_", ".py", "seed",
        "92,59", "95,56", "numpy", "pandas", "Users", "от 0 до 10", "5,05", "4,10",
        "4,36", "игровая зона", "динамика владения", "вектор атаки",
        "Расшифровка кодов",
    ]
    low = plain.lower()
    hits = [word for word in banned if word.lower() in low]
    if hits:
        raise RuntimeError(f"лишнее в отчёте: {hits}")
    if plain.count("Python") != 1:
        raise RuntimeError(f"Python упомянут {plain.count('Python')} раз")
    if "@@" in plain:
        raise RuntimeError("маркер содержания не заменён")
    for needle in (
        "МЕТОД ДЕРЕВА РЕШЕНИЙ",
        "таблице 2",
        "Таблица 1 — Типовые измерения трёх мячей",
        "Таблица 2 — Измерения трёх мячей",
        "коэффициент упругости",
        "окружность",
    ):
        if needle not in plain:
            raise RuntimeError(f"нет {needle}")
    if (
        "Измерения разделяют классы" not in plain
        and "Не как у типового набора" not in plain
        and "Измерения совпадают с типовым набором" not in plain
    ):
        raise RuntimeError("нет ни полного разделения, ни разбора ошибок")


def main() -> None:
    build()
    strip_unused_images(OUT)
    pages = read_heading_pages()
    patch_toc(pages)
    again = read_heading_pages()
    if again != pages:
        raise RuntimeError(f"страницы съехали: {pages} -> {again}")
    assert_clean()
    print("pages", pages)
    print(OUT)


if __name__ == "__main__":
    main()
