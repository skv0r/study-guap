# -*- coding: utf-8 -*-
import sys
from pathlib import Path

from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_common import (
    add_caption, add_code, add_contents_title, add_h, add_p, add_table, add_toc_item, measure_pages, new_doc,
)

BASE = Path(__file__).resolve().parent
OUT = BASE / "ЛР4" / "ЛР4_Отчет_Буренков_МиСП.docx"
TOC = ["1 Цель работы", "2 Задание", "3 Ход работы", "4 Выводы"]

SQL_A = """DROP SCHEMA IF EXISTS node4 CASCADE;
DROP SCHEMA IF EXISTS office CASCADE;

CREATE SCHEMA office;
CREATE SCHEMA node4;

CREATE TABLE office."Студент" (
  "ID" integer PRIMARY KEY,
  "ФИО" varchar(100) NOT NULL,
  "Факультет" smallint NOT NULL
);

CREATE TABLE node4."Студент" (
  LIKE office."Студент" INCLUDING ALL
);

CREATE TABLE office."Конференция" (
  "ID" integer PRIMARY KEY,
  "Название" varchar(150) NOT NULL
);

CREATE TABLE node4."Конференция" (
  LIKE office."Конференция" INCLUDING ALL
);

CREATE OR REPLACE FUNCTION office.copy_student()
RETURNS trigger
LANGUAGE plpgsql
AS $$
BEGIN
  IF NEW."Факультет" = 4 THEN
    INSERT INTO node4."Студент" ("ID", "ФИО", "Факультет")
    VALUES (NEW."ID", NEW."ФИО", NEW."Факультет");
  END IF;
  RETURN NEW;
END;
$$;

CREATE TRIGGER frag
AFTER INSERT ON office."Студент"
FOR EACH ROW EXECUTE FUNCTION office.copy_student();"""

SQL_B = """INSERT INTO office."Студент" ("ID", "ФИО", "Факультет")
VALUES (1, 'Иванов Иван', 1);

INSERT INTO office."Студент" ("ID", "ФИО", "Факультет")
VALUES (2, 'Сидорова Анна', 4);

SELECT 'центр' AS узел, "ID", "ФИО", "Факультет" FROM office."Студент" ORDER BY "ID";
SELECT 'узел_4' AS узел, "ID", "ФИО", "Факультет" FROM node4."Студент" ORDER BY "ID";"""

SQL_C = """INSERT INTO office."Конференция" ("ID", "Название") VALUES (7, 'Цифровой вуз');
INSERT INTO node4."Конференция" ("ID", "Название") VALUES (7, 'Цифровой вуз');

UPDATE office."Конференция" SET "Название" = 'Цифровой вуз, пленарное' WHERE "ID" = 7;
UPDATE node4."Конференция" SET "Название" = 'Цифровой вуз, секция' WHERE "ID" = 7;

SELECT 'центр' AS сторона, "Название" FROM office."Конференция" WHERE "ID" = 7
UNION ALL
SELECT 'узел_4', "Название" FROM node4."Конференция" WHERE "ID" = 7;

CREATE OR REPLACE FUNCTION office.check_conflict()
RETURNS void
LANGUAGE plpgsql
AS $$
DECLARE
  a text;
  b text;
BEGIN
  SELECT "Название" INTO a FROM office."Конференция" WHERE "ID" = 7;
  SELECT "Название" INTO b FROM node4."Конференция" WHERE "ID" = 7;
  IF a IS DISTINCT FROM b THEN
    RAISE EXCEPTION 'конфликт репликации: центр [%], узел [%]', a, b;
  END IF;
END;
$$;

SELECT office.check_conflict();"""

SQL_FIX = """UPDATE node4."Конференция"
   SET "Название" = (SELECT "Название" FROM office."Конференция" WHERE "ID" = 7)
 WHERE "ID" = 7;

SELECT office.check_conflict();

SELECT 'центр' AS сторона, "Название" FROM office."Конференция" WHERE "ID" = 7
UNION ALL
SELECT 'узел_4', "Название" FROM node4."Конференция" WHERE "ID" = 7;"""


def build(pages):
    doc = new_doc("4", "Репликация слиянием, фрагмент и конфликт")
    add_contents_title(doc)
    for title in TOC:
        add_toc_item(doc, title, pages[title])

    add_h(doc, "1 Цель работы", page_break=True)
    add_p(
        doc,
        "Получить навык фрагментации таблицы «Студент» по факультету и разбора конфликта "
        "по одной и той же конференции, изменённой в центре и на узле.",
    )

    add_h(doc, "2 Задание", page_break=True)
    add_p(
        doc,
        "Разделить данные варианта 5 так, чтобы на узел факультета попадали только его студенты. "
        "Смоделировать конфликт двух правок одной строки и показать, как сервер его объявляет и как правка выравнивается. "
        "Блоки вставляются в редактор SQL по очереди: следующий начинается, когда предыдущий уже выполнен.",
    )

    add_h(doc, "3 Ход работы", page_break=True)
    add_p(
        doc,
        "Среда работы — PostgreSQL. Центр находится в схеме office, узел факультета 4 — в схеме node4. "
        "Триггер frag копирует студента только при факультете 4. Строка другого факультета остаётся в центре. "
        "Конференция копируется на обе стороны заранее, затем названия расходятся. Сверка поднимает исключение, "
        "пока тексты разные, и проходит, когда на узле оставлено название центра.",
    )

    add_h(doc, "3.1 Блок A. Фрагмент")
    add_p(doc, "Блок создаёт обе схемы, таблицы студентов и конференций и триггер фрагмента. На первом запуске DROP сообщает, что схем ещё нет.")
    add_code(doc, SQL_A)
    add_p(doc, "Ответ сервера:")
    add_code(
        doc,
        "NOTICE:  schema \"node4\" does not exist, skipping\n"
        "NOTICE:  schema \"office\" does not exist, skipping\n"
        "DROP SCHEMA\n"
        "DROP SCHEMA\n"
        "CREATE SCHEMA\n"
        "CREATE SCHEMA\n"
        "CREATE TABLE\n"
        "CREATE TABLE\n"
        "CREATE TABLE\n"
        "CREATE TABLE\n"
        "CREATE FUNCTION\n"
        "CREATE TRIGGER",
    )

    heading = add_h(doc, "3.2 Блок B. Кто попал на узел")
    heading.paragraph_format.space_before = Pt(12)
    add_p(doc, "Иванов Иван относится к факультету 1, Сидорова Анна — к факультету 4. В центре после блока две строки. На узле одна: только факультет 4.")
    add_code(doc, SQL_B)
    add_p(doc, "Ответ сервера:")
    add_code(
        doc,
        "INSERT 0 1\n"
        "INSERT 0 1\n"
        "\n"
        " узел  | ID |      ФИО      | Факультет\n"
        "-------+----+---------------+-----------\n"
        " центр |  1 | Иванов Иван   |         1\n"
        " центр |  2 | Сидорова Анна |         4\n"
        "(2 строки)\n"
        "\n"
        "  узел  | ID |      ФИО      | Факультет\n"
        "--------+----+---------------+-----------\n"
        " узел_4 |  2 | Сидорова Анна |         4\n"
        "(1 строка)",
    )
    add_p(doc, "Те же строки сведены в таблице 1: Иванов Иван на узел не попадает, Сидорова Анна попадает.")
    add_caption(doc, "Таблица 1 — Фрагмент по факультету 4")
    add_table(
        doc,
        ["Сторона", "ID", "ФИО", "Факультет", "На узле"],
        [
            ["центр", "1", "Иванов Иван", "1", "нет"],
            ["центр и узел", "2", "Сидорова Анна", "4", "да"],
        ],
        [3.5, 1.5, 4.5, 3, 4],
    )

    heading = add_h(doc, "3.3 Блок C. Конфликт")
    heading.paragraph_format.space_before = Pt(12)
    add_p(doc, "Конференция 7 сначала одинакова: «Цифровой вуз». Затем центр записывает «Цифровой вуз, пленарное», а узел — «Цифровой вуз, секция». Сверка видит два текста и поднимает исключение. Это и есть ответ сервера на конфликт: оператор не проходит.")
    add_code(doc, SQL_C)
    add_p(doc, "Ответ сервера на сверку:")
    add_code(
        doc,
        "INSERT 0 1\n"
        "INSERT 0 1\n"
        "UPDATE 1\n"
        "UPDATE 1\n"
        "\n"
        " сторона |        Название\n"
        "---------+-------------------------\n"
        " центр   | Цифровой вуз, пленарное\n"
        " узел_4  | Цифровой вуз, секция\n"
        "(2 строки)\n"
        "\n"
        "CREATE FUNCTION\n"
        "ERROR:  конфликт репликации: центр [Цифровой вуз, пленарное], узел [Цифровой вуз, секция]\n"
        "КОНТЕКСТ:  PL/pgSQL function office.check_conflict() line 9 at RAISE",
    )
    add_p(doc, "Названия до разрешения сведены в таблице 2.")
    add_caption(doc, "Таблица 2 — Строка конференции до разрешения")
    add_table(
        doc,
        ["Сторона", "ID", "Название"],
        [["центр", "7", "Цифровой вуз, пленарное"], ["узел факультета 4", "7", "Цифровой вуз, секция"]],
        [5, 2, 9.5],
    )

    heading = add_h(doc, "3.4 Разрешение конфликта")
    heading.paragraph_format.space_before = Pt(12)
    add_p(doc, "Разрешение оставляет значение центра. Узел получает «Цифровой вуз, пленарное». Повторная сверка уже не поднимает исключение и возвращает пустую строку результата. Оба SELECT после этого показывают одно название.")
    add_code(doc, SQL_FIX)
    add_p(doc, "Ответ сервера:")
    add_code(
        doc,
        "UPDATE 1\n"
        "\n"
        " check_conflict\n"
        "----------------\n"
        " \n"
        "(1 строка)\n"
        "\n"
        " сторона |        Название\n"
        "---------+-------------------------\n"
        " центр   | Цифровой вуз, пленарное\n"
        " узел_4  | Цифровой вуз, пленарное\n"
        "(2 строки)",
    )

    add_h(doc, "4 Выводы", page_break=True)
    add_p(
        doc,
        "Фрагмент узла ограничен факультетом 4: Иванов Иван в выборке узла отсутствует, Сидорова Анна присутствует. "
        "Конфликт конференции 7 сервер объявил текстом «конфликт репликации: центр [Цифровой вуз, пленарное], узел [Цифровой вуз, секция]». "
        "После выравнивания на узле осталось название центра, повторная сверка прошла без ошибки, обе стороны вернули «Цифровой вуз, пленарное».",
    )
    doc.save(OUT)


def main():
    build({title: "…" for title in TOC})
    pages = measure_pages(OUT, TOC)
    build(pages)
    print(pages)


if __name__ == "__main__":
    main()
