# -*- coding: utf-8 -*-
"""Сборка отчёта из бланка прошлого года. Имя этого файла в отчёт не попадает."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

TEMPLATE = Path(
    r"C:\Users\skvor\OneDrive\Desktop\GitHub\study-guap\6_semester\МиCПИСиТ\lr1\Буренков_ВАР5_ЛР1.docx"
)
INDENT = Cm(1.25)


def set_run_font(run, name, size_pt, bold=None):
    run.bold = bold
    run.italic = False
    run.font.name = name
    run.font.size = Pt(size_pt)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)


def strip_body(doc):
    body = doc.element.body
    seen = 0
    remove_rest = False
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if remove_rest:
            body.remove(child)
            continue
        if child.tag == qn("w:p"):
            pPr = child.find(qn("w:pPr"))
            if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
                seen += 1
                if seen >= 2:
                    remove_rest = True
    if seen < 2:
        raise RuntimeError(f"title sections not found: {seen}")


def set_theme(doc, number: str, theme: str):
    cell = doc.tables[1].cell(0, 0)
    for run in cell.paragraphs[0].runs:
        if run.text.strip() in {"1", "2", "3", "4", "5", "6", "7", "8"}:
            run.text = number
            run.bold = None
    # тема — отдельный абзац, без жирного
    theme_par = cell.paragraphs[2]
    if theme_par.runs:
        theme_par.runs[0].text = theme
        theme_par.runs[0].bold = None
        for extra in theme_par.runs[1:]:
            extra.text = ""
    else:
        theme_par.add_run(theme)


def _base_paragraph(doc, align, first_indent, line, space_after=0, page_break=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if line == 1.5 else WD_LINE_SPACING.SINGLE
    if line == 1.5:
        pf.line_spacing = 1.5
    else:
        pf.line_spacing = 1.0
    pf.first_line_indent = first_indent
    pf.left_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if page_break:
        pf.page_break_before = True
    return p


def add_contents_title(doc):
    p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER, Cm(0), 1.5)
    run = p.add_run("СОДЕРЖАНИЕ")
    set_run_font(run, "Times New Roman", 14, bold=True)
    return p


def add_h(doc, text, page_break=False):
    p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.LEFT, INDENT, 1.5, page_break=page_break)
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", 14, bold=True)
    return p


def add_p(doc, text):
    p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.JUSTIFY, INDENT, 1.5)
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", 14, bold=False)
    return p


def text_width(doc):
    """Ширина текстового поля последней секции, без выхода за правое поле."""
    sec = doc.sections[-1]
    emu = int(sec.page_width) - int(sec.left_margin) - int(sec.right_margin)
    return Emu(emu)


def add_toc_item(doc, text, page):
    p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.LEFT, INDENT, 1.5)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.tab_stops.add_tab_stop(
        text_width(doc), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    set_run_font(p.add_run(text), "Times New Roman", 14, bold=False)
    set_run_font(p.add_run("\t" + str(page)), "Times New Roman", 14, bold=False)
    return p


def add_caption(doc, text):
    p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), 1.0, space_after=6)
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", 14, bold=False)
    return p


def add_code(doc, text):
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for line in lines:
        p = _base_paragraph(doc, WD_ALIGN_PARAGRAPH.LEFT, Cm(0), 1.0)
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Times New Roman", 14, bold=False)
    return


def _fill_cell(par, text, bold):
    parts = str(text).split("\n")
    run = par.add_run(parts[0] if parts[0] else " ")
    set_run_font(run, "Times New Roman", 14, bold=bold)
    for part in parts[1:]:
        run.add_break()
        run = par.add_run(part)
        set_run_font(run, "Times New Roman", 14, bold=bold)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        table.style = "Table Normal"
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        _fill_cell(par, h, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            _fill_cell(par, str(val), bold=False)
    if widths:
        table.autofit = False
        table.allow_autofit = False
        total = 0
        grid = table._tbl.find(qn("w:tblGrid"))
        for i, cm in enumerate(widths):
            tw = int(Cm(cm).twips)
            total += tw
            grid[i].set(qn("w:w"), str(tw))
            for row in table.rows:
                row.cells[i].width = Cm(cm)
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            table._tbl.tblPr.append(tbl_w)
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(total))
        margins = OxmlElement("w:tblCellMar")
        for edge, val in (("top", "28"), ("left", "40"), ("bottom", "28"), ("right", "40")):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), val)
            el.set(qn("w:type"), "dxa")
            margins.append(el)
        table._tbl.tblPr.append(margins)
    doc.add_paragraph()
    return table


def measure_pages(path: Path, titles: list[str]) -> dict[str, str]:
    import subprocess
    import tempfile

    listing = Path(tempfile.gettempdir()) / "misp_pages.txt"
    script = Path(tempfile.gettempdir()) / "misp_pages.ps1"
    wanted = set(titles)
    script.write_text(
        "\n".join(
            [
                "$ErrorActionPreference = 'Stop'",
                "$word = New-Object -ComObject Word.Application",
                "$word.Visible = $false",
                "$word.DisplayAlerts = 0",
                f"$doc = $word.Documents.Open('{path}', $false, $true)",
                "$doc.Repaginate()",
                "$lines = New-Object System.Collections.Generic.List[string]",
                "foreach ($p in $doc.Paragraphs) {",
                '  $t = $p.Range.Text.Replace("`r","").Replace("`a","").Trim()',
                '  if ($t.Length -gt 0 -and -not $t.Contains("`t")) {',
                "    $page = [int]$p.Range.Information(1)",
                '    $lines.Add("$page`t$t")',
                "  }",
                "}",
                "$doc.Close(0)",
                "$word.Quit()",
                f"[System.IO.File]::WriteAllLines('{listing}', $lines, [System.Text.UTF8Encoding]::new($false))",
            ]
        ),
        encoding="utf-8-sig",
    )
    subprocess.run(["powershell", "-NoProfile", "-File", str(script)], check=True)
    found: dict[str, str] = {}
    for line in listing.read_text(encoding="utf-8").splitlines():
        if "\t" not in line:
            continue
        page, text = line.split("\t", 1)
        text = text.strip()
        if text in wanted and text not in found:
            found[text] = page.strip()
    missing = [title for title in titles if title not in found]
    if missing:
        raise RuntimeError("Не найдены страницы: " + ", ".join(missing))
    return found


def tighten_initials(doc: Document) -> None:
    def fix(text: str) -> str:
        return text.replace("А. В.", "А.В.")

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text and "А. В." in run.text:
                run.text = fix(run.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.text and "А. В." in run.text:
                            run.text = fix(run.text)


def new_doc(number: str, theme: str) -> Document:
    doc = Document(str(TEMPLATE))
    strip_body(doc)
    set_theme(doc, number, theme)
    tighten_initials(doc)
    return doc
