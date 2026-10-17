# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР3. В текст отчёта не входит."""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from lxml import etree
from PIL import Image

import draw_uml

BASE = Path(__file__).resolve().parent
SRC = BASE.parent / "ЛР2" / "ЛР2_Отчет_Буренков_КИС.docx"
OUT = BASE / "ЛР3_Отчет_Буренков_КИС.docx"
FIG = BASE / "figures"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TITLE_REPLACEMENTS = {
    "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3",
    "Изучение методологии моделирования бизнес-процессов IDEF0":
        "Изучение методологии структурного моделирования UML",
}


def prepare_title(path: Path) -> None:
    shutil.copy(SRC, path)
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml")
    root = etree.fromstring(xml)
    body = root.find(f"{W}body")
    children = list(body)
    p11 = children[11]
    sect = p11.find(f"{W}pPr/{W}sectPr")
    if sect is None:
        raise RuntimeError("В титуле ЛР2 нет sectPr")
    sect = etree.fromstring(etree.tostring(sect))
    for child in children:
        body.remove(child)
    for child in children[:11]:
        body.append(child)
    body.append(sect)
    replaced = []
    for node in body.iter(f"{W}t"):
        if node.text in TITLE_REPLACEMENTS:
            replaced.append(node.text)
            node.text = TITLE_REPLACEMENTS[node.text]
    missing = [k for k in TITLE_REPLACEMENTS if k not in replaced]
    if missing:
        raise RuntimeError(f"Не найдены строки титула: {missing}")
    report = None
    for node in body.iter(f"{W}t"):
        if node.text and node.text.startswith("ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ"):
            report = node
            break
    rpr = report.getparent().find(f"{W}rPr")
    if rpr is not None and rpr.find(f"{W}b") is not None:
        raise RuntimeError("Строка отчёта на титуле стала жирной")
    xml_out = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = path.with_suffix(".tmp.zip")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            data = xml_out if item.filename == "word/document.xml" else zin.read(item.filename)
            zout.writestr(item, data)
    tmp.replace(path)


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(p, *, indent=True, align="justify", before=0, after=0, line=1.5, left=0):
    pf = p.paragraph_format
    if line == 1:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Cm(1.25 if indent else 0)
    pf.left_indent = Cm(left)
    pf.right_indent = Cm(0)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, bookmark, bid):
    p = doc.add_paragraph()
    format_paragraph(p, indent=True, align="left", before=12, after=6)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), bookmark)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    p_pr = p._p.find(qn("w:pPr"))
    if p_pr is not None:
        p_pr.addnext(start)
    else:
        p._p.insert(0, start)
    p._p.append(end)
    return p


def add_toc_line(doc, title, page, level):
    p = doc.add_paragraph()
    format_paragraph(p, indent=False, align="left", left=0.75 * level)
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    pPr.append(tabs)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t" + str(page)))
    return p


def add_caption_figure(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, indent=False, align="center", before=6, after=0)
    set_run_font(p.add_run(text))
    return p


def add_picture(doc, path, width_cm):
    p = doc.add_paragraph()
    format_paragraph(p, indent=False, align="center", before=6, after=0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    return p


def footer_rel_id(doc) -> str:
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/footer"):
            return rel.rId
    raise RuntimeError("В образце нет колонтитула с номером страницы")


def apply_section(section, *, landscape: bool, start: int | None, footer_id: str):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    sect = section._sectPr
    for tag in ("w:footerReference", "w:headerReference"):
        for el in list(sect.findall(qn(tag))):
            sect.remove(el)
    ref = OxmlElement("w:footerReference")
    ref.set(qn("w:type"), "default")
    ref.set(qn("r:id"), footer_id)
    sect.insert(0, ref)
    pg = sect.find(qn("w:pgNumType"))
    if start is None:
        if pg is not None:
            sect.remove(pg)
    else:
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sect.append(pg)
        pg.set(qn("w:start"), str(start))


def heading(doc, text, bookmark, bid, new_page=False):
    p = add_heading(doc, text, bookmark, bid)
    if new_page:
        p.paragraph_format.page_break_before = True
    return p


def fit_width(path: Path, max_h_cm: float, max_w_cm: float = 24.2) -> float:
    with Image.open(path) as im:
        w_px, h_px = im.size
    height_at_full = max_w_cm * h_px / w_px
    if height_at_full <= max_h_cm:
        return max_w_cm
    return max_h_cm * w_px / h_px


def slim_package(path: Path) -> None:
    with zipfile.ZipFile(path) as z:
        rels_xml = z.read("word/_rels/document.xml.rels")
        doc_xml = z.read("word/document.xml")
        files = {name: z.read(name) for name in z.namelist()}
    rels = etree.fromstring(rels_xml)
    root = etree.fromstring(doc_xml)
    used = set()
    for el in root.iter():
        for attr, value in el.attrib.items():
            if attr.endswith("}embed") or attr.endswith("}id") and value.startswith("rId"):
                used.add(value)
    keep_targets = set()
    for rel in list(rels):
        target = rel.get("Target") or ""
        rid = rel.get("Id")
        droppable = target.startswith("media/") or target.startswith("footer")
        if droppable and rid not in used:
            rels.remove(rel)
        else:
            keep_targets.add("word/" + target.replace("\\", "/"))
    rels_out = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
    ctypes = etree.fromstring(files["[Content_Types].xml"])
    ns = {"ct": "http://schemas.openxmlformats.org/package/2006/content-types"}
    for override in list(ctypes.findall("ct:Override", ns)):
        part = override.get("PartName", "")
        if part.startswith("/word/media/") or part.startswith("/word/footer"):
            if part.lstrip("/") not in keep_targets:
                ctypes.remove(override)
    files["word/_rels/document.xml.rels"] = rels_out
    files["[Content_Types].xml"] = etree.tostring(
        ctypes, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    drop = [
        name for name in files
        if (name.startswith("word/media/") or name.startswith("word/footer")) and name not in keep_targets
    ]
    for name in drop:
        del files[name]
    tmp = path.with_suffix(".tmp.zip")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


def enable_field_update(path: Path) -> None:
    with zipfile.ZipFile(path) as z:
        settings = z.read("word/settings.xml")
    root = etree.fromstring(settings)
    tag = f"{W}updateFields"
    if root.find(tag) is None:
        el = etree.Element(tag)
        el.set(qn("w:val"), "true")
        root.append(el)
    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    tmp = path.with_suffix(".tmp.zip")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for item in zin.infolist():
            payload = data if item.filename == "word/settings.xml" else zin.read(item.filename)
            zout.writestr(item, payload)
    tmp.replace(path)


SPECS = [
    (
        "3.1 Диаграмма вариантов использования",
        "z31",
        "fig01_usecase.png",
        "Рисунок 1 — Диаграмма вариантов использования",
        "На рисунке 1 граница системы — СервисЗадач. Заказчик сообщает требование и запрашивает статус задачи. "
        "Руководитель команды регистрирует срочную задачу, назначает разработчика и закрывает задачу. "
        "Разработчик принимает задачу и оформляет протокол проверки. Вторичный актор БазаДанных участвует в записи "
        "состояния: назначение и закрытие включают сохранение в хранилище, поэтому состояние задачи не «живёт» "
        "только у пользователя, а фиксируется в контуре учёта.",
        11.0,
    ),
    (
        "3.2 Диаграмма классов",
        "z32",
        "fig02_classes.png",
        "Рисунок 2 — Диаграмма классов",
        "На рисунке 2 выделены два архитектурных слоя. СервисЗадач (приложение) принимает команды учётного контура "
        "и использует БазаДанных (хранилище таблиц задач, модулей, назначений и протоколов). Доменный слой "
        "сохраняет сущности Задача, Модуль, Команда и ПротоколПроверки: задача относится к модулю, до назначения "
        "связана с командой кратностью 0..1, имеет не более одного протокола. Сервис управляет задачей и через "
        "хранилище меняет её состояние; доменные операции соответствуют переходам на диаграмме состояний.",
        10.8,
    ),
    (
        "3.3 Диаграмма последовательности",
        "z33",
        "fig03_sequence.png",
        "Рисунок 3 — Назначение разработчика на срочную задачу",
        "На рисунке 3 взаимодействие идёт через СервисЗадач к БазаДанных, а не напрямую от руководителя к разработчику. "
        "Руководитель регистрирует срочную задачу; сервис сохраняет её в БД со состоянием «Новая» и возвращает номер и срок. "
        "При назначении сервис проверяет срок сдачи, обновляет состояние в БД на «Назначена» и уведомляет разработчика. "
        "После принятия задачи сервис записывает в БД состояние «В работе» и сообщает его руководителю. "
        "Владелец переходов — сервис; источник истины по состоянию — хранилище.",
        11.2,
    ),
    (
        "3.4 Диаграмма состояний задачи",
        "z34",
        "fig04_states.png",
        "Рисунок 4 — Диаграмма состояний задачи",
        "На рисунке 4 у задачи ровно пять состояний. Состояние хранится в БазаДанных, переходы выполняет СервисЗадач. "
        "Регистрация создаёт «Новую», назначение — «Назначена», принятие — «В работе». Если проверка пройдена, "
        "задача переходит в «Проверена», иначе возвращается в «В работе» без введения отдельного статуса. "
        "Закрытие даёт «Закрыта».",
        10.0,
    ),
    (
        "3.5 Диаграмма деятельности",
        "z35",
        "fig05_activity.png",
        "Рисунок 5 — Закрытие задачи после проверки",
        "На рисунке 5 четыре дорожки отражают закрытие после проверки в архитектуре ИС: разработчик, СервисЗадач, "
        "БазаДанных и руководитель команды. Задача находится в «В работе». Разработчик фиксирует результат и передаёт "
        "протокол в сервис; сервис проверяет результат, БД сохраняет протокол и записывает состояние «Проверена». "
        "Руководитель сверяет протокол со сроком и закрывает задачу; сервис запрашивает закрытие, БД записывает "
        "«Закрыта». Проверка и запись в хранилище — обязательные шаги контура, а не «бумажный» финал вне системы.",
        11.0,
    ),
]


def build():
    draw_uml.main()
    prepare_title(OUT)
    doc = Document(str(OUT))
    footer_id = footer_rel_id(doc)
    body = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section(body, landscape=False, start=2, footer_id=footer_id)

    p = doc.add_paragraph()
    format_paragraph(p, indent=False, align="center", before=0, after=12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run("СОДЕРЖАНИЕ"), bold=True)

    toc = [
        ("1 Задание", 3, 0),
        ("2 Описание системы", 4, 0),
        ("3 Спецификации UML", 5, 0),
        ("3.1 Диаграмма вариантов использования", 6, 1),
        ("3.2 Диаграмма классов", 7, 1),
        ("3.3 Диаграмма последовательности", 8, 1),
        ("3.4 Диаграмма состояний задачи", 9, 1),
        ("3.5 Диаграмма деятельности", 10, 1),
        ("4 Выводы", 11, 0),
    ]
    for title, page, level in toc:
        add_toc_line(doc, title, page, level)

    heading(doc, "1 Задание", "z1", 1, True)
    add_body(
        doc,
        "Нужно изучить методологию структурного моделирования UML и показать её на одной информационной системе. "
        "Ниже дано краткое описание КИС «СпринтКонтур» как архитектурного контура учёта задач и пять нотаций: "
        "варианты использования, классы, последовательность, состояния и деятельность. "
        "Состояния задачи на этих диаграммах названы одинаково.",
    )

    heading(doc, "2 Описание системы", "z2", 2, True)
    add_body(
        doc,
        "КИС «СпринтКонтур» — информационная система учёта задач команды, разрабатывающей программный модуль. "
        "Архитектурная задача работы: спроектировать контур учёта с разделением на подсистемы — сервис контроля "
        "и учёта задач (СервисЗадач), хранилище (БазаДанных: задачи, модули, назначения, протоколы проверки), "
        "клиентские роли на границе системы (руководитель команды, разработчик, заказчик) и взаимодействие "
        "компонентов при назначении срочной задачи и при закрытии после проверки.",
    )
    add_body(
        doc,
        "Выбранное решение: пользователи обращаются к СервисЗадач; сервис владеет правилами переходов "
        "и записывает состояние задачи в БазаДанных; доменные сущности Задача, Модуль, Команда и ПротоколПроверки "
        "хранятся в БД и не существуют «вне» контура. Задача проходит пять состояний: «Новая», «Назначена», "
        "«В работе», «Проверена» и «Закрыта». Если проверка не пройдена, сервис возвращает задачу в «В работе». "
        "Учёт рабочего времени, договоры и обмен с бухгалтерией на этих диаграммах не рассматриваются.",
    )

    heading(doc, "3 Спецификации UML", "z3", 3, True)
    add_body(
        doc,
        "Пять диаграмм описывают один архитектурный контур срочной задачи. Варианты использования задают роли "
        "и границу СервисЗадач с участием БазаДанных. Классы показывают слой приложения, хранилище и домен. "
        "Последовательность проводит назначение через сервис к БД до состояния «В работе». "
        "Диаграмма состояний фиксирует пять статусов и владельца переходов. "
        "Деятельность закрывает задачу после проверки с записью протокола и состояния в БД.",
    )

    for i, (title, mark, filename, caption, text, max_h) in enumerate(SPECS, start=1):
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        apply_section(sec, landscape=True, start=None, footer_id=footer_id)
        heading(doc, title, mark, 10 + i, False)
        add_body(doc, text)
        picture = FIG / filename
        add_picture(doc, picture, fit_width(picture, max_h))
        add_caption_figure(doc, caption)

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_section(sec, landscape=False, start=None, footer_id=footer_id)
    heading(doc, "4 Выводы", "z4", 90, False)
    add_body(
        doc,
        "Для КИС «СпринтКонтур» составлены пять UML-диаграмм архитектурного контура учёта задач: СервисЗадач, "
        "БазаДанных и доменные сущности. Варианты использования отделяют роли пользователей от хранилища. "
        "Классы показывают, что сервис использует БД и управляет задачей; кратность 0..1 у команды и протокола "
        "означает, что до назначения и до проверки этих связей ещё нет.",
    )
    add_body(
        doc,
        "Состояния задачи совпадают на диаграммах последовательности, состояний и деятельности: "
        "«Новая», «Назначена», «В работе», «Проверена», «Закрыта». Переходы выполняет сервис, состояние хранится в БД. "
        "Назначение разработчика заканчивается «В работе». Если проверка не пройдена, задача возвращается "
        "в «В работе». Закрытие после проверки проходит через запись протокола и состояния «Проверена», "
        "затем «Закрыта».",
    )
    doc.save(str(OUT))
    slim_package(OUT)
    enable_field_update(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
