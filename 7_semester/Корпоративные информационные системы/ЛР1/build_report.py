#!/usr/bin/env python3
"""Отчёт ЛР1 по КИС: бланк ГУАП, Times New Roman 14, интервал 1,5."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

import generate_figures as gf
import project_plan as P

BASE = Path(__file__).resolve().parent
BLANK = BASE / "guap_blanks" / "lab.docx"
OUT = BASE / "ЛР1_Отчет_Буренков_КИС.docx"
FIG = BASE / "figures"

TOC_DEF = [
    ("ВВЕДЕНИЕ", "intro", 0),
    ("1 Краткое описание проектируемой информационной системы", "s1", 0),
    ("2 Описание используемого программного пакета", "s2", 0),
    ("2.1 Производитель и права использования", "s21", 1),
    ("2.2 Назначение и основные функции", "s22", 1),
    ("2.3 Порядок работы в лабораторной", "s23", 1),
    ("3 Список этапов и подэтапов проекта", "s3", 0),
    ("3.1 Ресурсы проекта", "s31", 1),
    ("3.2 Оценка затрат по работам", "s32", 1),
    ("4 Основные этапы управления проектом", "s4", 0),
    ("5 Развёрнутая диаграмма Ганта", "s5", 0),
    ("ЗАКЛЮЧЕНИЕ", "conc", 0),
    ("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", "src", 0),
]

SEARCH = {
    "intro": "ВВЕДЕНИЕ",
    "s1": "1 Краткое описание проектируемой информационной системы",
    "s2": "2 Описание используемого программного пакета",
    "s21": "2.1 Производитель и права использования",
    "s22": "2.2 Назначение и основные функции",
    "s23": "2.3 Порядок работы в лабораторной",
    "s3": "3 Список этапов и подэтапов проекта",
    "s31": "3.1 Ресурсы проекта",
    "s32": "3.2 Оценка затрат по работам",
    "s4": "4 Основные этапы управления проектом",
    "s5": "5 Развёрнутая диаграмма Ганта",
    "conc": "ЗАКЛЮЧЕНИЕ",
    "src": "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
}


def rub(n: float) -> str:
    return f"{int(round(n)):,}".replace(",", "\u00a0") + " ₽"


def set_run_font(run, size=14, bold=False, name="Times New Roman", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(
    p,
    *,
    first_indent=True,
    align="justify",
    space_after=0,
    space_before=0,
    left_indent=0,
    line_spacing=1.5,
):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]


def add_body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p)
    set_run_font(p.add_run(text))
    return p


def add_section_h(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_before=0, space_after=12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)


def add_struct(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_after=18)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)


def add_sub_h(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_before=10, space_after=6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)


def add_toc_line(doc, title, page, *, indent=0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=0.75 * indent)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    p.add_run("\t")
    set_run_font(p.add_run(str(page)))


def add_caption_table(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=10, space_after=4)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def add_figure(doc, path, caption, w=16.0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=10)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(w))
    c = doc.add_paragraph()
    format_paragraph(c, first_indent=False, align="center", space_before=6, space_after=10)
    set_run_font(c.add_run(caption))


def _set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_shading(cell, fill="D9D9D9"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, headers, rows, widths_cm=None, size=10):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    usable = Cm(16.5)
    if widths_cm:
        total = sum(widths_cm)
        widths = [int(usable * (w / total)) for w in widths_cm]
    else:
        widths = [int(usable / len(headers))] * len(headers)
    grid = t._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        if i < len(widths):
            child.set("w", str(widths[i]))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for cell in row.cells:
            _set_cell_border(cell)
            for i, w in enumerate(widths):
                if cell._tc is row.cells[i]._tc:
                    cell.width = Emu(w)
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        set_cell_shading(cell)
        p = cell.paragraphs[0]
        format_paragraph(p, first_indent=False, align="center", line_spacing=1.15)
        set_run_font(p.add_run(h), bold=True, size=size)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            align = "center" if j == 0 else "left"
            format_paragraph(p, first_indent=False, align=align, line_spacing=1.15)
            set_run_font(p.add_run(v), size=size)
    spacer = doc.add_paragraph()
    format_paragraph(spacer, first_indent=False, line_spacing=1.0, space_after=6)
    return t


def add_page_field(paragraph):
    run = paragraph.add_run()
    for kind, val in [("begin", None), ("instr", " PAGE "), ("end", None)]:
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    set_run_font(run)


def put_cell(cell, text, size=14):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r.text = ""
    if not p.runs:
        run = p.add_run(text)
    else:
        run = p.runs[0]
        run.text = text
    set_run_font(run, size=size)


def fill_title(doc: Document):
    for p in doc.paragraphs:
        if p.text.startswith("КАФЕДРА"):
            if len(p.runs) >= 2 and p.runs[1].text.strip("_") == "":
                p.runs[1].text = "42"
            else:
                p.clear()
                r = p.add_run("КАФЕДРА № 42")
                set_run_font(r, size=14)
        elif "20__" in p.text:
            for r in p.runs:
                if "20__" in r.text:
                    r.text = r.text.replace("20__", "2026")

    t_prep, t_work, t_stud = doc.tables
    put_cell(t_prep.rows[0].cells[0], "д-р техн. наук, профессор")
    put_cell(t_prep.rows[0].cells[4], "В. В. Фомин")
    put_cell(t_work.rows[0].cells[0], "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1")
    put_cell(t_work.rows[1].cells[0], "Изучение планирование и управление ресурсами")
    put_cell(t_work.rows[2].cells[0], "по курсу: Корпоративные информационные системы")
    put_cell(t_stud.rows[0].cells[1], "4321")
    put_cell(t_stud.rows[0].cells[5], "Г.В. Буренков")


def setup_body_section(doc: Document):
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    title_sec, body_sec = doc.sections[0], doc.sections[1]

    for sec in (title_sec, body_sec):
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin, sec.right_margin = Mm(30), Mm(15)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)

    title_sec.footer.is_linked_to_previous = False
    for p in title_sec.footer.paragraphs:
        p.clear()

    body_sec.footer.is_linked_to_previous = False
    body_sec.header.is_linked_to_previous = False

    sect_pr = body_sec._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "2")
    sect_pr.append(pg)

    fp = body_sec.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def task_rows():
    rows = []
    for code in P.ORDER:
        t = P.TASKS[code]
        if t.kind == "summary":
            typ = "этап верхнего уровня"
            dur = "—"
        elif t.kind == "milestone":
            typ = "промежуточная точка (веха)"
            dur = "0"
        else:
            typ = "подзадача" + (f" ({t.parallel_note})" if t.parallel_note else "")
            dur = f"{t.duration} раб. дн."
        pred = "—" if not t.preds else ", ".join(t.preds)
        rows.append([code, t.name, typ, dur, pred, P.fmt(t.es), P.fmt(t.ef)])
    return rows


def allocation_text(t: P.Task) -> str:
    if not t.alloc:
        return "—"
    parts = []
    for rid, frac in t.alloc:
        pct = int(round(frac * 100))
        parts.append(f"{P.RESOURCES[rid]['fio']} {pct}%")
    return ", ".join(parts)


def build_body(doc: Document, toc_pages: dict):
    add_struct(doc, "СОДЕРЖАНИЕ", new_page=False)
    for title, key, ind in TOC_DEF:
        add_toc_line(doc, title, toc_pages.get(key, "…"), indent=ind)

    add_struct(doc, "ВВЕДЕНИЕ")
    add_body(
        doc,
        "Корпоративная информационная система отражает архитектуру организации и "
        "сопровождает её многофункциональную деятельность. На современном этапе ядром "
        "КИС предприятий служат системы планирования ресурсов (ERP), а управление "
        "проектом создания такой системы само по себе требует календарного плана, "
        "учёта зависимостей работ и распределения трудовых и денежных ресурсов [1]. "
        "Целевая задача настоящей лабораторной работы — изучение технологий планирования "
        "и управления процессами и ресурсами: построение расписания с выделением "
        "критического пути и визуализация плана на диаграмме Ганта.",
    )
    add_body(
        doc,
        "По условию задания выбран собственный сценарий проекта и составлена иерархия "
        "не менее чем из восьми этапов и подэтапов. Часть работ выполняется параллельно, "
        "чтобы сократить общий срок. Тема проекта — разработка КИС «ЛифтКонтур», "
        "системы диспетчеризации заявок на техническое обслуживание лифтов. "
        "Календарный план рассчитан методом критического пути; результаты представлены "
        "диаграммой Ганта, сетевым графиком и профилями загрузки ресурсов.",
    )

    add_section_h(doc, "1 Краткое описание проектируемой информационной системы")
    add_body(
        doc,
        "Проектируемая корпоративная информационная система «ЛифтКонтур» предназначена "
        "для предприятия технического обслуживания лифтового хозяйства в жилом фонде "
        "Санкт-Петербурга. Система закрывает контур от поступления заявки жителя или "
        "диспетчера управляющей компании до закрытия наряда механиком и отражения "
        "расхода запасных частей. Клиентская часть включает веб-АРМ диспетчера и "
        "мобильное приложение выездного механика; серверная — сервисы заявок, договоров "
        "с ТСЖ и УК, склада ЗИП и обмена с «1С:Бухгалтерия».",
    )
    add_body(
        doc,
        "В отличие от «коробочного» складского учёта система опирается на сервисный "
        "цикл: нормативы SLA по авариям и плановым осмотрам, маршрутизация бригад по "
        "адресам домов, фотофиксация выполненных работ, резервирование узлов кабины и "
        "лебедки. Планирование разработки выполняется с 6 октября 2026 года по "
        f"{P.fmt(P.TASKS['4.4'].ef)} с учётом выходных, Дня народного единства "
        "(4 ноября 2026 года) и новогодних каникул 1–8 января 2027 года. Критический "
        "путь и параллельные ветви рассчитаны явно, чтобы показать сокращение срока "
        "относительно полностью последовательного выполнения всех работ.",
    )

    add_section_h(doc, "2 Описание используемого программного пакета")
    add_body(
        doc,
        "Для планирования использован метод диаграммы Ганта совместно с расчётом "
        "критического пути (CPM). Календарь проекта учитывает рабочие дни, выходные "
        "и праздники. На основе иерархии работ (WBS), длительностей и связей FS "
        "построены сетевой график, диаграмма Ганта и профили загрузки исполнителей. "
        "Иллюстрации включены в отчёт [2].",
    )

    add_sub_h(doc, "2.1 Производитель и права использования")
    add_body(
        doc,
        "Диаграмма Ганта — общепринятый способ календарного представления работ "
        "проекта, предложенный Г. Л. Гантом в начале XX века. Метод критического "
        "пути применяется для определения работ без резерва времени. Оба приёма "
        "являются открытыми методиками проектного управления и не требуют "
        "коммерческой лицензии на сам способ расчёта [2].",
    )
    add_body(
        doc,
        "В лабораторной работе лицензионный настольный планировщик не использовался: "
        "расписание и рисунки подготовлены по рассчитанным датам и включены в отчёт "
        "как графические материалы.",
    )

    add_sub_h(doc, "2.2 Назначение и основные функции")
    add_body(
        doc,
        "Средства планирования в работе закрывают следующие задачи:",
    )
    add_body(
        doc,
        "— производственный календарь: рабочие дни понедельник–пятница, исключение "
        "выходных и праздников Российской Федерации;",
    )
    add_body(
        doc,
        "— иерархический список работ (WBS) с этапами, подзадачами и вехами;",
    )
    add_body(
        doc,
        "— сетевой график «работа на вершине» (AON) и связи типа FS;",
    )
    add_body(
        doc,
        "— диаграмма Ганта с выделением критического пути;",
    )
    add_body(
        doc,
        "— назначение исполнителей, ставки и оценка затрат по работам;",
    )
    add_body(
        doc,
        "— контроль загрузки ресурсов по рабочим дням.",
    )
    add_body(
        doc,
        "Расчёт критического пути выполнен прямым и обратным проходом. Прямой проход "
        "задаёт ранние начала и окончания (ES, EF), обратный — поздние (LS, LF); "
        "полный резерв равен числу рабочих дней между ES и LS. Работы с нулевым "
        "резервом образуют критический путь. Связи приняты типа FS без лагов.",
    )

    add_sub_h(doc, "2.3 Порядок работы в лабораторной")
    add_body(
        doc,
        "Сначала зафиксированы тема продукта и календарь: рабочие дни — "
        "понедельник–пятница, плюс праздники, попадающие на горизонт плана. Затем "
        "составлена WBS из четырёх этапов верхнего уровня и семнадцати листьев и вех. "
        "Для листьев заданы длительности в рабочих днях, предшественники и доли "
        "загрузки исполнителей. По этим данным рассчитаны даты, резервы и стоимость, "
        "после чего построены диаграммы.",
    )
    add_body(
        doc,
        "Чтобы сократить срок, заранее заложены три зоны параллелизма: обследование "
        "диспетчерской и разбор Excel-журналов; проектирование модели данных, "
        "интерфейсов и обмена с 1С; кодирование сервера, веб-АРМ и мобильного клиента. "
        "Интеграционный контур 3.4 намеренно стоит после готовности API, иначе обмен с "
        "бухгалтерией пришлось бы писать «в пустоту». Пилот вынесен на жилые дома "
        "Выборгского района и пересекает новогодние каникулы — это видно на шкале Ганта "
        "как серый разрыв внутри полосы работы 4.2.",
    )

    add_section_h(doc, "3 Список этапов и подэтапов проекта")
    add_body(
        doc,
        "Перед детальным календарём выделены четыре задачи верхнего уровня. Подзадачи "
        "вложены в этапы. Вехи «Утверждение ТЗ» и «Сдача проекта заказчику» имеют "
        "нулевую длительность. Сводный перечень с датами после расчёта CPM приведён "
        "в таблице 1.",
    )
    add_caption_table(doc, "Таблица 1 — Этапы и подэтапы разработки КИС «ЛифтКонтур»")
    add_table(
        doc,
        ["Код", "Наименование", "Тип", "Длит.", "Предш.", "Начало", "Окончание"],
        task_rows(),
        widths_cm=[1.2, 4.4, 3.2, 1.8, 1.6, 2.1, 2.2],
        size=9,
    )
    add_body(
        doc,
        f"Календарная длительность принятого плана — {P.planned_duration()} рабочих дней "
        f"(с {P.fmt(P.TASKS['1.1'].es)} по {P.fmt(P.TASKS['4.4'].ef)}). Если выполнять "
        f"все листовые работы строго последовательно, срок составил бы "
        f"{P.sequential_duration()} рабочих дней. Распараллеливание экономит "
        f"{P.sequential_duration() - P.planned_duration()} рабочих дней. Критический путь: "
        "1.1 и 1.2 → 1.3 → 1.4 → 1.5 → 2.1 → 2.2 → 3.1 → 3.4 → 4.1 → 4.2 → 4.3 → 4.4. "
        "Ветви интерфейсов и мобильного клиента (2.3, 3.2, 3.3) имеют резерв 10 рабочих "
        "дней; проектирование обмена с 1С (2.4) — 17 дней, потому что контур 3.4 всё "
        "равно ждёт окончания серверной части.",
    )

    add_sub_h(doc, "3.1 Ресурсы проекта")
    add_body(
        doc,
        "Распределены человеческие ресурсы со ставками в рублях за рабочий день. Денежный "
        "ресурс выражен через фонд оплаты: стоимость работы равна сумме «ставка × "
        "длительность × доля загрузки» по всем назначенным исполнителям. Состав команды "
        "и контакты — в таблице 2.",
    )
    add_caption_table(doc, "Таблица 2 — Сотрудники проекта и дневные ставки")
    people = []
    for rid, r in P.RESOURCES.items():
        people.append([r["fio"], r["role"], f"{r['rate']:,}".replace(",", "\u00a0"), r["email"], r["phone"]])
    add_table(
        doc,
        ["ФИО", "Роль", "Ставка, ₽/день", "E-mail", "Телефон"],
        people,
        widths_cm=[3.0, 3.4, 2.4, 4.4, 3.3],
        size=9,
    )

    add_sub_h(doc, "3.2 Оценка затрат по работам")
    add_body(
        doc,
        "Оценка фонда оплаты по листовым работам приведена в таблице 3. "
        "Вехи стоимости не несут.",
    )
    cost_rows = []
    for code in P.LEAF_CODES:
        t = P.TASKS[code]
        if t.duration == 0:
            continue
        cost_rows.append(
            [
                code,
                t.name,
                allocation_text(t),
                str(t.duration),
                rub(P.task_cost(t)),
            ]
        )
    cost_rows.append(["", "Итого по плану", "", "", rub(P.total_cost())])
    add_caption_table(doc, "Таблица 3 — Оценка затрат по задачам (ставка × дни × загрузка)")
    add_table(
        doc,
        ["Код", "Задача", "Ресурсы", "Дней", "Затраты"],
        cost_rows,
        widths_cm=[1.2, 4.6, 6.2, 1.2, 3.3],
        size=9,
    )
    add_body(
        doc,
        f"Ориентировочный фонд оплаты труда составляет {rub(P.total_cost())} без накладных "
        "расходов, закупки серверов и командировок на пилот. Самые дорогие пакеты — "
        "проектирование интерфейсов и кодирование серверной части: там сочетаются высокая "
        "ставка архитектора или разработчика и длительный календарный интервал. Пилот "
        "на двух домах тоже заметен в смете, потому что одновременно заняты тестировщик, "
        "руководитель проекта и аналитик.",
    )

    add_section_h(doc, "4 Основные этапы управления проектом")
    add_body(
        doc,
        "Ниже приведены материалы планирования: календарь, список задач, структура WBS, "
        "сетевой график, диаграмма Ганта, вехи, состав команды, ставки, загрузка "
        "ресурсов и оценка затрат.",
    )

    add_figure(doc, FIG / "fig01_calendar.png", "Рисунок 1 — Настройка календаря проекта")
    add_body(
        doc,
        "При создании плана заданы название КИС «ЛифтКонтур», дата старта 06.10.2026 и "
        "нерабочие дни. Суббота и воскресенье исключены из длительности работ. Отдельно "
        "закрыты 04.11.2026 и интервал 01.01–08.01.2027. Благодаря этому пилот 4.2, "
        "стартующий 29 декабря, фактически «перепрыгивает» каникулы и заканчивается "
        "15 января.",
    )

    add_figure(doc, FIG / "fig03_task_list.png", "Рисунок 2 — Список задач и подзадач")
    add_body(
        doc,
        "Сформирован иерархический список: четыре этапа и вложенные работы анализа, "
        "проектирования, кодирования, испытаний. Для каждой строки видны тип, "
        "длительность и предшественники — это исходные данные прямого прохода CPM.",
    )

    add_figure(doc, FIG / "fig02_wbs.png", "Рисунок 3 — Иерархия работ (WBS)")
    add_body(
        doc,
        "Подзадачи отнесены к родительским этапам. Знак параллелизма на схеме отмечает "
        "ветки, которые стартуют от общего предшественника и не связаны жёсткой "
        "последовательностью друг с другом.",
    )

    add_figure(doc, FIG / "fig04_network.png", "Рисунок 4 — Связывание задач (сетевой график AON)")
    add_body(
        doc,
        "На сетевом графике задана последовательность «окончание — начало». Красные узлы "
        "лежат на критическом пути. Параллельные тройки 2.2 / 2.3 / 2.4 и 3.1 / 3.2 / 3.3 "
        "видны как ярусы под общим предком. Работа 3.4 имеет двух предшественников "
        "(проект обмена и готовый API), поэтому не может начаться раньше окончания 3.1.",
    )

    add_figure(doc, FIG / "fig05_gantt.png", "Рисунок 5 — Диаграмма Ганта с зависимостями")
    add_body(
        doc,
        "На шкале времени отображены полосы работ, сводные скобки этапов, ромбы вех и "
        "стрелки FS. Выходные показаны вертикальными серыми полосами: из-за них "
        "календарный интервал шире суммы рабочих дней. Критические работы выделены "
        "отдельно от «запасных» веток интерфейса и мобильного клиента.",
    )

    add_figure(doc, FIG / "fig06_milestones.png", "Рисунок 6 — Создание промежуточных точек")
    add_body(
        doc,
        "Контрольные события «Утверждение ТЗ» (27.10.2026) и «Сдача проекта» "
        f"({P.fmt(P.TASKS['4.4'].es)}) имеют нулевую длительность. Дополнительно на ось "
        "вынесены старт проекта, начало проектирования, начало кодирования и старт "
        "испытаний — это удобные точки для совещаний с заказчиком.",
    )

    add_figure(doc, FIG / "fig07_team.png", "Рисунок 7 — Карточки сотрудников")
    add_body(
        doc,
        "В план введены восемь человек. Для каждого заполнены имя, роль, ставка, "
        "электронная почта и телефон. Руководитель проекта Волков С.Н. присутствует "
        "на обследовании, согласовании ТЗ, архитектуре, пилоте и сдаче, но не кодирует "
        "сам — это снижает риск, что «менеджер рисует Гант и пишет API одновременно».",
    )

    add_figure(doc, FIG / "fig08_rates.png", "Рисунок 8 — Ресурсы: дневные ставки")
    add_body(
        doc,
        "Ставки заданы как пользовательская характеристика ресурса. Самая высокая — "
        "у архитектора (8 400 ₽/день), самая низкая — у тестировщика (5 300 ₽/день). "
        "Значения учебные, но внутренне согласованы: разрыв между ролями сохранён, "
        "чтобы смета реагировала на состав назначений, а не была «круглой цифрой».",
    )

    add_figure(doc, FIG / "fig09_resource_load.png", "Рисунок 9 — Занятость ресурсов")
    add_body(
        doc,
        "Тепловая карта показывает процент загрузки по рабочим дням. Лебедева И.А. "
        "занята на 100 % весь анализ; Чернов П.Д. переходит от разбора журналов к "
        "архитектуре и модели данных; Морозов А.В. после проектирования БД полностью "
        "уходит в API. Перегрузки свыше 100 % в плане нет: доли на параллельных работах "
        "2.2 и 2.3 для архитектора сложены как 80 % + 20 %. Беляев Р.И. появляется "
        "только на испытаниях — до кодирования держать тестировщика в простое незачем.",
    )

    add_figure(doc, FIG / "fig10_gantt_resources.png", "Рисунок 10 — Ресурсы на диаграмме Ганта")
    add_body(
        doc,
        "Рядом с полосами работ подписаны фамилии исполнителей. Так проверяется, что "
        "параллельные контуры кодирования действительно назначены разным людям: сервер — "
        "Морозов, веб-АРМ — Кузнецова, мобильный клиент — Григорьев. Если бы все три "
        "полосы висели на одном разработчике, формальный параллелизм плана был бы фикцией.",
    )

    add_figure(
        doc,
        FIG / "fig16_gantt_zoom_resources.png",
        "Рисунок 11 — Ресурсы на фрагменте параллельной разработки",
    )
    add_body(
        doc,
        "Укрупнение фазы разработки показывает, как 3.4 стартует только после 3.1, хотя "
        "3.2 и 3.3 к этому моменту уже закончены и имеют резерв. Именно поэтому "
        "критический путь идёт через API и обмен с 1С, а не через мобильный клиент.",
    )

    add_figure(doc, FIG / "fig12_costs.png", "Рисунок 12 — Стоимость задач")
    add_body(
        doc,
        "Столбцы затрат дублируют таблицу 3 в графической форме. Красным отмечены "
        "работы критического пути: задержка на них одновременно сдвигает сдачу и "
        "удерживает фонд оплаты «вправо» по календарю. Накладные и закупка оборудования "
        "в эту оценку не входят.",
    )

    add_section_h(doc, "5 Развёрнутая диаграмма Ганта")
    add_body(
        doc,
        "На рисунке 13 сведена полная картина плана: этапы, листовые работы, вехи, "
        "исполнители, выходные и праздники. Параллельные участки 1.1∥1.2, 2.2∥2.3∥2.4 и "
        "3.1∥3.2∥3.3 сокращают срок относительно строго последовательного изготовления "
        "продукта. Стрелки зависимостей для читаемости вынесены на рисунки 4 и 5; здесь "
        "акцент на шкале времени и назначениях.",
    )
    add_figure(doc, FIG / "fig13_gantt_full.png", "Рисунок 13 — Развёрнутая диаграмма Ганта", w=16.5)
    add_figure(
        doc,
        FIG / "fig11_parallel.png",
        "Рисунок 14 — Фрагмент Ганта: параллельное проектирование и кодирование",
    )
    add_figure(doc, FIG / "fig15_compare.png", "Рисунок 15 — Сравнение последовательного и параллельного сроков")
    add_body(
        doc,
        f"Численно: последовательный срок {P.sequential_duration()} раб. дн., принятый план "
        f"{P.planned_duration()} раб. дн., экономия "
        f"{P.sequential_duration() - P.planned_duration()} раб. дн. "
        f"({100 * (P.sequential_duration() - P.planned_duration()) / P.sequential_duration():.0f} %). "
        "Календарная сдача — 25 января 2027 года. Если убрать параллелизм кодирования, "
        "критический путь удлинился бы примерно на сумму 3.2 и 3.3 (ещё 24 рабочих дня) "
        "плюс последовательное проектирование UI и обмена с 1С.",
    )

    add_struct(doc, "ЗАКЛЮЧЕНИЕ")
    add_body(
        doc,
        "В ходе лабораторной работы изучены технологии планирования и управления "
        "процессами и ресурсами на примере разработки корпоративной информационной "
        "системы «ЛифтКонтур». Составлен календарный план с иерархией из четырёх этапов "
        "и более восьми подэтапов, настроен производственный календарь с выходными и "
        "праздниками, введены восемь человеческих ресурсов со ставками, выполнено "
        "назначение на работы и оценён фонд оплаты около "
        f"{rub(P.total_cost())}.",
    )
    add_body(
        doc,
        "Удалось распараллелить подзадачи основных этапов: обследование бригад и разбор "
        "журналов заявок; проектирование модели данных, интерфейсов и интеграции с 1С; "
        "кодирование серверной части, веб-АРМ диспетчера и мобильного клиента механика. "
        f"За счёт этого срок сокращён с {P.sequential_duration()} до {P.planned_duration()} "
        "рабочих дней. Критический путь проходит через требования, ТЗ, архитектуру, "
        "модель данных, API, обмен с бухгалтерией, испытания, пилот и сдачу: задержка "
        "на любой из этих работ сдвигает дату 25.01.2027.",
    )
    add_body(
        doc,
        "Цель лабораторной работы — изучение технологий планирования и управления "
        "ресурсами — достигнута. Получены практические навыки построения расписания с "
        "учётом критического пути, трудовых и денежных ресурсов и визуализации плана "
        "на диаграмме Ганта.",
    )

    add_struct(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    sources = [
        "Тематический конспект лекций «Корпоративные информационные системы» / "
        "кафедра № 42 ГУАП. – Санкт-Петербург, 2026. – Текст : непосредственный.",
        "ГОСТ Р 54869–2011. Проектный менеджмент. Требования к управлению проектом. – "
        "Введ. 2012-09-01. – Москва : Стандартинформ, 2011. – Текст : непосредственный.",
        "Задание 1. Изучение, планирование и управление ресурсами / "
        "кафедра № 42 ГУАП. – Санкт-Петербург, 2026. – Текст : непосредственный.",
    ]
    for i, src in enumerate(sources, 1):
        p = doc.add_paragraph()
        format_paragraph(p)
        set_run_font(p.add_run(f"{i}. {src}"))


def build(toc_pages: dict) -> Path:
    shutil.copy(BLANK, OUT)
    doc = Document(str(OUT))
    fill_title(doc)
    setup_body_section(doc)
    build_body(doc, toc_pages)
    doc.save(str(OUT))
    return OUT


def main():
    P.schedule()
    gf.main()
    build({k: "…" for k in SEARCH})
    print(OUT)
    print("planned", P.planned_duration(), "seq", P.sequential_duration(), "cost", round(P.total_cost()))


if __name__ == "__main__":
    main()
