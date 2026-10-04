#!/usr/bin/env python3
"""Расчётная таблица ПР2: формулы Excel, не готовые цифры."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "ПР2_Расчет_рейтинг_поставщиков.xlsx"


def style_range(ws, rows, cols, **kwargs):
    for r in rows:
        for c in cols:
            cell = ws.cell(r, c)
            for k, v in kwargs.items():
                setattr(cell, k, v)


def main():
    wb = Workbook()
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    fill_h = PatternFill("solid", fgColor="D9D9D9")
    font_h = Font(name="Times New Roman", bold=True, size=11)
    font_n = Font(name="Times New Roman", size=11)
    wrap = Alignment(wrap_text=True, vertical="center")

    # --- исходные ---
    ws = wb.active
    ws.title = "Исходные"
    headers = [
        "Параметр",
        "А. Поролон-Север (Вологда)",
        "Б. Теплолист СПб",
        "В. ПеноПласт-НН (Нижний Новгород)",
    ]
    ws.append(headers)
    rows = [
        ["Цена листа ППЭ 30 мм, ₽", 1180, 1410, 1050],
        ["Срок поставки, дни", 10, 2, 18],
        ["Минимальная партия, листов", 160, 40, 400],
        ["Предоплата, %", 40, 0, 100],
        ["Постоплата, дни", 0, 10, 0],
        ["Доставка, ₽/рейс", 29000, 6500, 48000],
        ["Бесплатная доставка от, листов", 0, 80, 0],
        ["Лет на рынке", 5, 12, 4],
        ["Конфликты по браку за год", 0, 1, 0],
        ["Возврат брака, дней", 0, 7, 0],
        ["Лимит брака, % партии", 1.5, 4, 2.5],
        ["Сертификация", "ЕАС", "ЕАС + ISO 9001", "ЕАС"],
        ["Потребность, листов/мес", 480, 480, 480],
    ]
    for row in rows:
        ws.append(row)

    ws["A15"] = "Стоимость партии, ₽"
    ws["B15"] = "=B2*$B$14"
    ws["C15"] = "=C2*$C$14"
    ws["D15"] = "=D2*$D$14"
    ws["A16"] = "Число рейсов в месяц"
    ws["B16"] = 1
    ws["C16"] = 1
    ws["D16"] = 1
    ws["A17"] = "Затраты на доставку, ₽"
    # Б: потребность 480 >= 80 → 0
    ws["B17"] = "=B7*B16"
    ws["C17"] = '=IF(C14>=C8,0,C7*C16)'
    ws["D17"] = "=D7*D16"
    ws["A18"] = "Итого за месяц, ₽"
    ws["B18"] = "=B15+B17"
    ws["C18"] = "=C15+C17"
    ws["D18"] = "=D15+D17"
    ws["A19"] = "Полная цена листа, ₽"
    ws["B19"] = "=B18/$B$14"
    ws["C19"] = "=C18/$C$14"
    ws["D19"] = "=D18/$D$14"

    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 28 if col == 1 else 32
    for r in range(1, 20):
        for c in range(1, 5):
            cell = ws.cell(r, c)
            cell.font = font_h if r == 1 else font_n
            cell.border = thin
            cell.alignment = wrap
            if r == 1:
                cell.fill = fill_h
            if r >= 15 and c >= 2:
                cell.number_format = "0.00"

    ws["A20"] = "JIT: партия Б 40 листов, доставка, ₽"
    ws["C20"] = "=IF(40>=C8,0,C7)"
    ws["A21"] = "JIT: полная цена листа Б, ₽"
    ws["C21"] = "=C2+C20/40"
    for r in (20, 21):
        for c in range(1, 5):
            cell = ws.cell(r, c)
            cell.font = font_n
            cell.border = thin
            cell.alignment = wrap
            if c >= 2:
                cell.number_format = "0.00"

    # --- рейтинг ---
    wr = wb.create_sheet("Рейтинг")
    wr.append(["Критерий", "Вес, %", "А, балл", "Б, балл", "В, балл", "А, вклад", "Б, вклад", "В, вклад", "Пояснение балла"])
    crit = [
        ["Цена с логистикой (ниже — лучше)", 25, 7, 4, 10, "В 1150 ₽/лист, А 1240,42 ₽, Б 1410 ₽"],
        ["Срок поставки", 20, 6, 10, 3, "Б 2 дня, А 10, В 18"],
        ["Условия оплаты", 15, 6, 10, 3, "Б постоплата 10 дней; А 40% предоплата; В 100%"],
        ["Минимальная партия", 10, 7, 10, 4, "Б 40 листов, А 160, В 400 при потребности 480"],
        ["Репутация и надёжность", 10, 6, 8, 5, "Б 12 лет и ISO, 1 конфликт; А 5 лет; В 4 года"],
        ["Возврат брака", 10, 6, 8, 5, "Б 7 дней и до 4%; А по согласованию; В замена в следующей партии"],
        ["Сертификация", 10, 7, 10, 7, "ЕАС у всех, ISO 9001 только у Б"],
    ]
    for i, row in enumerate(crit, start=2):
        wr.append(
            [
                row[0],
                row[1],
                row[2],
                row[3],
                row[4],
                f"=C{i}*B{i}/100",
                f"=D{i}*B{i}/100",
                f"=E{i}*B{i}/100",
                row[5],
            ]
        )
    wr.append(["Итого рейтинг", "=SUM(B2:B8)", "", "", "", "=SUM(F2:F8)", "=SUM(G2:G8)", "=SUM(H2:H8)", ""])
    wr["A11"] = "Победитель"
    wr["B11"] = '=IF(AND(G9>=F9,G9>=H9),"Б. Теплолист СПб",IF(F9>=H9,"А. Поролон-Север","В. ПеноПласт-НН"))'
    wr["A12"] = "Самый дешёвый по полной цене листа"
    wr["B12"] = "В. ПеноПласт-НН (1 150 ₽ против 1 240,42 ₽ и 1 410 ₽)"

    for col, w in enumerate([32, 10, 12, 12, 12, 12, 12, 12, 55], start=1):
        wr.column_dimensions[get_column_letter(col)].width = w
    for r in range(1, 13):
        for c in range(1, 10):
            cell = wr.cell(r, c)
            cell.font = font_h if r in (1, 9) else font_n
            cell.border = thin
            cell.alignment = wrap
            if r == 1:
                cell.fill = fill_h

    # --- чувствительность ---
    ws2 = wb.create_sheet("Чувствительность")
    ws2.append(["Критерий", "Вес, %", "А, балл", "Б, балл", "В, балл", "А, вклад", "Б, вклад", "В, вклад"])
    # цена 40%, остальные * 60/75 = 0.8
    sens = [
        ["Цена с логистикой", 40, 7, 4, 10],
        ["Срок поставки", 16, 6, 10, 3],
        ["Условия оплаты", 12, 6, 10, 3],
        ["Минимальная партия", 8, 7, 10, 4],
        ["Репутация и надёжность", 8, 6, 8, 5],
        ["Возврат брака", 8, 6, 8, 5],
        ["Сертификация", 8, 7, 10, 7],
    ]
    for i, row in enumerate(sens, start=2):
        ws2.append(
            [
                row[0],
                row[1],
                row[2],
                row[3],
                row[4],
                f"=C{i}*B{i}/100",
                f"=D{i}*B{i}/100",
                f"=E{i}*B{i}/100",
            ]
        )
    ws2.append(["Итого", "=SUM(B2:B8)", "", "", "", "=SUM(F2:F8)", "=SUM(G2:G8)", "=SUM(H2:H8)"])
    ws2["A11"] = "Победитель при весе цены 40%"
    ws2["B11"] = '=IF(AND(G9>=F9,G9>=H9),"Б. Теплолист СПб",IF(F9>=H9,"А. Поролон-Север","В. ПеноПласт-НН"))'
    ws2["A12"] = "Выбор изменился?"
    ws2["B12"] = "Нет: Б остаётся первым, разрыв сокращается"

    for col, w in enumerate([32, 10, 12, 12, 12, 12, 12, 12], start=1):
        ws2.column_dimensions[get_column_letter(col)].width = w
    for r in range(1, 13):
        for c in range(1, 9):
            cell = ws2.cell(r, c)
            cell.font = font_h if r in (1, 9) else font_n
            cell.border = thin
            cell.alignment = wrap
            if r == 1:
                cell.fill = fill_h

    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
