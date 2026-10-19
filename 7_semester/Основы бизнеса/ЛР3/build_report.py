#!/usr/bin/env python3
"""Практическая работа № 3. Основы бизнеса: оргструктура ООО «Жар-короб»."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
BLANK = Path(r"C:\Users\skvor\OneDrive\Desktop\GitHub\study-guap\7_semester\_титулы\ТИТУЛ_ОБ_ПР1.docx")
OUT = BASE / "ПР3_Отчет_Буренков_ОБ.docx"
PDF = BASE / "ПР3_Отчет_Буренков_ОБ.pdf"
FIG_ASIS = BASE / "figures" / "org_as_is.png"
FIG_TOBE = BASE / "figures" / "org_to_be.png"

BODY = 12
TBL = 10

TOC_DEF = [
    ("ВВЕДЕНИЕ", "intro", 0),
    ("1 Цель работы", "s1", 0),
    ("2 Действующая организационная структура", "s2", 0),
    ("3 Схема действующей организационной структуры", "s3", 0),
    ("4 Проблемы структуры", "s4", 0),
    ("5 Предлагаемая организационная структура", "s5", 0),
    ("6 Схема предлагаемой организационной структуры", "s6", 0),
    ("7 Расчёт изменения ФОТ", "s7", 0),
    ("8 Матрица RACI закупки листа ППЭ", "s8", 0),
    ("9 Ожидаемый эффект", "s9", 0),
    ("ЗАКЛЮЧЕНИЕ", "conc", 0),
]

SEARCH = {key: title for title, key, _ in TOC_DEF}


def set_run_font(run, size=BODY, bold=False, name="Times New Roman", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def format_paragraph(
    p,
    *,
    first_indent=True,
    align="justify",
    space_after=0,
    space_before=0,
    left_indent=0,
    line_spacing=1.5,
):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]


def add_body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p)
    set_run_font(p.add_run(text))
    return p


def add_struct(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_after=12)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)


def add_section_h(doc, text, *, new_page=True):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=True, align="left", space_before=0, space_after=6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)


def add_caption_table(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=8, space_after=2, line_spacing=1.0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def add_figure(doc, path, caption, w=17.5):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=8)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(w))
    c = doc.add_paragraph()
    format_paragraph(c, first_indent=False, align="center", space_before=6, space_after=8)
    set_run_font(c.add_run(caption))


def add_toc_line(doc, title, page, *, indent=0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=0.75 * indent)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    p.add_run("\t")
    set_run_font(p.add_run(str(page)))


def _set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_shading(cell, fill="D9D9D9"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_table(doc, headers, rows, widths_cm=None, size=TBL):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    usable = Cm(18.4)
    if widths_cm:
        total = sum(widths_cm)
        widths = [int(usable * (w / total)) for w in widths_cm]
    else:
        widths = [int(usable / len(headers))] * len(headers)
    grid = t._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        if i < len(widths):
            child.set("w", str(widths[i]))
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        for cell in row.cells:
            _set_cell_border(cell)
            for i, w in enumerate(widths):
                if cell._tc is row.cells[i]._tc:
                    cell.width = Emu(w)
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        set_cell_shading(cell)
        p = cell.paragraphs[0]
        format_paragraph(p, first_indent=False, align="center", line_spacing=1.0)
        set_run_font(p.add_run(h), bold=True, size=size)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            align = "center" if j > 0 and len(headers) > 4 else "left"
            format_paragraph(p, first_indent=False, align=align, line_spacing=1.0)
            set_run_font(p.add_run(v), size=size, bold=(i == len(rows) - 1 and "Итого" in str(row[0])))
    spacer = doc.add_paragraph()
    format_paragraph(spacer, first_indent=False, line_spacing=1.0, space_after=4)
    return t


def add_page_field(paragraph):
    run = paragraph.add_run()
    for kind, val in [("begin", None), ("instr", " PAGE "), ("end", None)]:
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = val
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)
    set_run_font(run)


def fill_title(doc: Document):
    t_work = doc.tables[1]
    t_work.rows[0].cells[0].paragraphs[0].runs[0].text = "ОТЧЕТ О ПРАКТИЧЕСКОЙ РАБОТЕ № 3"
    t_work.rows[1].cells[0].paragraphs[0].runs[0].text = "Диагностика и реструктуризация ООО «Жар-короб»"


def setup_body_section(doc: Document):
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    title_sec, body_sec = doc.sections[0], doc.sections[1]
    title_sec.page_width, title_sec.page_height = Mm(210), Mm(297)
    title_sec.left_margin, title_sec.right_margin = Mm(30), Mm(10)
    title_sec.top_margin, title_sec.bottom_margin = Mm(20), Mm(20)
    body_sec.page_width, body_sec.page_height = Mm(210), Mm(297)
    body_sec.left_margin = Mm(12.7)
    body_sec.right_margin = Mm(12.7)
    body_sec.top_margin = Mm(12.7)
    body_sec.bottom_margin = Mm(16)
    title_sec.footer.is_linked_to_previous = False
    for p in title_sec.footer.paragraphs:
        p.clear()
    body_sec.footer.is_linked_to_previous = False
    body_sec.header.is_linked_to_previous = False
    sect_pr = body_sec._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "2")
    sect_pr.append(pg)
    fp = body_sec.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)


def _box(ax, x, y, w, h, text, *, dash=False, fill="#ffffff", size=8.0):
    from matplotlib.patches import FancyBboxPatch

    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            facecolor=fill,
            edgecolor="black",
            linewidth=1.05,
            linestyle=(0, (3, 2)) if dash else "-",
        )
    )
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, linespacing=1.15)


def _arrow(ax, x1, y1, x2, y2):
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="-|>", color="black", lw=1.05, mutation_scale=9),
    )


def make_as_is():
    import matplotlib.pyplot as plt

    FIG_ASIS.parent.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(12.4, 8.2), dpi=170)
    ax.set_xlim(0, 12.4)
    ax.set_ylim(0, 8.4)
    ax.axis("off")

    _box(ax, 4.55, 7.35, 3.3, 0.78, "Генеральный директор\nтехнолог, доля 45 %", fill="#f2f2f2", size=8.6)
    gd_x, gd_y = 6.2, 7.35

    boxes = [
        (0.25, 5.45, 2.85, 0.85, "Начальник\nпроизводства"),
        (3.30, 5.45, 2.45, 0.85, "Главный\nбухгалтер"),
        (5.95, 5.45, 3.20, 0.85, "Коммерческий директор\nдоля 35 %"),
        (9.35, 5.45, 2.75, 0.85, "Офис-менеджер"),
    ]
    for x, y, w, h, t in boxes:
        _box(ax, x, y, w, h, t, size=8.0)
        _arrow(ax, gd_x, gd_y, x + w / 2, y + h)

    kids = [
        (0.10, 3.20, 1.70, 0.80, "Раскрой ППЭ\n2 чел."),
        (1.90, 3.20, 1.70, 0.80, "Сборка\n4 чел."),
        (3.70, 3.20, 1.70, 0.80, "Кладовщик\n1 чел."),
        (5.50, 3.20, 1.70, 0.80, "ОТК\n1 чел."),
        (3.45, 1.55, 2.15, 0.75, "Бухгалтер\n1 чел."),
        (5.85, 1.55, 2.35, 0.75, "Менеджер B2B\n1 чел."),
        (8.35, 1.55, 2.45, 0.75, "Мелкий опт\n1 чел."),
    ]
    for x, y, w, h, t in kids:
        fill = "#fff4cc" if "ОТК" in t else "#ffffff"
        _box(ax, x, y, w, h, t, fill=fill, size=7.6)

    np_cx = 0.25 + 2.85 / 2
    _arrow(ax, np_cx, 5.45, 0.95, 4.00)
    _arrow(ax, np_cx, 5.45, 2.75, 4.00)
    _arrow(ax, np_cx, 5.45, 4.55, 4.00)
    _arrow(ax, np_cx, 5.45, 6.35, 4.00)
    _arrow(ax, 4.525, 5.45, 4.525, 2.30)
    _arrow(ax, 7.55, 5.45, 7.02, 2.30)
    _arrow(ax, 7.55, 5.45, 9.57, 2.30)

    _box(
        ax,
        0.25,
        0.20,
        5.60,
        1.05,
        "Жёлтый блок: ОТК подчиняется производству.\nКонтролёр не может остановить приёмку брака.",
        dash=True,
        fill="#ffffff",
        size=7.2,
    )
    _box(
        ax,
        6.15,
        0.20,
        5.95,
        1.05,
        "Штат 16 чел. Логиста нет: рейсы ППЭ возит\nкоммерческий директор между созвонами.",
        dash=True,
        fill="#ffffff",
        size=7.2,
    )
    fig.tight_layout(pad=0.15)
    fig.savefig(FIG_ASIS, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def make_to_be():
    import matplotlib.pyplot as plt

    FIG_TOBE.parent.mkdir(parents=True, exist_ok=True)
    plt.rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(12.6, 8.4), dpi=170)
    ax.set_xlim(0, 12.6)
    ax.set_ylim(0, 8.6)
    ax.axis("off")

    _box(ax, 4.65, 7.55, 3.30, 0.75, "Генеральный директор", fill="#f2f2f2", size=8.8)
    gd_x, gd_y = 6.30, 7.55

    top = [
        (0.15, 5.70, 1.90, 0.80, "ОТК\nпрямо ГД", "#fff4cc"),
        (2.20, 5.70, 2.10, 0.80, "Начальник\nпроизводства", "#ffffff"),
        (4.45, 5.70, 1.90, 0.80, "Главный\nбухгалтер", "#ffffff"),
        (6.50, 5.70, 2.20, 0.80, "Коммерческий\nдиректор", "#ffffff"),
        (8.85, 5.70, 1.85, 0.80, "Логист-\nзакупщик", "#e8f4e8"),
        (10.85, 5.70, 1.55, 0.80, "Офис-\nменеджер", "#ffffff"),
    ]
    for x, y, w, h, t, fill in top:
        _box(ax, x, y, w, h, t, fill=fill, size=7.4)
        _arrow(ax, gd_x, gd_y, x + w / 2, y + h)

    kids = [
        (1.30, 3.45, 1.65, 0.75, "Раскрой\n2 чел."),
        (3.10, 3.45, 1.65, 0.75, "Сборка\n6 чел."),
        (4.90, 3.45, 1.65, 0.75, "Кладовщик\n1 чел."),
        (4.50, 1.80, 1.80, 0.70, "Бухгалтер\n1 чел."),
        (6.80, 3.45, 1.65, 0.75, "B2B\n2 чел."),
        (8.55, 3.45, 1.70, 0.75, "Мелкий опт\n1 чел."),
        (10.35, 3.45, 1.90, 0.75, "Маркетплейсы\n1 чел."),
    ]
    for x, y, w, h, t in kids:
        _box(ax, x, y, w, h, t, size=7.3)

    np_cx = 3.25
    _arrow(ax, np_cx, 5.70, 2.125, 4.20)
    _arrow(ax, np_cx, 5.70, 3.925, 4.20)
    _arrow(ax, np_cx, 5.70, 5.725, 4.20)
    _arrow(ax, 5.40, 5.70, 5.40, 2.50)
    kd_cx = 7.60
    _arrow(ax, kd_cx, 5.70, 7.625, 4.20)
    _arrow(ax, kd_cx, 5.70, 9.40, 4.20)
    _arrow(ax, kd_cx, 5.70, 11.30, 4.20)

    _box(
        ax,
        0.15,
        0.20,
        5.90,
        1.25,
        "Жёлтый: ОТК выведен из производства.\nЗелёный: логист сам закрывает партию ППЭ\nдо 800 000 ₽, гендиректор только в курсе.",
        dash=True,
        fill="#ffffff",
        size=7.1,
    )
    _box(
        ax,
        6.25,
        0.20,
        6.10,
        1.25,
        "Штат 21 чел. (+2 сборщика, логист,\nвторой B2B, маркетплейсы).\nПервичка: кладовщик + офис-менеджер.",
        dash=True,
        fill="#ffffff",
        size=7.1,
    )
    fig.tight_layout(pad=0.15)
    fig.savefig(FIG_TOBE, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def build_body(doc: Document, toc_pages: dict):
    add_struct(doc, "СОДЕРЖАНИЕ", new_page=False)
    for title, key, ind in TOC_DEF:
        add_toc_line(doc, title, toc_pages.get(key, "…"), indent=ind)

    add_struct(doc, "ВВЕДЕНИЕ")
    add_body(
        doc,
        "ООО «Жар-короб» шьёт двухконтурные термокороба на Парнасе. В практических "
        "работах 1 и 2 уже зафиксированы форма общества, трое участников (45 / 35 / 20) "
        "и поставщик листа ППЭ — «Теплолист СПб». Сейчас узкое место не сырьё, а то, "
        "как внутри устроены полномочия: кто утверждает закупку, кто принимает брак "
        "и кто отвечает агрегаторам.",
    )
    add_body(
        doc,
        "Организационная структура здесь — схема подчинённости и распределения функций, "
        "а не «список фамилий в штатном расписании». Линейно-функциональная схема "
        "означает: у каждого сотрудника один прямой начальник (линейная ветка) и есть "
        "функциональные блоки (производство, учёт, продажи). Ниже показаны текущая "
        "схема, шесть конкретных сбоев, новая схема, арифметика фонда оплаты труда "
        "и матрица ответственности за закупку листа ППЭ.",
    )

    add_section_h(doc, "1 Цель работы")
    add_body(
        doc,
        "Цель — найти разрывы в организационной структуре ООО «Жар-короб» и предложить "
        "схему, которая выдерживает рост выручки на 35 % и запуск витрины на маркетплейсах.",
    )
    add_body(doc, "Задачи работы:")
    add_body(doc, "1) описать текущую линейно-функциональную структуру и штат 16 человек;")
    add_body(doc, "2) построить схему действующей организационной структуры;")
    add_body(doc, "3) назвать шесть проблем: дублирование, очередь согласования, конфликт ОТК, дыра в логистике и продажах;")
    add_body(doc, "4) предложить целевую организационную структуру с введением новых должностей и изменением порядка подчинённости;")
    add_body(doc, "5) посчитать фонд оплаты труда до и после, показать разницу по шагам;")
    add_body(
        doc,
        "6) заполнить матрицу RACI для процесса «закупка листа ППЭ 30 мм» и оценить "
        "срок согласования, ответ агрегаторам и нагрузку на директора.",
    )

    add_section_h(doc, "2 Действующая организационная структура")
    add_body(
        doc,
        "Цех 340 м², план 400 коробов в месяц, 480 листов ППЭ, выручка 38,4 млн ₽ в год "
        "(38 400 000 / 12 = 3 200 000 ₽ в месяц). Штат 16 человек. Генеральный директор — "
        "технолог и участник с долей 45 %: он и режет лекала, и подписывает счета. "
        "Коммерческий директор — участник с долей 35 %, ведёт агрегаторов. Третий участник "
        "(20 %) в операционке не сидит, на схеме его нет.",
    )
    add_body(
        doc,
        "Под генеральным директором четыре ветки. Производство: начальник, два раскройщика "
        "листа, четыре сборщика коробов, кладовщик и контролёр ОТК. Учёт: главный бухгалтер "
        "и бухгалтер. Продажи: коммерческий директор, один менеджер B2B (заявки кухонь "
        "агрегаторов) и один менеджер мелкого опта (хорека без договора-рамки). Офис-менеджер "
        "подчиняется директору напрямую и занят пропусками, курьерами и записью на совещания.",
    )
    add_caption_table(doc, "Таблица 1 — Штат и оклады действующей структуры")
    add_table(
        doc,
        ["Должность", "Человек", "Оклад, ₽/мес", "ФОТ должности, ₽/мес"],
        [
            ["Генеральный директор", "1", "110 000", "110 000"],
            ["Начальник производства", "1", "85 000", "85 000"],
            ["Раскройщик ППЭ", "2", "58 000", "116 000"],
            ["Сборщик коробов", "4", "52 000", "208 000"],
            ["Кладовщик", "1", "48 000", "48 000"],
            ["Контролёр ОТК", "1", "55 000", "55 000"],
            ["Главный бухгалтер", "1", "78 000", "78 000"],
            ["Бухгалтер", "1", "48 000", "48 000"],
            ["Коммерческий директор", "1", "90 000", "90 000"],
            ["Менеджер B2B", "1", "50 000", "50 000"],
            ["Менеджер мелкого опта", "1", "45 000", "45 000"],
            ["Офис-менеджер", "1", "40 000", "40 000"],
            ["Итого", "16", "—", "973 000"],
        ],
        widths_cm=[6.4, 2.4, 4.2, 5.4],
    )
    add_body(
        doc,
        "ФОТ должности = оклад × число ставок. Пример: раскрой 58 000 × 2 = 116 000; "
        "сборка 52 000 × 4 = 208 000. Сумма 12 строк: 110 000 + 85 000 + 116 000 + "
        "208 000 + 48 000 + 55 000 + 78 000 + 48 000 + 90 000 + 50 000 + 45 000 + "
        "40 000 = 973 000 ₽ в месяц. Год: 973 000 × 12 = 11 676 000 ₽. Страховые взносы "
        "в эту таблицу не входят: сравниваем окладный фонд, как в задании.",
    )
    add_body(
        doc,
        "Планы на следующий год, под которые схему надо ломать сейчас: выручка +35 % "
        "(38,4 × 1,35 = 51,84 млн ₽), витрина коробов на Ozon и Wildberries, два "
        "дополнительных сборщика под рост сборки, внедрение CRM заявок агрегаторов. "
        "Поставщик листа уже выбран — «Теплолист СПб», срок 2 дня, при работе без "
        "большого склада рейсы чаще одного в месяц.",
    )

    add_section_h(doc, "3 Схема действующей организационной структуры")
    add_body(
        doc,
        "На рисунке 1 — текущая подчинённость. Жёлтый блок ОТК стоит под производством: "
        "это не украшение схемы, а источник конфликта в разделе 4. Логиста на рисунке нет, "
        "потому что должности нет.",
    )
    add_figure(doc, FIG_ASIS, "Рисунок 1 — Действующая организационная структура ООО «Жар-короб»")

    add_section_h(doc, "4 Проблемы структуры")
    add_body(
        doc,
        "1. Очередь на подпись гендиректора. Счёт на лист ППЭ и счёт на ножи раскроя "
        "подписывает только он. Директор по 3–4 дня на линии: счёт от «Теплолист СПб» "
        "лежит в ящике, пока он не вернётся в офис. Срок поставки 2 дня из практической "
        "работы 2 сгорает на согласовании. Это не «бюрократия вообще», а конкретная "
        "задержка до 4 рабочих дней на каждую закупку.",
    )
    add_body(
        doc,
        "2. Конфликт интересов у ОТК. Контролёр подчиняется начальнику производства. "
        "При поступлении партии листа с отклонением по толщине начальник производства "
        "требует принять продукцию во избежание срыва плана сборки 400 коробов. "
        "У контролёра отсутствует возможность эскалации вопроса генеральному директору. "
        "В результате брак либо направляется в сборку, либо скрывается до поступления "
        "следующей партии.",
    )
    add_body(
        doc,
        "3. Один менеджер B2B на всех агрегаторов. Заявка кухни «нужно 40 коробов к пятнице» "
        "висит до двух дней. Слот на неделю уходит другому поставщику коробов. Мелкий опт "
        "при этом живёт отдельно и не подхватывает B2B: разные каналы, разные договоры, "
        "перекинуть человека «на часок» нельзя без потери рамки.",
    )
    add_body(
        doc,
        "4. Бухгалтерия закрывает первичку склада. По оценке главного бухгалтера около "
        "30 % времени бухгалтера уходит на ТТН и УПД по 480 листам и отгрузке коробов. "
        "Это работа кладовщика и офис-менеджера по шаблону, а не работа человека, который "
        "считает УСН. Пока первичка сидит в учёте, отчёт закрывается на 3–4 дня позже.",
    )
    add_body(
        doc,
        "5. Нет логистики как функции. Коммерческий директор между созвонами с агрегаторами "
        "заказывает рейс и ловит водителя. После выбора поставщика Б рейс нужен чаще: "
        "партия 40 листов при «точно в срок» — это не «один раз в месяц между делом». "
        "Продажи и доставка сырья в одних руках дублируют внимание и рвут оба процесса.",
    )
    add_body(
        doc,
        "6. В организационной структуре отсутствует должность, ответственная за торговлю "
        "через маркетплейсы. Запуск витрины предусмотрен планом, однако соответствующая "
        "штатная единица не введена. Коммерческий директор относит размещение карточек "
        "товаров к неопределённому сроку. Ключевые показатели эффективности для менеджера "
        "B2B и менеджера мелкого опта не установлены; фактически применяется неформальный "
        "критерий выполнения месячного плана продаж. При отсутствии должностного лица, "
        "ответственного за канал, витрина на Ozon не будет сформирована, в том числе "
        "после внедрения CRM.",
    )

    add_section_h(doc, "5 Предлагаемая организационная структура")
    add_body(
        doc,
        "Новая схема остаётся линейно-функциональной: один прямой начальник у каждого. "
        "Меняется не «философия управления», а три вещи: кому подчиняется ОТК, кто "
        "имеет право закрыть закупку ППЭ без директора и какие каналы продаж имеют "
        "своего человека.",
    )
    add_body(
        doc,
        "ОТК выводится из производства и подчиняется генеральному директору. Контролёр "
        "может остановить приёмку листа, начальник производства об этом узнаёт, но не "
        "отменяет решение. Логист-закупщик подчиняется директору и закрывает счета "
        "до 800 000 ₽ сам. Месячная закупка у «Теплолист СПб» из практической работы 2 "
        "равна 676 800 ₽, то есть обычный рейс больше не стоит в очереди директора. "
        "Сверх 800 000 ₽ — подпись директора, логист готовит пакет за день, а не за четыре.",
    )
    add_body(
        doc,
        "В продажах появляются второй менеджер B2B и специалист по маркетплейсам. "
        "Оба подчиняются коммерческому директору. Второй B2B закрывает ночные и "
        "пиковые заявки кухонь; маркетплейсы — отдельный канал с комиссией площадки, "
        "его нельзя отдавать «между рамками агрегаторов». На сборке +2 человека: "
        "4 + 2 = 6, это как раз план роста выпуска, а не «на всякий случай».",
    )
    add_body(
        doc,
        "Первичка склада уходит из бухгалтерии: кладовщик заполняет шаблон ТТН, "
        "офис-менеджер проводит скан в папку, бухгалтер только проверяет комплект "
        "раз в день. Должность бухгалтера не сокращаем: при выручке 51,84 млн ₽ "
        "объём документов вырастет, высвобожденные 30 % уйдут на закрытие месяца.",
    )

    add_section_h(doc, "6 Схема предлагаемой организационной структуры")
    add_body(
        doc,
        "На рисунке 2 жёлтый блок ОТК стоит на одном уровне с производством, зелёный "
        "блок логиста — новая должность. Штат 21 человек: 16 прежних плюс два сборщика, "
        "логист, второй B2B и специалист маркетплейсов. 16 + 5 = 21.",
    )
    add_figure(doc, FIG_TOBE, "Рисунок 2 — Предлагаемая организационная структура ООО «Жар-короб»")

    add_section_h(doc, "7 Расчёт изменения ФОТ")
    add_body(
        doc,
        "Новые оклады берём в той же сетке, что таблица 1: сборщик 52 000, логист 60 000 "
        "(между начальником производства и кладовщиком: возит деньги и сроки, но не "
        "командует цехом), маркетплейсы 55 000, второй B2B 50 000 как первый.",
    )
    add_caption_table(doc, "Таблица 2 — Прирост фонда оплаты труда в предлагаемой структуре")
    add_table(
        doc,
        ["Должность", "Человек", "Оклад, ₽/мес", "Прирост ФОТ, ₽/мес"],
        [
            ["Сборщик коробов (доп.)", "2", "52 000", "104 000"],
            ["Логист-закупщик", "1", "60 000", "60 000"],
            ["Специалист маркетплейсов", "1", "55 000", "55 000"],
            ["Менеджер B2B (второй)", "1", "50 000", "50 000"],
            ["Итого прирост", "5", "—", "269 000"],
        ],
        widths_cm=[6.4, 2.4, 4.2, 5.4],
    )
    add_body(
        doc,
        "Прирост: 52 000 × 2 = 104 000; 104 000 + 60 000 = 164 000; 164 000 + 55 000 = "
        "219 000; 219 000 + 50 000 = 269 000 ₽ в месяц. Новый ФОТ: 973 000 + 269 000 = "
        "1 242 000 ₽ в месяц. Год по предлагаемой структуре: 1 242 000 × 12 = 14 904 000 ₽. Разница "
        "за год: 14 904 000 − 11 676 000 = 3 228 000 ₽, или 269 000 × 12 = 3 228 000 ₽ "
        "(проверка тем же числом).",
    )
    add_body(
        doc,
        "Относительный рост ФОТ: 269 000 / 973 000 ≈ 0,276, то есть +27,6 %. План "
        "выручки +35 %: 38,4 × 0,35 = 13,44 млн ₽ доп. выручки. Доп. ФОТ 3,228 млн ₽ "
        "съедает 3,228 / 13,44 ≈ 0,24, то есть около 24 % прироста выручки. Остаётся "
        "место под сырьё, комиссию маркетплейса и прибыль. Если прирост выручки не "
        "выйдет, реструктуризация всё равно устраняет очередь согласования и конфликт ОТК: "
        "введение указанных должностей направлено на устранение выявленных разрывов, "
        "а не на увеличение численности как самоцель.",
    )
    add_caption_table(doc, "Таблица 3 — Сводка ФОТ")
    add_table(
        doc,
        ["Показатель", "Действующая", "Предлагаемая", "Разница"],
        [
            ["Численность, чел.", "16", "21", "+5"],
            ["ФОТ в месяц, ₽", "973 000", "1 242 000", "+269 000"],
            ["ФОТ в год, ₽", "11 676 000", "14 904 000", "+3 228 000"],
            ["ФОТ к выручке", "11,676 / 38,4 ≈ 30,4 %", "14,904 / 51,84 ≈ 28,7 %", "доля падает"],
        ],
        widths_cm=[4.2, 5.0, 5.4, 3.8],
    )
    add_body(
        doc,
        "Доля ФОТ в выручке сейчас: 11 676 000 / 38 400 000 = 0,304, около 30,4 %. "
        "При плане: 14 904 000 / 51 840 000 = 0,287, около 28,7 %. Фонд растёт в рублях "
        "и не растёт как доля, если 35 % выручки выполняется.",
    )

    add_section_h(doc, "8 Матрица RACI закупки листа ППЭ")
    add_body(
        doc,
        "RACI — четыре роли на каждом шаге процесса. R (responsible) — кто делает работу. "
        "A (accountable) — кто отвечает за результат; на шаг назначается один A, иначе "
        "непонятно, с кого спрашивать. C (consulted) — с кем советуются до действия. "
        "I (informed) — кого ставят в известность после. Матрица ниже — для новой структуры "
        "и для процесса «закупка листа ППЭ 30 мм», не для абстрактного «снабжения».",
    )
    add_caption_table(doc, "Таблица 4 — RACI процесса «закупка листа ППЭ 30 мм»")
    add_table(
        doc,
        ["Шаг", "ГД", "НП", "Логист", "Кладовщик", "ОТК", "Главбух"],
        [
            ["1. Расчёт потребности 480 листов", "I", "R / A", "C", "C", "I", "I"],
            ["2. Заявка поставщику Б", "I", "C", "R / A", "I", "I", "I"],
            ["3. Проверка счёта до 800 000 ₽", "I", "C", "R / A", "I", "I", "C"],
            ["4. Оплата", "I", "I", "C", "I", "I", "R / A"],
            ["5. Приёмка количества", "I", "I", "C", "R / A", "C", "I"],
            ["6. Входной контроль толщины и брака", "I", "I", "I", "C", "R / A", "I"],
            ["7. Оприходование на склад", "I", "A", "I", "R", "I", "C"],
            ["8. Комплект закрывающих документов", "I", "I", "C", "C", "I", "R / A"],
        ],
        widths_cm=[5.2, 2.0, 2.0, 2.4, 2.6, 2.0, 2.2],
        size=9,
    )
    add_body(
        doc,
        "ГД — генеральный директор, НП — начальник производства. На шаге 1 потребность "
        "считает сам НП: 400 коробов × 1,2 листа = 480, поэтому он и R, и A; логист "
        "только уточняет кратность партии поставщика. На шаге 2 логист и "
        "делает заявку, и отвечает за неё: директор больше не A. На шаге 4 деньги "
        "проводкой отпускает главбух, логист только подтверждает, что счёт тот. "
        "На шаге 6 A — ОТК, не начальник производства: принять тонкий лист «чтобы линия "
        "не встала» теперь нельзя без контролёра. Если счёт выше 800 000 ₽, на шаге 3 "
        "A переходит к гендиректору, логист остаётся R: готовит пакет в тот же день.",
    )

    add_section_h(doc, "9 Ожидаемый эффект")
    add_body(
        doc,
        "Ожидаемый эффект оценивается по трём количественным показателям, соответствующим "
        "проблемам, изложенным в разделе 4, а также по снижению нагрузки на генерального "
        "директора.",
    )
    add_caption_table(doc, "Таблица 5 — Ожидаемый эффект реструктуризации")
    add_table(
        doc,
        ["Показатель", "Действующее значение", "Целевое значение", "За счёт чего"],
        [
            ["Согласование закупки ППЭ", "до 4 дней", "в день счёта", "лимит логиста 800 000 ₽"],
            ["Ответ агрегаторам", "до 2 дней", "в день заявки", "второй B2B и CRM"],
            ["Часы ГД на закупки в неделю", "около 8", "около 1", "функция у логиста"],
            ["Доля времени бухгалтера на ТТН", "около 30 %", "около 10 %", "кладовщик и офис"],
            ["Владелец витрины маркетплейсов", "нет", "есть, 1 ставка", "новая должность"],
        ],
        widths_cm=[5.0, 3.2, 3.4, 6.8],
    )
    add_body(
        doc,
        "Срок согласования: 4 дня очереди директора минус полномочие логиста. Для партии "
        "676 800 ₽ директор больше не нужен как подпись. Ответ агрегаторам: два менеджера "
        "делят кухни по сменам, CRM не даёт заявке лежать без ответственного. Восемь часов "
        "директора в неделю на закупки — это примерно один рабочий день; после передачи "
        "остаётся час на счета выше лимита и на конфликт по браку, который эскалирует ОТК.",
    )
    add_body(
        doc,
        "Продажи: второй B2B не увеличивает выручку сам по себе, он снимает потерянные "
        "слоты. Маркетплейс — новый канал, в 35 % роста он заложен как часть плана, "
        "а не как «гарантированные 13,44 млн ₽ с карточки». Если витрина запустится "
        "позже, рост держат агрегаторы и мелкий опт; структура всё равно закрывает "
        "очередь и ОТК.",
    )

    add_struct(doc, "ЗАКЛЮЧЕНИЕ")
    add_body(
        doc,
        "Текущая линейно-функциональная схема на 16 человек держит сборку 400 коробов, "
        "но ломается на закупке ППЭ, приёмке брака и заявках агрегаторов. Шесть разрывов "
        "сводятся к одному: полномочия сидят у людей, у которых нет на них времени или "
        "есть прямой интерес обойти контроль.",
    )
    add_body(
        doc,
        "Новая схема на 21 человека выводит ОТК к директору, отдаёт закупку логисту "
        "с лимитом 800 000 ₽, добавляет второго B2B, маркетплейсы и двух сборщиков. "
        "ФОТ растёт с 973 000 до 1 242 000 ₽ в месяц (+27,6 %) при плане выручки +35 %. "
        "Матрица RACI фиксирует, что заявка и входной контроль больше не висят на "
        "гендиректоре и начальнике производства. Согласование партии ППЭ сжимается "
        "с четырёх дней до дня счёта.",
    )


def build(toc_pages: dict) -> Path:
    make_as_is()
    make_to_be()
    shutil.copy(BLANK, OUT)
    doc = Document(str(OUT))
    fill_title(doc)
    setup_body_section(doc)
    build_body(doc, toc_pages)
    doc.save(str(OUT))
    return OUT


def convert():
    src = str(OUT.resolve()).replace("'", "''")
    dst = str(PDF.resolve()).replace("'", "''")
    ps = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
foreach ($d in @($word.Documents)) {{ $d.Close([ref]0) }}
$doc = $word.Documents.Open('{src}')
$wdFormatPDF = 17
$doc.SaveAs([ref] '{dst}', [ref]$wdFormatPDF)
$doc.Close([ref]0)
$word.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($word) | Out-Null
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


def measure() -> dict[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(str(PDF))
    pages: dict[str, int] = {}
    for i, page in enumerate(reader.pages):
        if i < 2:
            continue
        text = page.extract_text() or ""
        printed = i + 1
        for key, needle in SEARCH.items():
            if key not in pages and needle in text:
                pages[key] = printed
    return pages


def main():
    build({k: "…" for k in SEARCH})
    try:
        convert()
        pages = measure()
        print("pass1", pages)
        from pypdf import PdfReader

        print("PDF pages", len(PdfReader(str(PDF)).pages))
        if pages:
            build(pages)
            convert()
            print("pass2", measure())
            print("PDF pages", len(PdfReader(str(PDF)).pages))
    except Exception as exc:
        print("PDF skipped:", exc)
    print(OUT)


if __name__ == "__main__":
    main()
