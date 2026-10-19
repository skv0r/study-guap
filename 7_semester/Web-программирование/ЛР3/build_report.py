# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР3. В текст отчёта не входит."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"C:\Users\skvor\OneDrive\Desktop\GitHub\study-guap\7_semester\Web-программирование")
LAB = ROOT / "ЛР3"
SRC_DOC = ROOT / "ЛР1_Отчет_Буренков_PairWise.docx"
OUT = LAB / "ЛР3_Отчет_Буренков_PairWise.docx"
APP = Path(r"C:\Users\skvor\OneDrive\Desktop\GitHub\pairwise\app")
SHOTS = Path(r"C:\Users\skvor\AppData\Local\Temp\cursor\screenshots")
FIG = LAB / "fig"
BLACK = RGBColor(0, 0, 0)


def rfonts(run, name):
    run.font.name = name
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def style_run(run, size=14, bold=False, font="Times New Roman"):
    rfonts(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def set_spacing(paragraph, before=0, after=0, line=360):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line / 240.0


def set_indent(paragraph, first=None, left=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))
    if left is not None:
        ind.set(qn("w:left"), str(left))


def add_text(doc, text, *, first=709, before=0, after=0, line=360, center=False, bold=False, size=14, font="Times New Roman"):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before, after, line)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first:
            set_indent(paragraph, first=first)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, font=font)
    return paragraph


def add_heading(doc, text, *, center=False, break_before=False, bookmark=None, bid=None):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=12, line=360)
    paragraph.paragraph_format.keep_with_next = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph, first=709)
    if break_before:
        run_br = paragraph.add_run()
        style_run(run_br, bold=True)
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_br._r.append(br)
    run = paragraph.add_run(text)
    style_run(run, bold=True)
    if bookmark is not None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bid))
        start.set(qn("w:name"), bookmark)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bid))
        p_pr = paragraph._p.find(qn("w:pPr"))
        if p_pr is not None:
            p_pr.addnext(start)
        else:
            paragraph._p.insert(0, start)
        paragraph._p.append(end)
    return paragraph


def field_run(paragraph, kind, payload=None):
    run = paragraph.add_run()
    style_run(run, bold=False)
    for child in list(run._r):
        if child.tag == qn("w:t"):
            run._r.remove(child)
    if kind == "text":
        node = OxmlElement("w:t")
        node.text = payload
        run._r.append(node)
        return
    if kind == "instr":
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = payload
    else:
        node = OxmlElement("w:fldChar")
        node.set(qn("w:fldCharType"), kind)
    run._r.append(node)


def add_toc(doc, title, bookmark):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, line=360)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    ppr.append(tabs)
    run = paragraph.add_run(title)
    style_run(run, bold=False)
    paragraph.add_run().add_tab()
    field_run(paragraph, "begin")
    field_run(paragraph, "instr", f" PAGEREF {bookmark} \\h ")
    field_run(paragraph, "separate")
    field_run(paragraph, "text", "1")
    field_run(paragraph, "end")
    return paragraph


def shade_cell(cell):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcpr.append(shd)


def set_cell(cell, text, *, bold=False, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_spacing(paragraph, before=2, after=2, line=360)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    style_run(run, size=14, bold=bold)
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    if header:
        shade_cell(cell)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    total = sum(widths)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc = cell._tc
            tcpr = tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)
    jc = tblpr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblpr.append(jc)
    jc.set(qn("w:val"), "left")


def add_table(doc, caption, headers, rows, widths):
    content = 9355
    inset = max(0, (content - sum(widths)) // 2)
    add_text(
        doc,
        caption,
        first=0,
        before=6,
        after=6,
        line=240,
        center=False,
        bold=False,
        size=14,
    )
    # caption is the last paragraph
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = True
    set_indent(cap, first=0, left=inset)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, header=True)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], value)
    return table


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=6, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_text(doc, caption, first=0, before=6, after=10, line=360, center=True)


def add_listing(doc, caption, path):
    add_text(doc, caption, first=0, before=10, after=2, line=360, center=False)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    lines = path.read_text(encoding="utf-8").replace("\t", "    ").splitlines()
    for line in lines:
        paragraph = doc.add_paragraph()
        set_spacing(paragraph, line=240)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(line if line else " ")
        style_run(run, size=10, bold=False, font="Courier New")
        if line.startswith(" ") or line.endswith(" "):
            run._element.find(qn("w:t")).set(qn("xml:space"), "preserve")


def make_tree():
    FIG.mkdir(parents=True, exist_ok=True)
    lines = [
        "app",
        "├─ layout.tsx",
        "├─ gallery/page.tsx",
        "├─ gallery/gallery.module.css",
        "├─ gallery/_components/PlacesGallery.tsx",
        "├─ gallery/_components/PlacesGallery.module.css",
        "├─ gallery/_components/PhotoModal.tsx",
        "├─ gallery/_components/PhotoModal.module.css",
        "├─ gallery/_components/PhotoMap.tsx",
        "├─ gallery/_components/PhotoMap.module.css",
        "├─ ui/Header/Header.tsx",
        "├─ ui/BurgerMenu/BurgerMenu.tsx",
        "├─ ui/BurgerMenu/BurgerMenu.module.css",
        "├─ data/spots.ts",
        "├─ providers/DemoProvider.tsx",
        "└─ providers/useDemoAuth.ts",
        "public/images/spots",
        "├─ tochka-kipeniya.jpg",
        "├─ growup-ligovsky.jpg",
        "├─ bench-petrogradka.jpg",
        "├─ sevkabel.jpg",
        "├─ itmo-cowork.jpg",
        "└─ praktik-vo.jpg",
    ]
    font = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 28)
    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    width = max(probe.textlength(line, font=font) for line in lines) + 80
    height = 36 * len(lines) + 48
    image = Image.new("RGB", (int(width), height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width - 1, height - 1), outline=(180, 180, 180))
    y = 24
    for line in lines:
        draw.text((32, y), line, font=font, fill=(0, 0, 0))
        y += 36
    dest = FIG / "tree.png"
    image.save(dest)
    return dest


def replace_title(doc):
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 3",
        "Создание проекта Next.js с использованием React, TypeScript и App Router":
            "Реализация клиентской части веб-приложения: фотографии, модальные окна, бургер-меню, состояние пользователя и карта",
    }
    found = set()
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
            found.add(node.text)
    expected = set(replacements.values())
    if found != expected:
        raise SystemExit(f"Титул заменён не полностью: {found}")


def update_page_fields(path):
    import subprocess
    import tempfile

    script = Path(tempfile.gettempdir()) / "lr3_web_pages.ps1"
    target = str(path).replace("'", "''")
    script.write_text(
        "\n".join(
            [
                "$ErrorActionPreference = 'Stop'",
                "$word = New-Object -ComObject Word.Application",
                "$word.Visible = $false",
                "$word.DisplayAlerts = 0",
                f"$doc = $word.Documents.Open('{target}')",
                "$doc.Repaginate()",
                "$doc.Fields.Update() | Out-Null",
                "$doc.Repaginate()",
                "$doc.Save()",
                "$doc.Close()",
                "$word.Quit()",
            ]
        ),
        encoding="utf-8-sig",
    )
    subprocess.run(["powershell", "-NoProfile", "-File", str(script)], check=True)


def build():
    tree = make_tree()
    pictures = {
        "gallery": SHOTS / "lr3-gallery.png",
        "modal": SHOTS / "lr3-modal-guest.png",
        "burger": SHOTS / "lr3-burger.png",
        "guest": SHOTS / "lr3-modal-guest-action.png",
        "user": SHOTS / "lr3-modal-user.png",
        "map": SHOTS / "lr3-map-pins.png",
        "nokey": SHOTS / "lr3-map-nokey.png",
    }
    for path in pictures.values():
        if not path.exists():
            raise SystemExit(f"Нет снимка: {path}")

    doc = Document(str(SRC_DOC))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replace_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ", center=True, bookmark="toc", bid=20)
    toc = [
        ("1 Цель работы", "h1"),
        ("2 Клиентский сценарий", "h2"),
        ("3 Структура приложения", "h3"),
        ("4 Сущность и фотографии", "h4"),
        ("5 Галерея и модальное окно", "h5"),
        ("5.1 Состояние и хуки", "h51"),
        ("5.2 Доступность диалога", "h52"),
        ("6 Бургер-меню", "h6"),
        ("7 Карта", "h7"),
        ("8 Результаты проверки", "h8"),
        ("9 Ответы на контрольные вопросы", "h9"),
        ("ЗАКЛЮЧЕНИЕ", "h10"),
        ("ПРИЛОЖЕНИЕ А Листинги исходного кода", "h11"),
    ]
    for title, bookmark in toc:
        add_toc(doc, title, bookmark)

    add_heading(doc, "1 Цель работы", break_before=True, bookmark="h1", bid=21)
    add_text(
        doc,
        "Цель работы — добавить в PairWise клиентскую часть: фотографии площадок, галерею с модальным окном, бургер-меню, действие избранного для учебного пользователя и карту с метками через JavaScript API Яндекс Карт.",
    )

    add_heading(doc, "2 Клиентский сценарий", break_before=True, bookmark="h2", bid=22)
    add_text(
        doc,
        "Пользователь открывает раздел «Галерея и карта». На странице шесть карточек коворкингов Санкт-Петербурга. Щелчок по карточке открывает окно с фотографией, описанием и адресом. Гость в этом окне видит ссылку «Войдите, чтобы сохранить площадку». После учебного входа на том же месте появляется кнопка «Добавить в избранное». Повторное нажатие меняет подпись на «Убрать из избранного». Под сеткой написано, что избранное хранится только до обновления страницы.",
    )
    add_text(
        doc,
        "Шапка открывает бургер-меню. В нём пункты «Главная», «Каталог», «Галерея и карта», «О проекте» и «Кабинет». Текущий раздел подсвечен. У гостя внизу панели ссылка «Войти». У учебного пользователя там имя и кнопка «Выйти». Ссылка «Вход» в правой части шапки остаётся на всех экранах и ведёт на /auth/login.",
    )
    add_text(
        doc,
        "Под карточками расположен блок карты. Координаты площадок лежат в массиве данных. С ключом NEXT_PUBLIC_YANDEX_MAPS_API_KEY компонент создаёт карту Санкт-Петербурга и метку на каждую из шести площадок. Щелчок по метке открывает то же модальное окно, что и карточка: с метки открылась «ITMO Cowork». Если ключа в окружении нет, скрипт карт не запрашивается и на странице остаётся сообщение, что карта не подключена.",
    )

    add_heading(doc, "3 Структура приложения", break_before=True, bookmark="h3", bid=23)
    add_text(
        doc,
        "Страница галереи живёт в app/gallery/page.tsx и открывается по адресу /gallery. Директивы «use client» в этом файле нет. Интерактивность вынесена в PlacesGallery, PhotoModal, PhotoMap и BurgerMenu. Шапка Header подключает бургер. Учебный пользователь по-прежнему хранится в DemoAuthProvider. Назначение файлов приведено в таблице 1, состав каталогов показан на рисунке 1.",
    )
    add_table(
        doc,
        "Таблица 1 — Новые и изменённые части клиента",
        ["Назначение", "Файл", "Где работает"],
        [
            ["Страница галереи", "app/gallery/page.tsx", "/gallery"],
            ["Сетка карточек и избранное", "app/gallery/_components/PlacesGallery.tsx", "клиент"],
            ["Модальное окно", "app/gallery/_components/PhotoModal.tsx", "document.body"],
            ["Карта", "app/gallery/_components/PhotoMap.tsx", "блок под сеткой"],
            ["Бургер-меню", "app/ui/BurgerMenu/BurgerMenu.tsx", "шапка"],
            ["Площадки, фото, координаты", "app/data/spots.ts", "каталог и галерея"],
            ["Файлы изображений", "public/images/spots", "локальные jpg"],
        ],
        [3200, 4000, 2155],
    )
    add_picture(doc, tree, "Рисунок 1 — Файлы клиентской части и фотографии площадок", 14.5)

    add_heading(doc, "4 Сущность и фотографии", break_before=True, bookmark="h4", bid=24)
    add_text(
        doc,
        "Предметная сущность — площадка. К прежним полям добавлены путь к изображению, альтернативный текст, широта и долгота. Карта читает latitude и longitude. Поле coords сохранено: x совпадает с долготой, y — с широтой. В массиве шесть площадок с номерами от 0 до 5. Поля сведены в таблице 2.",
    )
    add_table(
        doc,
        "Таблица 2 — Свойства площадки",
        ["Свойство", "Тип", "Назначение"],
        [
            ["id", "number", "Номер площадки"],
            ["name", "string", "Название"],
            ["description", "string", "Краткое описание"],
            ["city", "string", "Город"],
            ["adress", "string", "Адрес"],
            ["stack", "string[]", "Стек людей на площадке"],
            ["image", "string", "Путь к файлу в public"],
            ["imageAlt", "string", "Текст, если фото не видно"],
            ["latitude", "number", "Широта для метки"],
            ["longitude", "number", "Долгота для метки"],
            ["coords", "{x, y}", "Та же точка: x — долгота, y — широта"],
            ["createdAt", "Date", "Дата появления в базе"],
        ],
        [2200, 2200, 4955],
    )
    add_text(
        doc,
        "Фотографии лежат в public/images/spots и уже использовались в каталоге. Компонент Image из next/image выводит их с fill и sizes. Альтернативный текст описывает кадр. Файлы, кадры и источники сведены в таблице 3.",
        before=8,
    )
    add_table(
        doc,
        "Таблица 3 — Локальные фотографии и их источники",
        ["Файл", "Что на снимке", "Источник"],
        [
            ["tochka-kipeniya.jpg", "Зал с паркетом, ряды стульев, спикер у сцены", "Иллюстрация каталога. Автор и лицензия в файле не записаны"],
            ["growup-ligovsky.jpg", "Длинные столы у окон и люди за мониторами", "Иллюстрация каталога. Автор и лицензия в файле не записаны"],
            ["bench-petrogradka.jpg", "Двое за ноутбуком у книжных полок", "Иллюстрация каталога. Автор и лицензия в файле не записаны"],
            ["sevkabel.jpg", "Фасад порта: мурал, лестница, столики на площади", "Wikimedia Commons, Sevkabel Port 01.jpg, автор MBH, CC BY-SA 4.0"],
            ["itmo-cowork.jpg", "Интерьер зала: столы, ноутбуки, доска", "Иллюстрация интерьера для карточки ITMO Cowork. Автор и лицензия в файле не записаны"],
            ["praktik-vo.jpg", "Переговорная на двоих и стеклянная стена", "Иллюстрация каталога. Автор и лицензия в файле не записаны"],
        ],
        [2500, 3100, 3755],
    )
    add_text(
        doc,
        "Файл sevkabel.jpg — копия снимка Sevkabel Port 01.jpg с Wikimedia Commons в размере 1280×960. Автор — MBH, дата съёмки 30 апреля 2024 года, лицензия CC BY-SA 4.0, страница файла https://commons.wikimedia.org/wiki/File:Sevkabel_Port_01.jpg. У остальных пяти файлов в каталоге нет записи об авторе и лицензии.",
        before=8,
    )

    add_heading(doc, "5 Галерея и модальное окно", break_before=True, bookmark="h5", bid=25)
    add_text(
        doc,
        "PlacesGallery перебирает массив spots и рисует кнопку на каждую площадку. Внутри кнопки стоит next/image. Кнопка, а не div с обработчиком, получает фокус и срабатывает с клавиатуры. Выбранная площадка записывается в selected. Пока selected равен null, окна нет. Когда в selected лежит объект, рядом монтируется PhotoModal.",
    )
    add_text(
        doc,
        "Окно выводится через createPortal во второй аргумент document.body. В дереве React модальное окно остаётся потомком галереи, а в DOM узел оказывается прямым потомком body и перекрывает страницу. Контекст DemoAuthProvider при этом доступен: портал меняет место узла в документе и не вырезает компонент из дерева React, поэтому useDemoAuth внутри окна читает того же пользователя.",
    )

    add_heading(doc, "5.1 Состояние и хуки", bookmark="h51", bid=26)
    add_text(
        doc,
        "useState возвращает пару: текущее значение и функцию, которая его меняет. После вызова этой функции React рисует компонент заново. В галерее три таких значения. selected хранит площадку или null. favoriteIds хранит Set со строковыми номерами выбранных площадок. returnFocusTo хранит кнопку карточки, с которой открыли окно, чтобы после закрытия вернуть на неё фокус. Открытие с карты передаёт null: метка на странице не является кнопкой карточки.",
    )
    add_text(
        doc,
        "useRef хранит узел, который не должен сам по себе вызывать повторный рисунок. У окна это контейнер диалога и кнопка закрытия. У меню это кнопка открытия, кнопка закрытия и панель. У карты это div, в который Яндекс Карты вставляют свой холст.",
    )
    add_text(
        doc,
        "useEffect в окне при открытии прячет прокрутку documentElement, ставит фокус на кнопку закрытия и слушает Escape. Функция очистки снимает слушатель, возвращает прежнее значение overflow и вызывает focus у кнопки карточки. Та же очистка выполняется, когда окно размонтируется и когда меняются зависимости эффекта. У карты очистка вызывает destroy у созданного экземпляра Map и помечает сборку как отменённую, чтобы поздний ответ скрипта не трогал уже снятый узел.",
    )

    add_heading(doc, "5.2 Доступность диалога", bookmark="h52", bid=27)
    add_text(
        doc,
        "Контейнер окна объявлен диалогом. Заголовок h2 имеет id spot-modal-title, и aria-labelledby указывает на него. Закрытие по подложке срабатывает только когда target совпадает с currentTarget, поэтому щелчок по фотографии и тексту окно не закрывает. Tab и Shift+Tab переходят по ссылкам и кнопкам внутри диалога. С последнего элемента фокус уходит на первый, с первого при Shift+Tab — на последний. Сводка требований приведена в таблице 4.",
    )
    add_table(
        doc,
        "Таблица 4 — Требования к модальному диалогу",
        ["Требование", "Как сделано"],
        [
            ["Роль диалога", "role=\"dialog\" и aria-modal=\"true\""],
            ["Имя окна", "aria-labelledby=\"spot-modal-title\""],
            ["Видимое закрытие", "На кнопке виден знак ×, доступное имя — «Закрыть окно»"],
            ["Клавиша Escape", "Слушатель keydown вызывает onClose"],
            ["Щелчок по подложке", "Закрытие при target === currentTarget"],
            ["Фокус при открытии", "Фокус на кнопке закрытия"],
            ["Цикл Tab", "С краёв списка фокус переносится на другой край"],
            ["Возврат фокуса", "focus() у кнопки карточки в очистке эффекта"],
            ["Блокировка прокрутки", "overflow документа ставится в hidden и восстанавливается"],
        ],
        [3200, 6155],
    )
    add_text(
        doc,
        "Гость видит ссылку на /auth/login. Учебный пользователь видит кнопку избранного. Набор идентификаторов живёт в useState и пропадает при обновлении страницы. Записи в localStorage нет. Скрытая у гостя кнопка меняет только разметку. Это условие интерфейса. Права на сервере оно не проверяет, потому что серверной операции сохранения нет: тот же Set можно было бы изменить только в памяти этой вкладки.",
        before=8,
    )

    add_heading(doc, "6 Бургер-меню", break_before=True, bookmark="h6", bid=28)
    add_text(
        doc,
        "Кнопка «Открыть меню» имеет aria-expanded и aria-controls=\"main-menu-panel\". Панель — это dialog с тем же id. Список ссылок собран через Link. usePathname сравнивает адрес с пунктом. Для «Главной» активны и /, и /start, потому что страница /start выполняет redirect на /. Остальные пункты активны на своём адресе и на вложенных сегментах. Текущая ссылка получает класс menu__link_current и aria-current=\"page\".",
    )
    add_text(
        doc,
        "Пока панель открыта, прокрутка документа заблокирована, фокус стоит на кнопке «Закрыть меню», Escape и щелчок по подложке закрывают панель, Tab ходит по её ссылкам и кнопкам. Переход по Link закрывает меню. После закрытия фокус возвращается на кнопку открытия. Гость видит «Войти». Если в контексте есть user, на этом месте имя и кнопка «Выйти»: она вызывает logout и закрывает панель.",
    )

    add_heading(doc, "7 Карта", break_before=True, bookmark="h7", bid=29)
    add_text(
        doc,
        "PhotoMap читает process.env.NEXT_PUBLIC_YANDEX_MAPS_API_KEY. Префикс NEXT_PUBLIC_ означает, что значение попадает в код, который выполняется в браузере. Его видно в собранной странице, поэтому такой ключ нельзя считать серверным секретом. Для карты это ключ доступа к клиентскому API, его и нужно отдавать браузеру вместе со скриптом.",
    )
    add_text(
        doc,
        "Скрипт https://api-maps.yandex.ru/2.1/ подключается один раз, с id pairwise-yandex-maps, параметрами apikey и lang=ru_RU. Загрузка асинхронная: документ не ждёт файл карт, чтобы нарисовать галерею. Обработчик onload вызывает ymaps.ready. Этот метод дожидается, пока API закончит свою инициализацию, и только потом разрешает создать Map и Placemark. Центр карты — [59.942, 30.3], масштаб 12. Для каждой площадки создаётся метка в точке [latitude, longitude]. В hintContent передаётся название, в balloonContent — адрес. Событие click метки вызывает onSelect, и галерея открывает то же окно, что и карточка.",
    )
    add_text(
        doc,
        "Если ключа нет, состояние сразу равно no-key и скрипт не запрашивается. На странице остаётся текст «Карта не подключена: в .env.local нет ключа NEXT_PUBLIC_YANDEX_MAPS_API_KEY». Если скрипт или инициализация завершились ошибкой, показывается «Не удалось загрузить карту».",
    )

    add_heading(doc, "8 Результаты проверки", break_before=True, bookmark="h8", bid=30)
    add_text(
        doc,
        "Команды npm run lint и npm run build завершились без ошибок. Сборка отдала страницу /gallery. Ниже экраны работающего приложения: сетка, окно гостя, бургер, действие гостя, действие после учебного входа, карта с метками и сообщение на случай, когда ключа нет.",
    )
    add_picture(doc, pictures["gallery"], "Рисунок 2 — Галерея площадок", 15.2)
    add_picture(doc, pictures["modal"], "Рисунок 3 — Модальное окно «Точка Кипения»", 15.2)
    add_picture(doc, pictures["burger"], "Рисунок 4 — Открытое бургер-меню", 15.2)
    add_picture(doc, pictures["guest"], "Рисунок 5 — Действие гостя в модальном окне", 15.2)
    add_picture(doc, pictures["user"], "Рисунок 6 — Кнопка избранного после учебного входа", 15.2)
    add_picture(doc, pictures["map"], "Рисунок 7 — Карта площадок с шестью метками", 15.2)
    add_picture(doc, pictures["nokey"], "Рисунок 8 — Сообщение, если ключ карты не задан", 15.2)
    add_text(
        doc,
        "На рисунке 2 видны шапка и карточки площадок. На рисунке 3 окно перекрывает страницу и показывает фотографию зала. На рисунке 4 текущий пункт — «Галерея и карта», внизу панели у гостя ссылка «Войти». На рисунке 5 гостю предложено войти. На рисунке 6 после входа под адресом стоит «Добавить в избранное». На рисунке 7 карта Санкт-Петербурга и шесть меток. Щелчок по метке открыл окно «ITMO Cowork». На рисунке 8 тот же блок без ключа: скрипт не загружается, виден текст про NEXT_PUBLIC_YANDEX_MAPS_API_KEY.",
    )

    add_heading(doc, "9 Ответы на контрольные вопросы", break_before=True, bookmark="h9", bid=31)
    questions = [
        (
            "Для чего требуется директива «use client»?",
            "Она помечает модуль, которому нужны состояние, события, эффекты или браузерные API. В этой работе директива стоит у галереи, окна, карты и бургер-меню.",
        ),
        (
            "Почему page.tsx галереи может оставаться Server Component?",
            "Сам файл только выводит заголовок и подключает PlacesGallery. Состояние и обработчики живут в клиентском компоненте, поэтому страница маршрута остаётся серверной.",
        ),
        (
            "Что возвращает useState?",
            "Пару из текущего значения и функции обновления. Вызов функции обновления приводит к новому рендеру. В галерее так устроены выбранная площадка, множество избранного и кнопка для возврата фокуса.",
        ),
        (
            "Для каких задач используется useRef?",
            "Для хранения DOM-узла между рендерами без лишнего обновления состояния. Через ref берутся диалог, кнопки меню и контейнер карты.",
        ),
        (
            "Когда выполняется cleanup useEffect?",
            "Перед повторным запуском эффекта и при размонтировании компонента. В окне очистка снимает Escape, возвращает прокрутку и фокус. На карте очистка уничтожает экземпляр Map.",
        ),
        (
            "Чем портал изменяет физическое размещение DOM?",
            "createPortal оставляет компонент в дереве React и переносит его узел в другой контейнер. Окно монтируется в document.body и рисуется поверх страницы.",
        ),
        (
            "Почему контекст остаётся доступен внутри портала?",
            "Провайдер ищется по дереву React. Портал это дерево не разрывает, поэтому PhotoModal по-прежнему читает DemoAuthProvider.",
        ),
        (
            "Какие атрибуты нужны контейнеру модального диалога?",
            "role=\"dialog\", aria-modal=\"true\" и aria-labelledby на заголовок окна. Рядом нужна видимая кнопка закрытия.",
        ),
        (
            "Как должны работать Tab, Shift+Tab и Escape?",
            "Tab и Shift+Tab циклически ходят по фокусируемым элементам внутри диалога. Escape закрывает окно.",
        ),
        (
            "Почему для карточки используется button, а не div с onClick?",
            "Кнопка по умолчанию попадает в порядок Tab и активируется с клавиатуры. У div этого нет, пока ему отдельно не задать роль и обработку клавиш.",
        ),
        (
            "Какие задачи решает next/image?",
            "Компонент выводит локальный файл, задаёт alt и sizes и подставляет размеры под ширину экрана. В карточках и в окне изображение занимает контейнер через fill.",
        ),
        (
            "Где хранятся локальные изображения?",
            "В каталоге public/images/spots. Путь в поле image начинается с /images/spots/.",
        ),
        (
            "Почему API карты загружается асинхронно?",
            "Скрипт карт нужен только блоку карты. Асинхронная вставка не задерживает разметку галереи. Страница остаётся рабочей, даже если скрипт ещё не пришёл или ключа нет.",
        ),
        (
            "Для чего используется ymaps.ready?",
            "Колбэк внутри ready выполняется после инициализации API. До этого момента конструкторы Map и Placemark ещё нельзя вызывать.",
        ),
        (
            "Как Placemark добавляется на карту?",
            "Для площадки создаётся ymaps.Placemark с координатами [широта, долгота], подсказкой и адресом. Метка добавляется через map.geoObjects.add. Щелчок вызывает выбор этой площадки.",
        ),
        (
            "Почему NEXT_PUBLIC_YANDEX_MAPS_API_KEY не является серверным секретом?",
            "Префикс NEXT_PUBLIC_ встраивает значение в клиентский код. Браузер передаёт ключ в адресе скрипта карт, и его можно прочитать на странице.",
        ),
        (
            "Почему скрытие действия для гостя не защищает данные?",
            "Условие user меняет кнопку на ссылку только в интерфейсе. Избранное и так хранится в памяти вкладки до обновления. Серверной проверки права в работе нет, поэтому скрытая кнопка не закрывает данные.",
        ),
    ]
    for title, answer in questions:
        add_text(doc, title, first=709, before=8, after=2, bold=True)
        add_text(doc, answer)

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", center=True, break_before=True, bookmark="h10", bid=32)
    add_text(
        doc,
        "В PairWise собрана клиентская часть галереи. У площадки есть локальная фотография, альтернативный текст и координаты. Страница /gallery показывает карточки, модальное окно и блок карты. Окно открывается в document.body, закрывается по Escape и по подложке, держит фокус внутри и возвращает его на карточку.",
    )
    add_text(
        doc,
        "Бургер-меню заменило прежний список в шапке и отмечает текущий раздел. Гость и учебный пользователь видят в окне разные действия. Избранное живёт до обновления страницы. Карта показывает шесть меток, щелчок по метке открывает окно площадки. Без ключа в окружении на этом месте остаётся сообщение, что карта не подключена. Команды проверки и производственной сборки прошли без ошибок.",
    )

    add_heading(doc, "ПРИЛОЖЕНИЕ А Листинги исходного кода", center=True, break_before=True, bookmark="h11", bid=33)
    listings = [
        ("Листинг А.1 — Файл app/data/spots.ts", APP / "data" / "spots.ts"),
        ("Листинг А.2 — Файл app/gallery/page.tsx", APP / "gallery" / "page.tsx"),
        ("Листинг А.3 — Файл app/gallery/_components/PlacesGallery.tsx", APP / "gallery" / "_components" / "PlacesGallery.tsx"),
        ("Листинг А.4 — Файл app/gallery/_components/PhotoModal.tsx", APP / "gallery" / "_components" / "PhotoModal.tsx"),
        ("Листинг А.5 — Файл app/gallery/_components/PhotoMap.tsx", APP / "gallery" / "_components" / "PhotoMap.tsx"),
        ("Листинг А.6 — Файл app/ui/BurgerMenu/BurgerMenu.tsx", APP / "ui" / "BurgerMenu" / "BurgerMenu.tsx"),
    ]
    for caption, path in listings:
        add_listing(doc, caption, path)

    doc.save(str(OUT))
    update_page_fields(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
