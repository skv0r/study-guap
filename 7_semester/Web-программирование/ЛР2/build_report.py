# -*- coding: utf-8 -*-
"""Сборка отчёта ЛР2. В текст отчёта не входит."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"C:\Users\skvor\OneDrive\Desktop\GitHub\study-guap\7_semester\Web-программирование")
LAB = ROOT / "ЛР2"
SRC_DOC = ROOT / "ЛР1_Отчет_Буренков_PairWise.docx"
OUT = LAB / "ЛР2_Отчет_Буренков_PairWise.docx"
APP = Path(r"C:\Users\skvor\OneDrive\Desktop\GitHub\pairwise\app")
SHOTS = Path(r"C:\Users\skvor\AppData\Local\Temp\cursor\screenshots")
FIG = LAB / "fig"
BLACK = RGBColor(0, 0, 0)


def rfonts(run, name):
    run.font.name = name
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def style_run(run, size=14, bold=False, font="Times New Roman"):
    rfonts(run, font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def set_spacing(paragraph, before=0, after=0, line=360):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    paragraph.paragraph_format.line_spacing = line / 240.0


def set_indent(paragraph, first=None, left=None):
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))
    if left is not None:
        ind.set(qn("w:left"), str(left))


def add_text(doc, text, *, first=709, before=0, after=0, line=360, center=False, bold=False, size=14, font="Times New Roman"):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before, after, line)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if first:
            set_indent(paragraph, first=first)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold, font=font)
    return paragraph


def add_heading(doc, text, *, center=False, break_before=False):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=0, after=12, line=360)
    paragraph.paragraph_format.keep_with_next = True
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_indent(paragraph, first=709)
    if break_before:
        run_br = paragraph.add_run()
        style_run(run_br, bold=True)
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        run_br._r.append(br)
    run = paragraph.add_run(text)
    style_run(run, bold=True)
    return paragraph


def add_toc(doc, title, token):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, line=360)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ppr = paragraph._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9071")
    tabs.append(tab)
    ppr.append(tabs)
    run = paragraph.add_run(title)
    style_run(run, bold=False)
    paragraph.add_run().add_tab()
    num = paragraph.add_run(token)
    style_run(num, bold=False)
    return paragraph


def shade_cell(cell):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcpr.append(shd)


def set_cell(cell, text, *, bold=False, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_spacing(paragraph, before=2, after=2, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    style_run(run, size=12, bold=bold)
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "60")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)
    if header:
        shade_cell(cell)


def set_table_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    total = sum(widths)
    tblw.set(qn("w:w"), str(total))
    tblw.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblpr.addnext(grid)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc = cell._tc
            tcpr = tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    tblpr.append(borders)
    jc = tblpr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblpr.append(jc)
    jc.set(qn("w:val"), "left")


def add_table(doc, caption, headers, rows, widths):
    content = 9355
    inset = max(0, (content - sum(widths)) // 2)
    add_text(
        doc,
        caption,
        first=0,
        before=6,
        after=6,
        line=240,
        center=False,
        bold=False,
        size=14,
    )
    # caption is the last paragraph
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.keep_with_next = True
    set_indent(cap, first=0, left=inset)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, header=True)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.rows[r + 1].cells[c], value)
    return table


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    set_spacing(paragraph, before=6, after=0, line=240)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_text(doc, caption, first=0, before=6, after=10, line=360, center=True)


def add_listing(doc, caption, path):
    add_text(doc, caption, first=0, before=10, after=2, line=360, center=False)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    lines = path.read_text(encoding="utf-8").replace("\t", "    ").splitlines()
    for line in lines:
        paragraph = doc.add_paragraph()
        set_spacing(paragraph, line=240)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.keep_together = True
        run = paragraph.add_run(line if line else " ")
        style_run(run, size=10, bold=False, font="Courier New")
        if line.startswith(" ") or line.endswith(" "):
            run._element.find(qn("w:t")).set(qn("xml:space"), "preserve")


def make_tree():
    FIG.mkdir(parents=True, exist_ok=True)
    lines = [
        "app",
        "├─ layout.tsx",
        "├─ page.tsx",
        "├─ not-found.tsx",
        "├─ globals.css",
        "├─ about/page.tsx",
        "├─ start/page.tsx",
        "├─ auth/login/page.tsx",
        "├─ auth/register/page.tsx",
        "├─ cabinet/page.tsx",
        "├─ cabinet/cabinet.tsx",
        "├─ catalog/page.tsx",
        "├─ catalog/[id]/page.tsx",
        "├─ data/spots.ts",
        "├─ data/features.ts",
        "├─ providers/DemoProvider.tsx",
        "├─ providers/useDemoAuth.ts",
        "├─ ui/Header/Header.tsx",
        "├─ ui/AppNav/AppNav.tsx",
        "├─ ui/LoginMenu/LoginMenu.tsx",
        "└─ ui/RegisterMenu/RegisterMenu.tsx",
    ]
    font = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 28)
    probe = ImageDraw.Draw(Image.new("RGB", (10, 10)))
    width = max(probe.textlength(line, font=font) for line in lines) + 80
    height = 36 * len(lines) + 48
    image = Image.new("RGB", (int(width), height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width - 1, height - 1), outline=(180, 180, 180))
    y = 24
    for line in lines:
        draw.text((32, y), line, font=font, fill=(0, 0, 0))
        y += 36
    dest = FIG / "tree.png"
    image.save(dest)
    return dest


def replace_title(doc):
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 1": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2",
        "Создание проекта Next.js с использованием React, TypeScript и App Router":
            "Создание структуры веб-приложения с помощью App Router. Проектирование авторизации",
    }
    found = set()
    for node in doc.element.body.iter(qn("w:t")):
        if node.text in replacements:
            node.text = replacements[node.text]
            found.add(node.text)
    expected = set(replacements.values())
    if found != expected:
        raise SystemExit(f"Титул заменён не полностью: {found}")


def build():
    tree = make_tree()
    pictures = {
        "home": SHOTS / "lr2-home.png",
        "catalog": SHOTS / "lr2-catalog.png",
        "spot": SHOTS / "lr2-spot.png",
        "missing": SHOTS / "lr2-notfound.png",
        "login": SHOTS / "page-2026-10-04T14-57-58-883Z.png",
        "cabinet": SHOTS / "lr2-cabinet-guest.png",
    }
    for path in pictures.values():
        if not path.exists():
            raise SystemExit(f"Нет снимка: {path}")

    doc = Document(str(SRC_DOC))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replace_title(doc)

    add_heading(doc, "СОДЕРЖАНИЕ", center=True, break_before=False)
    toc = [
        ("1 Цель работы", "P1"),
        ("2 Продолжение темы PairWise", "P2"),
        ("3 Маршруты приложения", "P3"),
        ("4 Реализация", "P4"),
        ("4.1 Страницы и корневой макет", "P41"),
        ("4.2 Динамическая карточка площадки", "P42"),
        ("4.3 Учебный вход и кабинет", "P43"),
        ("5 Результаты", "P5"),
        ("6 Ответы на контрольные вопросы", "P6"),
        ("ЗАКЛЮЧЕНИЕ", "P7"),
        ("ПРИЛОЖЕНИЕ А Листинги исходного кода", "P8"),
    ]
    for title, token in toc:
        add_toc(doc, title, token)

    add_heading(doc, "1 Цель работы", break_before=True)
    add_text(
        doc,
        "Цель работы — продолжить приложение PairWise и собрать его структуру на App Router: общие страницы, корневой макет, динамическую карточку площадки и учебные формы входа и регистрации.",
    )

    add_heading(doc, "2 Продолжение темы PairWise", break_before=True)
    add_text(
        doc,
        "PairWise подбирает коворкинг для очной встречи. Пользователь смотрит площадки по городу, адресу и стеку людей, которые там бывают. Главная сущность каталога — площадка. У неё есть номер, название, описание, город, адрес, список технологий, изображение и дата появления в базе.",
    )
    add_text(
        doc,
        "В учебном массиве шесть площадок Санкт-Петербурга: «Точка Кипения», GrowUp Ligovsky, Bench на Петроградке, Севкабель Порт, ITMO Cowork и «Практик на Васильевском». Номера идут от 0 до 5. Каталог показывает карточки этих площадок, а страница /catalog/[id] открывает одну из них.",
    )

    add_heading(doc, "3 Маршруты приложения", break_before=True)
    add_text(
        doc,
        "Адрес страницы собирается из вложенных каталогов app, в которых лежит page.tsx. Каталоги data, providers и ui адреса не создают: в них данные, состояние и компоненты. Корневой layout.tsx задаёт язык ru, шрифт Manrope, шапку и оборачивает все страницы в DemoAuthProvider.",
    )
    add_table(
        doc,
        "Таблица 1 — Назначение раздела, каталог app и URL",
        ["Назначение", "Каталог app", "URL"],
        [
            ["Главная", "app/page.tsx", "/"],
            ["Переход на главную", "app/start/page.tsx", "/start"],
            ["О проекте", "app/about/page.tsx", "/about"],
            ["Каталог площадок", "app/catalog/page.tsx", "/catalog"],
            ["Карточка площадки", "app/catalog/[id]/page.tsx", "/catalog/[id]"],
            ["Вход", "app/auth/login/page.tsx", "/auth/login"],
            ["Регистрация", "app/auth/register/page.tsx", "/auth/register"],
            ["Личный кабинет", "app/cabinet/page.tsx", "/cabinet"],
            ["Неизвестный адрес", "app/not-found.tsx", "notFound()"],
            ["Общий макет", "app/layout.tsx", "для всех URL"],
        ],
        [3200, 3600, 2555],
    )
    add_text(
        doc,
        "Страница /start выполняет redirect на /. Меню шапки собрано через Link. Текущий пункт определяется хуком usePathname и помечается атрибутом aria-current. В меню также есть «Карта» (/map) и «Избранное» (/saved). Отдельных page.tsx для этих адресов нет, поэтому открывается not-found.",
        before=8,
    )
    add_picture(doc, tree, "Рисунок 1 — Структура каталога app", 12.5)

    add_heading(doc, "4 Реализация", break_before=True)
    add_heading(doc, "4.1 Страницы и корневой макет", break_before=False)
    add_text(
        doc,
        "Файл page.tsx задаёт содержимое одного адреса. Файл layout.tsx адрес не добавляет: он рисует общую оболочку и остаётся на экране при переходе между вложенными страницами. Корневой макет один на всё приложение. В нём подключены шапка, кнопка прокрутки наверх и провайдер учебного пользователя.",
    )
    add_text(
        doc,
        "Страницы главной, каталога, карточки, проекта, входа, регистрации и кабинета серверные: директивы «use client» в их page.tsx нет. Страница входа только выводит LoginMenu, страница регистрации — RegisterMenu, страница кабинета — CabinetView. Интерактивные части вынесены в клиентские компоненты. Директива стоит у DemoAuthProvider, AppNav, LoginMenu, RegisterMenu и CabinetView, потому что им нужны состояние, контекст или события формы.",
    )

    add_heading(doc, "4.2 Динамическая карточка площадки", break_before=False)
    add_text(
        doc,
        "Каталог app/catalog/[id] создаёт динамический сегмент. Квадратные скобки в имени означают, что на место id в адрес подставляется конкретное значение: /catalog/0, /catalog/1 и дальше. Страница объявлена как async и получает params через await: в текущем App Router это Promise. Строка id переводится в число. Если число не получилось или площадки с таким номером нет, вызывается notFound().",
    )
    add_text(
        doc,
        "Карточка показывает фотографию, название, описание, город, адрес, дату попадания в базу и стек площадки. Ссылка «К каталогу» возвращает на /catalog. Из списка площадок переход идёт по адресу /catalog/ и номеру карточки. Открыты /catalog/0 («Точка Кипения») и /catalog/1 (GrowUp Ligovsky). Адрес /catalog/999 попадает в not-found: на экране текст «404 ха ха» и ссылка «Вернуться на главную».",
    )

    add_heading(doc, "4.3 Учебный вход и кабинет", break_before=False)
    add_text(
        doc,
        "DemoAuthProvider хранит учебного пользователя в useState. В объекте два поля: name и email. Функция login принимает почту и берёт имя из части адреса до символа @. Если такой части нет, имя становится «Пользователь». Функция register записывает переданные имя и почту. logout возвращает состояние к null. Хук useDemoAuth читает контекст и выбрасывает ошибку, если его вызвали вне провайдера.",
    )
    add_text(
        doc,
        "Пароль в login и register не передаётся и в состоянии не хранится. На форме входа есть флажок «Запомнить логин и пароль», но обработчик отправки его не читает. Подсказка поля пароля говорит «Не менее 8 символов». Длину форма входа не проверяет: пустое поле останавливает атрибут required. На регистрации сравниваются два введённых пароля. Если они различаются, под полями появляется текст «Пароли не совпадают», и переход в кабинет не выполняется. Дата рождения на форме есть, в провайдер она не уходит.",
    )
    add_text(
        doc,
        "Кабинет без пользователя показывает заголовок, фразу «Войдите или зарегистрируйтесь, чтобы открыть кабинет» и ссылки на вход и регистрацию. После учебного входа на странице приветствие с именем, почта и кнопка «Выйти». Шапка состояние не читает: ссылка «Вход» на всех экранах ведёт на /auth/login. Имя появляется в кабинете.",
    )
    add_table(
        doc,
        "Таблица 2 — Аутентификация, сессия и авторизация",
        ["Понятие", "Назначение", "Как сделано в PairWise"],
        [
            [
                "Аутентификация",
                "Установить, кто обращается к системе",
                "Форма принимает почту. Пароль ни с чем не сравнивается и в состояние не записывается",
            ],
            [
                "Сессия",
                "Сохранить результат входа между запросами",
                "Имя и почта лежат в памяти вкладки. Обновление страницы возвращает пустое состояние",
            ],
            [
                "Авторизация",
                "Решить, какое действие разрешено",
                "Проверки прав нет. Кабинет меняет текст, если пользователь уже записан",
            ],
        ],
        [2200, 3400, 3755],
    )
    add_text(
        doc,
        "Прототип показывает сценарий гостя и сценарий учебного пользователя. Отдельной надписи о том, что вход учебный, на формах нет. Подтверждения личности и выдачи прав в этом состоянии нет. Состояние создаётся заново при загрузке дерева React, в cookie и localStorage провайдер ничего не пишет. Условие «пользователь есть» меняет только разметку кабинета. Скрытая кнопка на странице не закрывает серверную операцию: ту же операцию можно вызвать запросом мимо интерфейса. Базы пользователей и серверной сессии в работе нет. После подключения PostgreSQL и Prisma на их место встанут учётная запись, хеш пароля, серверная сессия и проверка права перед действием. Формы можно оставить, отправлять их нужно будет уже на сервер.",
        before=8,
    )

    add_heading(doc, "5 Результаты", break_before=True)
    add_text(
        doc,
        "Ниже экраны работающего приложения. В package.json для статической проверки и производственной сборки заданы команды npm run lint и npm run build.",
    )
    add_picture(doc, pictures["home"], "Рисунок 2 — Главная страница PairWise", 15.2)
    add_picture(doc, pictures["catalog"], "Рисунок 3 — Каталог площадок", 15.2)
    add_picture(doc, pictures["spot"], "Рисунок 4 — Карточка «Точка Кипения», адрес /catalog/0", 15.2)
    add_picture(doc, pictures["missing"], "Рисунок 5 — Неизвестный идентификатор /catalog/999", 15.2)
    add_picture(doc, pictures["login"], "Рисунок 6 — Форма входа", 15.2)
    add_picture(doc, pictures["cabinet"], "Рисунок 7 — Личный кабинет до входа", 15.2)
    add_text(
        doc,
        "На главной виден текст о подборе коворкинга и блок функций. В каталоге карточки площадок со стеком и ссылкой «Перейти к площадке». На карточке /catalog/0 показаны город, адрес и дата. Для номера 999 открывается страница неизвестного адреса. Форма входа содержит почту, пароль, флажок и ссылку на регистрацию. Кабинет до входа предлагает войти или зарегистрироваться. В шапке на этих экранах остаётся пункт «Вход».",
    )

    add_heading(doc, "6 Ответы на контрольные вопросы", break_before=True)
    questions = [
        (
            "Как каталог app преобразуется в URL?",
            "Вложенный каталог с файлом page.tsx добавляет сегмент адреса. Файл app/catalog/page.tsx открывается как /catalog, а app/catalog/[id]/page.tsx — как /catalog/ и значение сегмента. Каталог без page.tsx, например data или ui, адресом не становится.",
        ),
        (
            "Чем page.tsx отличается от layout.tsx?",
            "page.tsx — содержимое конкретного адреса. layout.tsx — оболочка вокруг этого содержимого и вложенных страниц. При переходе по Link оболочка сохраняется, меняется страница внутри неё.",
        ),
        (
            "Для чего используются квадратные скобки в имени каталога?",
            "Скобки обозначают динамический сегмент. Имя внутри скобок, здесь id, становится полем params. Один файл страницы обслуживает все карточки площадок.",
        ),
        (
            "Почему params в актуальном App Router получают через await?",
            "В текущей версии params передаётся как Promise. Страница объявлена async и сначала дожидается этого значения, затем читает id.",
        ),
        (
            "Когда требуется директива «use client»?",
            "Она нужна компоненту с состоянием, контекстом, обработчиком события или хуком браузера. В PairWise это провайдер, меню, формы и кабинет. Серверная страница может подключить такой компонент и сама остаться серверной.",
        ),
        (
            "Какие преимущества даёт Link по сравнению с полной перезагрузкой?",
            "Link меняет адрес без полной перезагрузки документа. Общий макет и дерево React сохраняются, поэтому учебный пользователь остаётся в памяти, пока вкладку не обновят. Обычный переход с перезагрузкой собирает страницу заново.",
        ),
        (
            "Чем аутентификация отличается от авторизации?",
            "Аутентификация устанавливает, кто обращается. Авторизация решает, что этому человеку разрешено. Форма PairWise записывает почту в состояние и права на действия не проверяет.",
        ),
        (
            "Для чего нужна сессия?",
            "Сессия хранит уже установленный результат входа между запросами, чтобы пользователь не представлялся на каждой странице заново. В прототипе эту роль играет useState внутри вкладки.",
        ),
        (
            "Почему скрытая кнопка не защищает серверную операцию?",
            "Кнопка и условие в разметке работают в браузере. Серверная операция обязана проверять право сама. Если проверки на сервере нет, запрос можно отправить мимо спрятанной кнопки.",
        ),
        (
            "Почему DemoAuthProvider сбрасывается после обновления страницы?",
            "Состояние создаётся в памяти вкладки при монтировании провайдера. Начальное значение user равно null. Обновление страницы монтирует дерево заново. В хранилище браузера и в cookie провайдер запись не делает.",
        ),
        (
            "Какие части прототипа будут заменены после подключения PostgreSQL и Prisma?",
            "Вместо объекта в useState появятся запись пользователя и хеш пароля в базе, серверная проверка пароля, сессия между запросами и проверка права перед действием. Клиентские формы могут остаться, но отправка уйдёт на сервер.",
        ),
    ]
    for title, answer in questions:
        add_text(doc, title, first=709, before=8, after=2, bold=True)
        add_text(doc, answer)

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", center=True, break_before=True)
    add_text(
        doc,
        "В работе продолжен проект PairWise. На App Router собраны главная страница, страница о проекте, каталог, динамическая карточка площадки, вход, регистрация и личный кабинет. Общий макет держит шапку и провайдер учебного пользователя. Неизвестный номер площадки обрабатывается через notFound().",
    )
    add_text(
        doc,
        "Учебный вход хранит имя и почту в памяти вкладки. Пароль в это состояние не попадает. После обновления страницы записанный пользователь пропадает. Такой прототип показывает сценарий интерфейса и не заменяет проверку личности, сессию и права на сервере.",
    )

    add_heading(doc, "ПРИЛОЖЕНИЕ А Листинги исходного кода", center=True, break_before=True)
    listings = [
        ("Листинг А.1 — Файл app/data/spots.ts", APP / "data" / "spots.ts"),
        ("Листинг А.2 — Файл app/catalog/page.tsx", APP / "catalog" / "page.tsx"),
        ("Листинг А.3 — Файл app/catalog/[id]/page.tsx", APP / "catalog" / "[id]" / "page.tsx"),
        ("Листинг А.4 — Файл app/providers/DemoProvider.tsx", APP / "providers" / "DemoProvider.tsx"),
        ("Листинг А.5 — Файл app/providers/useDemoAuth.ts", APP / "providers" / "useDemoAuth.ts"),
        ("Листинг А.6 — Файл app/ui/LoginMenu/LoginMenu.tsx", APP / "ui" / "LoginMenu" / "LoginMenu.tsx"),
        ("Листинг А.7 — Файл app/ui/RegisterMenu/RegisterMenu.tsx", APP / "ui" / "RegisterMenu" / "RegisterMenu.tsx"),
        ("Листинг А.8 — Файл app/cabinet/cabinet.tsx", APP / "cabinet" / "cabinet.tsx"),
    ]
    for caption, path in listings:
        add_listing(doc, caption, path)

    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
