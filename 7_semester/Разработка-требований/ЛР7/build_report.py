# -*- coding: utf-8 -*-
"""Сборка отчётов ЛР7 и ЛР8. В текст отчётов не входит."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "ЛР2" / "ЛР2_Отчет_Буренков_Требования.docx"
FIG7 = ROOT / "ЛР7" / "figures"
FIG8 = ROOT / "ЛР8" / "figures"
USABLE = Cm(16.5)


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def format_paragraph(p, *, first_indent=True, align="justify", space_after=0, space_before=0, left_indent=0, single=False):
    pf = p.paragraph_format
    if single:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT}[align]


def add_body(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, *, new_page=False, level=1):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    if level == 0:
        format_paragraph(p, first_indent=False, align="center", space_after=8)
    else:
        format_paragraph(p, first_indent=True, align="left", space_before=8, space_after=4)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    return p


def add_toc_line(doc, title, page):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(str(page)))


def add_caption(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=8, space_after=6, single=True)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def add_fig(doc, path, caption, width=15.5):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="center", space_before=6, single=True)
    p.add_run().add_picture(str(path), width=Cm(width))
    cap = doc.add_paragraph()
    format_paragraph(cap, first_indent=False, align="center", space_before=4, space_after=8)
    set_run_font(cap.add_run(caption))


def add_code(doc, line):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=0.75, single=True)
    set_run_font(p.add_run(line if line else " "), size=12)


def _borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _cell(cell, text, *, bold=False, size=11):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        format_paragraph(p, first_indent=False, align="left", single=True)
        set_run_font(p.add_run(line if line else " "), size=size, bold=bold)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths)
    emu = [int(USABLE * (w / total)) for w in widths]
    grid = table._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        child.set(qn("w:w"), str(int(emu[i] / 635)))
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(USABLE / 635)))
    tblW.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Emu(emu[idx])
            _borders(cell)
    for j, header in enumerate(headers):
        _cell(table.rows[0].cells[j], header, bold=True)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            _cell(table.rows[i + 1].cells[j], value)
    return table


def clone(out: Path, number: str, topic: str) -> Document:
    shutil.copy(TEMPLATE, out)
    doc = Document(str(out))
    body = doc.element.body
    for child in list(body)[12:-1]:
        body.remove(child)
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2": f"ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № {number}",
        "Идентификация действующих лиц и вариантов использования": topic,
    }
    found = set()
    for node in body.iter(qn("w:t")):
        if node.text in replacements:
            found.add(node.text)
            node.text = replacements[node.text]
    if set(replacements) - found:
        raise SystemExit(f"Титул не заменён: {set(replacements) - found}")
    return doc


def measure(path: Path, headings: list[str]) -> dict[str, str]:
    listing = Path(tempfile.gettempdir()) / "req_pages.txt"
    script = Path(tempfile.gettempdir()) / "req_pages.ps1"
    out = str(path).replace("'", "''")
    listing_s = str(listing).replace("'", "''")
    script.write_text(
        "\n".join([
            "$ErrorActionPreference = 'Stop'",
            "$word = $null; $doc = $null",
            "try {",
            "  $word = New-Object -ComObject Word.Application",
            "  $word.Visible = $false; $word.DisplayAlerts = 0",
            f"  $doc = $word.Documents.Open('{out}', $false, $true)",
            "  $doc.Repaginate()",
            "  $lines = New-Object System.Collections.Generic.List[string]",
            "  foreach ($p in $doc.Paragraphs) {",
            "    $t = $p.Range.Text.Replace(\"`r\",\"\").Replace(\"`a\",\"\").Trim()",
            "    if ($t.Length -gt 0 -and -not $t.Contains(\"`t\")) {",
            "      $page = [int]$p.Range.Information(3)",
            "      $lines.Add(\"$page`t$t\")",
            "    }",
            "  }",
            f"  [System.IO.File]::WriteAllLines('{listing_s}', $lines, [System.Text.UTF8Encoding]::new($false))",
            "} finally {",
            "  if ($doc -ne $null) { $doc.Close(0) }",
            "  if ($word -ne $null) { $word.Quit() }",
            "}",
        ]),
        encoding="utf-8-sig",
    )
    subprocess.run(["powershell", "-NoProfile", "-File", str(script)], check=True)
    found = {}
    wanted = set(headings)
    for line in listing.read_text(encoding="utf-8").splitlines():
        if "\t" not in line:
            continue
        page, text = line.split("\t", 1)
        text = text.strip()
        if text in wanted and text not in found:
            found[text] = page.strip()
    missing = [title for title in headings if title not in found]
    if missing:
        raise SystemExit("Нет страниц: " + ", ".join(missing))
    return found


def finish(doc, path, headings, build):
    pages = {title: "…" for title in headings}
    for _ in range(3):
        doc = clone(path, build.number, build.topic)
        build(doc, pages)
        doc.save(str(path))
        measured = measure(path, headings)
        if measured == pages:
            print(path.name, measured)
            return
        pages = measured
    print(path.name, pages)


H7 = ["ВВЕДЕНИЕ", "1 Задание", "2 Два слоя", "3 Правила слоя данных", "4 Прогон", "ЗАКЛЮЧЕНИЕ"]
H8 = ["ВВЕДЕНИЕ", "1 Задание", "2 Матрицы отслеживания", "3 Тест-кейсы", "4 Прогон", "5 Проверка требований", "ЗАКЛЮЧЕНИЕ"]


def build7(doc, pages):
    add_heading(doc, "СОДЕРЖАНИЕ", level=0)
    for title in H7:
        add_toc_line(doc, title, pages[title])
    add_heading(doc, "ВВЕДЕНИЕ", new_page=True, level=0)
    add_body(doc, "Цель работы — запрограммировать интерфейс пользователя и слой доступа к данным АИС железной дороги. Экран показывает поля и ответ. Решение, можно ли записать станцию, билет или заявку, принимает слой данных. Отчёт описывает это разделение и результат прогона.")
    add_heading(doc, "1 Задание", new_page=True)
    add_body(doc, "Среда та же, что выбрана для прототипа: TypeScript и React. Новые библиотеки не подключались: слой данных написан средствами языка и вызывается из форм. В отчёте — технические решения и кадры прогона, по которым видно и успешную запись, и отказ.")
    add_heading(doc, "2 Два слоя", new_page=True)
    add_body(doc, "Интерфейс знает роль, текущий экран и текст ответа. Он не хранит список станций, рейсов и билетов и не решает, уникален ли код. Эти записи лежат в слое данных. Форма собирает поля и вызывает метод слоя. Метод возвращает ответ: запись прошла или нет, и текст, который экран показывает без переписывания.")
    add_body(doc, "На входе кассир или контролёр получают открытую смену, если станция в обслуживании. Остальные роли входят без смены. Выход закрывает смену. Справочники, маршрут, рейс, продажа, возврат, посадка, обслуживание, ремонт и инвентаризация читают и пишут только через слой. Под экраном справочника и продажи выведен список уже записанных объектов, чтобы было видно, что строка попала в данные, а не только в подпись кнопки.")
    add_heading(doc, "3 Правила слоя данных", new_page=True)
    add_body(doc, "Слой проверяет правило до записи. Повторный код станции не добавляет вторую строку. В маршруте нужны две разные станции. Рейс с тем же номером не публикуется второй раз, отменённый номер не открывается снова, повторная отмена не создаётся. Продажа требует открытую смену и открытый рейс. Место, на которое уже есть непроведённый возвратом билет, второй раз не продаётся. Черновик не создаёт билет и не даёт закрыть смену. Посадка ставится один раз. После посадки возврат не проходит. Заявку закрывает только инженер, закрытую той же неисправностью заново не открывают. Инвентаризация одного объекта на одну дату не повторяется.")
    add_body(doc, "Фрагмент продажи показывает, где сидит запрет второго билета. Экран сюда не заглядывает: он получает готовый текст.")
    for line in [
        "if (!stationOpen || !shiftOpen)",
        "    return отказ «Нет обслуживания станции или открытой смены.»",
        "if (!trip.open)",
        "    return отказ «Рейс закрыт. Билет не создан.»",
        "if (билет на это место рейса уже есть и не возвращён)",
        "    return отказ «Место 12 рейса 102 уже продано. Второй билет не создан.»",
        "записать билет и вернуть «Билет проведён, купон отправлен на печать.»",
    ]:
        add_code(doc, line)
    add_heading(doc, "4 Прогон", new_page=True)
    add_body(doc, "На рисунке 1 кассир ещё не вошёл: смена закрыта, станция в обслуживании. После входа смена открывается, и продажа места 12 рейса 102 проходит. Билет 102-12 появляется в списке проведённых (рисунок 2).")
    add_fig(doc, FIG7 / "login.png", "Рисунок 1 — Вход кассира до открытия смены")
    add_fig(doc, FIG7 / "sale-ok.png", "Рисунок 2 — Слой записал билет, экран показал список")
    add_body(doc, "Повтор той же продажи слой отвергает. В списке по-прежнему одна строка (рисунок 3). Та же схема у справочника: код LUG записывается, вторая попытка с тем же кодом не добавляет станцию (рисунки 4 и 5).")
    add_fig(doc, FIG7 / "sale-double.png", "Рисунок 3 — Отказ второго билета на место 12")
    add_fig(doc, FIG7 / "station-ok.png", "Рисунок 4 — Станция Луговая записана в слой")
    add_fig(doc, FIG7 / "station-dup.png", "Рисунок 5 — Повтор кода LUG не создал вторую станцию")
    add_heading(doc, "ЗАКЛЮЧЕНИЕ", new_page=True, level=0)
    add_body(doc, "Интерфейс АИС железной дороги показывает формы и ответ слоя данных. Слой хранит станции, схемы, маршруты, рейсы, билеты, заявки и акты инвентаризации и сам отказывает в повторной записи. На прогоне билет 102-12 записан, вторая продажа места 12 отвергнута, станция LUG записана один раз.")


def build8(doc, pages):
    add_heading(doc, "СОДЕРЖАНИЕ", level=0)
    for title in H8:
        add_toc_line(doc, title, pages[title])
    add_heading(doc, "ВВЕДЕНИЕ", new_page=True, level=0)
    add_body(doc, "Цель работы — проследить требования АИС железной дороги до экранов и тест-кейсов и проверить требования прогоном. Матрицы отслеживания связывают правило, вариант использования, экран и тест. Тест-кейсы собраны по сценариям вариантов: основной ход и отказ. Прогон показывает, какой ответ дал слой данных.")
    add_heading(doc, "1 Задание", new_page=True)
    add_body(doc, "Трассировка сделана матрицами по главе 29 Вигерса: от бизнес-правила к варианту использования, от варианта к экранной форме и от варианта к тесту. Тест-кейсы построены по статье о генерации тестов из вариантов использования: сценарий, условия, конкретные значения, ожидаемый результат. Проверка требований выполнена на прототипе, как приём из главы 17 Вигерса: тест ищет не только сбой экрана, но и пропуск в формулировке требования.")
    add_heading(doc, "2 Матрицы отслеживания", new_page=True)
    add_body(doc, "Таблица 1 связывает уже принятые правила с вариантами UC-1–UC-9. Таблица 2 доводит вариант до экрана прототипа. Таблица 3 доводит вариант до идентификатора теста, который разобран ниже.")
    add_caption(doc, "Таблица 1 — Правила и варианты использования")
    add_table(
        doc,
        ["Правило", "Вариант"],
        [
            ["Публикация рейса и фонд мест", "UC-1"],
            ["Продажа, в том числе запрет второго билета на место", "UC-2"],
            ["Отмена рейса закрывает новые продажи", "UC-3"],
            ["Посадка отмечается один раз", "UC-4"],
            ["Возврат до посадки", "UC-5"],
            ["Станции и схемы вагонов", "UC-6"],
            ["Маршрут из двух разных станций", "UC-7"],
            ["Открытие и закрытие смены", "UC-8"],
            ["Обслуживание станции, ремонт, инвентаризация", "UC-9"],
        ],
        [10, 6.5],
    )
    add_caption(doc, "Таблица 2 — Вариант и экран")
    add_table(
        doc,
        ["Вариант", "Экран", "Кто открывает"],
        [
            ["UC-1", "Опубликовать рейс, отменить рейс", "Диспетчер расписания"],
            ["UC-2", "Продать билет", "Кассир при открытой смене"],
            ["UC-3", "Отменить рейс, затем продажа", "Диспетчер, затем кассир"],
            ["UC-4", "Отметить посадку", "Контролёр"],
            ["UC-5", "Вернуть билет", "Кассир"],
            ["UC-6", "Справочник станций, схемы вагонов", "Администратор справочников"],
            ["UC-7", "Составить маршрут", "Диспетчер расписания"],
            ["UC-8", "Вход и закрыть смену", "Кассир, контролёр"],
            ["UC-9", "Обслуживание, заявка, инвентаризация", "Дежурный и инженеры"],
        ],
        [3, 7.5, 6],
    )
    add_caption(doc, "Таблица 3 — Вариант и тест")
    add_table(
        doc,
        ["Вариант", "Тесты"],
        [
            ["UC-1", "TC-1.1 публикация занятого номера, TC-1.2 повтор отмены, TC-1.3 номер после отмены"],
            ["UC-2", "TC-2.1 продажа места 12, TC-2.2 повтор места, TC-2.3 продажа без смены, TC-2.4 продажа закрытого рейса"],
            ["UC-3", "TC-1.2 и TC-2.4"],
            ["UC-4", "TC-4.1 первая посадка, TC-4.2 повтор"],
            ["UC-5", "TC-5.2 возврат до посадки, TC-5.1 возврат после посадки"],
            ["UC-6", "TC-6.1 новая станция, TC-6.2 повтор кода"],
            ["UC-7", "TC-7.1 маршрут из двух станций, TC-7.2 повтор станции"],
            ["UC-8", "TC-8.1 вход открывает смену, TC-8.2 черновик не даёт закрыть смену"],
            ["UC-9", "TC-9.1 дежурный не закрывает заявку, TC-9.2 инженер закрывает, TC-9.4 первая инвентаризация, TC-9.3 повтор"],
        ],
        [3, 13.5],
    )
    add_heading(doc, "3 Тест-кейсы", new_page=True)
    add_body(doc, "Для продажи разобраны сценарии, как в образце тест-кейсов. Сценарий 1 — основной ход. Сценарий 2 — место уже продано. Сценарий 3 — нет смены. Сценарий 4 — рейс закрыт. В матрице V означает допустимое значение, I — недопустимое, прочерк — условие к этому шагу не относится.")
    add_caption(doc, "Таблица 4 — Тест-кейсы продажи, UC-2")
    add_table(
        doc,
        ["ID", "Сценарий", "Рейс", "Место", "Смена", "Ожидаемый ответ"],
        [
            ["TC-2.1", "основной ход", "V 102 открыт", "V 12 свободно", "V открыта", "Билет проведён, купон отправлен на печать."],
            ["TC-2.2", "место занято", "V 102", "I 12 уже продано", "V открыта", "Место 12 рейса 102 уже продано. Второй билет не создан."],
            ["TC-2.3", "нет смены", "V 102", "V 15", "I закрыта", "Нет обслуживания станции или открытой смены."],
            ["TC-2.4", "рейс закрыт", "I 102 отменён", "V 16", "V открыта", "Рейс закрыт. Билет не создан."],
        ],
        [2.2, 2.8, 3, 3.2, 2.4, 2.9],
    )
    add_body(doc, "Таблица 5 задаёт остальные кейсы тем же способом: сценарий, пометка V или I и дословный ожидаемый ответ. Этот ответ записан до прогона. Коды M-13 и M-14 стоят в условии маршрута: M-13 допустим, M-14 с повторной станцией недопустим.")
    add_caption(doc, "Таблица 5 — Тест-кейсы остальных вариантов")
    add_table(
        doc,
        ["ID", "Сценарий", "Условие", "Ожидаемый ответ"],
        [
            ["TC-1.1", "номер уже открыт", "I рейс 102 уже открыт", "Рейс с этим номером уже открыт."],
            ["TC-1.2", "повтор отмены", "I рейс 102 уже отменён", "Рейс уже отменён. Повторная отмена не создана."],
            ["TC-1.3", "номер после отмены", "I номер 102 отменён", "Рейс отменён. Тем же номером не открывается."],
            ["TC-4.1", "первая посадка", "V билет 102-12, V смена открыта", "Посадка отмечена один раз."],
            ["TC-4.2", "повтор посадки", "I отметка уже есть", "Отметка посадки уже есть."],
            ["TC-5.2", "возврат до посадки", "V билет 102-15, V посадка не стоит", "Билет возвращён, место свободно."],
            ["TC-5.1", "возврат после посадки", "I посадка билета 102-12 отмечена", "Посадка уже отмечена. Возврат не проведён."],
            ["TC-6.1", "новая станция", "V код LUG свободен", "Станция Луговая сохранена. Код LUG уникален."],
            ["TC-6.2", "повтор кода", "I код LUG уже есть", "Код LUG уже есть. Вторая станция не записана."],
            ["TC-7.1", "две станции", "V код M-13, V «Северная — Южная»", "Маршрут сохранён: две станции, без повтора."],
            ["TC-7.2", "повтор станции", "I код M-14, I «Северная — Северная»", "Станция в маршруте повторяется. Маршрут не сохранён."],
            ["TC-8.1", "вход кассира", "V роль кассир, V станция в обслуживании", "Смена открыта."],
            ["TC-8.2", "смена при черновике", "I черновик продажи есть", "Есть черновик продажи. Смена не закрыта."],
            ["TC-9.1", "закрытие дежурным", "I роль дежурный", "Закрыть заявку может только инженер после отметки устранения."],
            ["TC-9.2", "закрытие инженером", "V роль инженер", "Заявка закрыта."],
            ["TC-9.4", "первая инвентаризация", "V касса 1, V дата 2026-10-20 свободна", "Инвентаризация за 2026-10-20 проведена. Повтор на эту дату закрыт."],
            ["TC-9.3", "повтор инвентаризации", "I касса 1 за 2026-10-20 уже есть", "Инвентаризация Касса 1 за 2026-10-20 уже есть. Повтор закрыт."],
        ],
        [2.2, 3.6, 5.2, 5.5],
    )
    add_heading(doc, "4 Прогон", new_page=True)
    add_body(doc, "Прогон выполнен на прототипе. Экран показывает тот текст, который вернул слой. Таблица 6 ставит рядом ответ прогона и итог. Для каждого теста этот ответ совпадает с ожидаемым текстом таблиц 4 и 5. После TC-7.1 в списке маршрутов есть M-13, после TC-7.2 кода M-14 нет. Рисунки 1–4 фиксируют успешную продажу места 12, отказ второго билета на это место, запись станции LUG и отказ повторного кода.")
    add_caption(doc, "Таблица 6 — Ответ прогона")
    add_table(
        doc,
        ["Тест", "Ответ", "Итог"],
        [
            ["TC-8.1", "Смена открыта.", "пройден"],
            ["TC-2.1", "Билет проведён, купон отправлен на печать.", "пройден"],
            ["TC-2.2", "Место 12 рейса 102 уже продано. Второй билет не создан.", "пройден"],
            ["TC-8.2", "Есть черновик продажи. Смена не закрыта.", "пройден"],
            ["TC-2.3", "Нет обслуживания станции или открытой смены.", "пройден"],
            ["TC-1.1", "Рейс с этим номером уже открыт.", "пройден"],
            ["TC-1.2", "Рейс уже отменён. Повторная отмена не создана.", "пройден"],
            ["TC-1.3", "Рейс отменён. Тем же номером не открывается.", "пройден"],
            ["TC-2.4", "Рейс закрыт. Билет не создан.", "пройден"],
            ["TC-4.1", "Посадка отмечена один раз.", "пройден"],
            ["TC-4.2", "Отметка посадки уже есть.", "пройден"],
            ["TC-5.2", "Билет возвращён, место свободно.", "пройден"],
            ["TC-5.1", "Посадка уже отмечена. Возврат не проведён.", "пройден"],
            ["TC-6.1", "Станция Луговая сохранена. Код LUG уникален.", "пройден"],
            ["TC-6.2", "Код LUG уже есть. Вторая станция не записана.", "пройден"],
            ["TC-7.1", "Маршрут сохранён: две станции, без повтора.", "пройден"],
            ["TC-7.2", "Станция в маршруте повторяется. Маршрут не сохранён.", "пройден"],
            ["TC-9.1", "Закрыть заявку может только инженер после отметки устранения.", "пройден"],
            ["TC-9.2", "Заявка закрыта.", "пройден"],
            ["TC-9.4", "Инвентаризация за 2026-10-20 проведена. Повтор на эту дату закрыт.", "пройден"],
            ["TC-9.3", "Инвентаризация Касса 1 за 2026-10-20 уже есть. Повтор закрыт.", "пройден"],
        ],
        [2.4, 11.2, 2.9],
    )
    add_fig(doc, FIG8 / "sale-ok.png", "Рисунок 1 — TC-2.1: билет записан и виден в списке")
    add_fig(doc, FIG8 / "sale-double.png", "Рисунок 2 — TC-2.2: второй билет на место 12 не создан")
    add_fig(doc, FIG8 / "station-ok.png", "Рисунок 3 — TC-6.1: станция LUG в списке")
    add_fig(doc, FIG8 / "station-dup.png", "Рисунок 4 — TC-6.2: повтор кода отклонён")
    add_heading(doc, "5 Проверка требований", new_page=True)
    add_body(doc, "По главе 17 Вигерса тест требования полезен ещё до спора о вёрстке: он показывает, можно ли по тексту правила получить однозначный ответ. Для продажи такое правило нашлось. Формулировки экрана говорили, что билет проводится, но не фиксировали, что то же место того же рейса второй раз проводить нельзя. Без этой строки тест TC-2.2 нельзя было ни написать, ни прогнать. Правило внесено в слой данных: второй билет не создаётся, в списке остаётся одна запись. Это и есть дефект требования, который закрыт до приёмки сценария.")
    add_body(doc, "Остальные ответы в таблице 6 читаются так же. У каждого есть одно ожидаемое предложение в таблицах 4 и 5, и прогон вернул его дословно. Противоречия между матрицами не всплыло: вариант, экран и тест ссылаются на один и тот же отказ. Пассажирского экрана в проверке нет, его нет и в границах системы.")
    add_heading(doc, "ЗАКЛЮЧЕНИЕ", new_page=True, level=0)
    add_body(doc, "Требования UC-1–UC-9 прослежены до экранов и тестов. Тест-кейсы задают сценарий, пометку V или I и ожидаемый текст. Прогон вернул эти тексты. Билет 102-15 возвращён до посадки, возврат билета 102-12 после посадки отклонён. Маршрут M-13 записан, маршрут M-14 с повторной станцией не записан. Первая инвентаризация кассы 1 на 20.10.2026 проходит, повтор отклоняется. Проверка требований добавила в слой явный запрет второй продажи того же места.")


def main():
    out7 = ROOT / "ЛР7" / "ЛР7_Отчет_Буренков_Требования.docx"
    out8 = ROOT / "ЛР8" / "ЛР8_Отчет_Буренков_Требования.docx"
    build7.number = "7"
    build7.topic = "Интерфейс пользователя и слой доступа к данным"
    build8.number = "8"
    build8.topic = "Трассировка и тестирование требований"
    finish(None, out7, H7, build7)
    finish(None, out8, H8, build8)


if __name__ == "__main__":
    main()
