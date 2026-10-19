#!/usr/bin/env python3
"""Сборка отчёта ЛР6. В текст отчёта не входит."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from lxml import etree

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE.parent / "ЛР2" / "ЛР2_Отчет_Буренков_Требования.docx"
OUT = BASE / "ЛР6_Отчет_Буренков_Требования.docx"
TOPIC = "Проект интерфейса пользователя"
USABLE = Cm(16.5)
PAGE_LIMIT = 15


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def format_paragraph(
    p,
    *,
    first_indent=True,
    align="justify",
    space_after=0,
    space_before=0,
    left_indent=0,
    line_spacing=1.5,
    single=False,
):
    pf = p.paragraph_format
    if single:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_body(doc, text, *, space_before=0):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, *, new_page=False, level=1):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    if level == 0:
        format_paragraph(p, first_indent=False, align="center", space_after=8)
    else:
        format_paragraph(
            p,
            first_indent=True,
            align="left",
            space_before=8 if level == 1 else 6,
            space_after=4 if level == 1 else 2,
        )
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    return p


def add_toc_line(doc, title, page, *, indent=0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=indent, line_spacing=1.5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(str(page)))


def add_table_caption(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=8, space_after=6, single=True)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def _borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _set_cell_text(cell, text, *, bold=False, align="left", size=12):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        format_paragraph(p, first_indent=False, align=align, single=True)
        set_run_font(p.add_run(line if line else " "), size=size, bold=bold)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths)
    emu = [int(USABLE * (w / total)) for w in widths]
    grid = table._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        child.set(qn("w:w"), str(int(emu[i] / 635)))
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(USABLE / 635)))
    tblW.set(qn("w:type"), "dxa")
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "left")
    ind = tblPr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tblPr.append(ind)
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Emu(emu[idx])
            _borders(cell)
    for j, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[j], header, bold=True, align="center")
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            _set_cell_text(table.rows[i + 1].cells[j], value, align="left")
    return table


def clone_title() -> Document:
    shutil.copy(TEMPLATE, OUT)
    doc = Document(str(OUT))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 6",
        "Идентификация действующих лиц и вариантов использования": TOPIC,
    }
    found = set()
    for node in body.iter(qn("w:t")):
        if node.text in replacements:
            found.add(node.text)
            node.text = replacements[node.text]
    missing = set(replacements) - found
    if missing:
        raise SystemExit(f"На титуле не найдены строки: {missing}")
    return doc


def strip_unused_images(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    rel_name = "word/_rels/document.xml.rels"
    root = etree.fromstring(files[rel_name])
    used = set()
    for name, data in files.items():
        if not (name.startswith("word/") and name.endswith(".xml") and "/_rels/" not in name):
            continue
        for key in (b'r:id="', b'r:embed="'):
            start = 0
            while True:
                pos = data.find(key, start)
                if pos < 0:
                    break
                end = data.find(b'"', pos + len(key))
                used.add(data[pos + len(key) : end].decode())
                start = end + 1
    removed = []
    for rel in list(root):
        if rel.get("Id") not in used and "image" in (rel.get("Type") or ""):
            removed.append("word/" + rel.get("Target").replace("\\", "/"))
            root.remove(rel)
    files[rel_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for name in removed:
        files.pop(name, None)
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


SHOTS = Path(r"C:\Users\skvor\AppData\Local\Temp\cursor\screenshots")
FIG = BASE / "figures"


def add_picture(doc, path, caption, width_cm):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, first_indent=False, align="center", space_before=8, space_after=0, single=True)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_p = doc.add_paragraph()
    format_paragraph(caption_p, first_indent=False, align="center", space_before=4, space_after=8, single=True)
    set_run_font(caption_p.add_run(caption))


def draw_graph():
    from PIL import Image, ImageDraw, ImageFont
    import math

    FIG.mkdir(parents=True, exist_ok=True)
    font = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", 22)
    image = Image.new("RGB", (1100, 1040), "white")
    draw = ImageDraw.Draw(image)

    def box(x, y, w, h, text):
        draw.rectangle((x, y, x + w, y + h), outline="black", width=2)
        bbox = draw.textbbox((0, 0), text, font=font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((x + (w - tw) / 2, y + (h - th) / 2 - 2), text, font=font, fill="black")

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill="black", width=2)
        angle = math.atan2(y2 - y1, x2 - x1)
        size = 11
        p1 = (x2 - size * math.cos(angle - 0.45), y2 - size * math.sin(angle - 0.45))
        p2 = (x2 - size * math.cos(angle + 0.45), y2 - size * math.sin(angle + 0.45))
        draw.polygon([p1, (x2, y2), p2], fill="black")

    box(40, 36, 250, 52, "Вход")
    box(390, 36, 280, 52, "Меню роли")
    box(760, 36, 290, 52, "Закрытие смены")
    arrow(290, 62, 390, 62)
    arrow(670, 62, 760, 62)
    draw.line((905, 36, 905, 16, 165, 16), fill="black", width=2)
    arrow(165, 16, 165, 36)

    forms = [
        "Справочник станций",
        "Схемы вагонов",
        "Маршрут",
        "Публикация рейса",
        "Отмена рейса",
        "Продажа билета",
        "Возврат билета",
        "Посадка",
        "Обслуживание станции",
        "Заявка на ремонт",
        "Инвентаризация",
    ]
    spine = 470
    top = 88
    last = 160 + (len(forms) - 1) * 74 + 26
    draw.line((spine, top, spine, last), fill="black", width=2)
    for i, name in enumerate(forms):
        y = 160 + i * 74
        box(520, y, 500, 52, name)
        arrow(spine, y + 26, 520, y + 26)
    dest = FIG / "graph.png"
    image.save(dest)
    return dest


HEADINGS = [
    ("ВВЕДЕНИЕ", 0),
    ("1 Задание", 0),
    ("2 Экранные формы", 0),
    ("3 Граф переходов", 0),
    ("4 Вид форм", 0),
    ("ЗАКЛЮЧЕНИЕ", 0),
]


def build_body(doc: Document, pages: dict[str, str]) -> None:
    graph = draw_graph()
    shots = {
        "login": SHOTS / "lr6-login.png",
        "menu": SHOTS / "lr6-menu.png",
        "route": SHOTS / "lr6-route.png",
        "sale": SHOTS / "lr6-sale.png",
        "service": SHOTS / "lr6-service.png",
        "repair": SHOTS / "lr6-repair.png",
    }
    for path in shots.values():
        if not path.exists():
            raise SystemExit(f"Нет снимка: {path}")

    add_heading(doc, "СОДЕРЖАНИЕ", level=0)
    for title, indent in HEADINGS:
        add_toc_line(doc, title, pages.get(title, "…"), indent=indent * 0.5)

    add_heading(doc, "ВВЕДЕНИЕ", new_page=True, level=0)
    add_body(
        doc,
        "Прототип собирает экранные формы автоматизированной информационной системы железной дороги "
        "для дальнейшей реализации. Среда — React и TypeScript. Формы обслуживают роли уже описанного контура: "
        "диспетчера, кассира, контролёра, администратора справочников, дежурного по станции и инженеров ремонта и инвентаризации. "
        "Экрана пассажира нет.",
    )

    add_heading(doc, "1 Задание", new_page=True)
    add_body(
        doc,
        "Нужно спроектировать интерфейс пользователя: экранные формы прототипа, граф переходов между ними и описание каждой формы. "
        "Правила те же, что в документе концепции, спецификации UC-1–UC-9 и дополнительной спецификации. "
        "Смена кассира и контролёра открывается только при обслуживании станции. Продажа, возврат и посадка требуют открытой смены. "
        "Черновик продажи не даёт закрыть смену. После отмены рейса новая продажа не проводится.",
    )

    add_heading(doc, "2 Экранные формы", new_page=True)
    add_body(doc, "Поля, роль и отказ каждой формы сведены в таблице 1. На рабочей форме есть «Сохранить» и «К меню».")
    add_table_caption(doc, "Таблица 1 — Экранные формы прототипа")
    add_table(
        doc,
        ["Форма", "Поля", "Роль, переход и отказ"],
        [
            ["Вход", "Роль, станция", "Все роли, далее меню. Смена кассы и посадки не открывается, если станция не в обслуживании"],
            ["Меню роли", "Кнопки форм этой роли, «Выйти»", "В форму роли. «Выйти» возвращает на вход и смену не закрывает"],
            ["Справочник станций", "Код, название", "Администратор, к меню. Повтор кода не сохраняется. Запись, занятую рейсом или маршрутом, не удаляют"],
            ["Схемы вагонов", "Тип вагона, номера мест", "Администратор, к меню. Повтор номера места не сохраняется. Схему, занятую открытым рейсом, не удаляют"],
            ["Маршрут", "Код, название, станции по порядку, время хода", "Диспетчер, к меню. Меньше двух станций, повтор станции или правка маршрута открытого рейса — отказ"],
            ["Публикация рейса", "Номер, дата, маршрут, схема, время отправления и прибытия", "Диспетчер, к меню. Нет маршрута или схемы, тот же номер уже открыт или рейс отменён — рейс не открывается"],
            ["Отмена рейса", "Рейс, причина", "Диспетчер, к меню. Повторная отмена не создаётся. Черновик продажи не проводится и отмену не блокирует"],
            ["Продажа билета", "Станции и дата, место, пассажир, цена из прайса, черновик", "Кассир при открытой смене, к меню. Нет смены, станция закрыта, рейс отменён, нет цены или место занято — билет не создаётся"],
            ["Возврат билета", "Номер билета", "Кассир при открытой смене, до отправления и до посадки, к меню. Посадка, погашенный билет, закрытая смена или истёкший срок — отказ. При отменённом рейсе место остаётся закрытым"],
            ["Посадка", "Номер билета или купона", "Контролёр при открытой смене, к меню. Билет не найден, возвращён, рейс отменён или отметка уже есть"],
            ["Обслуживание станции", "Кассы, платформы, зал, проходы, замечание", "Дежурный, к меню. Сутки не закрываются при открытой смене. Отказ, если станции нет в справочнике или сутки уже закрыты. Неготовый объект создаёт или обновляет заявку на ремонт"],
            ["Заявка на ремонт", "Объект, описание, статус", "Дежурный или инженер по ремонту, к меню. Закрывает заявку инженер. Закрытую той же неисправностью повторно не открывают"],
            ["Инвентаризация", "Дата, объект, факт", "Инженер по инвентаризации, к меню. Повтор на ту же дату закрыт"],
            ["Закрытие смены", "Подтверждение закрытия", "Кассир или контролёр, на вход. Черновик продажи смену не закрывает"],
        ],
        [3.8, 5.4, 7.3],
    )

    add_heading(doc, "3 Граф переходов", new_page=True)
    add_body(
        doc,
        "Рисунок 1 показывает вход в меню роли и формы, которые меню открывает. "
        "Закрытие смены возвращает на вход. Его выполняют только кассир и контролёр, и только если нет черновика продажи. "
        "Кнопка «Выйти» в шапке возвращает на вход и смену не закрывает. "
        "Остальные формы возвращаются в меню кнопкой «К меню». Меню показывает только формы текущей роли.",
    )
    add_picture(doc, graph, "Рисунок 1 — Граф переходов экранных форм", 16.2)

    add_heading(doc, "4 Вид форм", new_page=True)
    add_body(
        doc,
        "Рисунки 2–7 показывают вход, меню кассира, маршрут, продажу, обслуживание станции и заявку на ремонт. "
        "Справочник, схема вагона, публикация, отмена, возврат, посадка, инвентаризация и закрытие смены собраны тем же шаблоном: поля из таблицы 1, «Сохранить» и «К меню».",
    )
    add_picture(doc, shots["login"], "Рисунок 2 — Вход на рабочее место", 15.2)
    add_picture(doc, shots["menu"], "Рисунок 3 — Меню кассира при открытой смене", 15.2)
    add_picture(doc, shots["route"], "Рисунок 4 — Форма маршрута диспетчера", 15.2)
    add_picture(doc, shots["sale"], "Рисунок 5 — Продажа билета", 15.2)
    add_picture(doc, shots["service"], "Рисунок 6 — Обслуживание станции", 15.2)
    add_picture(doc, shots["repair"], "Рисунок 7 — Заявка на ремонт", 15.2)

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", new_page=True, level=0)
    add_body(
        doc,
        "В React и TypeScript собран прототип экранных форм АИС железной дороги. "
        "Граф ведёт от входа к меню роли и дальше к справочникам, маршруту, рейсу, кассе, посадке, обслуживанию станции, ремонту и инвентаризации. "
        "Смена и обслуживание станции связаны так же, как в спецификации вариантов. Пассажирского экрана в прототипе нет.",
    )


def save_doc(doc: Document) -> None:
    doc.save(str(OUT))
    strip_unused_images(OUT)


def measure_pages():
    import platform

    if platform.system() == "Darwin":
        return {title: "1" for title, _indent in HEADINGS}, 1
    listing = Path(tempfile.gettempdir()) / "lr6_req_pages.txt"
    script = Path(tempfile.gettempdir()) / "lr6_req_pages.ps1"
    out = str(OUT).replace("'", "''")
    listing_s = str(listing).replace("'", "''")
    script.write_text(
        "\n".join(
            [
                "$ErrorActionPreference = 'Stop'",
                "$word = $null",
                "$doc = $null",
                "try {",
                "  $word = New-Object -ComObject Word.Application",
                "  $word.Visible = $false",
                "  $word.DisplayAlerts = 0",
                f"  $doc = $word.Documents.Open('{out}', $false, $true)",
                "  $doc.Repaginate()",
                "  $total = $doc.ComputeStatistics(2)",
                "  $lines = New-Object System.Collections.Generic.List[string]",
                "  $lines.Add(\"TOTAL`t$total\")",
                "  foreach ($p in $doc.Paragraphs) {",
                "    $t = $p.Range.Text.Replace(\"`r\",\"\").Replace(\"`a\",\"\").Trim()",
                "    if ($t.Length -gt 0 -and -not $t.Contains(\"`t\")) {",
                "      $page = [int]$p.Range.Information(3)",
                "      $lines.Add(\"$page`t$t\")",
                "    }",
                "  }",
                f"  [System.IO.File]::WriteAllLines('{listing_s}', $lines, [System.Text.UTF8Encoding]::new($false))",
                "} finally {",
                "  if ($doc -ne $null) { $doc.Close(0) }",
                "  if ($word -ne $null) { $word.Quit() }",
                "}",
            ]
        ),
        encoding="utf-8-sig",
    )
    subprocess.run(["powershell", "-NoProfile", "-File", str(script)], check=True)
    found: dict[str, str] = {}
    total = 0
    wanted = {title for title, _indent in HEADINGS}
    for line in listing.read_text(encoding="utf-8").splitlines():
        if "\t" not in line:
            continue
        page, text = line.split("\t", 1)
        if page == "TOTAL":
            total = int(text.strip())
            continue
        text = text.strip()
        if text in wanted and text not in found:
            found[text] = page.strip()
    return found, total


def main() -> None:
    pages = {title: "…" for title, _indent in HEADINGS}
    total = 0
    for _ in range(3):
        doc = clone_title()
        build_body(doc, pages)
        save_doc(doc)
        measured, total = measure_pages()
        missing = [title for title, _indent in HEADINGS if title not in measured]
        if missing:
            raise SystemExit("Не найдены страницы: " + ", ".join(missing))
        if measured == pages:
            break
        pages = measured
    print("PAGES", total)
    print(pages)
    if total > PAGE_LIMIT:
        raise SystemExit(f"ЛР6 длиннее {PAGE_LIMIT} страниц: {total}")
    print(OUT)


if __name__ == "__main__":
    main()
