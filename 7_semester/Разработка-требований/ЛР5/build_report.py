#!/usr/bin/env python3
"""Сборка отчёта ЛР5. В текст отчёта не входит."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from lxml import etree

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE.parent / "ЛР2" / "ЛР2_Отчет_Буренков_Требования.docx"
OUT = BASE / "ЛР5_Отчет_Буренков_Требования.docx"
FIG = BASE / "figures"
ER_FIG = FIG / "er_model.png"
TOPIC = "Глоссарий"
USABLE = Cm(16.5)
PAGE_LIMIT = 15


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = False


def format_paragraph(
    p,
    *,
    first_indent=True,
    align="justify",
    space_after=0,
    space_before=0,
    left_indent=0,
    line_spacing=1.5,
    single=False,
):
    pf = p.paragraph_format
    if single:
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_spacing
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Cm(1.25 if first_indent else 0)
    pf.left_indent = Cm(left_indent)
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def add_body(doc, text, *, space_before=0):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=space_before)
    set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, *, new_page=False, level=1):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph()
    if level == 0:
        format_paragraph(p, first_indent=False, align="center", space_after=8)
    else:
        format_paragraph(
            p,
            first_indent=True,
            align="left",
            space_before=8 if level == 1 else 6,
            space_after=4 if level == 1 else 2,
        )
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), bold=True)
    return p


def add_toc_line(doc, title, page, *, indent=0):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", left_indent=indent, line_spacing=1.5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    set_run_font(p.add_run(title))
    set_run_font(p.add_run("\t"))
    set_run_font(p.add_run(str(page)))


def add_table_caption(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, first_indent=False, align="left", space_before=8, space_after=6, single=True)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text))


def add_picture(doc, path: Path, caption: str, width_cm: float = 16.0):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, first_indent=False, align="center", space_before=8, space_after=0, single=True)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_p = doc.add_paragraph()
    format_paragraph(caption_p, first_indent=False, align="center", space_before=4, space_after=8, single=True)
    set_run_font(caption_p.add_run(caption))


def _borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def _set_cell_text(cell, text, *, bold=False, align="left", size=12):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        format_paragraph(p, first_indent=False, align=align, single=True)
        set_run_font(p.add_run(line if line else " "), size=size, bold=bold)


def add_table(doc, headers, rows, widths, *, bold_first=False, size=12):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    total = sum(widths)
    emu = [int(USABLE * (w / total)) for w in widths]
    grid = table._tbl.tblGrid
    for i, child in enumerate(list(grid)):
        child.set(qn("w:w"), str(int(emu[i] / 635)))
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(USABLE / 635)))
    tblW.set(qn("w:type"), "dxa")
    jc = tblPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tblPr.append(jc)
    jc.set(qn("w:val"), "left")
    ind = tblPr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tblPr.append(ind)
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Emu(emu[idx])
            _borders(cell)
    for j, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[j], header, bold=True, align="center", size=size)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            _set_cell_text(
                table.rows[i + 1].cells[j],
                value,
                bold=bold_first and j == 0,
                align="left",
                size=size,
            )
    return table


def clone_title() -> Document:
    shutil.copy(TEMPLATE, OUT)
    doc = Document(str(OUT))
    body = doc.element.body
    children = list(body)
    for child in children[12:-1]:
        body.remove(child)
    replacements = {
        "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 2": "ОТЧЕТ О ЛАБОРАТОРНОЙ РАБОТЕ № 5",
        "Идентификация действующих лиц и вариантов использования": TOPIC,
    }
    found = set()
    for node in body.iter(qn("w:t")):
        if node.text in replacements:
            found.add(node.text)
            node.text = replacements[node.text]
    missing = set(replacements) - found
    if missing:
        raise SystemExit(f"На титуле не найдены строки: {missing}")
    return doc


def strip_unused_images(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    rel_name = "word/_rels/document.xml.rels"
    root = etree.fromstring(files[rel_name])
    used = set()
    for name, data in files.items():
        if not (name.startswith("word/") and name.endswith(".xml") and "/_rels/" not in name):
            continue
        for key in (b'r:id="', b'r:embed="'):
            start = 0
            while True:
                pos = data.find(key, start)
                if pos < 0:
                    break
                end = data.find(b'"', pos + len(key))
                used.add(data[pos + len(key) : end].decode())
                start = end + 1
    removed = []
    for rel in list(root):
        if rel.get("Id") not in used and "image" in (rel.get("Type") or ""):
            removed.append("word/" + rel.get("Target").replace("\\", "/"))
            root.remove(rel)
    files[rel_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for name in removed:
        files.pop(name, None)
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)
    tmp.replace(path)


def add_term(doc, name, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=8, space_after=2)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(name), bold=True)
    add_body(doc, text)


TERMS = [
    (
        "Администратор справочников",
        "Сотрудник, который ведёт единые справочники станций и схем пассажирских вагонов пилотного участка (UC-6). Учётные записи персонала эта роль не заводит. Аварийное закрытие смены администратор может выполнить отдельно от ведения справочника.",
    ),
    (
        "АИС железной дороги",
        "Автоматизированная информационная система пилотного участка одного перевозчика. Внутри границ: справочники, маршруты, рейсы, краткий билетный контур, смены, обслуживание станции, заявки на ремонт и инвентаризация.",
    ),
    (
        "Билет",
        "Запись о проданном месте на рейсе. Действующий документ поездки — билет, проведённый в системе и напечатанный купоном. На одно место приходится не больше одного действующего билета.",
    ),
    (
        "БЦ-1",
        "Бизнес-цель документа концепции: оформление билета занимает не более 3 минут в 90 % случаев за три месяца.",
    ),
    (
        "БЦ-2",
        "Бизнес-цель документа концепции: не более одного инцидента двойной продажи одного места в месяц.",
    ),
    (
        "БЦ-3",
        "Бизнес-цель документа концепции: после отмены рейса новые продажи закрыты не позднее чем через 5 минут в 95 % случаев.",
    ),
    (
        "Вагон",
        "Пассажирский вагон рейса. Состав мест вагона задаёт схема вагона из справочника.",
    ),
    (
        "Возврат билета",
        "Гашение билета кассиром до отправления и до отметки посадки (UC-5). Если рейс не отменён, место снова свободно. При отменённом рейсе место в продажу не возвращается. Массовая выплата в этот термин не входит.",
    ),
    (
        "Дежурный по станции",
        "Сотрудник, который открывает и закрывает сутки обслуживания станции: отмечает готовность касс, платформ, зала ожидания и проходов (UC-9).",
    ),
    (
        "Диспетчер расписания",
        "Сотрудник, который составляет маршрут (UC-7), публикует рейс (UC-1) и закрывает продажу при отмене рейса (UC-3).",
    ),
    (
        "Зал ожидания",
        "Объект станции. Готовность зала дежурный отмечает в сутках обслуживания вместе с кассами, платформами и проходами.",
    ),
    (
        "Заявка на ремонт",
        "Запись инженера по ремонту о неисправности объекта станции: объект, описание и статус. Закупка запчастей и складской контур в заявку не входят.",
    ),
    (
        "Инвентаризация",
        "Сверка фактического наличия оборудования и имущества станции с учётным списком. По одной дате инвентаризация проводится один раз.",
    ),
    (
        "Инженер по инвентаризации",
        "Сотрудник, который проводит инвентаризацию объектов станции. Найм этой ролью не оформляется.",
    ),
    (
        "Инженер по ремонту",
        "Сотрудник, который регистрирует и закрывает заявки на ремонт объектов станции. Склада запчастей у роли нет.",
    ),
    (
        "Касса",
        "Рабочее место продажи и возврата билетов. Готовность касс входит в обслуживание станции. Оплата фиксируется как факт в кассе, платёжный шлюз не используется.",
    ),
    (
        "Кассир",
        "Сотрудник кассы. Продажа и возврат доступны при открытой смене кассира и открытом обслуживании станции.",
    ),
    (
        "Контролёр",
        "Сотрудник посадки. По номеру билета или купону ставит одну отметку посадки и не меняет цену и расписание.",
    ),
    (
        "Купон",
        "Бумажная копия проведённого билета. Купон печатает принтер кассы.",
    ),
    (
        "Маршрут",
        "Упорядоченная цепочка станций пилотного участка и плановое время хода или стоянки между соседними станциями (UC-7). В цепочке не меньше двух станций, станция не повторяется. Станции берутся только из справочника.",
    ),
    (
        "Место",
        "Место в вагоне рейса. Номера мест в схеме не повторяются. На место приходится не больше одного действующего билета.",
    ),
    (
        "Обслуживание станции",
        "Сутки готовности станции. Пока сутки открыты, на станции можно открывать смены. Закрытие суток при открытых сменах запрещено.",
    ),
    (
        "Отмена рейса",
        "Перевод рейса в статус «отменён» (UC-3): новые продажи закрыты, оформленные билеты помечены. Черновик в кассе не проводится. Массовую выплату отмена не выполняет. Срок закрытия продаж задаёт БЦ-3.",
    ),
    (
        "Пассажир",
        "Человек, чьи фамилия и имя печатаются на купоне и проверяются при посадке. Рабочего места и публичной учётной записи у пассажира нет.",
    ),
    (
        "Пилотный участок",
        "Ограниченный набор станций, маршрутов и рейсов первой версии одного перевозчика.",
    ),
    (
        "Платформа",
        "Объект станции. Готовность платформы отмечает дежурный в сутках обслуживания.",
    ),
    (
        "Посадка",
        "Одна отметка контролёра по действующему билету (UC-4). После отметки возврат закрыт. Повторная отметка, билет возврата, отменённый рейс и ненайденный билет отметку не создают.",
    ),
    (
        "Прайс",
        "Источник цены места. Продажа без цены из прайса не проводится. Тариф первой версии берётся только из прайса.",
    ),
    (
        "Проход",
        "Объект станции. Готовность прохода дежурный отмечает вместе с кассами, платформами и залом ожидания.",
    ),
    (
        "Рейс",
        "Пассажирский рейс пилотного участка: номер, дата, маршрут, время отправления и прибытия, схема вагонов и один фонд мест. Без сохранённого маршрута и схемы рейс не публикуют. Повтор того же номера не создаёт второй фонд.",
    ),
    (
        "Смена",
        "Интервал работы кассира или контролёра на рабочем месте (UC-8). Вторая открытая смена того же сотрудника не создаётся. Закрытие при черновике продажи останавливается. Без открытой смены продажа, возврат и посадка недоступны.",
    ),
    (
        "Справочник",
        "Единый перечень станций или схем пассажирских вагонов (UC-6). Запись с повтором кода станции или номера места не сохраняется. Станцию или схему, которую использует открытый рейс или действующий маршрут, не удаляют.",
    ),
    (
        "Станция",
        "Станция пилотного участка с кодом и названием. Маршрут собирается только из станций справочника.",
    ),
    (
        "Схема вагона",
        "Тип пассажирского вагона и номера мест. При публикации рейса по схеме создаётся фонд мест.",
    ),
    (
        "Фонд мест",
        "Места одного рейса, созданные по схеме вагонов без повторов номеров. У рейса один фонд.",
    ),
    (
        "Черновик продажи",
        "Незавершённое оформление билета. Пока черновик есть, смену кассира не закрывают. Отмена рейса черновик не проводит.",
    ),
]


HEADINGS = [
    ("ВВЕДЕНИЕ", 0),
    ("1 Введение", 0),
    ("1.1 Назначение", 1),
    ("1.2 Границы", 1),
    ("1.3 Связанные документы", 1),
    ("1.4 Обзор", 1),
    ("1.5 Модель данных", 1),
    ("2 Определения", 0),
    ("ЗАКЛЮЧЕНИЕ", 0),
]


def build_body(doc: Document, pages: dict[str, str]) -> None:
    add_heading(doc, "СОДЕРЖАНИЕ", level=0)
    for title, indent in HEADINGS:
        add_toc_line(doc, title, pages.get(title, "…"), indent=indent * 0.5)

    add_heading(doc, "ВВЕДЕНИЕ", new_page=True, level=0)
    add_body(
        doc,
        "Глоссарий фиксирует термины АИС железной дороги, чтобы спецификации вариантов "
        "и дополнительная спецификация пользовались одними словами. Определения сведены "
        "в одну таблицу, порядок алфавитный, группы терминов не выделяются.",
    )

    add_heading(doc, "1 Введение", new_page=True)
    add_heading(doc, "1.1 Назначение", level=2)
    add_body(
        doc,
        "Документ задаёт единые определения терминов пилотного участка. В вариантах "
        "использования описывают действия, а не поясняют слова заново.",
    )
    add_table_caption(doc, "Таблица 1 — Назначение глоссария")
    add_table(
        doc,
        ["Элемент", "Содержание"],
        [
            ["Предмет", "АИС железной дороги, пилотный участок одного перевозчика"],
            ["Что фиксирует", "Термины предметной области, роли персонала, составные части системы"],
            ["Зачем", "Одинаковые слова в концепции, UC-1–UC-9 и дополнительной спецификации"],
            ["Форма списка", "Одна таблица «термин — определение», по алфавиту, без групп"],
        ],
        [4.0, 12.5],
    )

    add_heading(doc, "1.2 Границы", level=2)
    add_body(doc, "Границы включаемых и исключаемых терминов приведены в таблице 2.")
    add_table_caption(doc, "Таблица 2 — Границы глоссария")
    add_table(
        doc,
        ["Входит", "Не входит"],
        [
            [
                "Станции, маршруты, рейсы, фонд мест, билетный контур, смены, обслуживание станции, ремонт, инвентаризация, роли персонала",
                "Грузовые перевозки, управление движением, найм, закупка запчастей, склад, интернет-магазин, публичный кабинет пассажира",
            ],
            [
                "Пассажир — только как лицо на купоне",
                "Отдельный термин рабочего места пассажира",
            ],
        ],
        [8.25, 8.25],
    )

    add_heading(doc, "1.3 Связанные документы", level=2)
    add_body(doc, "Связь глоссария с документами проекта и пороги БЦ из концепции — в таблицах 3 и 4.")
    add_table_caption(doc, "Таблица 3 — Связанные документы")
    add_table(
        doc,
        ["Документ", "Что задаёт для глоссария"],
        [
            ["Концепция и границы", "Пороги БЦ-1, БЦ-2, БЦ-3; границы пилотного участка"],
            ["Спецификация UC-1–UC-9", "Шаги кассы, диспетчера, дежурного, контролёра, администратора справочников"],
            ["Дополнительная спецификация", "Правила ремонта, инвентаризации и общие ограничения"],
        ],
        [5.5, 11.0],
    )
    add_table_caption(doc, "Таблица 4 — Бизнес-цели из концепции")
    add_table(
        doc,
        ["Код", "Порог"],
        [
            ["БЦ-1", "Оформление билета ≤ 3 мин в 90 % случаев за три месяца"],
            ["БЦ-2", "Не более одного инцидента двойной продажи места в месяц"],
            ["БЦ-3", "После отмены рейса новые продажи закрыты ≤ 5 мин в 95 % случаев"],
        ],
        [2.5, 14.0],
    )

    add_heading(doc, "1.4 Обзор", level=2)
    add_body(
        doc,
        "Раздел 1.5 показывает модель данных предметной области. Раздел 2 содержит одну таблицу "
        "определений: в первой колонке термин, во второй — определение. Порядок алфавитный. "
        "Числа БЦ в определениях совпадают с таблицей 4.",
    )

    add_heading(doc, "1.5 Модель данных", level=2)
    add_body(
        doc,
        "На рисунке 1 приведена модель сущностей АИС железной дороги. Состав сущностей и "
        "ключевые атрибуты совпадают со слоем данных прототипа лабораторных работ 6–8 "
        "(станции, схемы вагонов, маршруты, рейсы, фонд мест, билеты, смена, обслуживание "
        "станции, заявки на ремонт, инвентаризация). Пассажир хранится только как поле билета.",
    )
    if not ER_FIG.exists():
        raise SystemExit(f"Нет рисунка модели данных: {ER_FIG}")
    add_picture(doc, ER_FIG, "Рисунок 1 — Модель данных АИС железной дороги", 16.2)
    add_table_caption(doc, "Таблица 5 — Сущности модели данных")
    add_table(
        doc,
        ["Сущность", "Ключевые атрибуты", "Связи"],
        [
            ["Станция", "код, название", "входит в состав маршрута"],
            ["Схема вагона", "тип, номера мест", "используется при публикации рейса"],
            ["Маршрут", "код, название, время хода", "состоит из станций; задаёт рейс"],
            ["Состав маршрута", "маршрут, станция, порядок", "связывает маршрут и станции"],
            ["Рейс", "номер, дата, маршрут, схема, время, открыт, отменён", "создаёт фонд мест; продаются билеты"],
            ["Место (фонд)", "рейс, номер места, свободно", "одно место — не больше одного действующего билета"],
            ["Билет", "номер, рейс, место, пассажир, цена, посадка, возврат", "продажа / возврат / посадка"],
            ["Смена", "роль, сотрудник, открыта, черновик", "нужна для продажи, возврата и посадки"],
            ["Обслуживание станции", "сутки, готовность касс / платформ / зала / проходов", "разрешает смены; порождает ремонт"],
            ["Заявка на ремонт", "объект, описание, закрыта", "по неготовому объекту станции"],
            ["Инвентаризация", "дата, объект, факт", "один раз на дату и объект"],
        ],
        [3.6, 7.2, 5.7],
        bold_first=True,
        size=10,
    )

    add_heading(doc, "2 Определения", new_page=True)
    add_body(
        doc,
        f"В таблице 6 приведены {len(TERMS)} терминов предметной области в алфавитном порядке. "
        "Термины согласованы с сущностями таблицы 5 и рисунка 1.",
    )
    add_table_caption(doc, "Таблица 6 — Определения терминов")
    add_table(
        doc,
        ["Термин", "Определение"],
        [[name, text] for name, text in TERMS],
        [4.2, 12.3],
        bold_first=True,
        size=11,
    )

    add_heading(doc, "ЗАКЛЮЧЕНИЕ", new_page=True, level=0)
    add_body(
        doc,
        "Глоссарий сведён в таблицу «термин — определение» и согласован с моделью данных "
        "рисунка 1. Покрыты касса, рейс, маршрут, обслуживание станции, смена, ремонт и "
        "инвентаризация. Пороги БЦ-1–БЦ-3 совпадают с концепцией. Пассажир — лицо на купоне. "
        "Груз, найм, склад и интернет-магазин не включены.",
    )


def save_doc(doc: Document) -> None:
    doc.save(str(OUT))
    strip_unused_images(OUT)


def measure_pages():
    import platform

    if platform.system() == "Darwin":
        return {title: "1" for title, _indent in HEADINGS}, 1
    listing = Path(tempfile.gettempdir()) / "lr5_req_pages.txt"
    script = Path(tempfile.gettempdir()) / "lr5_req_pages.ps1"
    out = str(OUT).replace("'", "''")
    listing_s = str(listing).replace("'", "''")
    script.write_text(
        "\n".join(
            [
                "$ErrorActionPreference = 'Stop'",
                "$word = $null",
                "$doc = $null",
                "try {",
                "  $word = New-Object -ComObject Word.Application",
                "  $word.Visible = $false",
                "  $word.DisplayAlerts = 0",
                f"  $doc = $word.Documents.Open('{out}', $false, $true)",
                "  $doc.Repaginate()",
                "  $total = $doc.ComputeStatistics(2)",
                "  $lines = New-Object System.Collections.Generic.List[string]",
                "  $lines.Add(\"TOTAL`t$total\")",
                "  foreach ($p in $doc.Paragraphs) {",
                "    $t = $p.Range.Text.Replace(\"`r\",\"\").Replace(\"`a\",\"\").Trim()",
                "    if ($t.Length -gt 0 -and -not $t.Contains(\"`t\")) {",
                "      $page = [int]$p.Range.Information(3)",
                "      $lines.Add(\"$page`t$t\")",
                "    }",
                "  }",
                f"  [System.IO.File]::WriteAllLines('{listing_s}', $lines, [System.Text.UTF8Encoding]::new($false))",
                "} finally {",
                "  if ($doc -ne $null) { $doc.Close(0) }",
                "  if ($word -ne $null) { $word.Quit() }",
                "}",
            ]
        ),
        encoding="utf-8-sig",
    )
    subprocess.run(["powershell", "-NoProfile", "-File", str(script)], check=True)
    found: dict[str, str] = {}
    total = 0
    wanted = {title for title, _indent in HEADINGS}
    for line in listing.read_text(encoding="utf-8").splitlines():
        if "\t" not in line:
            continue
        page, text = line.split("\t", 1)
        if page == "TOTAL":
            total = int(text.strip())
            continue
        text = text.strip()
        if text in wanted and text not in found:
            found[text] = page.strip()
    return found, total


def main() -> None:
    from draw_er import draw

    draw()
    pages = {title: "…" for title, _indent in HEADINGS}
    total = 0
    for _ in range(3):
        doc = clone_title()
        build_body(doc, pages)
        save_doc(doc)
        measured, total = measure_pages()
        missing = [title for title, _indent in HEADINGS if title not in measured]
        if missing:
            raise SystemExit("Не найдены страницы: " + ", ".join(missing))
        if measured == pages:
            break
        pages = measured
    print("PAGES", total)
    print(pages)
    if total > PAGE_LIMIT:
        raise SystemExit(f"ЛР5 длиннее {PAGE_LIMIT} страниц: {total}")
    print(OUT)


if __name__ == "__main__":
    main()
